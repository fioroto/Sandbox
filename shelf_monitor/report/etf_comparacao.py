"""Planilha da comparacao entre a captura e a exportacao da Comdinheiro.

"Resumo": uma linha por campo, com quantos batem e por que divergem.
"Detalhe": uma linha por ETF e campo, com os dois valores e a diferenca.
"Cobertura": os ETFs que so um dos lados tem.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from shelf_monitor.market.etf_bruto.comparacao import (
    DIFERENTE,
    IGUAL,
    PROXIMO,
    SO_CAPTURA,
    SO_COMDINHEIRO,
    Comparacao,
    resumo,
)
from shelf_monitor.report.shared.fontes import FONTES_ETF_BRUTO
from shelf_monitor.texto import data_br

FONT = "Arial"
_TITULO = Font(name=FONT, size=13, bold=True)
_NOTA = Font(name=FONT, size=9, italic=True, color="595959")
_CORPO = Font(name=FONT, size=10)
_CABECALHO = Font(name=FONT, size=10, bold=True, color="FFFFFF")
_FUNDO_CABECALHO = PatternFill("solid", fgColor="1F3864")
_LINK = Font(name=FONT, size=10, color="1A5FB4", underline="single")
_CORES = {
    IGUAL: PatternFill("solid", fgColor="E2EFDA"),
    PROXIMO: PatternFill("solid", fgColor="FFF2CC"),
    DIFERENTE: PatternFill("solid", fgColor="FCE4D6"),
}


def _cabecalho(ws, linha: int, colunas: tuple[tuple[str, int], ...]) -> None:
    for indice, (titulo, largura) in enumerate(colunas, start=1):
        celula = ws.cell(row=linha, column=indice, value=titulo)
        celula.font, celula.fill = _CABECALHO, _FUNDO_CABECALHO
        celula.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(indice)].width = largura


def _valor(valor: object) -> object:
    if isinstance(valor, float):
        return round(valor, 6)
    return valor


def escrever(c: Comparacao, caminho: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    analise = data_br(c.exportacao.data_analise) if c.exportacao.data_analise else "-"
    ws.cell(row=1, column=1, value="Captura x Comdinheiro").font = _TITULO
    ws.cell(row=2, column=1, value=(
        f"Data da análise da Comdinheiro: {analise}; os números da captura foram refeitos nessa data, com a janela "
        "de dias úteis de cada linha da Comdinheiro (calendário ANBIMA: 24/12 e 31/12 contam como dia útil). "
        f"ETFs nos dois lados: {c.em_ambos}; só na Comdinheiro: {len(c.so_comdinheiro)}; só na captura "
        f"(listados): {len(c.so_captura)}."
    )).font = _NOTA
    ws.cell(row=3, column=1, value="Arquivos: " + "; ".join(Path(a).name for a in c.exportacao.arquivos)).font = _NOTA
    _cabecalho(ws, 5, (("Seção", 18), ("Campo", 30), ("Comparáveis", 12), ("Iguais", 9), ("Próximos", 10),
                       ("Diferentes", 11), ("Só Comdinheiro", 14), ("Só captura", 11), ("Diferença mediana", 14),
                       ("Unidade", 9), ("Por que pode divergir", 80)))
    for numero, linha in enumerate(resumo(c), start=6):
        valores = (linha["secao"], linha["campo"], linha["comparaveis"], linha[IGUAL], linha[PROXIMO], linha[DIFERENTE],
                   linha[SO_COMDINHEIRO], linha[SO_CAPTURA], _valor(linha["diferenca_mediana"]), linha["unidade"], linha["porque"])
        for indice, valor in enumerate(valores, start=1):
            celula = ws.cell(row=numero, column=indice, value=valor)
            celula.font = _CORPO
        if linha["comparaveis"]:
            ws.cell(row=numero, column=4).fill = _CORES[IGUAL]
            if linha[DIFERENTE]:
                ws.cell(row=numero, column=6).fill = _CORES[DIFERENTE]
    ws.freeze_panes = "A6"

    detalhe = wb.create_sheet("Detalhe")
    _cabecalho(detalhe, 1, (("ETF", 10), ("Seção", 18), ("Campo", 28), ("Captura", 30), ("Comdinheiro", 30),
                            ("Diferença (captura − Comdinheiro)", 16), ("Situação", 14), ("Observação", 60)))
    for numero, d in enumerate(sorted(c.diferencas, key=lambda x: (x.ticker, x.secao, x.campo)), start=2):
        valores = (d.ticker, d.secao, d.campo, _valor(d.captura), _valor(d.comdinheiro), _valor(d.diferenca), d.situacao, d.observacao)
        for indice, valor in enumerate(valores, start=1):
            detalhe.cell(row=numero, column=indice, value=valor).font = _CORPO
        if d.situacao in _CORES:
            detalhe.cell(row=numero, column=7).fill = _CORES[d.situacao]
    detalhe.freeze_panes = "B2"
    detalhe.auto_filter.ref = f"A1:H{max(1, len(c.diferencas) + 1)}"

    cobertura = wb.create_sheet("Cobertura")
    _cabecalho(cobertura, 1, (("ETF", 12), ("Onde está", 18), ("Por quê", 70)))
    linha = 2
    for rotulo, lista in (("só Comdinheiro", c.so_comdinheiro), ("só captura", c.so_captura)):
        for ticker, motivo in lista:
            for indice, valor in enumerate((ticker, rotulo, motivo), start=1):
                cobertura.cell(row=linha, column=indice, value=valor).font = _CORPO
            linha += 1

    fontes = wb.create_sheet("Fontes")
    _cabecalho(fontes, 1, (("Fonte", 52), ("O que entra", 70), ("Endereço para consulta", 90)))
    for numero, fonte in enumerate(FONTES_ETF_BRUTO, start=2):
        nome = fontes.cell(row=numero, column=1, value=fonte.nome)
        nome.hyperlink, nome.font = fonte.url, _LINK
        fontes.cell(row=numero, column=2, value=fonte.uso).font = _CORPO
        url = fontes.cell(row=numero, column=3, value=fonte.url)
        url.hyperlink, url.font = fonte.url, _LINK

    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho
