"""Relatórios e exportação de dados de captação semanal de ETFs.

Gera tabelas ricas no terminal via `rich` e cria planilhas Excel formatadas (.xlsx)
com matrizes dinâmicas por categoria, por ETF e por semana ISO.
"""
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rich.console import Console
from rich.table import Table

from shelf_monitor.market.etf import EtfMetadata, catalog_counts
from shelf_monitor.market.etf_captacao_mensal import (
    DETALHE,
    ORDEM_DA_LEGENDA,
    ORIGENS,
    ROTULOS,
    SIMBOLOS,
    legenda,
    prazos_captacao,
    reais_com_sinal,
    rotulo,
)
from shelf_monitor.market.etf_competicao import curto, nome_com_ticker, pct, pp, prazos_etf, reais
from shelf_monitor.market.etf_flow import CategoryWeeklySummary, EtfWeeklyFlow, FlowCoverage
from shelf_monitor.market.fie_balancete import EtfMonthlyPl, PlCoverage
from shelf_monitor.normalize import format_cnpj
from shelf_monitor.report.etf_captacao import CORES, mes_curto
from shelf_monitor.report.shared.fontes import FONTES_ETF, Fonte
from shelf_monitor.texto import contagem, data_br

_KIND_LABEL = {"etf": "ETF", "feeder": "Feeder"}

FONT = "Arial"
_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
_TITLE_FONT = Font(name=FONT, size=13, bold=True)
_BODY_FONT = Font(name=FONT, size=10)
_NOTE_FONT = Font(name=FONT, size=9, italic=True, color="595959")

_GREEN_FILL = PatternFill("solid", fgColor="E2EFDA")
_RED_FILL = PatternFill("solid", fgColor="FCE4D6")
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(bottom=_THIN)

_MONEY = "#,##0;(#,##0);-"
_PCT = '0.00"%";-;-'
_INT = "0;-;-"


def _format_rentab(flow: EtfWeeklyFlow) -> str:
    """"suspeito" e "-" são coisas diferentes: a primeira é um número que
    existia e foi descartado por implausível; a segunda é ausência de cota
    para calcular. Misturar as duas esconderia justamente o alerta que a
    guarda de sanidade existe para dar.
    """
    if flow.data_suspeita:
        return "suspeito"
    if flow.rentabilidade_semana is None:
        return "-"
    return f"{flow.rentabilidade_semana:+.2f}%".replace(".", ",")


def format_money_millions(val: float, sinal: bool = True) -> str:
    """Formata valor em milhões (ex: R$ +12,4M ou R$ -3,1M).

    `sinal=False` para estoque e magnitude (PL, captacao bruta, resgates): o "+"
    num PL de R$ 14.090,0M faz ler variacao onde ha saldo.
    """
    val_m = val / 1_000_000.0
    if abs(val_m) < 0.01:
        return "R$ 0,0M"
    sign = "+" if val_m > 0 and sinal else ""
    return f"R$ {sign}{val_m:,.1f}M".replace(",", "X").replace(".", ",").replace("X", ".")


def print_etf_summary(
    weekly_flows: list[EtfWeeklyFlow],
    category_summaries: list[CategoryWeeklySummary],
    catalog: dict[str, EtfMetadata],
    coverage: FlowCoverage | None = None,
    source_label: str = "INF_DIARIO",
) -> None:
    """Exibe no terminal a visão resumida da captação de ETFs.

    O catálogo (quantos ETFs existem, por categoria) sai sempre - é a resposta
    a "o monitor enxerga todos os ETFs do mercado?" - mesmo quando não há
    fluxo semanal para mostrar. Sem essa separação, uma falha no download do
    INF_DIARIO parecia "não há ETF nenhum", quando na verdade o catálogo
    estava completo e só a captação é que não veio.

    `source_label` diz de onde o fluxo veio (INF_DIARIO ou Fundos.NET): a
    ressalva de cobertura do INF_DIARIO só faz sentido para ele.
    """
    console = Console()
    is_inf_diario = source_label.upper() == "INF_DIARIO"

    counts = catalog_counts(catalog)
    console.print()
    console.print("[bold cyan]=== CATÁLOGO DE ETFs NA CVM ===[/]")
    console.print(
        f"[dim]{counts['etfs']} ETFs (FIIM) cadastrados e ativos, "
        f"{counts['feeders']} feeders (FIF que apenas replicam um ETF)[/]"
    )
    if counts["por_categoria"]:
        by_cat = sorted(counts["por_categoria"].items(), key=lambda kv: -kv[1])
        console.print("[dim]Por categoria: " + ", ".join(f"{cat} {n}" for cat, n in by_cat) + "[/]")
    console.print()

    if coverage is not None and coverage.etfs_no_catalogo:
        pct = coverage.etf_coverage_ratio * 100
        style = "yellow" if pct < 50 else "green"
        line = (
            f"[{style}]Cobertura de fluxo ({source_label}): {coverage.etfs_com_fluxo} de "
            f"{coverage.etfs_no_catalogo} ETFs ({pct:.0f}%).[/]"
        )
        if is_inf_diario:
            line += (
                " [dim]O informe diário cobre FI/FIF, não FIIM - a captação de ETF de "
                "verdade vem do Fundos.NET (`etf --probe-sources`).[/]"
            )
        console.print(line)
        if coverage.feeders_no_catalogo:
            console.print(
                f"[dim]Feeders com fluxo: {coverage.feeders_com_fluxo} de "
                f"{coverage.feeders_no_catalogo}[/]"
            )
        if coverage.semanas_suspeitas:
            console.print(
                f"[yellow]{contagem(coverage.semanas_suspeitas, 'semana', 'semanas')} com "
                "rentabilidade descartada por implausível - ver 'Dado suspeito' na "
                "planilha.[/]"
            )
        console.print()

    if not weekly_flows:
        console.print(f"[yellow]Nenhum dado de captação semanal encontrado (fluxo: {source_label}).[/]")
        return

    # Descobrir a última semana disponível
    latest_week = max(f.week_key for f in weekly_flows)
    latest_flows = [f for f in weekly_flows if f.week_key == latest_week]
    latest_cats = [c for c in category_summaries if c.week_key == latest_week]

    console.print(f"[bold cyan]=== FLUXO SEMANAL (Semana {latest_week}) ===[/]")
    console.print(f"[dim]ETFs com movimento na semana: {len(latest_flows)}[/]\n")

    # Tabela 1: Resumo por Categoria
    cat_table = Table(title=f"Captação por Categoria de ETF ({latest_week})", title_style="bold green")
    cat_table.add_column("Categoria")
    cat_table.add_column("ETFs", justify="right")
    cat_table.add_column("Captação Bruta", justify="right")
    cat_table.add_column("Resgates", justify="right")
    # So o Fundos.NET tem o ajuste (o dia em que o declarado nao fecha com a
    # variacao de cotas); no INF_DIARIO a coluna sairia zerada.
    com_ajuste = any(c.ajuste_pl for c in latest_cats)
    if com_ajuste:
        cat_table.add_column("Ajuste pela dif. de PL", justify="right")
    cat_table.add_column("Captação Líquida", justify="right", style="bold")
    cat_table.add_column("PL Total", justify="right")

    for cat in sorted(latest_cats, key=lambda c: c.captacao_liquida, reverse=True):
        style = "green" if cat.captacao_liquida > 0 else ("red" if cat.captacao_liquida < 0 else "dim")
        cat_table.add_row(
            cat.category,
            str(cat.qtd_etfs),
            format_money_millions(cat.captacao_bruta, sinal=False),
            format_money_millions(cat.resgates, sinal=False),
            *([format_money_millions(cat.ajuste_pl)] if com_ajuste else []),
            f"[{style}]{format_money_millions(cat.captacao_liquida)}[/]",
            format_money_millions(cat.pl_total, sinal=False),
        )

    console.print(cat_table)

    # Tabela 2: Top 10 Captação Líquida Positiva
    top_inflow = sorted([f for f in latest_flows if f.captacao_liquida > 0], key=lambda f: f.captacao_liquida, reverse=True)[:10]
    if top_inflow:
        in_table = Table(title=f"\nTop 10 Maior Captação Líquida ({latest_week})", title_style="bold green")
        in_table.add_column("Ticker / Nome", style="cyan")
        in_table.add_column("Tipo")
        in_table.add_column("Categoria")
        in_table.add_column("Gestora")
        in_table.add_column("Captação Líquida", justify="right", style="bold green")
        in_table.add_column("PL Fim", justify="right")
        in_table.add_column("Cotistas", justify="right")
        in_table.add_column("Retorno Semanal", justify="right")

        for f in top_inflow:
            name = f.ticker or f.etf_name[:30]
            in_table.add_row(
                name,
                _KIND_LABEL.get(f.kind, f.kind),
                f.category,
                (f.manager_name or "-")[:25],
                format_money_millions(f.captacao_liquida),
                format_money_millions(f.pl_fim, sinal=False),
                f"{f.cotistas_fim:,}".replace(",", "."),
                _format_rentab(f),
            )
        console.print(in_table)

    # Tabela 3: Top 5 Resgates
    top_outflow = sorted([f for f in latest_flows if f.captacao_liquida < 0], key=lambda f: f.captacao_liquida)[:5]
    if top_outflow:
        out_table = Table(title=f"\nTop 5 Maiores Resgates ({latest_week})", title_style="bold red")
        out_table.add_column("Ticker / Nome", style="cyan")
        out_table.add_column("Tipo")
        out_table.add_column("Categoria")
        out_table.add_column("Gestora")
        out_table.add_column("Captação Líquida", justify="right", style="bold red")
        out_table.add_column("PL Fim", justify="right")
        out_table.add_column("Cotistas", justify="right")
        out_table.add_column("Retorno Semanal", justify="right")

        for f in top_outflow:
            name = f.ticker or f.etf_name[:30]
            out_table.add_row(
                name,
                _KIND_LABEL.get(f.kind, f.kind),
                f.category,
                (f.manager_name or "-")[:25],
                format_money_millions(f.captacao_liquida),
                format_money_millions(f.pl_fim, sinal=False),
                f"{f.cotistas_fim:,}".replace(",", "."),
                _format_rentab(f),
            )
        console.print(out_table)

    _print_valorizacao_ranking(console, latest_flows, latest_week)
    _print_cotistas_ranking(console, latest_flows, latest_week)


def _print_valorizacao_ranking(
    console: Console, flows: list[EtfWeeklyFlow], week: str
) -> None:
    """Ranking de valorizacao: quanto do PL veio de mercado, e nao de captacao.

    Duas medidas lado a lado de proposito. O retorno da cota e a performance por
    unidade (comparavel entre ETFs de tamanhos diferentes); a valorizacao em reais
    e o quanto do patrimonio o mercado somou ou tirou. Um ETF pequeno pode liderar
    a primeira e nao aparecer na segunda - sao perguntas distintas.
    """
    ranked = [f for f in flows if f.rentabilidade_semana is not None]
    if not ranked:
        return
    ranked.sort(key=lambda f: f.rentabilidade_semana or 0.0, reverse=True)

    console.print()
    table = Table(
        title=f"Ranking de Valorização ({week}) - retorno da cota",
        title_style="bold",
    )
    table.add_column("Ticker / Nome", style="cyan")
    table.add_column("Categoria")
    table.add_column("Retorno da cota", justify="right")
    table.add_column("Valorização (R$)", justify="right")
    table.add_column("Captação líquida", justify="right")
    table.add_column("Variação de PL", justify="right")

    def row(flow: EtfWeeklyFlow) -> None:
        style = "green" if (flow.rentabilidade_semana or 0) > 0 else "red"
        table.add_row(
            (flow.ticker or flow.etf_name[:30]) + (" *" if flow.decomposicao_suspeita else ""),
            flow.category,
            f"[{style}]{flow.rentabilidade_semana:+.2f}%[/]".replace(".", ","),
            format_money_millions(flow.valorizacao),
            format_money_millions(flow.captacao_liquida),
            format_money_millions(flow.variacao_pl),
        )

    for flow in ranked[:10]:
        row(flow)
    if len(ranked) > 12:
        table.add_section()
        for flow in ranked[-3:]:
            row(flow)

    console.print(table)
    # A cota do informe e cota patrimonial ex-provento: num ETF que distribui, o
    # retorno medido aqui fica abaixo do retorno total do investidor. O informe
    # diario nao traz provento pago, entao nao ha como reincorporar sem uma
    # segunda fonte - dizer isso e mais util que publicar o numero sem ressalva.
    console.print(
        "[dim]Retorno da cota, ex-provento: o informe diário não publica provento "
        "pago, então um ETF que distribui aparece abaixo do retorno total.[/]"
    )
    duvidosos = [f for f in flows if f.decomposicao_suspeita]
    if duvidosos:
        console.print(
            f"[yellow]* {len(duvidosos)} de {contagem(len(flows), 'ETF', 'ETFs')} com "
            "decomposição suspeita:[/] a valorização por diferença não bate com a que a "
            "cota implica - semana com dia faltando na coleta, não movimento real."
        )


def _print_cotistas_ranking(
    console: Console, flows: list[EtfWeeklyFlow], week: str
) -> None:
    """Quem ganhou e quem perdeu cotistas na semana."""
    # Registros com cotistas desmentidos pelo PL ficam FORA do ranking e do
    # saldo, como o retorno implausivel ja ficava: publicar -51 mil cotistas de um
    # erro de preenchimento e pior que nao publicar numero nenhum.
    descartados = [f for f in flows if f.cotistas_suspeito]
    moved = [f for f in flows if f.variacao_cotistas and not f.cotistas_suspeito]
    if not moved:
        return
    moved.sort(key=lambda f: f.variacao_cotistas, reverse=True)

    console.print()
    table = Table(title=f"Evolução de Cotistas ({week})", title_style="bold")
    table.add_column("Ticker / Nome", style="cyan")
    table.add_column("Categoria")
    table.add_column("Cotistas início", justify="right")
    table.add_column("Cotistas fim", justify="right")
    table.add_column("Variação", justify="right")

    def row(flow: EtfWeeklyFlow) -> None:
        style = "green" if flow.variacao_cotistas > 0 else "red"
        table.add_row(
            flow.ticker or flow.etf_name[:30],
            flow.category,
            f"{flow.cotistas_inicio:,}".replace(",", "."),
            f"{flow.cotistas_fim:,}".replace(",", "."),
            f"[{style}]{flow.variacao_cotistas:+,}[/]".replace(",", "."),
        )

    for flow in moved[:8]:
        row(flow)
    perdas = [f for f in moved if f.variacao_cotistas < 0]
    if perdas:
        table.add_section()
        for flow in perdas[-3:]:
            row(flow)

    total = sum(f.variacao_cotistas for f in moved)
    console.print(table)
    console.print(
        f"[bold]Saldo de cotistas na semana:[/] {total:+,}".replace(",", ".")
        + f" em {contagem(len(moved), 'ETF', 'ETFs')} com movimento"
    )
    if descartados:
        console.print(
            f"[yellow]{contagem(len(descartados), 'ETF descartado', 'ETFs descartados')} "
            "do saldo:[/] variação de cotistas que o PL desmente (erro de preenchimento "
            "na origem)."
        )


def export_etf_csv(weekly_flows: list[EtfWeeklyFlow], output_path: Path) -> Path:
    """Exporta o histórico completo semanal de captação de ETFs para CSV."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    headers = [
        "CNPJ", "Tipo", "Ticker", "Nome ETF", "Categoria", "Gestora", "Semana ISO",
        "Data Fim Semana", "Captação Bruta (RS)", "Resgates (RS)",
        "Ajuste pela diferença de PL (RS)",
        "Captação Líquida (RS)", "PL Início (RS)", "PL Fim (RS)",
        "Variação PL (RS)", "Valorização (RS)",
        "Cotistas Início", "Cotistas Fim", "Variação Cotistas",
        "Rentabilidade Semanal (%)", "Dado Suspeito", "Decomposição Suspeita",
        "Cotistas Suspeito",
    ]

    with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(headers)
        for f in weekly_flows:
            writer.writerow([
                format_cnpj(f.cnpj),
                _KIND_LABEL.get(f.kind, f.kind),
                f.ticker or "",
                f.etf_name,
                f.category,
                f.manager_name or "",
                f.week_key,
                f.week_end_date.strftime("%Y-%m-%d"),
                f.captacao_bruta,
                f.resgates,
                f.ajuste_pl,
                f.captacao_liquida,
                f.pl_inicio,
                f.pl_fim,
                f.variacao_pl,
                f.valorizacao,
                f.cotistas_inicio,
                f.cotistas_fim,
                f.variacao_cotistas,
                f.rentabilidade_semana if f.rentabilidade_semana is not None else "",
                "sim" if f.data_suspeita else "",
                "sim" if f.decomposicao_suspeita else "",
                "sim" if f.cotistas_suspeito else "",
            ])

    return output_path


def export_etf_excel(
    weekly_flows: list[EtfWeeklyFlow],
    category_summaries: list[CategoryWeeklySummary],
    catalog: dict[str, EtfMetadata],
    output_path: Path,
    pl_series: list[EtfMonthlyPl] | None = None,
    fonte: str = "inf_diario",
    competicao: dict | None = None,
    captacao: dict | None = None,
) -> Path:
    """Exporta a planilha Excel completa com matriz semanal por ETF e por categoria.

    `fonte` ("fnet" ou "inf_diario") vai no titulo das abas de fluxo. O padrao
    e o conservador: sem dizer a fonte, a planilha se apresenta como fluxo de
    feeder - antes ela chamava de "Captacao Liquida de ETFs" um fluxo em que
    nenhum ETF aparecia (0 de 249 no INF_DIARIO em 19/09).
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    de_feeder = fonte != "fnet"
    rotulo_fonte = "Fundos.NET" if not de_feeder else "INF_DIARIO"
    nota_feeder = (
        "O INF_DIARIO não cobre ETF (0 de 249 em 19/09): estes números são dos "
        "feeders - FIF que replicam um ETF -, não do ETF. O fluxo do ETF vem do "
        "Fundos.NET (etf --flow-source fnet)."
    )

    # --- Aba 1: Resumo Categorias x Semanas
    ws_cat = wb.active
    ws_cat.title = "Resumo Categorias"

    weeks = sorted({f.week_key for f in weekly_flows})
    categories = sorted({f.category for f in weekly_flows})

    titulo_cat = (
        "Captação Líquida dos FEEDERS de ETF por Categoria (Semana a Semana)"
        if de_feeder
        else "Captação Líquida de ETFs por Categoria (Semana a Semana)"
    )
    ws_cat.cell(row=1, column=1, value=f"{titulo_cat} - {rotulo_fonte}").font = _TITLE_FONT
    ws_cat.cell(row=2, column=1, value=f"Gerado em {datetime.now():%d/%m/%Y %H:%M}").font = _NOTE_FONT
    if de_feeder:
        ws_cat.cell(row=3, column=1, value=nota_feeder).font = _NOTE_FONT

    headers_cat = ["Categoria", *weeks, "Total Acumulado (R$)"]
    for col_idx, h in enumerate(headers_cat, start=1):
        ws_cat.cell(row=4, column=col_idx, value=h)
        ws_cat.column_dimensions[get_column_letter(col_idx)].width = 24 if col_idx == 1 else 18

    # Estilizar cabeçalho
    for col_idx in range(1, len(headers_cat) + 1):
        cell = ws_cat.cell(row=4, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center" if col_idx > 1 else "left")

    cat_map: dict[tuple[str, str], float] = {(c.category, c.week_key): c.captacao_liquida for c in category_summaries}

    row_idx = 5
    for cat in categories:
        ws_cat.cell(row=row_idx, column=1, value=cat).font = _BODY_FONT
        for w_idx, w in enumerate(weeks, start=2):
            val = cat_map.get((cat, w), 0.0)
            cell = ws_cat.cell(row=row_idx, column=w_idx, value=val)
            cell.font = _BODY_FONT
            cell.number_format = _MONEY
            if val > 0:
                cell.fill = _GREEN_FILL
            elif val < 0:
                cell.fill = _RED_FILL

        # Fórmula de Total Acumulado por Linha
        col_start = get_column_letter(2)
        col_end = get_column_letter(len(weeks) + 1)
        tot_cell = ws_cat.cell(row=row_idx, column=len(weeks) + 2, value=f"=SUM({col_start}{row_idx}:{col_end}{row_idx})")
        tot_cell.font = Font(name=FONT, size=10, bold=True)
        tot_cell.number_format = _MONEY
        row_idx += 1

    # --- Aba 2: Matriz Semanal por ETF
    ws_etf = wb.create_sheet("Matriz ETFs")
    titulo_etf = (
        "Captação Líquida por FEEDER de ETF (Semana a Semana)"
        if de_feeder
        else "Captação Líquida por ETF (Semana a Semana)"
    )
    ws_etf.cell(row=1, column=1, value=f"{titulo_etf} - {rotulo_fonte}").font = _TITLE_FONT
    if de_feeder:
        ws_etf.cell(row=2, column=1, value=nota_feeder).font = _NOTE_FONT

    headers_etf = ["Ticker", "Tipo", "Nome", "Categoria", "Gestora", *weeks, "Total Acumulado (R$)"]
    _widths_etf = {1: 14, 2: 10, 3: 36}
    for col_idx, h in enumerate(headers_etf, start=1):
        ws_etf.cell(row=3, column=col_idx, value=h)
        ws_etf.column_dimensions[get_column_letter(col_idx)].width = _widths_etf.get(col_idx, 18)

    for col_idx in range(1, len(headers_etf) + 1):
        cell = ws_etf.cell(row=3, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    etf_cnpjs = sorted(
        {f.cnpj for f in weekly_flows},
        key=lambda c: catalog[c].ticker or catalog[c].name if c in catalog else c,
    )

    etf_week_map: dict[tuple[str, str], float] = {(f.cnpj, f.week_key): f.captacao_liquida for f in weekly_flows}
    etf_kind_map: dict[str, str] = {f.cnpj: f.kind for f in weekly_flows}

    row_idx = 4
    for cnpj in etf_cnpjs:
        meta = catalog.get(cnpj)
        ticker = meta.ticker if meta else "-"
        name = meta.name if meta else cnpj
        category = meta.category if meta else "Outros"
        manager = meta.manager_name if meta else "-"
        kind = _KIND_LABEL.get(etf_kind_map.get(cnpj, "etf"), "ETF")

        ws_etf.cell(row=row_idx, column=1, value=ticker).font = _BODY_FONT
        ws_etf.cell(row=row_idx, column=2, value=kind).font = _BODY_FONT
        ws_etf.cell(row=row_idx, column=3, value=name).font = _BODY_FONT
        ws_etf.cell(row=row_idx, column=4, value=category).font = _BODY_FONT
        ws_etf.cell(row=row_idx, column=5, value=manager).font = _BODY_FONT

        for w_idx, w in enumerate(weeks, start=6):
            val = etf_week_map.get((cnpj, w), 0.0)
            cell = ws_etf.cell(row=row_idx, column=w_idx, value=val)
            cell.font = _BODY_FONT
            cell.number_format = _MONEY
            if val > 0:
                cell.fill = _GREEN_FILL
            elif val < 0:
                cell.fill = _RED_FILL

        col_start = get_column_letter(6)
        col_end = get_column_letter(len(weeks) + 5)
        tot_cell = ws_etf.cell(row=row_idx, column=len(weeks) + 6, value=f"=SUM({col_start}{row_idx}:{col_end}{row_idx})")
        tot_cell.font = Font(name=FONT, size=10, bold=True)
        tot_cell.number_format = _MONEY
        row_idx += 1

    ws_etf.freeze_panes = "F4"

    # --- Aba 3: Detalhamento Completo
    ws_flat = wb.create_sheet("Detalhamento Completo")
    headers_flat = [
        "CNPJ", "Tipo", "Ticker", "Nome ETF", "Categoria", "Gestora", "Semana ISO",
        "Data Fim", "Captação Bruta (R$)", "Resgates (R$)", "Ajuste pela diferença de PL (R$)",
        "Captação Líquida (R$)", "PL Fim (R$)", "Variação PL (R$)", "Valorização (R$)",
        "Cotistas Fim", "Variação Cotistas", "Rentabilidade Semanal (%)",
        "Dado Suspeito", "Decomposição Suspeita", "Cotistas Suspeito",
    ]
    widths_flat = [18, 10, 12, 38, 16, 26, 12, 14, 18, 18, 18, 20, 18, 18, 18, 14, 16, 20, 14, 20, 18]

    for col_idx, (h, w) in enumerate(zip(headers_flat, widths_flat), start=1):
        ws_flat.cell(row=1, column=col_idx, value=h)
        ws_flat.column_dimensions[get_column_letter(col_idx)].width = w

    for col_idx in range(1, len(headers_flat) + 1):
        cell = ws_flat.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    row_idx = 2
    for f in weekly_flows:
        values = [
            format_cnpj(f.cnpj),
            _KIND_LABEL.get(f.kind, f.kind),
            f.ticker or "",
            f.etf_name,
            f.category,
            f.manager_name or "",
            f.week_key,
            f.week_end_date.strftime("%d/%m/%Y"),
            f.captacao_bruta,
            f.resgates,
            f.ajuste_pl,
            f.captacao_liquida,
            f.pl_fim,
            f.variacao_pl,
            f.valorizacao,
            f.cotistas_fim,
            f.variacao_cotistas,
            (f.rentabilidade_semana / 100.0) if f.rentabilidade_semana is not None else None,
            "sim" if f.data_suspeita else "",
            "sim" if f.decomposicao_suspeita else "",
            "sim" if f.cotistas_suspeito else "",
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws_flat.cell(row=row_idx, column=col_idx, value=val)
            cell.font = _BODY_FONT
            cell.border = _BORDER

        # 9-15 sao reais (captacao bruta, resgates, ajuste pela diferenca de PL,
        # liquida, PL fim, variacao de PL, valorizacao); 16-17 contagem de
        # cotistas; 18 percentual.
        for money_col in range(9, 16):
            ws_flat.cell(row=row_idx, column=money_col).number_format = _MONEY
        ws_flat.cell(row=row_idx, column=16).number_format = _INT
        ws_flat.cell(row=row_idx, column=17).number_format = _INT
        ws_flat.cell(row=row_idx, column=18).number_format = _PCT
        row_idx += 1

    ws_flat.freeze_panes = "A2"
    if row_idx > 2:
        ws_flat.auto_filter.ref = f"A1:{get_column_letter(len(headers_flat))}{row_idx - 1}"

    _write_catalog_sheet(wb, catalog)
    if pl_series:
        _write_pl_sheet(wb, pl_series)
    _write_fontes_sheet(wb, FONTES_ETF)
    # A competicao vai para a frente da planilha: e a leitura que decide, e as
    # abas de fluxo semanal viram o detalhe dela.
    if competicao and competicao.get("available"):
        _write_competicao_sheets(wb, competicao)
        wb.active = 0
    # A captacao mes a mes abre a planilha: e o monitoramento de fluxo, agora
    # mensal; o PL da competicao e o fluxo semanal vem depois, como detalhe.
    if captacao and captacao.get("available"):
        _write_captacao_sheets(wb, captacao)
        wb.active = 0

    wb.save(output_path)
    return output_path


def _write_fontes_sheet(wb: Workbook, fontes: tuple[Fonte, ...]) -> None:
    """Aba "Fontes": de onde vem cada número da planilha, com link para consulta."""
    ws = wb.create_sheet("Fontes")
    ws.cell(row=1, column=1, value="Fontes").font = _TITLE_FONT
    ws.cell(
        row=2, column=1,
        value="Onde conferir cada número desta planilha - clique no nome para abrir.",
    ).font = _NOTE_FONT
    for col, (titulo, largura) in enumerate(
        (("Fonte", 48), ("O que entra", 60), ("Endereço para consulta", 90)), start=1
    ):
        cell = ws.cell(row=4, column=col, value=titulo)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        ws.column_dimensions[get_column_letter(col)].width = largura
    link = Font(name=FONT, size=10, color="1A5FB4", underline="single")
    for linha, fonte in enumerate(fontes, start=5):
        nome = ws.cell(row=linha, column=1, value=fonte.nome)
        nome.hyperlink, nome.font = fonte.url, link
        ws.cell(row=linha, column=2, value=fonte.uso).font = _BODY_FONT
        url = ws.cell(row=linha, column=3, value=fonte.url)
        url.hyperlink, url.font = fonte.url, link


def _write_catalog_sheet(wb: Workbook, catalog: dict[str, EtfMetadata]) -> None:
    """Aba "Catálogo FIIM": todo ETF cadastrado na CVM, tenha fluxo ou não.

    É a resposta direta a "o monitor traz todos os ETFs do mercado, incluindo
    os de renda fixa?" - sem depender do INF_DIARIO, que não cobre FIIM (ver
    docstring de market/etf_flow.py). Feeders entram numa seção separada, no
    fim da aba: são relevantes para a prateleira, mas não são o ETF em si.
    """
    ws = wb.create_sheet("Catálogo FIIM")
    ws.cell(row=1, column=1, value="Catálogo de ETFs (FIIM) registrados na CVM").font = _TITLE_FONT
    ws.cell(row=2, column=1, value=f"Gerado em {datetime.now():%d/%m/%Y %H:%M}").font = _NOTE_FONT

    headers = [
        "CNPJ", "Ticker", "Nome", "Categoria", "Benchmark", "Gestora",
        "Situação", "PL do Cadastro (R$)",
    ]
    widths = [18, 12, 42, 16, 20, 30, 20, 20]
    header_row = 4
    for col_idx, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    def write_rows(items: list[EtfMetadata], start_row: int) -> int:
        row_idx = start_row
        for meta in sorted(items, key=lambda m: (m.category, m.ticker or m.name)):
            values = [
                format_cnpj(meta.cnpj),
                meta.ticker or "",
                meta.name,
                meta.category,
                meta.benchmark or "",
                meta.manager_name or "",
                meta.situation or "",
                meta.net_worth,
            ]
            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = _BODY_FONT
                cell.border = _BORDER
            ws.cell(row=row_idx, column=8).number_format = _MONEY
            row_idx += 1
        return row_idx

    etfs = [m for m in catalog.values() if m.is_etf]
    feeders = [m for m in catalog.values() if not m.is_etf]

    row_idx = write_rows(etfs, header_row + 1)
    ws.freeze_panes = f"A{header_row + 1}"
    if row_idx > header_row + 1:
        ws.auto_filter.ref = f"A{header_row}:H{row_idx - 1}"

    if feeders:
        row_idx += 2
        ws.cell(
            row=row_idx, column=1,
            value="Feeders (FIF que apenas replicam um ETF - não somam com o ETF acima)",
        ).font = _NOTE_FONT
        row_idx += 1
        feeder_header_row = row_idx
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=feeder_header_row, column=col_idx, value=h)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
        write_rows(feeders, feeder_header_row + 1)


def print_etf_pl(
    series: list[EtfMonthlyPl],
    coverage: PlCoverage,
    missing: list[str] | None = None,
) -> None:
    """Exibe o PL mensal dos ETFs vindo do balancete de FIE.

    Fica separado de print_etf_summary de proposito: sao fontes com cadencia e
    semantica diferentes (semanal/captacao contra mensal/PL) e juntar as duas
    numa tabela convidaria a somar numeros que nao se somam.
    """
    console = Console()

    if not series:
        console.print(
            "[yellow]Balancete de FIE sem nenhum ETF do catálogo.[/] "
            "Nenhuma competência disponível no período pedido."
        )
        return

    latest = coverage.competencia
    current = [s for s in series if s.competencia == latest]
    comparable = [s for s in current if s.variacao_pl is not None]
    comparable.sort(key=lambda s: -abs(s.variacao_pl or 0))

    label = f"{latest[4:]}/{latest[:4]}" if latest else "?"
    table = Table(
        title=f"PL dos ETFs no balancete de FIE - competência {label}",
        title_style="bold",
    )
    table.add_column("ETF")
    table.add_column("Categoria")
    table.add_column("PL", justify="right")
    table.add_column("Variação no mês", justify="right")
    table.add_column("%", justify="right")

    for item in comparable[:15]:
        variacao = item.variacao_pl or 0.0
        style = "green" if variacao > 0 else "red" if variacao < 0 else ""
        pct = f"{item.variacao_pct:+.1f}%".replace(".", ",") if item.variacao_pct is not None else "-"
        table.add_row(
            (item.ticker and f"{item.ticker} - {item.etf_name}" or item.etf_name)[:44],
            item.category[:18],
            format_money_millions(item.pl, sinal=False),
            f"[{style}]{format_money_millions(variacao)}[/]" if style else format_money_millions(variacao),
            f"[{style}]{pct}[/]" if style else pct,
        )

    console.print()
    console.print(table)
    console.print(
        f"[bold]Cobertura:[/] {coverage.etfs_com_pl} de {coverage.etfs_no_catalogo} "
        f"ETFs do catálogo ({coverage.coverage_ratio:.0%}), "
        f"PL somado {reais(coverage.pl_total)}"
    )
    # A ressalva sai junto do numero, sempre: sem ela a variacao de PL seria lida
    # como captacao - e ela inclui marcacao a mercado. Ver market/fie_balancete.py.
    console.print(
        "[dim]PL = Ativo menos Passivo do balancete, na data-base da competência. "
        "Variação de PL não é captação líquida: inclui valorização. A captação "
        "publicada vem do informe diário do Fundos.NET.[/]"
    )
    if coverage.identidades_suspeitas:
        console.print(
            f"[yellow]{contagem(coverage.identidades_suspeitas, 'ETF', 'ETFs')} com "
            "identidade contábil fora de linha[/] (Ativo menos Passivo não é o grupo 6 "
            "mais receitas e despesas) - sinal de que o plano de contas mudou; confira "
            "market/fie_balancete.py."
        )
    if missing:
        console.print(
            f"[dim]Competências sem balancete publicado: {', '.join(missing)} "
            "(a CVM publica com defasagem).[/]"
        )


def _write_pl_sheet(wb: Workbook, series: list[EtfMonthlyPl]) -> None:
    """Aba "PL mensal (FIE)": a serie de PL por ETF e competencia."""
    ws = wb.create_sheet("PL mensal (FIE)")

    ws.cell(row=1, column=1, value="PL mensal dos ETFs - balancete de Fundos Estruturados (CVM)").font = _TITLE_FONT
    ws.cell(
        row=2, column=1,
        value=(
            "PL = Ativo menos Passivo, na data-base de cada competência (o grupo 6 "
            "do balancete deixa de fora o resultado do exercício em curso). Variação de "
            "PL NÃO é captação líquida: inclui valorização; a captação publicada vem do "
            "informe diário do Fundos.NET."
        ),
    ).font = _NOTE_FONT

    headers = [
        "Competência", "Ticker", "ETF", "CNPJ", "Categoria", "Gestora",
        "PL (R$)", "PL anterior (R$)", "Variação (R$)", "Variação (%)",
    ]
    header_row = 4
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        ws.column_dimensions[get_column_letter(col_idx)].width = 44 if col_idx == 3 else 18

    row_idx = header_row + 1
    for item in series:
        values = [
            item.competencia,
            item.ticker or "",
            item.etf_name,
            format_cnpj(item.cnpj),
            item.category,
            item.manager_name or "",
            item.pl,
            item.pl_anterior,
            item.variacao_pl,
            item.variacao_pct,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = _BODY_FONT
            if col_idx in (7, 8, 9):
                cell.number_format = _MONEY
            elif col_idx == 10 and value is not None:
                cell.number_format = "0.0"
            # Reatribuir cell.fill a si mesmo levanta TypeError (StyleProxy nao
            # e hashable), entao a cor so e tocada quando ha sinal a mostrar.
            if col_idx == 9 and isinstance(value, (int, float)) and value:
                cell.fill = _GREEN_FILL if value > 0 else _RED_FILL
        row_idx += 1

    ws.freeze_panes = f"A{header_row + 1}"
    if row_idx > header_row + 1:
        ws.auto_filter.ref = f"A{header_row}:J{row_idx - 1}"


# --- competicao --------------------------------------------------------------

_SHARE = '0.0%;-0.0%;-'
_PP = '+0.0" pp";-0.0" pp";0.0" pp"'
_LINK_FONT = Font(name=FONT, size=10, color="1A5FB4", underline="single")


def print_etf_competicao(bloco: dict) -> None:
    """Quem lidera cada mercado de ETF, e onde a casa em foco compete."""
    console = Console()
    console.print()
    console.print("[bold]Competição em ETFs[/] - gestoras e segmentos, pelo PL do balancete")
    if not bloco.get("available"):
        console.print(f"   [yellow]{bloco.get('reason')}[/]")
        return
    for prazo in prazos_etf(bloco):
        console.print(f"   [dim]{prazo}[/]")
    for frase in bloco["leitura"]:
        console.print(f"   - {frase}")

    tabela = Table(title="Gestoras de ETF", title_style="bold")
    for titulo, alinhamento in (
        ("#", "right"), ("Gestora", "left"), ("ETFs", "right"), ("PL", "right"),
        ("Participação", "right"), ("Variação (pp)", "right"), ("Cotistas", "right"),
        ("Lidera", "left"),
    ):
        tabela.add_column(titulo, justify=alinhamento)
    for g in bloco["gestoras"][:12]:
        estilo = "bold" if g["foco"] else None
        tabela.add_row(
            str(g["posicao"]), g["gestora"], str(g["etfs"]), reais(g["pl"]),
            pct(g["share"]), pp(g["delta_share_pp"]), pct(g["share_cotistas"]),
            ", ".join(g["lidera"][:3]) + ("..." if len(g["lidera"]) > 3 else ""),
            style=estilo,
        )
    console.print(tabela)

    foco = bloco.get("foco") or "casa em foco"
    tabela = Table(title=f"Segmentos - onde a {foco} está", title_style="bold")
    for titulo, alinhamento in (
        ("Segmento", "left"), ("PL", "right"), ("Líder", "left"), (foco, "right"),
        ("Posição", "right"), ("Variação (pp)", "right"), ("Lançamentos", "left"),
    ):
        tabela.add_column(titulo, justify=alinhamento)
    for s in bloco["segmentos"]:
        if s["pl"] <= 0 and not s["lancando"]:
            continue
        lancamentos = sorted({curto(p["gestora"]) for p in s["lancando"]})
        tabela.add_row(
            s["segmento"], reais(s["pl"]) if s["pl"] else "-",
            f"{curto(s['lider'])} {pct(s['share_lider'], 0)}" if s["lider"] else "-",
            pct(s["foco_share"]) if s["foco_share"] is not None else _STATUS[s["foco_status"]],
            str(s["foco_posicao"] or "-"), pp(s["foco_delta_pp"]),
            ", ".join(lancamentos) or "-",
        )
    console.print(tabela)


_STATUS = {
    "lidera": "lidera",
    "presente": "presente",
    "entrando": "lançando",
    "ausente": "ausente",
}


def _tabela(
    ws, linha: int, colunas: list[tuple[str, int, str | None]], linhas: list[list],
    filtro: bool = True,
) -> int:
    """Cabecalho estilizado + linhas com formato por coluna. Devolve a proxima linha."""
    for col, (titulo, largura, _) in enumerate(colunas, start=1):
        cell = ws.cell(row=linha, column=col, value=titulo)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        atual = ws.column_dimensions[get_column_letter(col)].width or 0
        ws.column_dimensions[get_column_letter(col)].width = max(largura, atual)
    cabecalho = linha
    linha += 1
    for valores in linhas:
        for col, (valor, (_, _, formato)) in enumerate(zip(valores, colunas), start=1):
            cell = ws.cell(row=linha, column=col, value=valor)
            cell.font = _BODY_FONT
            cell.border = _BORDER
            if formato:
                cell.number_format = formato
        linha += 1
    if filtro and linha > cabecalho + 1:
        ws.auto_filter.ref = f"A{cabecalho}:{get_column_letter(len(colunas))}{linha - 1}"
    return linha + 1


def _cabecalho_competicao(ws, titulo: str, bloco: dict) -> int:
    ws.cell(row=1, column=1, value=titulo).font = _TITLE_FONT
    linha = 2
    for prazo in prazos_etf(bloco):
        ws.cell(row=linha, column=1, value=prazo[0].upper() + prazo[1:]).font = _NOTE_FONT
        linha += 1
    return linha + 1


def _write_competicao_sheets(wb: Workbook, bloco: dict) -> None:
    """Abas de competicao, na frente da planilha: e a leitura que decide.

    "Competição" (leitura + gestoras), "Segmentos", "Oportunidades e defesa",
    "Lançamentos" e "ETFs por segmento". Cada aba abre com o prazo de cada
    numero - o PL e mensal, a comparacao e entre competencias e cotistas e fluxo
    sao da janela de informes do Fundos.NET.
    """
    if not bloco.get("available"):
        return
    foco = bloco.get("foco") or "casa em foco"
    janela = bloco.get("fluxo") or {}
    rotulo_fluxo = (
        f"Captação líquida de {janela['de'][8:10]}/{janela['de'][5:7]} a "
        f"{janela['ate'][8:10]}/{janela['ate'][5:7]} (R$)"
        if janela.get("de") else "Captação líquida no Fundos.NET (R$)"
    )
    inicio = bloco["competencia_inicio"]
    rotulo_inicio = f"{inicio[4:]}/{inicio[:4]}"

    # --- Competição: leitura + gestoras
    ws = wb.create_sheet("Competição", 0)
    linha = _cabecalho_competicao(ws, f"Competição em ETFs - onde a {foco} está", bloco)
    for frase in bloco["leitura"]:
        ws.cell(row=linha, column=1, value=frase).font = _BODY_FONT
        linha += 1
    linha += 1
    ws.cell(row=linha, column=1, value="Gestoras de ETF").font = _TITLE_FONT
    linha += 1
    colunas = [
        ("Posição", 9, _INT), ("Gestora", 26, None), ("ETFs operando", 12, _INT),
        ("Pré-operacionais", 14, _INT), ("PL (R$)", 18, _MONEY),
        ("Participação no PL", 14, _SHARE), (f"Participação em {rotulo_inicio}", 14, _SHARE),
        ("Variação de participação (pp)", 16, _PP), ("Variação do PL", 12, _SHARE),
        ("Cotistas (soma por ETF)", 16, _INT), ("Participação nos cotistas", 14, _SHARE),
        ("PL por cotista (R$)", 16, _MONEY), (rotulo_fluxo, 20, _MONEY),
        ("Segmentos que lidera", 60, None),
    ]
    linha = _tabela(ws, linha, colunas, [
        [
            g["posicao"], g["gestora"], g["etfs"], g["pre_operacionais"], g["pl"],
            g["share"], g["share_inicio"], g["delta_share_pp"], g["variacao_pct"],
            g["cotistas"], g["share_cotistas"], g["pl_por_cotista"], g["fluxo"],
            ", ".join(g["lidera"]),
        ]
        for g in bloco["gestoras"]
    ])
    mensal = bloco.get("participacao_mensal") or {}
    competencias = mensal.get("competencias") or []
    if len(competencias) > 1:
        ws.cell(row=linha, column=1, value="Participação no PL, competência a competência").font = _TITLE_FONT
        linha += 1
        _tabela(ws, linha, [
            ("Gestora", 26, None),
            *[(f"{c[4:]}/{c[:4]}", 12, _SHARE) for c in competencias],
            (f"Variação {competencias[0][4:]}/{competencias[0][:4]} a "
             f"{competencias[-1][4:]}/{competencias[-1][:4]} (pp)", 18, _PP),
        ], [
            [
                g["gestora"], *[g["share"][c] for c in competencias],
                (
                    (g["share"][competencias[-1]] - g["share"][competencias[0]]) * 100
                    if g["share"][competencias[-1]] is not None
                    and g["share"][competencias[0]] is not None else None
                ),
            ]
            for g in mensal["gestoras"]
        ], filtro=False)
    ws.freeze_panes = "A2"

    # --- Segmentos
    ws = wb.create_sheet("Segmentos", 1)
    linha = _cabecalho_competicao(ws, "Segmentos de ETF - o que cada ETF replica", bloco)
    ws.cell(
        row=linha - 1, column=1,
        value=(
            "Segmento lido do nome: a CVM deixa vazios o benchmark e as classificações "
            "dos ETFs. Dentro do segmento a valorização é a mesma para todos, então a "
            "variação de participação é fluxo - menos exata onde o segmento junta "
            "índices diferentes (Inflação mistura IMA-B e NTN-B de prazos distintos)."
        ),
    ).font = _NOTE_FONT
    linha += 1
    colunas = [
        ("Grupo", 16, None), ("Segmento", 34, None), ("PL (R$)", 18, _MONEY),
        ("Variação do PL", 12, _SHARE), ("ETFs operando", 12, _INT), ("Gestoras", 10, _INT),
        ("Líder", 22, None), ("Participação do líder", 14, _SHARE), ("2ª gestora", 22, None),
        ("Participação da 2ª", 14, _SHARE), (f"{foco}: situação", 14, None),
        (f"{foco}: PL (R$)", 16, _MONEY), (f"{foco}: participação", 14, _SHARE),
        (f"{foco}: posição", 10, _INT), (f"{foco}: variação (pp)", 14, _PP),
        ("Pré-operacionais", 40, None), ("Estreias", 40, None),
        ("Começaram depois do balancete", 40, None),
        ("Cotistas (soma por ETF)", 16, _INT), (rotulo_fluxo, 20, _MONEY),
    ]
    _tabela(ws, linha, colunas, [
        [
            s["grupo"], s["segmento"], s["pl"], s["variacao_pct"], s["etfs"], s["gestoras"],
            s["lider"], s["share_lider"], s["segundo"], s["share_segundo"],
            _STATUS[s["foco_status"]], s["foco_pl"] or None, s["foco_share"],
            s["foco_posicao"], s["foco_delta_pp"],
            "; ".join(f"{p['gestora']}: {nome_com_ticker(p['nome'], p.get('ticker'))}"
                      for p in s["pre_operacionais"]),
            "; ".join(f"{e['gestora']}: {nome_com_ticker(e['nome'], e.get('ticker'))}"
                      for e in s["estreias"]),
            "; ".join(
                f"{r['gestora']}: {nome_com_ticker(r['nome'], r.get('ticker'))} "
                f"({reais(r['pl'])} em {data_br(r['data'])})"
                for r in s["recentes"]
            ),
            s["cotistas"], s["fluxo"],
        ]
        for s in bloco["segmentos"]
    ])
    ws.freeze_panes = "C2"

    # --- Oportunidades e defesa
    ws = wb.create_sheet("Oportunidades e defesa", 2)
    linha = _cabecalho_competicao(ws, f"Onde a {foco} não compete, e onde lidera", bloco)
    limiares = bloco.get("limiares") or {}
    ws.cell(row=linha, column=1, value=(
        f"Oportunidades: segmentos com mais de {reais(limiares.get('pl_oportunidade'))} "
        f"em que a {foco} tem menos de {pct(limiares.get('share_pequena'), 0)} do PL, "
        "do maior para o menor."
    )).font = _TITLE_FONT
    linha += 1
    linha = _tabela(ws, linha, [
        ("Segmento", 34, None), ("PL (R$)", 18, _MONEY), ("Variação do PL", 12, _SHARE),
        ("Líder", 22, None), ("Participação do líder", 14, _SHARE),
        (f"{foco}: situação", 14, None), (f"{foco}: participação", 14, _SHARE),
        (f"{foco}: posição", 10, _INT), ("Concorrentes lançando", 40, None),
    ], [
        [
            s["segmento"], s["pl"], s["variacao_pct"], s["lider"], s["share_lider"],
            _STATUS[s["foco_status"]], s["foco_share"], s["foco_posicao"],
            ", ".join(sorted({p["gestora"] for p in s["lancando"]})),
        ]
        for s in bloco["oportunidades"]
    ], filtro=False)
    ws.cell(row=linha, column=1, value=(
        f"Defesa: segmentos que a {foco} lidera, com quem vem atrás e quem está lançando."
    )).font = _TITLE_FONT
    linha += 1
    _tabela(ws, linha, [
        ("Segmento", 34, None), ("PL (R$)", 18, _MONEY), (f"{foco}: participação", 14, _SHARE),
        (f"{foco}: variação (pp)", 14, _PP), ("2ª gestora", 22, None),
        ("Participação da 2ª", 14, _SHARE), ("Concorrentes lançando", 40, None),
    ], [
        [
            s["segmento"], s["pl"], s["foco_share"], s["foco_delta_pp"], s["segundo"],
            s["share_segundo"],
            ", ".join(sorted({p["gestora"] for p in s["lancando"] if p["gestora"] != foco})),
        ]
        for s in bloco["defesa"]
    ], filtro=False)

    # --- Lançamentos
    ws = wb.create_sheet("Lançamentos", 3)
    linha = _cabecalho_competicao(ws, "Lançamentos - pré-operacionais e estreias", bloco)
    ws.cell(row=linha - 1, column=1, value=(
        "Pré-operacional: registrado na CVM e ainda sem operar. Estreia: sem PL no "
        "balancete da primeira competência e com PL na última. Depois do balancete: "
        "sem PL em nenhum balancete e com PL no Fundos.NET depois da data-base - o PL "
        "é o do último informe, na data ao lado. A data de registro não serve para "
        "isso - a migração para a RCVM 175 a reescreveu (o iShares Ibovespa, de 2008, "
        "aparece registrado em 2025)."
    )).font = _NOTE_FONT
    linha += 1
    _tabela(ws, linha, [
        ("Tipo", 30, None), ("Gestora", 22, None), ("Grupo", 16, None),
        ("Segmento", 34, None), ("Ticker", 10, None), ("ETF", 70, None), ("PL (R$)", 16, _MONEY),
        ("Data do PL", 12, None),
    ], [
        [
            l["tipo"], l["gestora"], l["grupo"], l["segmento"], l.get("ticker"), l["nome"], l["pl"] or None,
            data_br(l["data_pl"]) if l["data_pl"] else None,
        ]
        for l in bloco["lancamentos"]
    ])
    ws.freeze_panes = "A2"

    # --- ETFs por segmento
    ws = wb.create_sheet("ETFs por segmento", 4)
    linha = _cabecalho_competicao(ws, "ETFs por segmento", bloco)
    _tabela(ws, linha, [
        ("Grupo", 16, None), ("Segmento", 34, None), ("Gestora", 22, None),
        ("Ticker", 10, None), ("ETF", 70, None), ("Situação", 22, None), ("PL (R$)", 18, _MONEY),
        (f"PL em {rotulo_inicio} (R$)", 18, _MONEY), ("Cotistas", 12, _INT),
        (rotulo_fluxo, 20, _MONEY), ("CNPJ", 20, None),
    ], [
        [
            e["grupo"], e["segmento"], e["gestora"], e.get("ticker"), e["nome"], e["situacao"],
            e["pl"] or None, e["pl_inicio"] or None, e["cotistas"], e["fluxo"],
            format_cnpj(e["cnpj"]),
        ]
        for e in bloco["etfs"]
    ])
    ws.freeze_panes = "A2"


# --- captacao mes a mes -----------------------------------------------------

_ESTILO_DO_ESTADO = {
    "acelerando": "bold green",
    "captando": "green",
    "sem_movimento": "dim",
    "virou_saida": "yellow",
    "saida_persistente": "bold red",
    "sem_dado": "magenta",
}


def _simbolo_rich(estado: str, largura: int = 0) -> str:
    """O simbolo do estado na cor dele; `largura` alinha a grade no console."""
    simbolo = SIMBOLOS.get(estado, "").ljust(largura)
    estilo = _ESTILO_DO_ESTADO.get(estado)
    return f"[{estilo}]{simbolo}[/]" if estilo and simbolo.strip() else simbolo


def print_captacao_mensal(bloco: dict) -> None:
    """Captacao mes a mes no console: leitura, gestoras, segmentos, oportunidades e defesa."""
    console = Console()
    console.print()
    fonte = (
        "pelo informe diário do Fundos.NET (captação, resgates e diferença de PL)"
        if bloco.get("metodo") == "informe" else "pela variação de cotas do Fundos.NET"
    )
    console.print(f"[bold]Captação de ETF mês a mês[/] - {fonte}")
    if not bloco.get("available"):
        console.print(f"   [yellow]{bloco.get('reason')}[/]")
        return
    for prazo in prazos_captacao(bloco):
        console.print(f"   [dim]{prazo}[/]")
    console.print(f"   [dim]legenda: {legenda()}[/]")
    for frase in bloco["leitura"]:
        console.print(f"   - {frase}")

    meses = bloco["meses"]
    ultimo = meses[-1]
    janela = f"{mes_curto(meses[0])} → {mes_curto(ultimo)}"

    def estados(por_mes: dict | None) -> str:
        return " ".join(_simbolo_rich((por_mes or {}).get(m, "fora"), largura=2) for m in meses)

    tabela = Table(title="Captação por gestora", title_style="bold")
    for titulo, alinhamento in (
        ("#", "right"), ("Gestora", "left"), ("12 meses", "right"), ("% da captação", "right"),
        (janela, "left"), (f"Em {mes_curto(ultimo)}", "right"),
    ):
        tabela.add_column(titulo, justify=alinhamento)
    gestoras = bloco["gestoras"][:12]
    if not any(g["foco"] for g in gestoras):
        gestoras += [g for g in bloco["gestoras"] if g["foco"]]
    for g in gestoras:
        tabela.add_row(
            str(g["posicao"]), g["gestora"], reais_com_sinal(g["doze_meses"]),
            pct(g["share_doze_meses"]), estados(g["estado"]), reais_com_sinal(g["fluxo"][ultimo]),
            style="bold" if g["foco"] else None,
        )
    console.print(tabela)

    foco = bloco.get("foco") or "casa em foco"
    tabela = Table(title="Captação por segmento", title_style="bold")
    for titulo, alinhamento in (
        ("Segmento", "left"), ("12 meses", "right"), (janela, "left"),
        ("Quem mais captou", "left"), (foco, "right"),
    ):
        tabela.add_column(titulo, justify=alinhamento)
    for s_ in bloco["segmentos"]:
        tabela.add_row(
            s_["segmento"], reais_com_sinal(s_["doze_meses"]), estados(s_["estado"]),
            curto(s_["lider_captacao"]) if s_["lider_captacao"] else "-",
            reais_com_sinal(s_["foco_doze_meses"]) if s_["foco_doze_meses"] is not None else "sem ETF",
        )
    console.print(tabela)
    if bloco["oportunidades"]:
        console.print(f"[bold]Oportunidades[/] - onde o dinheiro entrou e a {foco} ficou de fora ou pequena:")
        for s_ in bloco["oportunidades"][:6]:
            console.print(
                f"   {s_['segmento']}: {reais_com_sinal(s_['doze_meses'])} em 12 meses, "
                f"{curto(s_['lider_captacao'])} à frente; {foco}: "
                + (f"{reais_com_sinal(s_['foco_doze_meses'])}" if s_["foco_tem_etf"] else "sem ETF")
            )
    alertas = [s_ for s_ in bloco["defesa"] if s_.get("alertas")]
    if alertas:
        console.print(f"[bold]Defesa[/] - segmentos que a {foco} lidera em PL com alerta:")
        for s_ in alertas:
            console.print(f"   {s_['segmento']}: {'; '.join(s_['alertas'])}")


def _cabecalho_captacao(ws, titulo: str, bloco: dict) -> int:
    """Titulo, prazos e legenda de cores. Devolve a proxima linha livre."""
    ws.cell(row=1, column=1, value=titulo).font = _TITLE_FONT
    linha = 2
    for prazo in prazos_captacao(bloco):
        ws.cell(row=linha, column=1, value=prazo[0].upper() + prazo[1:]).font = _NOTE_FONT
        linha += 1
    ws.cell(row=linha, column=1, value="Cor do mês:").font = _NOTE_FONT
    for coluna, estado in enumerate(ORDEM_DA_LEGENDA, start=2):
        fundo, texto = CORES[estado]
        cell = ws.cell(row=linha, column=coluna, value=ROTULOS[estado])
        cell.fill = PatternFill("solid", fgColor=fundo)
        cell.font = Font(name=FONT, size=9, color=texto)
        cell.alignment = Alignment(horizontal="center")
    ws.cell(row=linha, column=len(ORDEM_DA_LEGENDA) + 2, value="em branco: antes da estreia").font = _NOTE_FONT
    return linha + 2


def _pintar_meses(ws, primeira_linha: int, primeira_coluna: int, meses: list[str], estados: list[dict]) -> None:
    """Cor de fundo e de texto de cada mes pelo estado dele."""
    for i, por_mes in enumerate(estados):
        for j, mes in enumerate(meses):
            estado = (por_mes or {}).get(mes, "fora")
            if estado not in CORES:
                continue
            fundo, texto = CORES[estado]
            cell = ws.cell(row=primeira_linha + i, column=primeira_coluna + j)
            cell.fill = PatternFill("solid", fgColor=fundo)
            cell.font = Font(name=FONT, size=10, color=texto)


def _colunas_de_meses(meses: list[str]) -> list[tuple[str, int, str]]:
    return [(mes_curto(m), 11, _MONEY) for m in meses]


# As colunas do informe diario do Fundos.NET em 12 meses, na ordem de DETALHE.
_COLUNAS_DO_DETALHE = (
    ("Captação declarada 12 meses (R$)", 18, _MONEY),
    ("Resgates declarados 12 meses (R$)", 18, _MONEY),
    ("Ajuste pela diferença de PL 12 meses (R$)", 18, _MONEY),
    ("Sem o informe do mês 12 meses (R$)", 18, _MONEY),
    ("Diferença de PL 12 meses (R$)", 18, _MONEY),
)
_LINHAS_DO_DETALHE = (
    ("Captação declarada (informe)", "captacao"),
    ("Resgates declarados (informe)", "resgates"),
    ("Ajuste pela diferença de PL", "ajuste"),
    ("Sem o informe do mês (fechamentos)", "sem_informe"),
    ("Diferença de PL", "diferenca_pl"),
)


def _detalhe_12_meses(item: dict) -> list:
    """Captacao, resgates, ajuste e diferenca de PL em 12 meses de uma linha do bloco."""
    doze = (item.get("detalhe") or {}).get("doze_meses") or {}
    return [doze.get(coluna) for coluna in DETALHE]


def _write_informe_mes_a_mes(wb: Workbook, bloco: dict) -> None:
    """Aba "Captação - informe mês a mês": cada ETF em cada mês, com as colunas
    do informe diario do Fundos.NET e as pontas de PL."""
    ws = wb.create_sheet("Captação - informe mês a mês", 4)
    linha = _cabecalho_captacao(ws, "Captação de cada ETF em cada mês, pelo informe diário do Fundos.NET", bloco)
    ws.cell(row=linha, column=1, value=(
        "Captação líquida = captação declarada - resgates declarados + ajuste pela diferença de PL "
        "(mês sem o informe: a variação de cotas entre fechamentos). "
        "Valorização = diferença de PL - captação líquida."
    )).font = _NOTE_FONT
    linha += 2
    corpo = []
    for e in bloco["etfs"]:
        detalhe = e.get("detalhe") or {}
        for m in bloco["meses"]:
            ponta = (e.get("mes_a_mes") or {}).get(m)
            if ponta is None or e["estado"][m] == "fora":
                continue
            liquida = e["fluxo"][m]
            diferenca = (detalhe.get("diferenca_pl") or {}).get(m)
            do_informe = ponta["fonte"] == "informe"
            corpo.append([
                e["grupo"], e["segmento"], e["gestora"], e.get("ticker"), e["nome"], mes_curto(m),
                *[(detalhe.get(c) or {}).get(m) if do_informe else None for c in ("captacao", "resgates", "ajuste")],
                liquida, ponta["pl_inicio"], ponta["pl_fim"], diferenca,
                diferenca - liquida if diferenca is not None and liquida is not None else None,
                ponta["dias"] or None, ponta["dias_pelo_pl"] or None, rotulo(e["estado"][m]),
                ORIGENS.get(ponta["origem"], "informe diário" if ponta["fonte"] == "informe" else "fechamentos"),
                format_cnpj(e["cnpj"]),
            ])
    _tabela(ws, linha, [
        ("Grupo", 16, None), ("Segmento", 30, None), ("Gestora", 22, None), ("Ticker", 10, None),
        ("ETF", 50, None), ("Mês", 9, None), ("Captação declarada (R$)", 17, _MONEY),
        ("Resgates declarados (R$)", 17, _MONEY), ("Ajuste pela diferença de PL (R$)", 17, _MONEY),
        ("Captação líquida (R$)", 17, _MONEY), ("PL no início (R$)", 18, _MONEY),
        ("PL no fim (R$)", 18, _MONEY), ("Diferença de PL (R$)", 17, _MONEY),
        ("Valorização (R$)", 17, _MONEY), ("Dias com informe", 10, _INT),
        ("Dias pela diferença de PL", 10, _INT), ("Estado", 16, None), ("De onde saiu", 30, None),
        ("CNPJ", 20, None),
    ], corpo)
    ws.freeze_panes = ws.cell(row=linha + 1, column=6)


def _write_captacao_sheets(wb: Workbook, bloco: dict) -> None:
    """Abas de captacao mes a mes, na frente da planilha.

    "Captação - Competição" (mercado e gestoras), "Captação - Segmentos",
    "Captação - Oport. e defesa" e "Captação - por ETF": as visoes do bloco de
    competicao, agora pela captacao de cada mes, cada celula de mes na cor do
    estado dele.
    """
    foco = bloco.get("foco") or "casa em foco"
    meses = bloco["meses"]
    ultimo = mes_curto(meses[-1])
    col_meses = _colunas_de_meses(meses)

    # --- Competição: mercado e gestoras
    ws = wb.create_sheet("Captação - Competição", 0)
    linha = _cabecalho_captacao(ws, "Captação de ETF mês a mês - competição entre gestoras", bloco)
    for frase in bloco["leitura"]:
        ws.cell(row=linha, column=1, value=frase).font = _BODY_FONT
        linha += 1
    linha += 1
    colunas = [
        ("Posição", 9, _INT), ("Gestora", 26, None), ("ETFs com PL", 11, _INT),
        ("PL no fim (R$)", 18, _MONEY), ("Participação no PL", 13, _SHARE),
        ("Captação 12 meses (R$)", 18, _MONEY), ("Participação na captação 12 meses", 15, _SHARE),
        ("Captação no ano (R$)", 18, _MONEY), *col_meses,
        (f"Estado em {ultimo}", 18, None), ("Meses captando", 10, _INT), ("Meses em saída", 10, _INT),
        *_COLUNAS_DO_DETALHE,
    ]
    mercado = bloco["mercado"]
    linhas = [[
        None, "Todos os ETFs", bloco["cobertura"]["com_dado"].get(meses[-1]), mercado["pl"], 1.0,
        mercado["doze_meses"], 1.0, mercado["no_ano"], *[mercado["fluxo"][m] for m in meses],
        rotulo(mercado["estado_atual"]), mercado["meses_captando"], mercado["meses_em_saida"],
        *_detalhe_12_meses(mercado),
    ]]
    estados = [mercado["estado"]]
    historico = mercado.get("historico") or {}
    vazias = [None] * (3 + len(_COLUNAS_DO_DETALHE))
    linhas.append([
        None, "Das quais estreias", None, None, None,
        sum((historico.get(m) or {}).get("estreias") or 0.0 for m in meses) or None, None,
        sum((historico.get(m) or {}).get("estreias") or 0.0 for m in bloco["meses_ano"]) or None,
        *[((historico.get(m) or {}).get("estreias") or None) for m in meses], *vazias,
    ])
    estados.append({})
    if bloco.get("metodo") == "informe":
        # As colunas do informe diario, na mesma grade: liquida = captacao -
        # resgates + ajuste; a diferenca de PL carrega tambem a valorizacao.
        detalhe = mercado["detalhe"]
        for titulo, coluna in _LINHAS_DO_DETALHE:
            if coluna == "sem_informe" and not any(detalhe[coluna].values()):
                continue
            linhas.append([
                None, titulo, None, None, None, detalhe["doze_meses"][coluna], None,
                detalhe["no_ano"][coluna], *[detalhe[coluna][m] for m in meses], *vazias,
            ])
            estados.append({})
    primeira_gestora = len(linhas) + 1
    for g in bloco["gestoras"]:
        linhas.append([
            g["posicao"], g["gestora"], g["etfs"], g["pl"], g["share_pl"], g["doze_meses"],
            g["share_doze_meses"], g["no_ano"], *[g["fluxo"][m] for m in meses],
            rotulo(g["estado_atual"]), g["meses_captando"], g["meses_em_saida"],
            *_detalhe_12_meses(g),
        ])
        estados.append(g["estado"])
    cabecalho = linha
    _tabela(ws, linha, colunas, linhas)
    _pintar_meses(ws, cabecalho + 1, 9, meses, estados)
    ws.cell(row=cabecalho + 1, column=2).font = Font(name=FONT, size=10, bold=True)
    for i, g in enumerate(bloco["gestoras"], start=primeira_gestora):
        if g["foco"]:
            ws.cell(row=cabecalho + i, column=2).font = Font(name=FONT, size=10, bold=True)
    ws.freeze_panes = ws.cell(row=cabecalho + 1, column=3)

    # --- Segmentos
    ws = wb.create_sheet("Captação - Segmentos", 1)
    linha = _cabecalho_captacao(ws, "Captação por segmento, mês a mês", bloco)
    colunas = [
        ("Grupo", 16, None), ("Segmento", 34, None), ("PL no fim (R$)", 18, _MONEY),
        ("Captação 12 meses (R$)", 18, _MONEY), ("Captação no ano (R$)", 18, _MONEY), *col_meses,
        (f"Estado em {ultimo}", 18, None), ("Quem mais captou em 12 meses", 24, None),
        ("Captação dela (R$)", 16, _MONEY), (f"{foco}: captação 12 meses (R$)", 18, _MONEY),
        (f"{foco}: parte da captação", 13, _SHARE), (f"{foco}: estado em {ultimo}", 18, None),
        (f"{foco} lidera o PL", 10, None), *_COLUNAS_DO_DETALHE,
    ]
    cabecalho = linha
    proxima = _tabela(ws, linha, colunas, [
        [
            s_["grupo"], s_["segmento"], s_["pl"], s_["doze_meses"], s_["no_ano"],
            *[s_["fluxo"][m] for m in meses], rotulo(s_["estado_atual"]),
            s_["lider_captacao"], s_["lider_captacao_valor"] if s_["lider_captacao"] else None,
            s_["foco_doze_meses"], s_["foco_share"],
            rotulo(s_["foco_estado_atual"]) if s_["foco_estado_atual"] else "sem ETF",
            "sim" if s_["foco_lidera_pl"] else "", *_detalhe_12_meses(s_),
        ]
        for s_ in bloco["segmentos"]
    ])
    _pintar_meses(ws, cabecalho + 1, 6, meses, [s_["estado"] for s_ in bloco["segmentos"]])
    com_foco = [s_ for s_ in bloco["segmentos"] if s_["foco_estado"]]
    if com_foco:
        ws.cell(row=proxima, column=1, value=f"A {foco} em cada segmento, mês a mês").font = _TITLE_FONT
        linha = proxima + 1
        cabecalho = linha
        _tabela(ws, linha, [
            ("Grupo", 16, None), ("Segmento", 34, None), ("PL no fim (R$)", 18, _MONEY),
            ("Captação 12 meses (R$)", 18, _MONEY), ("Captação no ano (R$)", 18, _MONEY), *col_meses,
            (f"Estado em {ultimo}", 18, None),
        ], [
            [
                s_["grupo"], s_["segmento"], None, s_["foco_doze_meses"], s_["foco_no_ano"],
                *[(s_["foco_fluxo"] or {}).get(m) for m in meses], rotulo(s_["foco_estado_atual"]),
            ]
            for s_ in com_foco
        ], filtro=False)
        _pintar_meses(ws, cabecalho + 1, 6, meses, [s_["foco_estado"] for s_ in com_foco])
    ws.freeze_panes = "C2"

    # --- Oportunidades e defesa
    ws = wb.create_sheet("Captação - Oport. e defesa", 2)
    linha = _cabecalho_captacao(ws, f"Captação: oportunidades e defesa da {foco}", bloco)
    limiares = bloco.get("limiares") or {}
    ws.cell(row=linha, column=1, value=(
        f"Oportunidades: segmentos que captaram mais de {reais(limiares.get('captacao_oportunidade'))} "
        f"em 12 meses e em que a {foco} ficou de fora ou com menos de "
        f"{pct(limiares.get('share_pequena'), 0)} da captação. A grade é a captação do segmento."
    )).font = _TITLE_FONT
    linha += 1
    cabecalho = linha
    linha = _tabela(ws, linha, [
        ("Segmento", 34, None), ("Captação 12 meses (R$)", 18, _MONEY), ("Captação no ano (R$)", 18, _MONEY),
        *[(t, w, f) for t, w, f in col_meses], (f"Estado em {ultimo}", 18, None),
        ("Quem captou (12 meses)", 50, None), (f"{foco}: captação 12 meses (R$)", 18, _MONEY),
        (f"{foco}: parte da captação", 13, _SHARE), (f"{foco} tem ETF", 10, None),
    ], [
        [
            s_["segmento"], s_["doze_meses"], s_["no_ano"], *[s_["fluxo"][m] for m in meses],
            rotulo(s_["estado_atual"]),
            "; ".join(f"{r['gestora']}: {reais_com_sinal(r['doze_meses'])}" for r in s_["rivais_captando"]),
            s_["foco_doze_meses"], s_["foco_share"], "sim" if s_["foco_tem_etf"] else "não",
        ]
        for s_ in bloco["oportunidades"]
    ], filtro=False)
    _pintar_meses(ws, cabecalho + 1, 4, meses, [s_["estado"] for s_ in bloco["oportunidades"]])
    ws.cell(row=linha, column=1, value=(
        f"Defesa: segmentos que a {foco} lidera em PL. A grade é a captação da {foco} no segmento."
    )).font = _TITLE_FONT
    linha += 1
    cabecalho = linha
    _tabela(ws, linha, [
        ("Segmento", 34, None), ("PL no fim (R$)", 18, _MONEY), (f"{foco}: captação 12 meses (R$)", 18, _MONEY),
        *[(t, w, f) for t, w, f in col_meses], (f"{foco}: estado em {ultimo}", 18, None),
        ("Quem mais captou além dela", 24, None), ("Captação dela (R$)", 16, _MONEY), ("Alertas", 60, None),
    ], [
        [
            s_["segmento"], s_["pl"], s_["foco_doze_meses"], *[(s_["foco_fluxo"] or {}).get(m) for m in meses],
            rotulo(s_["foco_estado_atual"]),
            s_["rivais_captando"][0]["gestora"] if s_["rivais_captando"] else None,
            s_["rivais_captando"][0]["doze_meses"] if s_["rivais_captando"] else None,
            "; ".join(s_.get("alertas") or []),
        ]
        for s_ in bloco["defesa"]
    ], filtro=False)
    _pintar_meses(ws, cabecalho + 1, 4, meses, [s_["foco_estado"] or {} for s_ in bloco["defesa"]])

    # --- por ETF
    ws = wb.create_sheet("Captação - por ETF", 3)
    linha = _cabecalho_captacao(ws, "Captação de cada ETF, mês a mês, com o seu segmento", bloco)
    cabecalho = linha
    _tabela(ws, linha, [
        ("Grupo", 16, None), ("Segmento", 34, None), ("Gestora", 22, None), ("Ticker", 10, None),
        ("ETF", 60, None), ("PL no fim (R$)", 18, _MONEY), ("Captação 12 meses (R$)", 18, _MONEY),
        ("Captação no ano (R$)", 18, _MONEY), *col_meses, (f"Estado em {ultimo}", 18, None),
        *_COLUNAS_DO_DETALHE, ("Dias com informe (12 meses)", 12, _INT),
        ("Dias pela diferença de PL (12 meses)", 12, _INT),
        ("Observações", 50, None), ("CNPJ", 20, None),
    ], [
        [
            e["grupo"], e["segmento"], e["gestora"], e.get("ticker"), e["nome"], e["pl"] or None,
            e["doze_meses"],
            e["no_ano"], *[e["fluxo"][m] for m in meses], rotulo(e["estado_atual"]),
            *_detalhe_12_meses(e), e.get("dias_com_informe") or None, e.get("dias_pelo_pl") or None,
            "; ".join(f"{mes_curto(m)}: {o}" for m, o in sorted(e["origem"].items())),
            format_cnpj(e["cnpj"]),
        ]
        for e in bloco["etfs"]
    ])
    _pintar_meses(ws, cabecalho + 1, 9, meses, [e["estado"] for e in bloco["etfs"]])
    for i, e in enumerate(bloco["etfs"], start=1):
        if e["foco"]:
            for coluna in (4, 5):
                ws.cell(row=cabecalho + i, column=coluna).font = Font(name=FONT, size=10, bold=True)
    ws.freeze_panes = ws.cell(row=cabecalho + 1, column=6)

    if bloco.get("metodo") == "informe":
        _write_informe_mes_a_mes(wb, bloco)
    if bloco.get("anbima"):
        _write_conferencia_anbima(wb, bloco["anbima"])


def _write_conferencia_anbima(wb: Workbook, conferencia: dict) -> None:
    """Aba "Captação - ANBIMA": monitor x boletim da ANBIMA, mes a mes, na janela dele."""
    ws = wb.create_sheet("Captação - ANBIMA", 4)
    ws.cell(row=1, column=1, value="Captação de ETF: monitor x boletim da ANBIMA").font = _TITLE_FONT
    notas = [
        f"Anexo do boletim: {conferencia['arquivo']}. A ANBIMA soma o que os administradores "
        "associados informam; o monitor, todos os ETFs da CVM, pelo informe diário do Fundos.NET.",
        "O monitor conta o PL de estreia como captação; a ANBIMA nem sempre - a coluna de "
        "estreias mostra quanto disso há em cada mês.",
        "Renda fixa no monitor é o grupo Renda Fixa dos segmentos; o resto é renda variável.",
    ]
    for linha, nota in enumerate(notas, start=2):
        ws.cell(row=linha, column=1, value=nota).font = _NOTE_FONT
    linha = len(notas) + 3
    corpo = [
        [
            mes_curto(l["mes"]), l["monitor"], l["estreias"] or None, l["anbima"], l["diferenca"],
            l["pl_monitor"], l["pl_anbima"], l["cobertura"],
            l["renda_fixa"]["monitor"], l["renda_fixa"]["anbima"],
            l["renda_variavel"]["monitor"], l["renda_variavel"]["anbima"],
        ]
        for l in conferencia["linhas"]
    ]
    for rotulo_, soma in (("12 meses", conferencia["doze_meses"]), ("No ano", conferencia.get("no_ano"))):
        if soma:
            corpo.append([
                rotulo_, soma["monitor"], soma["estreias"] or None, soma["anbima"],
                soma["monitor"] - soma["anbima"], None, None, None, None, None, None, None,
            ])
    _tabela(ws, linha, [
        ("Mês", 10, None), ("Monitor (R$)", 18, _MONEY), ("Das quais estreias (R$)", 18, _MONEY),
        ("ANBIMA (R$)", 18, _MONEY), ("Diferença (R$)", 18, _MONEY), ("PL monitor (R$)", 18, _MONEY),
        ("PL ANBIMA (R$)", 18, _MONEY), ("PL: ANBIMA / monitor", 12, _SHARE),
        ("Renda fixa: monitor (R$)", 18, _MONEY), ("Renda fixa: ANBIMA (R$)", 18, _MONEY),
        ("Renda variável: monitor (R$)", 18, _MONEY), ("Renda variável: ANBIMA (R$)", 18, _MONEY),
    ], corpo, filtro=False)


def export_captacao_csv(bloco: dict, output_path: Path) -> Path:
    """Captacao de cada ETF em cada mes, em formato longo (para outras ferramentas)."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow([
            "competencia", "cnpj", "ticker", "etf", "gestora", "grupo", "segmento", "captacao",
            "estado", "origem", "captacao_declarada", "resgates_declarados", "ajuste_diferenca_pl",
            "pl_inicio", "pl_fim", "diferenca_pl", "dias_com_informe", "dias_pela_diferenca_pl",
        ])

        def numero(valor) -> str | float:
            return "" if valor is None else round(valor, 2)

        for e in bloco.get("etfs") or []:
            detalhe = e.get("detalhe") or {}
            for m in bloco["meses"]:
                if e["estado"][m] == "fora":
                    continue
                ponta = (e.get("mes_a_mes") or {}).get(m) or {}
                # "captacao" e a liquida do mes; "origem" diz de onde saiu (o
                # informe do mes, os fechamentos ou um caso de borda).
                origem = e["origem"].get(m) or (
                    "informe" if ponta.get("fonte") == "informe" else "cotas"
                )
                declarado = ponta.get("fonte") == "informe"
                writer.writerow([
                    m, e["cnpj"], e.get("ticker") or "", e["nome"], e["gestora"], e["grupo"], e["segmento"],
                    numero(e["fluxo"][m]), e["estado"][m], origem,
                    numero((detalhe.get("captacao") or {}).get(m)) if declarado else "",
                    numero((detalhe.get("resgates") or {}).get(m)) if declarado else "",
                    numero((detalhe.get("ajuste") or {}).get(m)),
                    numero(ponta.get("pl_inicio")), numero(ponta.get("pl_fim")),
                    numero((detalhe.get("diferenca_pl") or {}).get(m)),
                    ponta.get("dias") or "", ponta.get("dias_pelo_pl") or "",
                ])
    return output_path
