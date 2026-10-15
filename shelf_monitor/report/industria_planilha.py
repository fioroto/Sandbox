"""A planilha do monitor da industria: leitura, funil e o explorador.

Uma aba por pergunta, e o explorador com uma linha por produto do universo,
com autofiltro, para quem quer filtrar e ordenar a industria inteira. Valores
em reais vao em milhoes (R$ mi); fracoes, em percentual. Cada aba diz o
proprio prazo no topo e a aba "Período de cada aba" junta todos.

Padrao visual: o executivo da XP (`docs/linguagem-visual.md`) - cabecalho preto
com texto branco, a linha do total com o fundo amarelo claro. Nenhuma casa em
destaque: o monitor e neutro.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from shelf_monitor.config import Config
from shelf_monitor.industria.cadastro import ROTULOS_PUBLICO
from shelf_monitor.industria.curadoria import TIPOS_DE_GESTORA
from shelf_monitor.industria.estrutura import PAPEIS
from shelf_monitor.industria.janelas import ORDEM, ROTULOS as ROTULOS_JANELA
from shelf_monitor.industria.taxonomia import N1_ORDEM
from shelf_monitor.report.shared.fontes import fontes_da_industria
from shelf_monitor.report.spreadsheet import _salvar
from shelf_monitor.report.shared.numeros import pct, reais
from shelf_monitor.texto import competencia_br, data_br, milhar

FONT = "Arial"
_PRETO, _AMARELO_CLARO = "0A0A0A", "FFF6CC"
_TITULO = Font(name=FONT, size=13, bold=True)
_NOTA = Font(name=FONT, size=9, italic=True, color="595959")
_CORPO = Font(name=FONT, size=10)
_NEGRITO = Font(name=FONT, size=10, bold=True)
_CABECALHO = Font(name=FONT, size=10, bold=True, color="FFFFFF")
_FUNDO_CABECALHO = PatternFill("solid", fgColor=_PRETO)
_FUNDO_TOTAL = PatternFill("solid", fgColor=_AMARELO_CLARO)
_LINK = Font(name=FONT, size=10, color="1A5FB4", underline="single")
_BORDA = Border(bottom=Side(style="thin", color="D9D9D9"))

MI = "#,##0;-#,##0;-"
MI_SINAL = "+#,##0;-#,##0;0"
PCT = "0.0%;-0.0%;-"
PCT_SINAL = "+0.0%;-0.0%;0.0%"
PP = '+0.0" pp";-0.0" pp";0.0" pp"'
INT = "#,##0;-#,##0;-"


def _mi(valor) -> float | None:
    return None if valor is None else round(valor / 1e6, 1)


class _Aba:
    """Escrita de uma aba com titulo, notas e tabelas empilhadas."""

    def __init__(self, workbook: Workbook, nome: str, titulo: str, primeira: bool = False):
        self.ws = workbook.active if primeira else workbook.create_sheet(nome)
        self.ws.title = nome
        self.linha = 1
        self.escrever(titulo, fonte=_TITULO)

    def escrever(self, texto: str, fonte: Font = _CORPO, coluna: int = 1) -> None:
        self.ws.cell(row=self.linha, column=coluna, value=texto).font = fonte
        self.linha += 1

    def pular(self, n: int = 1) -> None:
        self.linha += n

    def tabela(
        self,
        colunas: list[tuple[str, int, str | None]],
        linhas: list[list],
        *,
        total: list | None = None,
        filtro: bool = False,
        congelar: bool = False,
    ) -> None:
        inicio = self.linha
        for c, (titulo, largura, _) in enumerate(colunas, start=1):
            celula = self.ws.cell(row=self.linha, column=c, value=titulo)
            celula.font, celula.fill = _CABECALHO, _FUNDO_CABECALHO
            celula.alignment = Alignment(wrap_text=True, vertical="center")
            atual = self.ws.column_dimensions[get_column_letter(c)].width or 0
            self.ws.column_dimensions[get_column_letter(c)].width = max(largura, atual)
        self.linha += 1
        for valores in linhas + ([total] if total else []):
            eh_total = total is not None and valores is total
            for c, (valor, (_, _, formato)) in enumerate(zip(valores, colunas), start=1):
                celula = self.ws.cell(row=self.linha, column=c, value=valor)
                celula.font = _NEGRITO if eh_total else _CORPO
                celula.border = _BORDA
                if formato:
                    celula.number_format = formato
                if eh_total:
                    celula.fill = _FUNDO_TOTAL
            self.linha += 1
        if filtro and linhas:
            self.ws.auto_filter.ref = f"A{inicio}:{get_column_letter(len(colunas))}{self.linha - 1}"
        if congelar:
            self.ws.freeze_panes = self.ws.cell(row=inicio + 1, column=3)
        self.pular()


def _prazo(bloco: dict, chave: str = "12m") -> str:
    janela = bloco["janelas"][chave]
    return f"{janela['rotulo']}: de {data_br(janela['de'])} a {data_br(janela['ate'])}"


def _fluxo(item: dict, janela: str, campo: str = "fluxo_investidor"):
    dados = (item.get("janelas") or {}).get(janela) or {}
    return dados.get(campo)


def _aba_leitura(wb: Workbook, bloco: dict, gerado: datetime) -> None:
    aba = _Aba(wb, "Leitura", f"Monitor da indústria de fundos — {bloco['rotulo_competencia']}", primeira=True)
    aba.escrever(
        f"Gerado em {gerado:%d/%m/%Y %H:%M}. Universo disponível para PF e PJ, na camada do investidor "
        "(sem o dinheiro que um fundo aplica em outro). Nenhuma casa em destaque.",
        _NOTA,
    )
    aba.pular()
    for frase in bloco.get("leitura") or ():
        aba.escrever(frase.replace("**", ""))
    aba.pular()
    aba.escrever("Prazos", _NEGRITO)
    for prazo in bloco.get("prazos") or ():
        aba.escrever(prazo, _NOTA)
    if bloco.get("avisos"):
        aba.pular()
        aba.escrever("Avisos", _NEGRITO)
        for aviso in bloco["avisos"]:
            aba.escrever(aviso, _NOTA)
    aba.ws.column_dimensions["A"].width = 140


def _aba_funil(wb: Workbook, bloco: dict) -> None:
    aba = _Aba(wb, "Funil", "Funil do universo: quem sai e por quê")
    aba.escrever(
        f"Cadastro de {bloco['qualidade']['retrato_do_cadastro'][4:]}/{bloco['qualidade']['retrato_do_cadastro'][:4]}; "
        f"PL do informe diário em {data_br(bloco['informe']['ate'])}; estruturados e ETF com o PL do cadastro. "
        "Cada produto (subclasse, ou classe sem subclasse) cai no primeiro motivo que se aplicar.",
        _NOTA,
    )
    linhas = [[d["rotulo"], d["produtos"], d["classes"], _mi(d["pl"])] for d in bloco["funil"]]
    universo = linhas.pop()
    aba.tabela(
        [("Degrau", 42, None), ("Produtos", 12, INT), ("Classes", 12, INT), ("PL (R$ mi)", 16, MI)],
        linhas, total=universo,
    )


def _aba_estrategias(wb: Workbook, bloco: dict) -> None:
    aba = _Aba(wb, "Estratégias", "Estratégias: tamanho, captação e rentabilidade por peer group")
    aba.escrever("Camada do investidor. " + " · ".join(bloco.get("prazos") or ()), _NOTA)
    estr = bloco["estrategias"]
    colunas = [("Macroclasse", 22, None), ("Classes", 10, INT), ("PL (R$ mi)", 14, MI), ("Fatia do PL", 11, PCT)]
    colunas += [(f"Captação {ROTULOS_JANELA[j]} (R$ mi)", 14, MI_SINAL) for j in ORDEM]
    colunas += [("Captação 12m / PL de abertura", 14, PCT_SINAL)]
    linhas = [
        [i["n1"], i["classes"], _mi(i["pl_investidor"]), i.get("fatia")]
        + [_mi(_fluxo(i, j)) for j in ORDEM] + [_fluxo(i, "12m", "sobre_pl")]
        for i in estr["n1"]
    ]
    universo = bloco["mercado"]["janelas"]
    total = ["Universo", sum(i["classes"] for i in estr["n1"]), _mi(sum(i["pl_investidor"] for i in estr["n1"])), 1.0]
    total += [_mi(universo[j]["fluxo_investidor"]) for j in ORDEM] + [universo["12m"]["sobre_pl"]]
    aba.tabela(colunas, linhas, total=total)

    colunas = [("Macroclasse", 22, None), ("Peer group", 44, None), ("Classes", 10, INT),
               ("PL (R$ mi)", 14, MI), ("Fatia do PL", 11, PCT)]
    colunas += [(f"Captação {ROTULOS_JANELA[j]} (R$ mi)", 14, MI_SINAL) for j in ORDEM]
    colunas += [("Captação 12m / PL de abertura", 14, PCT_SINAL),
                ("Retorno 12m: 1º quartil", 12, PCT), ("Retorno 12m: mediana", 12, PCT),
                ("Retorno 12m: 3º quartil", 12, PCT), ("Mediana 12m em % do CDI", 12, PCT),
                ("Estratégias acima do CDI em 12m", 12, PCT), ("Retorno 3m: mediana", 12, PCT),
                ("Mediana 3m em % do CDI", 12, PCT),
                ("Estratégias com retorno", 12, INT), ("Amostra pequena", 10, None)]
    linhas = []
    for i in estr["n2"]:
        r12 = (i.get("rentabilidade") or {}).get("12m") or {}
        r3 = (i.get("rentabilidade") or {}).get("3m") or {}
        linhas.append(
            [i["n1"], i["n2"], i["classes"], _mi(i["pl_investidor"]), i.get("fatia")]
            + [_mi(_fluxo(i, j)) for j in ORDEM]
            + [_fluxo(i, "12m", "sobre_pl"), r12.get("q1"), r12.get("mediana"), r12.get("q3"),
               r12.get("mediana_sobre_cdi"), r12.get("acima_do_cdi"), r3.get("mediana"), r3.get("mediana_sobre_cdi"),
               r12.get("estrategias"), "sim" if r12.get("amostra_pequena") else ""]
        )
    cdi = {j: ((bloco.get("desempenho") or {}).get(j) or {}).get("cdi") for j in ("12m", "3m")}
    aba.escrever(
        "Retorno: mediana e quartis por estratégia-raiz (uma vez por master), base no fechamento da "
        "competência anterior; não é retorno total nem líquido de impostos. % do CDI: a mediana sobre o CDI "
        f"acumulado da mesma janela (12 meses: {pct(cdi['12m'])}; 3 meses: {pct(cdi['3m'])}; SGS do Banco Central).",
        _NOTA,
    )
    aba.tabela(colunas, linhas, filtro=True)


def _aba_mapa_de_calor(wb: Workbook, bloco: dict) -> None:
    calor = bloco["estrategias"].get("mapa_de_calor") or {}
    aba = _Aba(wb, "Mapa de calor", "Rotação de estratégias: captação do mês em % do PL do mês anterior")
    aba.escrever("Camada do investidor, por peer group. " + _prazo(bloco), _NOTA)
    meses = calor.get("meses") or []
    colunas = [("Macroclasse", 22, None), ("Peer group", 44, None)] + [
        (competencia_br(m), 10, PCT_SINAL) for m in meses
    ]
    aba.tabela(colunas, [[l["n1"], l["n2"], *l["valores"]] for l in calor.get("linhas") or []])


def _aba_familias(wb: Workbook, bloco: dict) -> None:
    familias = bloco["familias"]
    aba = _Aba(wb, "Famílias", "Estratégias que mais captam e mais perdem (master + feeders + espelhos)")
    aba.escrever("Camada do investidor. " + _prazo(bloco), _NOTA)
    colunas = [("Estratégia (raiz)", 50, None), ("Gestora", 24, None), ("Macroclasse", 20, None),
               ("Peer group", 36, None), ("Veículos", 10, INT), ("PL (R$ mi)", 14, MI),
               ("Captação 12m (R$ mi)", 14, MI_SINAL), ("Captação / PL de abertura", 14, PCT_SINAL)]
    for titulo, chave in (("Que mais captam", "que_mais_captam"), ("Que mais perdem", "que_mais_perdem")):
        aba.escrever(titulo, _NEGRITO)
        aba.tabela(colunas, [
            [f["nome"], f["gestora"], f["n1"], f["n2"], f["veiculos"], _mi(f["pl_investidor_fim"]),
             _mi(f["fluxo_investidor"]), f.get("sobre_pl")]
            for f in familias.get(chave) or []
        ])


def _aba_gestoras(wb: Workbook, bloco: dict) -> None:
    gestoras = bloco["gestoras"]
    aba = _Aba(wb, "Gestoras", "Quem ganha: grupos gestores da estratégia")
    aba.escrever(
        "Grupo econômico da gestora da raiz (master). Share do PL da camada do investidor no fim e "
        f"12 meses antes; {_prazo(bloco)}. Grupo curado cobre {pct(gestoras.get('cobertura_curada'))} do PL.",
        _NOTA,
    )
    colunas = [("Grupo", 30, None), ("Tipo", 16, None), ("PL (R$ mi)", 14, MI), ("Share", 10, PCT),
               ("Share 12m antes", 10, PCT), ("Variação", 12, PP), ("Captação 12m (R$ mi)", 14, MI_SINAL),
               ("Captação / PL de abertura", 14, PCT_SINAL)]
    aba.tabela(colunas, [
        [g["grupo"], TIPOS_DE_GESTORA.get(g["tipo"], g["tipo"]), _mi(g["pl_investidor"]), g.get("share"),
         g.get("share_12m_antes"), g.get("variacao_pp"), _mi(g.get("fluxo_investidor_12m")), g.get("sobre_pl_12m")]
        for g in gestoras.get("grupos") or []
    ], filtro=True)
    aba.escrever("Por tipo de gestora", _NEGRITO)
    aba.tabela(colunas, [
        [t.get("rotulo") or t["tipo"], "", _mi(t["pl_investidor"]), t.get("share"), t.get("share_12m_antes"),
         t.get("variacao_pp"), _mi(t.get("fluxo_investidor_12m")), t.get("sobre_pl_12m")]
        for t in gestoras.get("tipos") or []
    ])


def _aba_dinamica(wb: Workbook, bloco: dict) -> None:
    dinamica = bloco.get("dinamica") or {}
    aba = _Aba(wb, "Dinâmica", "Dinâmica da captação: giro, concentração, entrantes e fluxo × retorno")
    aba.escrever(f"Tudo na janela de {_prazo(bloco)}; PL de abertura no fechamento de "
                 f"{data_br(bloco['janelas']['12m']['data_base'])}.", _NOTA)

    giro = dinamica.get("giro") or {}
    aba.escrever("Giro: quanto entra e quanto sai (fluxo declarado, com o dinheiro de fundos)", _NEGRITO)
    colunas = [("Macroclasse", 24, None), ("Captação bruta (R$ mi)", 16, MI), ("Resgates (R$ mi)", 16, MI),
               ("Líquido (R$ mi)", 14, MI_SINAL), ("PL médio (R$ mi)", 16, MI),
               ("Entrada (× PL médio)", 12, "0.00"), ("Giro (× PL médio)", 12, "0.00")]

    def linha_giro(rotulo: str, item: dict) -> list:
        return [rotulo, _mi(item.get("captacao")), _mi(item.get("resgates")), _mi(item.get("captacao_liquida")),
                _mi(item.get("pl_medio")), item.get("entrada"), item.get("giro")]

    total = giro.get("total") or {}
    aba.tabela(colunas, [linha_giro(i["n1"], i) for i in giro.get("n1") or []],
               total=linha_giro("Universo", total) if total else None)

    concentracao = dinamica.get("concentracao") or {}
    aba.escrever("Concentração por peer group: grupos gestores da estratégia, pelo PL do investidor", _NEGRITO)
    aba.escrever("HHI de 0 a 10.000 (soma dos quadrados das fatias). Variação contra 12 meses antes.", _NOTA)
    colunas = [("Macroclasse", 20, None), ("Peer group", 40, None), ("PL (R$ mi)", 14, MI), ("Gestoras", 10, INT),
               ("Gestoras 12m antes", 10, INT), ("HHI", 10, "#,##0"), ("HHI 12m antes", 10, "#,##0"),
               ("Variação do HHI", 10, "+#,##0;-#,##0;0"), ("5 maiores", 10, PCT), ("5 maiores 12m antes", 10, PCT),
               ("Líder", 30, None), ("Fatia do líder", 10, PCT)]

    def linha_conc(n1: str, n2: str, item: dict) -> list:
        return [n1, n2, _mi(item.get("pl_investidor")), item.get("gestoras"), item.get("gestoras_12m_antes"),
                item.get("hhi"), item.get("hhi_12m_antes"), item.get("variacao_hhi"), item.get("top5"),
                item.get("top5_12m_antes"), item.get("lider"), item.get("fatia_do_lider")]

    universo = concentracao.get("universo") or {}
    aba.tabela(colunas, [linha_conc(i["n1"], i["n2"], i) for i in concentracao.get("n2") or []],
               total=linha_conc("Universo", "", universo) if universo else None, filtro=True)

    gestoras = dinamica.get("gestoras") or {}
    entraram, sairam = gestoras.get("entraram") or {}, gestoras.get("sairam") or {}
    aba.escrever(
        f"Grupos gestores que entraram e saíram: {gestoras.get('grupos_na_base', 0)} no universo na base, "
        f"{gestoras.get('grupos_no_fim', 0)} no fim; entraram {entraram.get('grupos', 0)}, saíram "
        f"{sairam.get('grupos', 0)}", _NEGRITO,
    )
    aba.escrever("Entrou: sem produto no universo na base e com produto dentro da janela. Saiu: tinha produto na "
                 "base ou na janela e não tem no fim. Os 20 maiores de cada lado.", _NOTA)
    aba.tabela(
        [("Grupo gestor", 40, None), ("Tipo", 16, None), ("Primeira competência", 14, None),
         ("Peer group principal", 34, None), ("PL do investidor no fim (R$ mi)", 16, MI),
         ("Captação 12m (R$ mi)", 14, MI_SINAL)],
        [[i["grupo"], TIPOS_DE_GESTORA.get(i.get("tipo"), i.get("tipo")), competencia_br(i["primeira"]), i.get("n2"),
          _mi(i.get("pl_investidor")), _mi(i.get("fluxo_investidor"))] for i in entraram.get("lista") or []],
    )
    aba.tabela(
        [("Grupo gestor", 40, None), ("Tipo", 16, None), ("Última competência", 14, None),
         ("PL do investidor ao sair (R$ mi)", 16, MI)],
        [[i["grupo"], TIPOS_DE_GESTORA.get(i.get("tipo"), i.get("tipo")), competencia_br(i["ultima"]),
          _mi(i.get("pl_investidor"))] for i in sairam.get("lista") or []],
    )

    aba.escrever("O dinheiro vai atrás do retorno? Captação por quartil de retorno de 12 meses no peer group", _NEGRITO)
    aba.escrever("Mesma janela nos dois eixos: mostra associação, não causa. Quartil 1 é o de cima.", _NOTA)
    aba.tabela(
        [("Quartil", 30, None), ("Classes no fim", 12, INT), ("PL de abertura (R$ mi)", 16, MI),
         ("PL no fim (R$ mi)", 16, MI), ("Captação 12m (R$ mi)", 16, MI_SINAL), ("Captação / PL de abertura", 14, PCT_SINAL)],
        [[q["rotulo"], q["classes_fim"], _mi(q["pl_investidor_abertura"]), _mi(q["pl_investidor_fim"]),
          _mi(q["fluxo_investidor"]), q.get("sobre_pl")] for q in dinamica.get("fluxo_x_retorno") or []],
    )


def _aba_macro(wb: Workbook, bloco: dict) -> None:
    macro = bloco.get("macro") or {}
    aba = _Aba(wb, "Macro", "O macro ao lado da captação: CDI, IPCA, Selic e dólar")
    if not macro.get("available"):
        aba.escrever(macro.get("reason") or "Sem o SGS do Banco Central.", _NOTA)
        return
    aba.escrever(f"SGS do Banco Central: CDI (série 12), meta Selic (432), IPCA (433), dólar PTAX venda (1). "
                 f"{macro.get('convencao', '')}.", _NOTA)
    colunas = [("Janela", 14, None), ("De", 12, None), ("Até", 12, None), ("CDI", 10, "0.00%"),
               ("IPCA", 10, "0.00%"), ("Selic na base (% a.a.)", 12, "0.00"), ("Selic no fim (% a.a.)", 12, "0.00"),
               ("Dólar na base (R$)", 12, "0.0000"), ("Dólar no fim (R$)", 12, "0.0000"), ("Variação do dólar", 12, PCT_SINAL)]
    linhas = []
    for chave in ORDEM:
        janela, item = bloco["janelas"][chave], (macro.get("janelas") or {}).get(chave) or {}
        linhas.append([janela["rotulo"], data_br(janela.get("data_base")), data_br(janela.get("ate")), item.get("cdi"),
                       item.get("ipca"), item.get("selic_base"), item.get("selic_fim"), item.get("dolar_base"),
                       item.get("dolar_fim"), item.get("dolar_variacao")])
    aba.tabela(colunas, linhas)
    meses = macro.get("meses") or []
    n1s = [n1 for n1 in N1_ORDEM if any(n1 in (m.get("por_n1") or {}) for m in meses)]
    colunas = [("Competência", 12, None), ("Captação do investidor (R$ mi)", 16, MI_SINAL)]
    colunas += [(f"{n1} (R$ mi)", 14, MI_SINAL) for n1 in n1s]
    colunas += [("CDI no mês", 10, "0.00%"), ("IPCA no mês", 10, "+0.00%;-0.00%;0.00%"),
                ("Selic no fim (% a.a.)", 12, "0.00"), ("Dólar no fim (R$)", 12, "0.0000")]
    aba.tabela(colunas, [
        [competencia_br(m["competencia"]), _mi(m["fluxo_investidor"])]
        + [_mi((m.get("por_n1") or {}).get(n1)) for n1 in n1s]
        + [m.get("cdi"), m.get("ipca"), m.get("selic"), m.get("dolar")]
        for m in meses
    ])


def _aba_dono(wb: Workbook, bloco: dict) -> None:
    aba = _Aba(wb, "Dono do dinheiro", "De quem é o dinheiro do universo, e por onde ele chega")
    aba.escrever(
        f"Perfil mensal da CVM em {bloco['rotulo_competencia']} e 12 meses antes. PL bruto do universo "
        "(a fatia de fundos é a dupla contagem que a camada do investidor tira).", _NOTA,
    )
    aba.tabela(
        [("Segmento", 30, None), ("PL (R$ mi)", 14, MI), ("Fatia", 10, PCT),
         ("Fatia 12m antes", 10, PCT), ("Variação", 12, PP)],
        [[s["rotulo"], _mi(s["pl"]), s.get("fatia"), s.get("fatia_12m_antes"), s.get("variacao_pp")]
         for s in bloco["segmentos"]],
    )
    plataforma = bloco.get("plataformizacao") or {}
    colunas = [("", 30, None), ("PL (R$ mi)", 14, MI), ("Conta e ordem", 12, PCT),
               ("12m antes", 12, PCT), ("Variação", 12, PP)]
    aba.escrever("Plataformização: fatia do PL por conta e ordem", _NEGRITO)
    total = plataforma.get("total") or {}
    aba.tabela(colunas, [
        [i["n1"], _mi(i["pl"]), i.get("conta_e_ordem"), i.get("conta_e_ordem_12m_antes"), i.get("variacao_pp")]
        for i in plataforma.get("por_n1") or []
    ], total=["Universo", _mi(total.get("pl")), total.get("conta_e_ordem"),
              total.get("conta_e_ordem_12m_antes"), total.get("variacao_pp")])
    aba.tabela(colunas, [
        [i["gestora"], _mi(i["pl"]), i.get("conta_e_ordem"), i.get("conta_e_ordem_12m_antes"), i.get("variacao_pp")]
        for i in plataforma.get("por_gestora") or []
    ])
    segmentos = (bloco.get("dinamica") or {}).get("captacao_por_segmento") or {}
    total = segmentos.get("total")
    if not total:
        return
    aba.escrever(f"Quem trouxe o dinheiro: captação implícita por tipo de cotista, {_prazo(bloco)}", _NEGRITO)
    aba.escrever(
        "Por classe e mês: PL(t) × fatia(t) − PL(t−1) × fatia(t−1) × (1 + retorno); sem o mês anterior, o fluxo "
        "declarado × fatia(t). Os tipos do investidor mais o resíduo dão a captação do investidor; \"fundos\" é a "
        f"dupla contagem que a camada do investidor tira. Ruído do arredondamento do perfil no universo: "
        f"{reais(total.get('ruido'))}; valor abaixo de {reais(total.get('piso'))} sai marcado.", _NOTA,
    )
    chaves = [s["chave"] for s in total["segmentos"]]
    rotulos = {s["chave"]: s["rotulo"] for s in total["segmentos"]}
    colunas = [("Macroclasse", 22, None), ("Peer group", 36, None)]
    colunas += [(f"{rotulos[c]} (R$ mi)", 14, MI_SINAL) for c in chaves]
    colunas += [("Resíduo (R$ mi)", 14, MI_SINAL), ("Captação do investidor (R$ mi)", 16, MI_SINAL),
                ("Fundos, fora da camada (R$ mi)", 16, MI_SINAL), ("Ruído (R$ mi)", 12, MI),
                ("Tipos dentro do ruído", 30, None)]

    def linha(n1: str, n2: str, item: dict) -> list:
        valores = {s["chave"]: s for s in item.get("segmentos") or []}
        return (
            [n1, n2] + [_mi((valores.get(c) or {}).get("captacao")) for c in chaves]
            + [_mi(item.get("residuo")), _mi(item.get("fluxo_investidor")), _mi(item.get("fundos")),
               _mi(item.get("ruido")), ", ".join(s["rotulo"] for s in item.get("segmentos") or [] if s.get("no_ruido"))]
        )

    aba.tabela(colunas, [linha(i["n1"], "", i) for i in segmentos.get("n1") or []]
               + [linha(i["n1"], i["n2"], i) for i in segmentos.get("n2") or []],
               total=linha("Universo", "", total), filtro=True)


def _aba_estrutura(wb: Workbook, bloco: dict) -> None:
    estrutura = bloco["estrutura"]
    aba = _Aba(wb, "Estrutura", "Master, feeder, espelho e FoF; e para quem cada veículo foi feito")
    aba.escrever(
        "Papel pela carteira da CVM (" + ", ".join(competencia_br(c) for c in estrutura.get("competencias_da_carteira") or [])
        + ") e pelo perfil mensal. Casa dedicada: espelho de grupo distribuidor ou marca de outra casa no nome.",
        _NOTA,
    )
    aba.tabela(
        [("Papel", 28, None), ("Produtos", 10, INT), ("PL (R$ mi)", 14, MI), ("Fatia", 10, PCT),
         ("Captação 12m (R$ mi)", 14, MI_SINAL)],
        [[p["rotulo"], p["produtos"], _mi(p["pl_investidor"]), p.get("fatia"), _mi(p.get("fluxo_investidor_12m"))]
         for p in estrutura.get("papeis") or []],
    )
    aba.escrever("Casas dedicadas", _NEGRITO)
    aba.tabela(
        [("Casa", 28, None), ("Veículos", 10, INT), ("PL (R$ mi)", 14, MI), ("Fatia", 10, PCT),
         ("Captação 12m (R$ mi)", 14, MI_SINAL)],
        [[c["casa"], c["produtos"], _mi(c["pl_investidor"]), c.get("fatia"), _mi(c.get("fluxo_investidor_12m"))]
         for c in estrutura.get("casas_dedicadas") or []],
    )
    aba.escrever("Parcerias: casa dedicada × gestora da estratégia", _NEGRITO)
    aba.tabela(
        [("Casa", 28, None), ("Gestora da estratégia", 28, None), ("Veículos", 10, INT), ("PL (R$ mi)", 14, MI)],
        [[p["casa"], p["gestora"], p["produtos"], _mi(p["pl_investidor"])] for p in estrutura.get("parcerias") or []],
        filtro=True,
    )


def _aba_publico(wb: Workbook, bloco: dict) -> None:
    aba = _Aba(wb, "Público", "Captação por público-alvo")
    aba.escrever("Camada do investidor. " + " · ".join(bloco.get("prazos") or ()), _NOTA)
    colunas = [("Público", 20, None), ("Produtos", 10, INT), ("PL (R$ mi)", 14, MI), ("Fatia", 10, PCT)]
    colunas += [(f"Captação {ROTULOS_JANELA[j]} (R$ mi)", 14, MI_SINAL) for j in ORDEM]
    aba.tabela(colunas, [
        [ROTULOS_PUBLICO.get(p["publico"], p["publico"]), p["produtos"], _mi(p["pl_investidor"]), p.get("fatia")]
        + [_mi(_fluxo(p, j)) for j in ORDEM]
        for p in bloco.get("publico") or []
    ])


_COLUNAS_EXPLORADOR: list[tuple[str, str, int, str | None]] = [
    ("cnpj", "CNPJ", 20, None), ("subclasse", "Subclasse", 18, None), ("nome", "Nome", 60, None),
    ("n1", "Macroclasse", 20, None), ("n2", "Peer group", 36, None), ("confianca", "Confiança", 10, None),
    ("papel", "Papel", 16, None), ("raiz_nome", "Estratégia (raiz)", 44, None), ("gestora", "Gestora da estratégia", 24, None),
    ("tipo_gestora", "Tipo de gestora", 16, None), ("gestora_do_veiculo", "Gestora do veículo", 24, None),
    ("casa_dedicada", "Casa dedicada", 16, None), ("publico", "Público", 14, None),
    ("faixa_liquidez", "Liquidez", 14, None), ("aplicacao_minima", "Aplicação mínima (R$)", 14, INT),
    ("taxa_adm", "Taxa de administração (% a.a.)", 12, "0.00"),
    ("pl", "PL (R$ mi)", 14, MI), ("pl_investidor", "PL do investidor (R$ mi)", 14, MI),
    ("fatia_fundos", "Detido por fundos", 10, PCT), ("cotistas", "Cotistas", 12, INT),
    ("fluxo_1m", "Captação 1 mês (R$ mi)", 14, MI_SINAL), ("fluxo_3m", "Captação 3 meses (R$ mi)", 14, MI_SINAL),
    ("fluxo_ano", "Captação no ano (R$ mi)", 14, MI_SINAL), ("fluxo_12m", "Captação 12 meses (R$ mi)", 14, MI_SINAL),
    ("sobre_pl_12m", "Captação 12m / PL de abertura", 12, PCT_SINAL),
    ("retorno_12m", "Retorno 12m", 10, PCT), ("quartil_12m", "Quartil 12m no peer group", 10, INT),
    ("retorno_3m", "Retorno 3m", 10, PCT), ("quartil_3m", "Quartil 3m", 10, INT),
    ("conta_e_ordem", "Conta e ordem", 10, PCT), ("pf_pj", "PF e PJ", 10, PCT),
]
_EM_MILHOES = {"pl", "pl_investidor", "fluxo_1m", "fluxo_3m", "fluxo_ano", "fluxo_12m"}


def _valor_do_explorador(campo: str, linha: dict):
    valor = linha.get(campo)
    if campo in _EM_MILHOES:
        return _mi(valor)
    if campo == "taxa_adm" and valor is not None:
        return valor
    if campo == "papel":
        return PAPEIS.get(valor, valor)
    if campo == "tipo_gestora":
        return TIPOS_DE_GESTORA.get(valor, valor)
    if campo == "publico":
        return ROTULOS_PUBLICO.get(valor, valor)
    return valor


def _aba_explorador(wb: Workbook, bloco: dict) -> None:
    aba = _Aba(wb, "Explorador", "Explorador: um produto do universo por linha")
    aba.escrever(
        "Filtre por macroclasse, peer group, papel, gestora, público ou casa dedicada. Quartil 1 é o de "
        "cima. " + " · ".join(bloco.get("prazos") or ()), _NOTA,
    )
    aba.tabela(
        [(titulo, largura, formato) for _, titulo, largura, formato in _COLUNAS_EXPLORADOR],
        [[_valor_do_explorador(campo, linha) for campo, *_ in _COLUNAS_EXPLORADOR]
         for linha in bloco.get("_explorador") or []],
        filtro=True, congelar=True,
    )


def _aba_lancamentos(wb: Workbook, bloco: dict) -> None:
    aba = _Aba(wb, "Lançamentos", "O que entrou e o que saiu do universo")
    aba.escrever(
        "Entrou: primeira competência no universo dentro da janela, com o cadastro começando no máximo 45 dias "
        "antes. Saiu: estava no universo na janela e não está na competência de referência.", _NOTA,
    )
    colunas = [("Nome", 56, None), ("CNPJ", 16, None), ("Competência", 12, None), ("Macroclasse", 20, None),
               ("Peer group", 34, None), ("Gestora", 24, None), ("Papel", 16, None), ("PL do investidor (R$ mi)", 14, MI)]
    for chave in ("1m", "12m"):
        dados = (bloco.get("lancamentos") or {}).get(chave) or {}
        for lado, titulo in (("entrou", "Entraram"), ("saiu", "Saíram")):
            info = dados.get(lado) or {}
            aba.escrever(
                f"{titulo} em {_prazo(bloco, chave)}: {info.get('produtos', 0)} produtos", _NEGRITO,
            )
            aba.tabela(colunas, [
                [i["nome"], i["cnpj"], competencia_br(i["competencia"]), i["n1"], i["n2"], i["gestora"],
                 PAPEIS.get(i["papel"], i["papel"]), _mi(i["pl_investidor"])]
                for i in info.get("maiores") or []
            ])


def _aba_adjacentes(wb: Workbook, bloco: dict) -> None:
    aba = _Aba(wb, "Camadas ao lado", "Previdência e ETF: ao lado do universo, nunca na soma")
    colunas = [("Camada", 18, None), ("Fonte", 50, None), ("PL (R$ mi)", 14, MI)]
    colunas += [(f"Captação {ROTULOS_JANELA[j]} (R$ mi)", 14, MI_SINAL) for j in ORDEM]
    linhas = []
    for item in (bloco.get("adjacentes") or {}).values():
        if item.get("available") is False:
            linhas.append([item["rotulo"], item.get("reason"), None] + [None] * len(ORDEM))
            continue
        linhas.append(
            [item["rotulo"], item.get("fonte"), _mi(item.get("pl_investidor", item.get("pl")))]
            + [_mi(((item.get("janelas") or {}).get(j) or {}).get("fluxo_investidor")) for j in ORDEM]
        )
    aba.tabela(colunas, linhas)


def _aba_qualidade(wb: Workbook, bloco: dict) -> None:
    q = bloco["qualidade"]
    estrutura = bloco.get("estrutura") or {}
    estrategias = bloco.get("estrategias") or {}
    aba = _Aba(wb, "Qualidade", "Qualidade dos dados: o que sustenta cada número")
    linhas = [
        ["Perfil mensal do próprio mês", pct(q.get("cobertura_perfil")), "fração do PL do universo"],
        ["Grupo gestor curado", pct((bloco.get("gestoras") or {}).get("cobertura_curada")), "fração do PL"],
        ["Estratégia com confiança alta", pct((estrategias.get("confianca") or {}).get("alta")), "fração do PL"],
        ["PL sem carteira da CVM", reais(estrutura.get("pl_sem_carteira")), ""],
        ["PL com perfil e carteira divergindo", reais(estrutura.get("pl_com_divergencia")), "acima de 25 pontos"],
        ["Produtos sem cadastro (no informe)", milhar(q["sem_cadastro"]["produtos"]), reais(q["sem_cadastro"]["pl"])],
        ["Eventos na cota no mês", milhar(q.get("eventos_na_cota") or 0), "salto acima de 50%: fica fora do retorno"],
        ["Resíduo da decomposição / variação do PL", pct(q.get("residuo_sobre_variacao")), "soma dos meses da série"],
    ]
    aba.tabela([("Indicador", 46, None), ("Valor", 16, None), ("Nota", 50, None)], linhas)
    aba.escrever(
        "Competências com perfil: " + ", ".join(competencia_br(m) for m in q.get("meses_com_perfil") or []), _NOTA,
    )


def _aba_periodos(wb: Workbook, bloco: dict) -> None:
    aba = _Aba(wb, "Período de cada aba", "O prazo de cada número desta planilha")
    linhas = [["Funil, Dono do dinheiro, Estrutura, Explorador (PL)", f"fim de {bloco['rotulo_competencia']} ({data_br(bloco['informe']['ate'])})"]]
    for chave in ORDEM:
        janela = bloco["janelas"][chave]
        linhas.append([f"Captação {janela['rotulo']}", f"de {data_br(janela['de'])} a {data_br(janela['ate'])}; "
                       f"PL de abertura no fechamento de {data_br(janela['data_base'])}"])
    linhas.append(["Retorno 12m e 3m", "base no fechamento da competência anterior à janela (convenção de lâmina)"])
    linhas.append(["Dinâmica (giro, concentração, entrantes, fluxo × retorno) e captação por tipo de cotista",
                   f"12 meses: de {data_br(bloco['janelas']['12m']['de'])} a {data_br(bloco['janelas']['12m']['ate'])}"])
    linhas.append(["Macro", "CDI e IPCA acumulados em cada janela; Selic e dólar no último dia útil da base e do fim"])
    aba.tabela([("Aba ou número", 50, None), ("Prazo", 90, None)], linhas)


def _aba_fontes(wb: Workbook, bloco: dict) -> None:
    aba = _Aba(wb, "Fontes", "Fontes")
    aba.escrever("Onde conferir cada número desta planilha - clique no nome para abrir.", _NOTA)
    inicio = aba.linha
    aba.tabela([("Fonte", 48, None), ("O que entra", 70, None), ("Endereço para consulta", 70, None)], [])
    aba.linha = inicio + 1
    com_carteira = (bloco.get("estrutura") or {}).get("com_carteira", True)
    com_macro = bool((bloco.get("macro") or {}).get("available"))
    for fonte in fontes_da_industria(com_carteira=com_carteira, com_macro=com_macro):
        nome = aba.ws.cell(row=aba.linha, column=1, value=fonte.nome)
        nome.hyperlink, nome.font = fonte.url, _LINK
        aba.ws.cell(row=aba.linha, column=2, value=fonte.uso).font = _CORPO
        url = aba.ws.cell(row=aba.linha, column=3, value=fonte.url)
        url.hyperlink, url.font = fonte.url, _LINK
        aba.linha += 1


def caminho_da_planilha(config: Config, bloco: dict) -> Path:
    return config.reports_dir / f"industria-{bloco['competencia'][:4]}-{bloco['competencia'][4:]}.xlsx"


def gravar_planilha(config: Config, bloco: dict, gerado: datetime | None = None) -> Path:
    """Grava a planilha da competencia e devolve o caminho de fato escrito."""
    gerado = gerado or datetime.now()
    wb = Workbook()
    _aba_leitura(wb, bloco, gerado)
    _aba_funil(wb, bloco)
    _aba_estrategias(wb, bloco)
    _aba_mapa_de_calor(wb, bloco)
    _aba_familias(wb, bloco)
    _aba_gestoras(wb, bloco)
    _aba_dinamica(wb, bloco)
    _aba_dono(wb, bloco)
    _aba_estrutura(wb, bloco)
    _aba_publico(wb, bloco)
    _aba_explorador(wb, bloco)
    _aba_lancamentos(wb, bloco)
    _aba_adjacentes(wb, bloco)
    _aba_macro(wb, bloco)
    _aba_qualidade(wb, bloco)
    _aba_periodos(wb, bloco)
    _aba_fontes(wb, bloco)
    wb.active = 0
    destino = caminho_da_planilha(config, bloco)
    destino.parent.mkdir(parents=True, exist_ok=True)
    return _salvar(wb, destino)
