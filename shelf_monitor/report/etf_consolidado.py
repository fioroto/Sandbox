"""Planilha consolidada de ETF: uma linha por ETF, todas as colunas numa aba so.

A aba "ETFs" tem o dado; a do dicionario diz de onde vem cada coluna e como
foi calculada; "Fontes" tem o link de cada fonte. O CSV ao lado tem as mesmas
colunas, para quem cruza com outra base.
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from shelf_monitor.market.etf_bruto import armazem
from shelf_monitor.market.etf_bruto.consolidacao import (
    NAO_SE_APLICA,
    PENDENTE_CADASTRO,
    PENDENTE_PDF,
    SEM_COMDINHEIRO,
    Consolidado,
)
from shelf_monitor.market.etf_bruto.mapa import SECOES
from shelf_monitor.report.shared.fontes import FONTES_ETF_BRUTO
from shelf_monitor.texto import competencia_br, contagem, data_br

FONT = "Arial"
_TITULO = Font(name=FONT, size=13, bold=True)
_NOTA = Font(name=FONT, size=9, italic=True, color="595959")
_CORPO = Font(name=FONT, size=10)
_CABECALHO = Font(name=FONT, size=10, bold=True, color="FFFFFF")
_FUNDO_CABECALHO = PatternFill("solid", fgColor="1F3864")
_SECAO = Font(name=FONT, size=10, bold=True, color="1F3864")
_FUNDO_SECAO = PatternFill("solid", fgColor="D9E1F2")
_LINK = Font(name=FONT, size=10, color="1A5FB4", underline="single")
# Pendente de leitura de PDF em amarelo: e o que uma pessoa ainda precisa ler.
_FUNDO_PDF = PatternFill("solid", fgColor="FFF2CC")
_PENDENTE = Font(name=FONT, size=10, italic=True, color="7F6000")
_AUSENTE = Font(name=FONT, size=10, italic=True, color="808080")

_FORMATOS = {
    "pct": "0.00%",
    "pp": '+0.00" pp";-0.00" pp";0.00" pp"',
    "preco": "#,##0.00",
    "dinheiro": "#,##0",
    "inteiro": "#,##0",
    "decimal": "0.00",
    "data": "dd/mm/yyyy",
}
_MARCAS_AUSENTES = {PENDENTE_CADASTRO, SEM_COMDINHEIRO, NAO_SE_APLICA, "não calculado", "não coletado (site do fundo)"}

LINHA_DAS_SECOES = 5
LINHA_DO_CABECALHO = 6


def _notas(c: Consolidado) -> tuple[str, str]:
    total = len(c.linhas)
    negociam = sum(1 for l in c.linhas if l.get("situacao") == "negocia na B3")
    j = c.janelas
    partes = []
    if j.ultimo_pregao:
        partes.append(f"último pregão em disco {data_br(j.ultimo_pregao)}")
    if j.liquidez:
        partes.append(f"liquidez: 63 pregões de {data_br(j.liquidez[0])} a {data_br(j.liquidez[1])}")
    if j.doze_meses:
        partes.append(f"12 meses: de {data_br(j.doze_meses[0])} a {data_br(j.doze_meses[1])}")
    if j.cota_12m:
        partes.append(f"cota: fechamento de {competencia_br(j.cota_12m[0])} a {competencia_br(j.cota_12m[1])}")
    if j.carteira:
        partes.append(f"carteira da CVM de {competencia_br(j.carteira)}")
    primeira = (
        f"Uma linha por ETF: {contagem(total, 'ETF', 'ETFs')}, {negociam} negociando na B3. Gerado em "
        f"{c.gerado_em:%d/%m/%Y %H:%M} a partir de data/etf/bruto/. \"{PENDENTE_PDF}\": o dado só existe no "
        f"regulamento ou nas demonstrações financeiras (PDF) e ninguém o conferiu ainda; \"{PENDENTE_CADASTRO}\": "
        "preencher em data/etf/manual/."
    )
    segunda = "Janelas - " + "; ".join(partes) + ". Rentabilidade de preço é fechamento contra fechamento, sem proventos."
    return primeira, segunda


def _aba_etfs(wb: Workbook, c: Consolidado) -> None:
    ws = wb.active
    ws.title = "ETFs"
    ws.cell(row=1, column=1, value="ETFs - consolidação dos dados brutos").font = _TITULO
    primeira, segunda = _notas(c)
    ws.cell(row=2, column=1, value=primeira).font = _NOTA
    ws.cell(row=3, column=1, value=segunda).font = _NOTA

    # Faixa das secoes: uma celula mesclada por secao, sobre as colunas dela.
    inicio = 1
    for secao in SECOES:
        largura = sum(1 for col in c.colunas if col.secao == secao)
        if not largura:
            continue
        celula = ws.cell(row=LINHA_DAS_SECOES, column=inicio, value=secao)
        celula.font, celula.fill = _SECAO, _FUNDO_SECAO
        celula.alignment = Alignment(horizontal="left", vertical="center")
        for coluna in range(inicio + 1, inicio + largura):
            ws.cell(row=LINHA_DAS_SECOES, column=coluna).fill = _FUNDO_SECAO
        if largura > 1:
            ws.merge_cells(start_row=LINHA_DAS_SECOES, start_column=inicio, end_row=LINHA_DAS_SECOES,
                           end_column=inicio + largura - 1)
        inicio += largura

    for indice, coluna in enumerate(c.colunas, start=1):
        celula = ws.cell(row=LINHA_DO_CABECALHO, column=indice, value=coluna.titulo)
        celula.font, celula.fill = _CABECALHO, _FUNDO_CABECALHO
        celula.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(indice)].width = coluna.largura
    ws.row_dimensions[LINHA_DO_CABECALHO].height = 66

    for numero, linha in enumerate(c.linhas, start=LINHA_DO_CABECALHO + 1):
        for indice, coluna in enumerate(c.colunas, start=1):
            valor = linha.get(coluna.chave)
            celula = ws.cell(row=numero, column=indice)
            if valor is None or valor == "":
                continue
            if coluna.formato == "link":
                celula.value = "abrir"
                celula.hyperlink = str(valor)
                celula.font = _LINK
                continue
            if isinstance(valor, datetime):
                valor = valor.date()
            celula.value = valor
            if valor == PENDENTE_PDF:
                celula.font, celula.fill = _PENDENTE, _FUNDO_PDF
            elif isinstance(valor, str) and valor in _MARCAS_AUSENTES:
                celula.font = _AUSENTE
            else:
                celula.font = _CORPO
                if coluna.formato in _FORMATOS and isinstance(valor, (int, float, date)):
                    celula.number_format = _FORMATOS[coluna.formato]
    ultima = LINHA_DO_CABECALHO + len(c.linhas)
    ws.freeze_panes = ws.cell(row=LINHA_DO_CABECALHO + 1, column=2)
    ws.auto_filter.ref = f"A{LINHA_DO_CABECALHO}:{get_column_letter(len(c.colunas))}{max(ultima, LINHA_DO_CABECALHO)}"


def _aba_dicionario(wb: Workbook, c: Consolidado) -> None:
    ws = wb.create_sheet("Dicionário")
    ws.cell(row=1, column=1, value="Dicionário das colunas").font = _TITULO
    ws.cell(row=2, column=1, value="De onde vem cada coluna da aba ETFs e como foi preenchida.").font = _NOTA
    cabecalhos = (("Seção", 30), ("Coluna", 48), ("Indicador do mapeamento", 46), ("Como foi preenchida", 90))
    for indice, (titulo, largura) in enumerate(cabecalhos, start=1):
        celula = ws.cell(row=4, column=indice, value=titulo)
        celula.font, celula.fill = _CABECALHO, _FUNDO_CABECALHO
        ws.column_dimensions[get_column_letter(indice)].width = largura
    for numero, coluna in enumerate(c.colunas, start=5):
        for indice, valor in enumerate((coluna.secao, coluna.titulo, coluna.indicador, coluna.como), start=1):
            celula = ws.cell(row=numero, column=indice, value=valor)
            celula.font = _CORPO
            celula.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A5"


def _aba_fontes(wb: Workbook) -> None:
    ws = wb.create_sheet("Fontes")
    ws.cell(row=1, column=1, value="Fontes").font = _TITULO
    ws.cell(row=2, column=1, value="Onde conferir cada número da planilha - clique no nome para abrir.").font = _NOTA
    for indice, (titulo, largura) in enumerate((("Fonte", 52), ("O que entra", 70), ("Endereço para consulta", 90)), start=1):
        celula = ws.cell(row=4, column=indice, value=titulo)
        celula.font, celula.fill = _CABECALHO, _FUNDO_CABECALHO
        ws.column_dimensions[get_column_letter(indice)].width = largura
    for numero, fonte in enumerate(FONTES_ETF_BRUTO, start=5):
        nome = ws.cell(row=numero, column=1, value=fonte.nome)
        nome.hyperlink, nome.font = fonte.url, _LINK
        ws.cell(row=numero, column=2, value=fonte.uso).font = _CORPO
        url = ws.cell(row=numero, column=3, value=fonte.url)
        url.hyperlink, url.font = fonte.url, _LINK


def escrever_planilha(c: Consolidado, caminho: Path) -> Path:
    wb = Workbook()
    _aba_etfs(wb, c)
    _aba_dicionario(wb, c)
    _aba_fontes(wb)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho


def _texto_csv(valor) -> object:
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    return valor


def escrever_csv(c: Consolidado, caminho: Path) -> Path:
    """As mesmas colunas, com data ISO e numero com ponto - para cruzar com outra base."""
    registros = [{col.titulo: _texto_csv(linha.get(col.chave)) for col in c.colunas} for linha in c.linhas]
    armazem.gravar_csv(caminho, [col.titulo for col in c.colunas], registros)
    return caminho
