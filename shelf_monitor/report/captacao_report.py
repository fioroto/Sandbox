"""Relatorios standalone do monitor de captacao (subcomando `captacao`).

Console via `rich` e uma planilha dedicada (`captacao-<data>.xlsx`), no mesmo
padrao de `report/etf_report.py`: workbook e constantes de estilo proprios,
sem depender de `report/spreadsheet.py` (que traz as MESMAS abas, mas como
parte de `prateleiras-<data>.xlsx` - ver `_add_captacao_sheets` la).
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rich.console import Console
from rich.table import Table

from shelf_monitor.captacao import (
    NOTA_FUNDO_NIVEL,
    ResultadoCaptacao,
    analise_cruzada,
    consolidar_por_classificacao,
    ranking_consolidado,
    ranking_por_plataforma,
)
from shelf_monitor.config import PLATFORMS

FONT = "Arial"
_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
_TITLE_FONT = Font(name=FONT, size=13, bold=True)
_BODY_FONT = Font(name=FONT, size=10)
_NOTE_FONT = Font(name=FONT, size=9, italic=True, color="595959")
_ALERT_FONT = Font(name=FONT, size=10, bold=True, color="C00000")
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(bottom=_THIN)

_MONEY = "#,##0;(#,##0);-"
_PCT = '0.00"%";-;-'
_INT = "0;-;-"

_LIMITE_CONSOLE = 10


def _platform_label(platform_id: str) -> str:
    return PLATFORMS[platform_id].label if platform_id in PLATFORMS else platform_id


def _format_money(value: float | None) -> str:
    """Reais cheios, formato BR - ao contrario de `format_money_millions` (feito
    para ETFs na casa de dezenas de milhoes), captacao de um fundo individual
    costuma ficar na casa dos milhares: dividir por milhao arredondaria tudo
    para "R$ 0,0M" e esconderia o numero real."""
    if value is None:
        return "-"
    sign = "-" if value < 0 else ""
    return f"R$ {sign}{abs(value):,.0f}".replace(",", "X").replace(".", ",").replace("X", ".")


def print_captacao_summary(resultado: ResultadoCaptacao, janela: str = "semana") -> None:
    """Resumo no terminal: cobertura, nota de metodologia e top/bottom da janela pedida."""
    console = Console()
    console.print()
    console.print("[bold cyan]=== CAPTACAO LIQUIDA DOS FUNDOS DAS VITRINES ===[/]")
    console.print(f"[yellow]{NOTA_FUNDO_NIVEL}[/]")

    cob = resultado.cobertura
    console.print(
        f"[dim]Universo: {resultado.rotulo_universo} | vitrines: "
        f"{', '.join(_platform_label(p) for p in resultado.plataformas) or 'nenhuma'} | "
        f"{cob.total_com_fluxo} de {cob.total_fundos} fundos com fluxo no INF_DIARIO | "
        f"ultima data: {cob.ultima_data.isoformat() if cob.ultima_data else '-'}[/]"
    )
    if cob.meses_faltantes:
        console.print(f"[yellow]Meses sem dado no cache local: {', '.join(cob.meses_faltantes)}[/]")
    if resultado.excluded:
        out = "; ".join(f"{e['label']} ({e['reason']})" for e in resultado.excluded)
        console.print(f"[dim]Fora do universo: {out}[/]")
    console.print()

    if not resultado.plataformas:
        console.print("[yellow]Nenhuma vitrine entrou no universo desta execucao.[/]")
        return

    janela_obj = next((j for j in resultado.janelas if j.chave == janela), None)
    if janela_obj is None:
        console.print(f"[yellow]Janela '{janela}' nao calculada nesta execucao.[/]")
        return
    console.print(f"[bold]Janela: {janela_obj.rotulo}[/]\n")

    def _tabela(titulo: str, fundos: list[dict], style: str) -> Table | None:
        if not fundos:
            return None
        table = Table(title=titulo, title_style=f"bold {style}")
        table.add_column("Fundo", style="cyan")
        table.add_column("Classificacao")
        table.add_column("Nº vitrines", justify="right")
        table.add_column("Vitrines")
        table.add_column("Captacao liquida", justify="right", style=f"bold {style}")
        for f in fundos:
            table.add_row(
                f["nome"][:40], f["classificacao"], str(f["n_vitrines"]),
                ", ".join(_platform_label(p) for p in f["plataformas"]),
                _format_money(f["captacao_liquida"]),
            )
        return table

    consolidado = ranking_consolidado(resultado, janela, limite=_LIMITE_CONSOLE)
    top_table = _tabela("Top 10 captacao liquida (consolidado, todas as vitrines)", consolidado["top"], "green")
    if top_table:
        console.print(top_table)
    bottom_table = _tabela("Bottom 10 (maiores saidas liquidas, consolidado)", consolidado["bottom"], "red")
    if bottom_table:
        console.print(bottom_table)

    por_plataforma = ranking_por_plataforma(resultado, janela, limite=_LIMITE_CONSOLE)
    for platform_id in resultado.plataformas:
        bloco = por_plataforma.get(platform_id)
        if not bloco:
            continue
        label = _platform_label(platform_id)
        top = _tabela(f"{label} - Top 10", bloco["top"], "green")
        if top:
            console.print(top)
        bottom = _tabela(f"{label} - Bottom 10", bloco["bottom"], "red")
        if bottom:
            console.print(bottom)


def _write_headers(sheet, headers: list[tuple[str, int]]) -> None:
    for index, (name, width) in enumerate(headers, start=1):
        sheet.cell(row=1, column=index, value=name)
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL


def _fundo_values(fundo: dict, extra: list[object] | None = None) -> list[object]:
    values = [
        fundo["nome"], fundo["cnpj_display"], fundo["classificacao"], fundo["nivel_classificacao"],
        fundo["n_vitrines"], ", ".join(_platform_label(p) for p in fundo["plataformas"]),
        fundo["captacao_bruta"], fundo["resgates"], fundo["captacao_liquida"],
        fundo["pl_inicio"], fundo["pl_fim"], fundo["captacao_liquida_pct_pl"],
    ]
    return (extra or []) + values


def export_captacao_excel(resultado: ResultadoCaptacao | None, output_path: Path) -> Path:
    """Workbook dedicado do monitor de captacao: uma aba de rankings por
    janela, classificacao, cruzada, cobertura e metodologia.

    Sem resultado (nenhuma vitrine, nenhum fluxo encontrado), o arquivo ainda
    sai - com a aba de rankings suprimida e o motivo, nunca com numero
    inventado.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)

    sheet = workbook.create_sheet("Rankings")
    if resultado is None or not resultado.plataformas or resultado.cobertura.total_com_fluxo == 0:
        sheet.column_dimensions["A"].width = 100
        sheet.cell(row=1, column=1, value="Captacao liquida dos fundos das vitrines").font = _TITLE_FONT
        motivo = (
            "nenhum resultado calculado" if resultado is None
            else "nenhuma vitrine no universo" if not resultado.plataformas
            else "nenhum fundo da vitrine encontrado no INF_DIARIO"
        )
        sheet.cell(row=3, column=1, value=f"Sem dados: {motivo}.").font = _ALERT_FONT
        workbook.create_sheet("Metodologia").cell(row=1, column=1, value=NOTA_FUNDO_NIVEL).font = _NOTE_FONT
        workbook.save(output_path)
        return output_path

    headers = [
        ("Janela", 30), ("Plataforma", 22), ("Tipo", 8), ("#", 5),
        ("Fundo", 46), ("CNPJ", 20), ("Classificacao", 30), ("Nivel", 12),
        ("Nº vitrines", 11), ("Vitrines", 30),
        ("Captacao bruta (R$)", 18), ("Resgates (R$)", 18), ("Captacao liquida (R$)", 20),
        ("PL inicio (R$)", 18), ("PL fim (R$)", 18), ("Liquida / PL (%)", 15),
    ]
    _write_headers(sheet, headers)

    row = 1
    for janela in resultado.janelas:
        consolidado = ranking_consolidado(resultado, janela.chave, limite=25)
        por_plataforma = ranking_por_plataforma(resultado, janela.chave, limite=25)

        def _write(tipo: str, plataforma_label: str, fundos: list[dict]) -> None:
            nonlocal row
            for posicao, fundo in enumerate(fundos, start=1):
                row += 1
                values = [janela.rotulo, plataforma_label, tipo, posicao] + _fundo_values(fundo)
                for index, value in enumerate(values, start=1):
                    cell = sheet.cell(row=row, column=index, value=value)
                    cell.font = _BODY_FONT
                    cell.border = _BORDER
                for col in (11, 12, 13, 14, 15):
                    sheet.cell(row=row, column=col).number_format = _MONEY
                sheet.cell(row=row, column=16).number_format = _PCT

        _write("Top", "Todas as vitrines", consolidado["top"])
        _write("Bottom", "Todas as vitrines", consolidado["bottom"])
        for platform_id, bloco in sorted(por_plataforma.items()):
            label = _platform_label(platform_id)
            _write("Top", label, bloco["top"])
            _write("Bottom", label, bloco["bottom"])

    if row > 1:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row}"
    sheet.freeze_panes = "A2"

    # --- Classificacao -------------------------------------------------
    class_sheet = workbook.create_sheet("Classificacao")
    class_headers = [
        ("Janela", 30), ("Plataforma", 22), ("Classificacao", 30),
        ("Nº fundos", 11), ("Com fluxo", 11),
        ("Captacao bruta (R$)", 18), ("Resgates (R$)", 18), ("Captacao liquida (R$)", 20),
    ]
    _write_headers(class_sheet, class_headers)
    row = 1
    for janela in resultado.janelas:
        for linha in consolidar_por_classificacao(resultado, janela.chave):
            row += 1
            plataforma = "Todas as vitrines" if linha["platform_id"] == "__todas__" else linha["label"]
            values = [
                janela.rotulo, plataforma, linha["classificacao"], linha["qtd_fundos"], linha["qtd_com_fluxo"],
                linha["captacao_bruta"], linha["resgates"], linha["captacao_liquida"],
            ]
            for index, value in enumerate(values, start=1):
                cell = class_sheet.cell(row=row, column=index, value=value)
                cell.font = _BODY_FONT
                cell.border = _BORDER
            for col in (6, 7, 8):
                class_sheet.cell(row=row, column=col).number_format = _MONEY
    if row > 1:
        class_sheet.auto_filter.ref = f"A1:{get_column_letter(len(class_headers))}{row}"
    class_sheet.freeze_panes = "A2"

    # --- Cruzada (secoes empilhadas, semana e ano) ----------------------
    cruz_sheet = workbook.create_sheet("Cruzada")
    cruz_sheet.column_dimensions["A"].width = 46
    for col in "BCDE":
        cruz_sheet.column_dimensions[col].width = 24
    r = 1
    for janela in resultado.janelas:
        if janela.chave not in ("semana", "ano"):
            continue
        cruzada = analise_cruzada(resultado, chave=janela.chave, limite=25)
        cruz_sheet.cell(row=r, column=1, value=f"Janela: {janela.rotulo}").font = _TITLE_FONT
        r += 2

        for titulo, fundos in (
            ("Maiores captadores em 2+ vitrines", cruzada["compartilhados_top"]),
            ("Maiores captadores exclusivos de uma vitrine", cruzada["exclusivos_top"]),
        ):
            cruz_sheet.cell(row=r, column=1, value=titulo).font = _HEADER_FONT
            r += 1
            sub_headers = ["Fundo", "CNPJ", "Classificacao", "Vitrines", "Captacao liquida (R$)"]
            for index, name in enumerate(sub_headers, start=1):
                cruz_sheet.cell(row=r, column=index, value=name).font = _HEADER_FONT
            r += 1
            if not fundos:
                cruz_sheet.cell(row=r, column=1, value="(nenhum)").font = _NOTE_FONT
                r += 1
            for f in fundos:
                values = [f["nome"], f["cnpj_display"], f["classificacao"],
                          ", ".join(_platform_label(p) for p in f["plataformas"]), f["captacao_liquida"]]
                for index, value in enumerate(values, start=1):
                    cruz_sheet.cell(row=r, column=index, value=value).font = _BODY_FONT
                r += 1
            r += 1

        ausentes = cruzada["ausentes_na_baseline"]
        if ausentes.get("available"):
            cruz_sheet.cell(
                row=r, column=1,
                value=f"Maiores captadores fora da vitrine da {ausentes['baseline_label']}",
            ).font = _HEADER_FONT
            r += 1
            for f in ausentes["fundos"]:
                values = [f["nome"], f["cnpj_display"], f["classificacao"],
                          ", ".join(_platform_label(p) for p in f["plataformas"]), f["captacao_liquida"]]
                for index, value in enumerate(values, start=1):
                    cruz_sheet.cell(row=r, column=index, value=value).font = _BODY_FONT
                r += 1
        else:
            cruz_sheet.cell(
                row=r, column=1, value=f"Ausentes da baseline: {ausentes.get('reason', 'nao disponivel')}.",
            ).font = _NOTE_FONT
            r += 1
        r += 2

    # --- Cobertura -------------------------------------------------------
    cov_sheet = workbook.create_sheet("Cobertura")
    cov_headers = [("Plataforma", 24), ("Vitrine total", 14), ("Com CNPJ", 12), ("Com fluxo no informe", 18)]
    _write_headers(cov_sheet, cov_headers)
    cob = resultado.cobertura
    row = 1
    for platform_id in resultado.plataformas:
        info = cob.por_plataforma.get(platform_id, {})
        row += 1
        values = [
            _platform_label(platform_id), info.get("vitrine_total", 0),
            info.get("com_cnpj", 0), info.get("com_fluxo", 0),
        ]
        for index, value in enumerate(values, start=1):
            cov_sheet.cell(row=row, column=index, value=value).font = _BODY_FONT
    cov_sheet.freeze_panes = "A2"
    note_row = row + 3
    cov_sheet.cell(
        row=note_row, column=1,
        value=(
            f"Total: {cob.total_com_fluxo} de {cob.total_fundos} fundos com fluxo. "
            f"Meses lidos: {', '.join(cob.meses_lidos) or 'nenhum'}."
        ),
    ).font = _NOTE_FONT
    if cob.meses_faltantes:
        note_row += 1
        cov_sheet.cell(
            row=note_row, column=1,
            value=f"Meses sem dado no cache local: {', '.join(cob.meses_faltantes)}.",
        ).font = _ALERT_FONT

    # --- Metodologia -------------------------------------------------------
    meta_sheet = workbook.create_sheet("Metodologia")
    meta_sheet.column_dimensions["A"].width = 34
    meta_sheet.column_dimensions["B"].width = 100
    meta_sheet.cell(row=1, column=1, value="Metodologia e proveniencia").font = _TITLE_FONT
    linhas = [
        ("Nivel do dado", NOTA_FUNDO_NIVEL),
        ("Universo", f"{resultado.rotulo_universo} - {', '.join(_platform_label(p) for p in resultado.plataformas)}"),
        ("Janelas", "; ".join(f"{j.rotulo}" + (f" (dado desde {j.dados_desde})" if j.dados_desde else "") for j in resultado.janelas)),
        ("Fonte", "CVM INF_DIARIO (dados.cvm.gov.br), lido por classe/fundo, deduplicado por (CNPJ, data)."),
        ("Mapeamento de CNPJ", "Classe RCVM175 casa direto; envelope expande para as classes via data/market/mercado.csv e soma (agregado_por_fundo); sem base, usa o proprio CNPJ da oferta."),
        ("Classificacao", "ANBIMA oficial (join CVM) > CVM > heuristica de nome/categoria - ver coluna 'Nivel' na aba Rankings."),
    ]
    row = 3
    for key, value in linhas:
        cell = meta_sheet.cell(row=row, column=1, value=key)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        meta_sheet.cell(row=row, column=2, value=value).font = _BODY_FONT
        row += 1
    if resultado.excluded:
        out = "; ".join(f"{e['label']} ({e['reason']})" for e in resultado.excluded)
        meta_sheet.cell(row=row, column=1, value="Fora do universo").font = _HEADER_FONT
        meta_sheet.cell(row=row, column=1).fill = _HEADER_FILL
        meta_sheet.cell(row=row, column=2, value=out).font = _BODY_FONT

    workbook.save(output_path)
    return output_path
