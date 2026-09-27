"""Comparacao da captura com uma exportacao da Comdinheiro, campo a campo.

Entende os dois formatos que a Comdinheiro exporta para ETF (medidos em
2026-09-25 nos arquivos COMDINHEIRO_XP_Fundos*.xlsx): o cadastral (coluna
"Identificador", CNPJ, prestadores, publico-alvo, taxas, DY) e o de mercado
("TICKER", "Data da Analise", precos, volume, retornos por janela,
volatilidade, beta, VaR, DY). O formato e reconhecido pelo cabecalho.

Os numeros da captura sao refeitos na data da analise da Comdinheiro, com a
mesma janela que ela informa em cada linha ("Delta du 12m" = 250...). A
contagem de dias uteis e a da ANBIMA, nao a do pregao: 24/12 e 31/12 nao tem
pregao na B3 mas contam como dia util - sem isso os "250 du" de 12 meses caem
dois pregoes antes e a diferenca media no retorno de 12 meses era 0,22 pp
(com eles, o preco-base implicito do BOVA11 e o fechamento exato).

Do lado da captura, o que vem do regulamento sai da planilha de conferencia
(`manual/regulamento.csv`): o valor conferido por uma pessoa ou, na falta, a
sugestao tirada do texto do PDF - marcada como sugestao.
"""
from __future__ import annotations

import re
import statistics
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from shelf_monitor.config import Config
from shelf_monitor.market.etf_bruto import manual
from shelf_monitor.market.etf_bruto.consolidacao import (
    Insumos,
    beta,
    eventos_de_cota,
    fator_ate,
    fechamento_ate,
    retornos_diarios,
    volatilidade_anual,
)
from shelf_monitor.normalize import normalize_cnpj, normalize_text, parse_number
from shelf_monitor.texto import data_br

IGUAL = "igual"
PROXIMO = "próximo"
DIFERENTE = "diferente"
SO_COMDINHEIRO = "só Comdinheiro"
SO_CAPTURA = "só captura"


def _cabecalho(valor: object) -> str:
    return " ".join(str(valor or "").split())


def _data(valor: object) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto[:10], formato).date()
        except ValueError:
            continue
    return None


@dataclass
class Exportacao:
    """As linhas de cada formato, por ticker."""

    cadastral: dict[str, dict] = field(default_factory=dict)
    mercado: dict[str, dict] = field(default_factory=dict)
    data_analise: date | None = None
    arquivos: list[str] = field(default_factory=list)


def ler_exportacoes(caminhos: list[Path]) -> Exportacao:
    exportacao = Exportacao()
    for caminho in caminhos:
        livro = load_workbook(caminho, read_only=True, data_only=True)
        linhas = list(livro.worksheets[0].iter_rows(values_only=True))
        livro.close()
        if not linhas:
            continue
        cabecalho = [_cabecalho(c) for c in linhas[0]]
        registros = []
        for linha in linhas[1:]:
            registro: dict = {}
            for nome, valor in zip(cabecalho, linha):
                # Coluna repetida: fica a primeira. A Comdinheiro exporta "Prazo cotizacao resgate" duas vezes e
                # "Delta du 03m" duas vezes - a segunda e a de 60 meses (1.254 dias uteis).
                registro.setdefault(nome, valor)
            registros.append(registro)
        if "Identificador" in cabecalho and "CNPJ do Fundo" in cabecalho:
            for r in registros:
                if r.get("Identificador"):
                    exportacao.cadastral[str(r["Identificador"]).strip().upper()] = r
        elif "TICKER" in cabecalho and "Data da Análise" in cabecalho:
            for r in registros:
                if r.get("TICKER"):
                    exportacao.mercado[str(r["TICKER"]).strip().upper()] = r
                    exportacao.data_analise = exportacao.data_analise or _data(r.get("Data da Análise"))
        else:
            raise ValueError(f"{Path(caminho).name}: formato de exportação da Comdinheiro não reconhecido")
        exportacao.arquivos.append(str(caminho))
    return exportacao


# --- calendario ------------------------------------------------------------------------


def calendario_anbima(pregoes: list[date]) -> list[date]:
    """Pregoes da B3 mais os dias uteis sem pregao (24/12 e 31/12 de cada ano no intervalo)."""
    if not pregoes:
        return []
    dias = set(pregoes)
    for ano in range(pregoes[0].year, pregoes[-1].year + 1):
        for dia in (date(ano, 12, 24), date(ano, 12, 31)):
            if pregoes[0] <= dia <= pregoes[-1] and dia.weekday() < 5:
                dias.add(dia)
    return sorted(dias)


# --- comparacao --------------------------------------------------------------------------


@dataclass
class Diferenca:
    ticker: str
    secao: str
    campo: str
    captura: object
    comdinheiro: object
    diferenca: object
    situacao: str
    observacao: str = ""


@dataclass
class Campo:
    secao: str
    nome: str
    unidade: str = ""   # "pp", "R$", "%"...
    porque: str = ""    # por que os dois lados podem divergir


CAMPOS = (
    Campo("Identificação", "CNPJ"),
    Campo("Identificação", "ISIN"),
    Campo("Identificação", "Gestor", porque="Razão social escrita de formas diferentes"),
    Campo("Identificação", "Administrador", porque="Razão social escrita de formas diferentes"),
    Campo("Identificação", "Custodiante", porque="Razão social escrita de formas diferentes"),
    Campo("Identificação", "Auditor", porque="Razão social escrita de formas diferentes"),
    Campo("Identificação", "Constituição",
          porque="A captura usa a data mais antiga entre fundo e classe no cadastro da CVM"),
    Campo("Custo", "Taxa de administração", "%",
          "Captura: valor conferido no regulamento ou sugestão tirada do PDF, não conferida"),
    Campo("Custo", "Taxa máxima", "%", "Idem; em ETF de ETF a Comdinheiro soma a taxa do fundo investido"),
    Campo("Custo", "Taxa de performance", "%", "Idem"),
    Campo("Custo", "Mínimo mensal (R$)", "R$", "Idem; a Comdinheiro guarda em \"Taxa de Administração Financeira\""),
    Campo("Regras", "Público-alvo", porque="Captura: classificação do trecho do regulamento"),
    Campo("Preço e negociação", "Preço de fechamento", "R$"),
    Campo("Preço e negociação", "Volume do dia", "R$"),
    Campo("Preço e negociação", "Negócios do dia", porque="A Comdinheiro conta alguns negócios a menos que o arquivo da B3"),
    Campo("Retorno", "Retorno 1 mês", "pp", "A Comdinheiro ajusta por provento; a captura, só por desdobramento"),
    Campo("Retorno", "Retorno 3 meses", "pp", "Idem"),
    Campo("Retorno", "Retorno 6 meses", "pp", "Idem"),
    Campo("Retorno", "Retorno 12 meses", "pp", "Idem"),
    Campo("Retorno", "Retorno no ano", "pp", "Idem"),
    Campo("Retorno", "Retorno no mês", "pp", "Idem"),
    Campo("Retorno", "Retorno do Ibovespa 12 meses", "pp", "Série do índice na B3"),
    Campo("Risco", "Volatilidade 12 meses", "pp", "Retornos diários; a captura só usa pregões consecutivos"),
    Campo("Risco", "Beta 12 meses", porque="Contra o Ibovespa; a captura só usa pregões consecutivos"),
    Campo("Retorno", "DY 12 meses", "%", "A captura não tem série de proventos"),
    Campo("Risco", "VaR 95%", "%", "A captura não calcula VaR"),
)
_SECAO = {c.nome: c.secao for c in CAMPOS}

# Janelas de retorno do formato de mercado: (campo, coluna do retorno, coluna do delta de dias uteis).
_RETORNOS = (
    ("Retorno 1 mês", "Retorno 1 mês", "Delta du 01m Data Análise"),
    ("Retorno 3 meses", "Retorno 3 meses", "Delta du 03m Data Análise"),
    ("Retorno 6 meses", "Retorno 6 meses", "Delta du 06m Data Análise"),
    ("Retorno 12 meses", "Retorno 12 meses", "Delta du 12m Data Análise"),
    ("Retorno no ano", "Retorno Ano Atual", "Delta du entre Prim DU do ano Data Análise"),
    ("Retorno no mês", "Retorno Mês Atual", "Delta du entre Prim DU do mês Data Análise"),
)


def _situacao_numerica(captura: float | None, comdinheiro: float | None, igual: float, proximo: float,
                       relativa: bool = False) -> tuple[str, float | None]:
    if captura is None and comdinheiro is None:
        return "", None
    if captura is None:
        return SO_COMDINHEIRO, None
    if comdinheiro is None:
        return SO_CAPTURA, None
    diferenca = captura - comdinheiro
    medida = abs(diferenca) / abs(comdinheiro) if relativa and comdinheiro else abs(diferenca)
    if medida <= igual:
        return IGUAL, diferenca
    if medida <= proximo:
        return PROXIMO, diferenca
    return DIFERENTE, diferenca


def _situacao_texto(captura: str | None, comdinheiro: str | None) -> str:
    a, b = normalize_text(captura), normalize_text(comdinheiro)
    if not a and not b:
        return ""
    if not a:
        return SO_COMDINHEIRO
    if not b:
        return SO_CAPTURA
    if a == b:
        return IGUAL
    # "BANCO BNP PARIBAS BRASIL S A" x "BNP PARIBAS": uma contida na outra, ou as tres primeiras palavras iguais.
    if a in b or b in a or a.split()[:3] == b.split()[:3]:
        return PROXIMO
    return DIFERENTE


def _percentual(texto: object) -> float | None:
    """ "0,30%" / "nao cobra" / 0.3 -> 0.3 (em %)."""
    if texto is None:
        return None
    if isinstance(texto, (int, float)):
        return float(texto)
    t = str(texto).strip().lower()
    if not t:
        return None
    if "não cobra" in t or "nao cobra" in t:
        return 0.0
    achado = re.search(r"(\d+(?:[.,]\d+)?)\s*%", t)
    return parse_number(achado.group(1)) if achado else None


def _reais(texto: object) -> float | None:
    achado = re.search(r"R\$\s*([\d.]+(?:,\d+)?)", str(texto or ""))
    if not achado:
        return None
    return parse_number(achado.group(1).replace(".", "").replace(",", "."))


def _publico(texto: object, qualificado: object = None) -> str | None:
    t = normalize_text(str(texto or ""))
    if normalize_text(str(qualificado or "")) == "SIM" or "QUALIFICAD" in t and "NAO" not in t.split("QUALIFICAD")[0][-12:]:
        return "qualificados"
    if "PROFISSION" in t:
        return "profissionais"
    if t:
        return "em geral"
    return None


class _Regulamento:
    """Taxas e publico-alvo do lado da captura: o conferido ou a sugestao do PDF."""

    def __init__(self, config: Config):
        self.por_chave: dict[tuple[str, str], tuple[str, str]] = {}
        for linha in manual.ler(config, "regulamento.csv"):
            chave = ((linha.get("chave") or "").strip(), (linha.get("campo") or "").strip())
            conferido = (linha.get("valor_conferido") or "").strip()
            if conferido and (linha.get("conferido_por") or "").strip():
                self.por_chave[chave] = (conferido, "conferido")
            elif (linha.get("candidato") or "").strip():
                self.por_chave[chave] = (linha["candidato"].strip(), "sugestão do PDF")

    def valor(self, chave: str, campo: str) -> tuple[str | None, str]:
        return self.por_chave.get((chave, campo), (None, ""))


@dataclass
class Comparacao:
    diferencas: list[Diferenca]
    exportacao: Exportacao
    # (ticker, motivo) de quem so um dos lados tem.
    so_comdinheiro: list[tuple[str, str]]
    so_captura: list[tuple[str, str]]
    em_ambos: int


def comparar(config: Config, exportacao: Exportacao, hoje: date | None = None) -> Comparacao:
    hoje = hoje or date.today()
    ins = Insumos(config, hoje)
    regulamento = _Regulamento(config)
    por_ticker = {(l.get("ticker") or "").upper(): l for l in ins.universo if l.get("ticker")}
    por_cnpj = {l["cnpj"]: l for l in ins.universo if l.get("cnpj")}
    calendario = calendario_anbima(ins.calendario)
    analise = exportacao.data_analise
    tickers_cd = set(exportacao.cadastral) | set(exportacao.mercado)
    listados = {t for t, l in por_ticker.items() if l.get("listado") == "sim"}
    diferencas: list[Diferenca] = []

    def anotar(ticker, campo, captura, comdinheiro, situacao, diferenca=None, observacao=""):
        if situacao:
            diferencas.append(Diferenca(ticker, _SECAO[campo], campo, captura, comdinheiro, diferenca, situacao, observacao))

    for ticker in sorted(tickers_cd):
        etf = por_ticker.get(ticker)
        cad = exportacao.cadastral.get(ticker, {})
        mer = exportacao.mercado.get(ticker, {})
        if etf is None:
            continue
        chave, cnpj = etf.get("chave", ""), etf.get("cnpj") or ""
        cvm = ins.cvm.get(cnpj, {})
        dy = parse_number(mer.get("DY_12M")) if mer else parse_number(cad.get("DY(12m,,)"))
        obs_provento = f"distribui proventos (DY 12 meses {dy:.2f}%)".replace(".", ",") if dy else ""

        if cad:
            cnpj_cd = normalize_cnpj(cad.get("CNPJ do Fundo"))
            anotar(ticker, "CNPJ", cnpj or None, cnpj_cd, _situacao_texto(cnpj, cnpj_cd))
            anotar(ticker, "ISIN", etf.get("isin") or None, cad.get("ISIN"), _situacao_texto(etf.get("isin"), cad.get("ISIN")))
            for campo, nosso, coluna in (("Gestor", cvm.get("gestor"), "Razão Social do Gestor"),
                                         ("Administrador", cvm.get("administrador"), "Razão Social do Administrador"),
                                         ("Custodiante", cvm.get("custodiante"), "Razão Social do Custodiante"),
                                         ("Auditor", cvm.get("auditor"), "Razão Social do Auditor")):
                anotar(ticker, campo, nosso or None, cad.get(coluna), _situacao_texto(nosso, cad.get(coluna)))
            constituicao_cd = _data(cad.get("Data da Constituição do Fundo"))
            constituicao = cvm.get("constituicao")
            if constituicao or constituicao_cd:
                situacao = (IGUAL if constituicao == constituicao_cd else DIFERENTE) if constituicao and constituicao_cd else (
                    SO_CAPTURA if constituicao else SO_COMDINHEIRO)
                anotar(ticker, "Constituição", data_br(constituicao) if constituicao else None,
                       data_br(constituicao_cd) if constituicao_cd else None, situacao)
            for campo, campo_reg, coluna in (("Taxa de administração", "taxa_administracao", "Taxa de Administracao"),
                                             ("Taxa máxima", "taxa_maxima", "Taxa de Administracao Máxima"),
                                             ("Taxa de performance", "taxa_performance", "Taxa de Performance")):
                texto, origem = regulamento.valor(chave, campo_reg)
                nosso, deles = _percentual(texto), parse_number(cad.get(coluna))
                situacao, dif = _situacao_numerica(nosso, deles, 0.0005, 0.05)
                anotar(ticker, campo, texto, deles, situacao, dif, origem)
            texto, origem = regulamento.valor(chave, "custos_extras")
            nosso, deles = _reais(texto), parse_number(cad.get("Taxa de Administração Financeira"))
            situacao, dif = _situacao_numerica(nosso, deles, 0.5, 0.01 * (deles or 1), relativa=False)
            anotar(ticker, "Mínimo mensal (R$)", texto if nosso is not None else None, deles, situacao, dif, origem)
            texto, origem = regulamento.valor(chave, "publico_alvo")
            deles = _publico(cad.get("Público Alvo"), cad.get("É para investidores qualificados?"))
            nosso = texto if texto and texto != "ver trecho" else None
            anotar(ticker, "Público-alvo", nosso, deles, _situacao_texto(nosso, deles), observacao=origem)

        if mer and analise:
            bruta = ins.por_ticker.get(etf.get("ticker", "").upper(), {})
            brutos = sorted((d, r[0]) for d, r in bruta.items() if r[0])
            eventos = eventos_de_cota(brutos)
            serie = {d: ((r[0] * fator_ate(eventos, d)) if r[0] else r[0], r[1], r[2]) for d, r in bruta.items()}
            obs_evento = "; ".join(f"{rotulo} em {data_br(d)}" for d, _, rotulo in eventos)
            dia = bruta.get(analise)
            situacao, dif = _situacao_numerica(dia[0] if dia else None, parse_number(mer.get("Preço de Fechamento")), 0.005, 0.05)
            anotar(ticker, "Preço de fechamento", dia[0] if dia else None, parse_number(mer.get("Preço de Fechamento")), situacao, dif)
            volume_cd = parse_number(mer.get("Volume Negociado"))
            volume_cd = volume_cd * 1e9 if volume_cd is not None else None  # a Comdinheiro exporta em R$ bilhoes
            situacao, dif = _situacao_numerica(dia[2] if dia else None, volume_cd, 0.001, 0.01, relativa=True)
            anotar(ticker, "Volume do dia", dia[2] if dia else None, volume_cd, situacao, dif)
            situacao, dif = _situacao_numerica(dia[1] if dia else None, parse_number(mer.get("NEGOCIOS_DIA")), 0, 0.01, relativa=True)
            anotar(ticker, "Negócios do dia", dia[1] if dia else None, parse_number(mer.get("NEGOCIOS_DIA")), situacao, dif)

            fim = fechamento_ate(serie, analise, ins.pregoes)
            if analise in calendario:
                indice_analise = calendario.index(analise)
                for campo, coluna, coluna_delta in _RETORNOS:
                    deles = parse_number(mer.get(coluna))
                    delta = parse_number(mer.get(coluna_delta))
                    nosso = None
                    if fim and delta is not None and indice_analise - int(delta) >= 0:
                        base = fechamento_ate(serie, calendario[indice_analise - int(delta)], ins.pregoes)
                        if base:
                            nosso = (fim[1] / base[1] - 1) * 100
                    situacao, dif = _situacao_numerica(nosso, deles, 0.05, 0.5)
                    observacao = "; ".join(x for x in (obs_provento if situacao == DIFERENTE else "", obs_evento) if x)
                    anotar(ticker, campo, nosso, deles, situacao, dif, observacao)

                delta12 = parse_number(mer.get("Delta du 12m Data Análise"))
                if delta12 is not None and indice_analise - int(delta12) >= 0:
                    inicio = calendario[indice_analise - int(delta12)]
                    niveis = ins.indices.get("IBOV", {})
                    base_ibov = max((d for d in niveis if d <= inicio), default=None)
                    fim_ibov = max((d for d in niveis if d <= analise), default=None)
                    nosso_ibov = (niveis[fim_ibov] / niveis[base_ibov] - 1) * 100 if base_ibov and fim_ibov else None
                    situacao, dif = _situacao_numerica(nosso_ibov, parse_number(mer.get("Retorno IBOV 12 meses")), 0.05, 0.5)
                    anotar(ticker, "Retorno do Ibovespa 12 meses", nosso_ibov, parse_number(mer.get("Retorno IBOV 12 meses")),
                           situacao, dif)
                    janela = sorted((d, p) for d, p in ((d, r[0]) for d, r in serie.items() if r[0]) if inicio <= d <= analise)
                    retornos = retornos_diarios(janela, ins.calendario)
                    vol = volatilidade_anual([r for _, r in retornos]) * 100 if len(retornos) >= 200 else None
                    situacao, dif = _situacao_numerica(vol, parse_number(mer.get("Volatilidade Anualizada 12 meses")), 0.1, 1.0)
                    anotar(ticker, "Volatilidade 12 meses", vol, parse_number(mer.get("Volatilidade Anualizada 12 meses")),
                           situacao, dif, obs_evento)
                    niveis_janela = sorted((d, v) for d, v in niveis.items() if inicio <= d <= analise)
                    b = beta(dict(retornos), dict(retornos_diarios(niveis_janela)))[0] if len(retornos) >= 200 else None
                    situacao, dif = _situacao_numerica(b, parse_number(mer.get("Beta 12 meses")), 0.02, 0.05)
                    anotar(ticker, "Beta 12 meses", b, parse_number(mer.get("Beta 12 meses")), situacao, dif, obs_evento)
            anotar(ticker, "DY 12 meses", None, dy, SO_COMDINHEIRO if dy is not None else "")
            var = parse_number(mer.get("VAR95"))
            anotar(ticker, "VaR 95%", None, var, SO_COMDINHEIRO if var is not None else "")

    # Quem so a Comdinheiro tem: a chave e o CNPJ. Em 2026-09-25, 10 dos 12
    # tickers sem par eram o ticker antigo de ETF renomeado (5GTK11 -> CHIP11,
    # SAET11 -> BOVS11...), que a Comdinheiro mantem em outra linha, sem preco.
    so_comdinheiro = []
    for ticker in sorted(t for t in tickers_cd if t not in por_ticker):
        cnpj = normalize_cnpj((exportacao.cadastral.get(ticker) or {}).get("CNPJ do Fundo"))
        mesmo = por_cnpj.get(cnpj) if cnpj else None
        if mesmo and mesmo.get("ticker"):
            motivo = f"ticker antigo: o mesmo CNPJ negocia hoje como {mesmo['ticker']}"
        elif mesmo:
            motivo = "sem ticker na B3 (pré-operacional na CVM)"
        else:
            motivo = "fora do universo da captura"
        so_comdinheiro.append((ticker, motivo))
    so_captura = [(t, "listado na B3; fora da lista exportada da Comdinheiro")
                  for t in sorted(t for t in listados if t not in tickers_cd)]
    em_ambos = len([t for t in tickers_cd if t in por_ticker])
    return Comparacao(diferencas, exportacao, so_comdinheiro, so_captura, em_ambos)


def resumo(comparacao: Comparacao) -> list[dict]:
    """Uma linha por campo: quantos batem, quantos divergem e a diferenca tipica."""
    linhas = []
    for campo in CAMPOS:
        itens = [d for d in comparacao.diferencas if d.campo == campo.nome]
        contagem = {s: sum(1 for d in itens if d.situacao == s) for s in (IGUAL, PROXIMO, DIFERENTE, SO_COMDINHEIRO, SO_CAPTURA)}
        numericas = [abs(d.diferenca) for d in itens if isinstance(d.diferenca, (int, float))]
        linhas.append({
            "secao": campo.secao, "campo": campo.nome, "comparaveis": contagem[IGUAL] + contagem[PROXIMO] + contagem[DIFERENTE],
            **contagem, "diferenca_mediana": statistics.median(numericas) if numericas else None,
            "unidade": campo.unidade, "porque": campo.porque,
        })
    return linhas
