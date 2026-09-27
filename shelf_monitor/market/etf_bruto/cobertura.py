"""Cobertura: para cada indicador do mapeamento, quantos ETFs ja tem o insumo em disco.

Nao calcula indicador nenhum - mede se o dado bruto que cada via do
mapeamento pede esta capturado, ETF a ETF. O denominador e o universo de ETFs
listados na B3. Cada medidor devolve o conjunto de chaves do universo cobertas;
os insumos de uma via se cruzam (E) e, entre as alternativas, sai a melhor.

Roda sem rede (`etf-bruto --cobertura`): le so o que esta em data/etf/.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable

from rich.console import Console
from rich.table import Table

from shelf_monitor.config import Config
from shelf_monitor.market.etf_bruto import armazem, manual
from shelf_monitor.market.etf_bruto.b3_arquivos import meses_antes, pasta_negocios
from shelf_monitor.market.etf_bruto.composicao import pasta_cda, pasta_gestoras
from shelf_monitor.market.etf_bruto.comdinheiro import pasta as pasta_comdinheiro
from shelf_monitor.market.etf_bruto.fnet import ler_indice, pasta_informe, pasta_regulamentos
from shelf_monitor.market.etf_bruto.indices_b3 import pasta as pasta_indices
from shelf_monitor.market.etf_bruto.mapa import INDICADORES, ORIGENS, SECOES, Estado, Indicador, Via
from shelf_monitor.market.etf_bruto.universo import (
    arquivos_datados,
    listados,
    pasta_cvm_cadastro,
    pasta_instrumentos,
    ultimo_datado,
)
from shelf_monitor.normalize import normalize_cnpj
from shelf_monitor.texto import data_br, periodo_br

PREGOES_LIQUIDEZ = 63
# Fracao dos pregoes de uma janela que pode faltar em disco (minimo de 2).
PREGOES_TOLERADOS = 0.03
DIAS_RECENTES = 7
# Informes diarios em 12 meses para a serie de cota contar como completa
# (~250 pregoes; a folga cobre informe que faltou num dia).
INFORMES_12M = 240
MESES_12M = 13

COLUNAS_CSV = (
    "secao", "indicador", "cliente", "estado", "recomendada", "cobertos_recomendada", "alternativa",
    "cobertos_alternativa", "universo", "janela", "observacao", "links", "medido_em",
)


@dataclass
class Medida:
    cobertos: set[str] = field(default_factory=set)
    janela: str | None = None
    detalhe: str = ""


class Disco:
    """O que esta em data/etf/, lido uma vez e reaproveitado pelos medidores."""

    def __init__(self, config: Config, universo: list[dict], hoje: date):
        self.config = config
        self.hoje = hoje
        self.universo = universo
        self.listados = {linha["chave"] for linha in listados(universo)}
        self.por_ticker: dict[str, str] = {}
        self.por_cnpj: dict[str, str] = {}
        for linha in universo:
            if linha.get("ticker"):
                self.por_ticker[linha["ticker"].strip().upper()] = linha["chave"]
            if linha.get("cnpj"):
                self.por_cnpj[linha["cnpj"]] = linha["chave"]
        self._negocios: list[tuple[date, list[dict]]] | None = None
        self._sem_pregao: set[str] | None = None
        self._cadastros: list[tuple[date, set[str]]] | None = None

    def chaves_de_tickers(self, tickers) -> set[str]:
        return {self.por_ticker[t] for t in (x.strip().upper() for x in tickers) if t in self.por_ticker}

    def chaves_de_cnpjs(self, cnpjs) -> set[str]:
        return {self.por_cnpj[c] for c in (normalize_cnpj(x) for x in cnpjs) if c and c in self.por_cnpj}

    # --- B3 negocios ---

    @property
    def negocios(self) -> list[tuple[date, list[dict]]]:
        if self._negocios is None:
            self._negocios = [
                (dia, armazem.ler_csv(caminho))
                for dia, caminho in arquivos_datados(pasta_negocios(self.config), ".csv")
            ]
        return self._negocios

    @property
    def sem_pregao(self) -> set[str]:
        if self._sem_pregao is None:
            caminho = self.config.etf_bruto_dir / "b3_arquivos" / "_sem_pregao.json"
            self._sem_pregao = set(armazem.ler_json(caminho, []) or [])
        return self._sem_pregao

    @property
    def cadastros(self) -> list[tuple[date, set[str]]]:
        """(dia, chaves listadas) de cada cadastro de instrumentos em disco, do mais antigo ao mais novo.

        A idade de um ETF vem daqui, e nao do primeiro negocio (ETF pouco
        negociado passa dias sem negocio) nem do TradgStartDt (a B3 o reinicia:
        o BOVA11, de 2008, estava com 16/07/2026 em 23/09/2026).
        """
        if self._cadastros is None:
            self._cadastros = [
                (dia, self.chaves_de_tickers(
                    l.get("x_ticker_base") or l.get("TckrSymb") or ""
                    for l in armazem.ler_csv(caminho) if l.get("x_tipo_linha") == "padrao"))
                for dia, caminho in arquivos_datados(pasta_instrumentos(self.config), ".csv")
            ]
        return self._cadastros

    def listados_em(self, dia: date) -> set[str] | None:
        """Chaves listadas no cadastro mais recente ate `dia`; None sem cadastro tao antigo."""
        anteriores = [chaves for data_cadastro, chaves in self.cadastros if data_cadastro <= dia]
        return anteriores[-1] if anteriores else None

    def estreantes(self) -> dict[str, date]:
        """Chave -> cadastro em que o ETF aparece pela primeira vez, para os que
        nao estavam no cadastro mais antigo em disco (os que estrearam depois)."""
        if len(self.cadastros) < 2:
            return {}
        primeiro = self.cadastros[0][1]
        estreia: dict[str, date] = {}
        for dia, chaves in self.cadastros[1:]:
            for chave in chaves - primeiro:
                estreia.setdefault(chave, dia)
        return estreia

    def pregoes_ate(self, fim: date, quantos: int | None = None, desde: date | None = None) -> list[date]:
        """Dias uteis ate `fim` que nao estao marcados como sem pregao, do mais recente para tras."""
        dias = []
        dia = fim
        limite = desde or (fim - timedelta(days=int((quantos or 0) * 1.6) + 15))
        while dia >= limite and (quantos is None or len(dias) < quantos):
            if dia.weekday() < 5 and dia.isoformat() not in self.sem_pregao:
                dias.append(dia)
            dia -= timedelta(days=1)
        return dias


# --- medidores ---------------------------------------------------------------------


def _b3_instrumentos(d: Disco) -> Medida:
    ultimo = ultimo_datado(pasta_instrumentos(d.config), ".csv")
    if ultimo is None:
        return Medida(detalhe="nenhum cadastro de instrumentos em disco")
    tickers = {(l.get("x_ticker_base") or l.get("TckrSymb") or "") for l in armazem.ler_csv(ultimo[1])
               if l.get("x_tipo_linha") == "padrao"}
    return Medida(d.chaves_de_tickers(tickers), janela=f"cadastro de {data_br(ultimo[0])}")


def _b3_fundos(d: Disco) -> Medida:
    ultimo = ultimo_datado(d.config.etf_bruto_dir / "b3_fundos", ".json")
    cobertos = {l["chave"] for l in d.universo if l.get("tipo_b3")}
    return Medida(cobertos, janela=f"foto de {data_br(ultimo[0])}" if ultimo else None,
                  detalhe="" if ultimo else "nenhuma foto de fundos listados em disco")


def _cvm_cadastro(d: Disco) -> Medida:
    pasta = pasta_cvm_cadastro(d.config)
    dias = sorted(p for p in pasta.iterdir() if p.is_dir()) if pasta.exists() else []
    if not dias:
        return Medida(detalhe="nenhum extrato do cadastro da CVM em disco")
    cnpjs = set()
    for nome, coluna in (("registro_classe_etf.csv", "CNPJ_Classe"), ("registro_fundo_etf.csv", "CNPJ_Fundo")):
        cnpjs.update(l.get(coluna) for l in armazem.ler_csv(dias[-1] / nome))
    return Medida(d.chaves_de_cnpjs(cnpjs), janela=f"extrato de {data_br(dias[-1].name)}")


def _b3_preco(d: Disco) -> Medida:
    recentes = d.negocios[-5:]
    if not recentes:
        return Medida(detalhe="nenhum pregão de negócios em disco")
    tickers = {
        l.get("x_ticker_base") or l.get("TckrSymb") or ""
        for _, linhas in recentes for l in linhas
        if l.get("x_tipo_linha") == "padrao" and (l.get("LastPric") or "").strip()
    }
    return Medida(d.chaves_de_tickers(tickers), janela=periodo_br(recentes[0][0], recentes[-1][0]),
                  detalhe="fechamento em algum dos últimos 5 pregões em disco")


def _janela_completa(d: Disco, pregoes: list[date]) -> list[date]:
    """Os pregoes da lista que faltam em disco."""
    em_disco = {dia for dia, _ in d.negocios}
    return [dia for dia in pregoes if dia not in em_disco]


def _listados_no_ultimo_cadastro(d: Disco) -> set[str]:
    return d.cadastros[-1][1] if d.cadastros else set()


def _janela_de_pregoes(d: Disco, pregoes: list[date], rotulo: str) -> Medida:
    """Cobertura de uma janela: em disco (com as pontas e ate PREGOES_TOLERADOS de
    folga), e o ETF ja listado no comeco dela.

    A folga existe porque a B3 publica, as vezes, o arquivo do pregao sem o
    mercado a vista (6 em 266 pregoes ate 24/09/2026): exigir todos zeraria a
    janela de 12 meses por um dia sem dado. Sem cadastro de instrumentos tao
    antigo quanto o comeco da janela, a idade nao da para conferir e vale o
    ultimo cadastro (o detalhe avisa).
    """
    fim = pregoes[0]
    inicio = pregoes[-1]
    faltam = _janela_completa(d, pregoes)
    janela = f"{rotulo}, {periodo_br(inicio, fim)}"
    folga = max(2, int(len(pregoes) * PREGOES_TOLERADOS))
    if faltam and (len(faltam) > folga or fim in faltam or inicio in faltam):
        return Medida(janela=janela, detalhe=f"janela incompleta: faltam {len(faltam)} pregões em disco")
    aviso = (f"faltam {len(faltam)} pregões em disco ({', '.join(data_br(x) for x in sorted(faltam))})"
             if faltam else "")
    atuais = _listados_no_ultimo_cadastro(d)
    no_inicio = d.listados_em(inicio)
    if no_inicio is None:
        detalhes = ["sem cadastro de instrumentos do começo da janela: idade não conferida", aviso]
        return Medida(atuais, janela=janela, detalhe="; ".join(x for x in detalhes if x))
    cobertos = atuais & no_inicio
    novos = len(atuais - no_inicio)
    detalhes = [f"{novos} listados depois do começo da janela" if novos else "", aviso]
    return Medida(cobertos, janela=janela, detalhe="; ".join(x for x in detalhes if x))


def _b3_janela_63(d: Disco) -> Medida:
    if not d.negocios:
        return Medida(detalhe="nenhum pregão de negócios em disco")
    pregoes = d.pregoes_ate(d.negocios[-1][0], quantos=PREGOES_LIQUIDEZ)
    if len(pregoes) < PREGOES_LIQUIDEZ:
        return Medida(detalhe=f"só {len(pregoes)} pregões no calendário em disco")
    return _janela_de_pregoes(d, pregoes, "63 pregões")


def _b3_janela_12m(d: Disco) -> Medida:
    if not d.negocios:
        return Medida(detalhe="nenhum pregão de negócios em disco")
    fim = d.negocios[-1][0]
    pregoes = d.pregoes_ate(fim, desde=meses_antes(fim, 12))
    return _janela_de_pregoes(d, pregoes, "12 meses")


def _b3_historico(d: Disco) -> Medida:
    """Serie de 12 meses, ou desde a estreia para quem estreou depois do comeco
    do trecho continuo de pregoes em disco."""
    doze = _b3_janela_12m(d)
    if not d.negocios:
        return doze
    fim = d.negocios[-1][0]
    em_disco = {dia for dia, _ in d.negocios}
    inicio_continuo = fim
    for dia in d.pregoes_ate(fim, desde=d.negocios[0][0]):
        if dia not in em_disco:
            break
        inicio_continuo = dia
    estreantes = {chave for chave, dia in d.estreantes().items() if dia >= inicio_continuo}
    return Medida(doze.cobertos | (estreantes & _listados_no_ultimo_cadastro(d)),
                  janela=periodo_br(inicio_continuo, fim), detalhe=doze.detalhe)


def _informes(d: Disco) -> dict[str, set[date]]:
    """CNPJ -> datas com informe (PL ou cota), do bruto e da coleta do `etf`."""
    datas: dict[str, set[date]] = {}

    def somar(cnpj, dia_texto, *valores):
        cnpj = normalize_cnpj(cnpj)
        if not cnpj or not any((v or "").strip() for v in valores):
            return
        for formato in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                dia = datetime.strptime((dia_texto or "")[:10], formato).date()
                break
            except ValueError:
                continue
        else:
            return
        datas.setdefault(cnpj, set()).add(dia)

    for _, caminho in arquivos_datados(pasta_informe(d.config), ".csv"):
        for l in armazem.ler_csv(caminho):
            somar(l.get("CNPJ_FDO") or l.get("cnpj_arquivo"), l.get("DT_COMPT") or l.get("data_referencia"),
                  l.get("PATRIM_LIQ"), l.get("VL_QUOTA"))
    for _, caminho in arquivos_datados(d.config.etf_fnet_dir, ".csv"):
        for l in armazem.ler_csv(caminho):
            somar(l.get("cnpj"), l.get("data"), l.get("pl"), l.get("cota"))
    return datas


def _fnet_informe(d: Disco) -> Medida:
    datas = _informes(d)
    limite = d.hoje - timedelta(days=DIAS_RECENTES)
    recentes = {cnpj for cnpj, dias in datas.items() if max(dias) >= limite}
    return Medida(d.chaves_de_cnpjs(recentes), janela=periodo_br(limite, d.hoje),
                  detalhe="informe com PL ou cota nos últimos 7 dias")


def _fnet_informe_12m(d: Disco) -> Medida:
    datas = _informes(d)
    inicio = meses_antes(d.hoje, 12)
    completos = {cnpj for cnpj, dias in datas.items() if sum(1 for x in dias if x >= inicio) >= INFORMES_12M}
    return Medida(d.chaves_de_cnpjs(completos), janela=periodo_br(inicio, d.hoje),
                  detalhe=f"{INFORMES_12M} informes diários ou mais em 12 meses")


def _fnet_mensal_12m(d: Disco) -> Medida:
    pasta = d.config.etf_fnet_mensal_dir
    meses = sorted(p.stem for p in pasta.glob("*.csv") if p.stem.isdigit()) if pasta.exists() else []
    ultimos = meses[-MESES_12M:]
    if len(ultimos) < MESES_12M:
        return Medida(detalhe=f"{len(ultimos)} fechamentos mensais em disco (precisa {MESES_12M})")
    presenca: dict[str, list[bool]] = {}
    for i, mes in enumerate(ultimos):
        cnpjs = {normalize_cnpj(l.get("cnpj")) for l in armazem.ler_csv(pasta / f"{mes}.csv") if l.get("pl") or l.get("cota")}
        for cnpj in cnpjs - {None}:
            presenca.setdefault(cnpj, [False] * len(ultimos))[i] = True
    # Completo: presente em todos os meses desde a primeira aparicao (ETF novo conta).
    completos = {cnpj for cnpj, marcas in presenca.items() if all(marcas[marcas.index(True):])}
    return Medida(d.chaves_de_cnpjs(completos), janela=f"{ultimos[0]} a {ultimos[-1]}",
                  detalhe="fechamento em todos os meses desde a estreia")


def _fnet_regulamento(d: Disco) -> Medida:
    pasta = pasta_regulamentos(d.config)
    cnpjs = {p.parent.name for p in pasta.glob("*/*.pdf")} if pasta.exists() else set()
    return Medida(d.chaves_de_cnpjs(cnpjs))


def _fnet_demonstracoes(d: Disco) -> Medida:
    cnpjs = {l.get("cnpj_catalogo") for l in ler_indice(d.config)
             if l.get("categoria") != "Regulamento" and l.get("cnpj_catalogo")}
    return Medida(d.chaves_de_cnpjs(cnpjs), detalhe="link no índice de documentos do FNET")


def _cvm_cda(d: Disco) -> Medida:
    arquivos = sorted(pasta_cda(d.config).glob("cda_fie_etf_*.csv")) if pasta_cda(d.config).exists() else []
    if not arquivos:
        return Medida(detalhe="nenhuma carteira mensal da CVM em disco")
    ultimo = arquivos[-1]
    cnpjs = {l.get("CNPJ_FUNDO_CLASSE") or l.get("CNPJ_FUNDO") for l in armazem.ler_csv(ultimo)}
    return Medida(d.chaves_de_cnpjs(cnpjs), janela=f"competência {ultimo.stem[-6:]}")


def _indice(d: Disco, carteira: bool) -> Medida:
    codigos = manual.codigos_b3(d.config, d.universo)
    pasta = pasta_indices(d.config) / ("carteira" if carteira else "serie")
    limite = d.hoje - timedelta(days=DIAS_RECENTES)
    com_dado = set()
    for codigo in set(codigos.values()):
        alvo = pasta / codigo
        if not alvo.exists():
            continue
        if carteira:
            if any(dia >= limite for dia, _ in arquivos_datados(alvo, ".json")):
                com_dado.add(codigo)
        elif (alvo / f"{d.hoje.year}.json").exists():
            com_dado.add(codigo)
    cobertos = {chave for chave, codigo in codigos.items() if codigo in com_dado}
    return Medida(cobertos, detalhe=f"{len(codigos)} ETFs com índice B3 apontado ou sugerido")


def _gestora_arquivo(d: Disco) -> Medida:
    pasta = pasta_gestoras(d.config)
    limite = d.hoje - timedelta(days=DIAS_RECENTES)
    tickers = set()
    if pasta.exists():
        for alvo in pasta.iterdir():
            if not alvo.is_dir() or alvo.name.startswith("_"):
                continue
            if any(dia >= limite for dia, _ in _datados_qualquer(alvo)):
                tickers.add(alvo.name)
        entrada = pasta / "_entrada"
        if entrada.exists():
            for pasta_dia in entrada.iterdir():
                try:
                    dia = date.fromisoformat(pasta_dia.name)
                except ValueError:
                    continue
                if dia >= limite:
                    tickers.update(re.split(r"[_\-. ]", p.name, maxsplit=1)[0].upper()
                                   for p in pasta_dia.rglob("*") if p.is_file())
    return Medida(d.chaves_de_tickers(tickers), detalhe="arquivo da gestora nos últimos 7 dias")


def _datados_qualquer(pasta: Path) -> list[tuple[date, Path]]:
    achados = []
    for caminho in pasta.iterdir():
        try:
            achados.append((date.fromisoformat(caminho.name[:10]), caminho))
        except ValueError:
            continue
    return achados


def _texto_de(caminho: Path) -> str:
    if caminho.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook

            livro = load_workbook(caminho, read_only=True, data_only=True)
            partes = [str(celula) for aba in livro.worksheets for linha in aba.iter_rows(values_only=True)
                      for celula in linha if celula is not None]
            livro.close()
            return " ".join(partes)
        except Exception:  # noqa: BLE001 - planilha ilegivel conta como vazia
            return ""
    try:
        return caminho.read_bytes().decode("utf-8", errors="ignore")
    except OSError:
        return ""


def _comdinheiro(d: Disco) -> Medida:
    base = pasta_comdinheiro(d.config)
    limite = d.hoje - timedelta(days=DIAS_RECENTES)
    textos = []
    for sub in ("api", "entrada"):
        pasta = base / sub
        if not pasta.exists():
            continue
        for pasta_dia in pasta.iterdir():
            try:
                dia = date.fromisoformat(pasta_dia.name)
            except ValueError:
                continue
            if dia >= limite:
                textos.extend(_texto_de(p) for p in pasta_dia.rglob("*")
                              if p.is_file() and not p.name.endswith(".erro.txt"))
    if not textos:
        return Medida(detalhe="nada da Comdinheiro nos últimos 7 dias")
    achados = set(re.findall(r"\b[A-Z0-9]{4}\d{2}\b", " ".join(textos).upper()))
    return Medida(d.chaves_de_tickers(achados), detalhe="ticker presente nas respostas ou exportações recentes")


def _tributacao(d: Disco) -> Medida:
    """A regra pede a categoria B3 de todo ETF e, so nos de renda fixa, o PMRC conferido."""
    categoria = _b3_instrumentos(d)
    renda_fixa = {l["chave"] for l in d.universo if (l.get("categoria_b3") or "").startswith("FIXED INCOME")}
    pmrc = _regulamento_campo("pmrc")(d).cobertos
    cobertos = (categoria.cobertos - renda_fixa) | (categoria.cobertos & renda_fixa & pmrc)
    faltam = len(categoria.cobertos & renda_fixa - pmrc)
    return Medida(cobertos, janela=categoria.janela,
                  detalhe=f"{faltam} ETFs de renda fixa sem PMRC conferido no regulamento" if faltam else "")


def _nome(d: Disco) -> Medida:
    return Medida({l["chave"] for l in d.universo if l.get("nome")}, detalhe="triagem pelo nome")


def _regulamento_campo(campo: str) -> Callable[[Disco], Medida]:
    def medir(d: Disco) -> Medida:
        linhas = [l for l in manual.ler(d.config, "regulamento.csv") if l.get("campo") == campo]
        conferidos = {l.get("chave") for l in linhas
                      if (l.get("valor_conferido") or "").strip() and (l.get("conferido_por") or "").strip()}
        com_candidato = sum(1 for l in linhas if (l.get("candidato") or "").strip()
                            and l.get("chave") in d.listados)
        return Medida(set(conferidos) - {None}, detalhe=f"conferido por uma pessoa; {com_candidato} com candidato a conferir")
    return medir


def _manual(arquivo: str, coluna: str) -> Callable[[Disco], Medida]:
    def medir(d: Disco) -> Medida:
        return Medida(manual.preenchidos(d.config, arquivo, coluna), detalhe=f"preenchido em manual/{arquivo}")
    return medir


MEDIDORES: dict[str, Callable[[Disco], Medida]] = {
    "b3_instrumentos": _b3_instrumentos,
    "b3_fundos": _b3_fundos,
    "cvm_cadastro": _cvm_cadastro,
    "b3_preco": _b3_preco,
    "b3_janela_63": _b3_janela_63,
    "b3_janela_12m": _b3_janela_12m,
    "b3_historico": _b3_historico,
    "fnet_informe": _fnet_informe,
    "fnet_informe_12m": _fnet_informe_12m,
    "fnet_mensal_12m": _fnet_mensal_12m,
    "fnet_regulamento": _fnet_regulamento,
    "fnet_demonstracoes": _fnet_demonstracoes,
    "cvm_cda": _cvm_cda,
    "indice_serie": lambda d: _indice(d, carteira=False),
    "indice_carteira": lambda d: _indice(d, carteira=True),
    "gestora_arquivo": _gestora_arquivo,
    "comdinheiro": _comdinheiro,
    "tributacao": _tributacao,
    "nome": _nome,
}


def medidor(chave: str) -> Callable[[Disco], Medida]:
    if chave in MEDIDORES:
        return MEDIDORES[chave]
    if chave.startswith("regulamento:"):
        return _regulamento_campo(chave.split(":", 1)[1])
    if chave.startswith("manual:"):
        _, arquivo, coluna = chave.split(":", 2)
        return _manual(f"{arquivo}.csv", coluna)
    raise KeyError(chave)


# --- cobertura por indicador -----------------------------------------------------------


@dataclass
class LinhaCobertura:
    indicador: Indicador
    universo: int
    recomendada: int | None
    alternativa: Via | None
    cobertos_alternativa: int | None
    janela: str
    observacao: str
    links: str


def _medir_via(via: Via, medidas: dict[str, Medida], listados_: set[str]) -> tuple[int | None, list[str], list[str]]:
    """(cobertos, janelas, detalhes) de uma via; None quando ela nao e coletada."""
    if not via.coletada or not via.insumos:
        return None, [], []
    cobertos = set(listados_)
    janelas, detalhes = [], []
    for insumo in via.insumos:
        medida = medidas[insumo]
        cobertos &= medida.cobertos
        if medida.janela:
            janelas.append(medida.janela)
        if medida.detalhe:
            detalhes.append(medida.detalhe)
    return len(cobertos), janelas, detalhes


def medir(config: Config, universo: list[dict], hoje: date | None = None) -> list[LinhaCobertura]:
    hoje = hoje or date.today()
    disco = Disco(config, universo, hoje)
    chaves = set()
    for item in INDICADORES:
        for via in ((item.recomendada,) if item.recomendada else ()) + item.alternativas:
            if via.coletada:
                chaves.update(via.insumos)
    medidas = {chave: medidor(chave)(disco) for chave in sorted(chaves)}
    total = len(disco.listados)
    linhas = []
    for item in INDICADORES:
        rec_n, rec_janelas, rec_detalhes = (None, [], [])
        if item.recomendada is not None:
            rec_n, rec_janelas, rec_detalhes = _medir_via(item.recomendada, medidas, disco.listados)
        melhor, melhor_n, alt_janelas = None, None, []
        for via in item.alternativas:
            n, janelas, _ = _medir_via(via, medidas, disco.listados)
            if n is not None and (melhor_n is None or n > melhor_n):
                melhor, melhor_n, alt_janelas = via, n, janelas
        if melhor is None and item.alternativas:
            melhor = item.alternativas[0]
        observacoes = []
        if item.estado == Estado.REGRA and item.regra:
            observacoes.append(f"regra: {item.regra}")
        if item.recomendada is not None and not item.recomendada.coletada:
            observacoes.append(f"não coletado: {item.recomendada.descricao}")
        observacoes.extend(rec_detalhes)
        if item.captura:
            observacoes.append(item.captura)
        if item.nota:
            observacoes.append(f"nota do mapeamento: {item.nota}")
        vias = ((item.recomendada,) if item.recomendada else ()) + item.alternativas
        links = []
        for via in vias:
            origem = ORIGENS.get(via.origem)
            if origem and origem.fonte and origem.fonte.url not in links:
                links.append(origem.fonte.url)
        linhas.append(LinhaCobertura(
            item, total, rec_n, melhor, melhor_n,
            "; ".join(dict.fromkeys(rec_janelas or alt_janelas)),
            " | ".join(observacoes), " ".join(links),
        ))
    return linhas


def _n(valor: int | None, total: int) -> str:
    return "—" if valor is None else f"{valor}/{total}"


def _estado(item: Indicador) -> str:
    return item.estado.value + (f", {item.qualificador}" if item.qualificador else "")


def gravar(config: Config, linhas: list[LinhaCobertura], hoje: date) -> Path:
    registros = []
    for linha in linhas:
        item = linha.indicador
        registros.append({
            "secao": item.secao, "indicador": item.nome, "cliente": item.cliente, "estado": _estado(item),
            "recomendada": item.recomendada.descricao if item.recomendada else "—",
            "cobertos_recomendada": "" if linha.recomendada is None else linha.recomendada,
            "alternativa": linha.alternativa.descricao if linha.alternativa else "",
            "cobertos_alternativa": "" if linha.cobertos_alternativa is None else linha.cobertos_alternativa,
            "universo": linha.universo, "janela": linha.janela, "observacao": linha.observacao,
            "links": linha.links, "medido_em": hoje.isoformat(),
        })
    armazem.gravar_csv(config.etf_cobertura_path, COLUNAS_CSV, registros)
    return config.etf_cobertura_path


def _cor(valor: int | None, total: int) -> str:
    if valor is None or not total:
        return "dim"
    if valor >= 0.9 * total:
        return "green"
    return "yellow" if valor else "red"


def imprimir(linhas: list[LinhaCobertura], console: Console | None = None) -> None:
    console = console or Console()
    if not linhas:
        return
    total = linhas[0].universo
    for secao in SECOES:
        tabela = Table(title=f"{secao} (universo: {total} ETFs listados)", title_justify="left", show_lines=False)
        tabela.add_column("Indicador")
        tabela.add_column("Estado")
        tabela.add_column("Recomendada", justify="right")
        tabela.add_column("Melhor alternativa", justify="right")
        tabela.add_column("Janela")
        for linha in (l for l in linhas if l.indicador.secao == secao):
            item = linha.indicador
            rec = _n(linha.recomendada, total)
            alt = _n(linha.cobertos_alternativa, total)
            alt_texto = f"[{_cor(linha.cobertos_alternativa, total)}]{alt}[/]"
            if linha.alternativa is not None:
                alt_texto += f" [dim]{linha.alternativa.descricao[:40]}[/]"
            tabela.add_row(
                item.nome, _estado(item), f"[{_cor(linha.recomendada, total)}]{rec}[/]", alt_texto,
                linha.janela[:45],
            )
        console.print(tabela)
