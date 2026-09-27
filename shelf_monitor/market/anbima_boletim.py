"""Boletim de Fundos de Investimento da ANBIMA: a captacao e o PL de ETF, para conferir.

O anexo do boletim mensal (`Anexo_boletim_fundos_investimento_<mes>_valor_*.xlsx`,
de data.anbima.com.br) traz, por tipo ANBIMA e por mes, o PL ("Pag. 5 - PL por
Tipo") e a captacao liquida ("Pag. 9 - Cap. Liq. por Tipo"), em R$ milhoes, e o
numero de classes por categoria ("Pag. 13 - N de Fundos"). As linhas de ETF sao
"ETF", "ETF Renda Fixa" e "ETF Renda Variavel".

Por que o monitor e a ANBIMA nao dao o mesmo numero - medido em 2026-09-24 com o
boletim de 07/2026, na janela dele (08/2025 a 07/2026):

- o total: +R$ 58,4 bi no monitor, +R$ 53,3 bi na ANBIMA; no ano, +R$ 35,4 bi e
  +R$ 34,1 bi. Mes a mes a curva e a mesma: em 7 dos 12 meses a diferenca fica
  abaixo de R$ 0,3 bi;
- a base: a ANBIMA soma o que os administradores associados informam - R$ 120,4
  bi de PL em ETF em 07/2026, 93% dos R$ 129,5 bi do balancete da CVM, com 210
  classes contra 213 com PL na CVM. A diferenca de PL fica toda em renda
  variavel (84% a 87% de cobertura); em renda fixa a ANBIMA tem ate R$ 1,7 bi a
  mais;
- a estreia: o monitor conta o PL do primeiro fechamento como captacao. A
  ANBIMA nem sempre - em 10/2025 os ETFs da B-Index estrearam com R$ 2,36 bi e a
  diferenca do mes foi de R$ 2,39 bi; em 12/2025 os R$ 3,0 bi da estreia do IT
  NOW Tesouro Selic entraram na captacao dela.

A conferencia mostra, mes a mes, as duas captacoes, a parte do monitor que e
estreia e a cobertura de PL - para que a diferenca seja lida, nao adivinhada.
"""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from shelf_monitor.config import Config

ABA_PL = "Pág. 5 - PL por Tipo"
ABA_CAPTACAO = "Pág. 9 - Cap. Líq. por Tipo"
ABA_FUNDOS = "Pág. 13 - N° de Fundos"
TIPOS = {"ETF": "total", "ETF Renda Fixa": "renda_fixa", "ETF Renda Variável": "renda_variavel"}
_PADRAO_DO_ARQUIVO = "Anexo_boletim*.xlsx"
_MESES = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez")
# Diferenca mensal que merece explicacao na conferencia.
DIFERENCA_RELEVANTE = 1_000_000_000.0


def localizar_anexo(config: Config, caminho: str | Path | None = None) -> Path | None:
    """O anexo a usar: o arquivo dado, o mais recente da pasta dada, ou o mais
    recente de `data/etf/anbima/`."""
    alvo = Path(caminho) if caminho else config.etf_dir / "anbima"
    if alvo.is_file():
        return alvo
    if alvo.is_dir():
        candidatos = sorted(alvo.glob(_PADRAO_DO_ARQUIVO), key=lambda p: p.stat().st_mtime)
        return candidatos[-1] if candidatos else None
    return None


def _mes_do_cabecalho(valor: object) -> str | None:
    """datetime(2026, 6, 1) -> "202606"; "jul-26" -> "202607"; o resto, None."""
    if isinstance(valor, (datetime, date)):
        return f"{valor.year}{valor.month:02d}"
    if isinstance(valor, str):
        achado = re.fullmatch(r"\s*([a-zç]{3})[-/ ]?(\d{2}|\d{4})\s*", valor.lower())
        if achado and achado.group(1) in _MESES:
            ano = int(achado.group(2))
            ano = ano + 2000 if ano < 100 else ano
            return f"{ano}{_MESES.index(achado.group(1)) + 1:02d}"
    return None


def _serie_por_tipo(planilha) -> dict[str, dict[str, float]]:
    """Por tipo de ETF: mes -> valor em reais (o anexo vem em R$ milhoes)."""
    linhas = list(planilha.iter_rows(values_only=True))
    cabecalho = next(
        (row for row in linhas if row and any(isinstance(v, str) and v.strip() == "Tipos ANBIMA" for v in row)),
        None,
    )
    if cabecalho is None:
        return {}
    meses = {j: m for j, v in enumerate(cabecalho) if (m := _mes_do_cabecalho(v))}
    saida: dict[str, dict[str, float]] = {}
    for row in linhas:
        nome = next((v.strip() for v in row[:3] if isinstance(v, str) and v.strip() in TIPOS), None) if row else None
        if nome is None or TIPOS[nome] in saida:
            continue
        saida[TIPOS[nome]] = {
            m: float(row[j]) * 1e6 for j, m in meses.items()
            if j < len(row) and isinstance(row[j], (int, float))
        }
    return saida


def _fundos(planilha) -> dict[str, int]:
    linhas = list(planilha.iter_rows(values_only=True))
    cabecalho = next((row for row in linhas if row and "ETF" in [str(v).strip() for v in row if v]), None)
    if cabecalho is None:
        return {}
    coluna = [str(v).strip() if v else "" for v in cabecalho].index("ETF")
    saida = {}
    for row in linhas:
        mes = _mes_do_cabecalho(row[0]) if row else None
        if mes and coluna < len(row) and isinstance(row[coluna], (int, float)):
            saida[mes] = int(row[coluna])
    return saida


def ler_anexo(path: Path) -> dict | None:
    """PL, captacao e numero de ETFs do anexo do boletim. None se nao for um."""
    from openpyxl import load_workbook

    try:
        livro = load_workbook(path, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - arquivo que nao e o anexo nao derruba o relatorio
        return None
    # Em modo so-leitura o openpyxl segura o arquivo ate o close(): no Windows,
    # isso trava o anexo na pasta da pessoa.
    try:
        abas = set(livro.sheetnames)
        if not {ABA_PL, ABA_CAPTACAO} <= abas:
            return None
        pl = _serie_por_tipo(livro[ABA_PL])
        captacao = _serie_por_tipo(livro[ABA_CAPTACAO])
        fundos = _fundos(livro[ABA_FUNDOS]) if ABA_FUNDOS in abas else {}
    finally:
        livro.close()
    if not captacao.get("total"):
        return None
    return {"arquivo": path.name, "pl": pl, "captacao": captacao, "fundos": fundos}


def conciliar(bloco: dict, anexo: dict) -> dict | None:
    """Monitor x ANBIMA, mes a mes, na janela que os dois cobrem."""
    historico = (bloco.get("mercado") or {}).get("historico") or {}
    grupos = {g["grupo"]: g.get("historico") or {} for g in bloco.get("grupos") or []}
    anbima = anexo["captacao"]
    comuns = sorted(m for m in historico if m in anbima.get("total", {}) and historico[m]["fluxo"] is not None)
    if not comuns:
        return None

    def do_monitor(mes: str, renda_fixa: bool) -> dict:
        fluxo = pl = 0.0
        for grupo, serie in grupos.items():
            if (grupo == "Renda Fixa") != renda_fixa or mes not in serie:
                continue
            fluxo += serie[mes]["fluxo"] or 0.0
            pl += serie[mes]["pl"] or 0.0
        return {"fluxo": fluxo, "pl": pl}

    linhas = []
    for mes in comuns:
        h = historico[mes]
        pl_anbima = anexo["pl"].get("total", {}).get(mes)
        rf, rv = do_monitor(mes, True), do_monitor(mes, False)
        linhas.append({
            "mes": mes,
            "monitor": h["fluxo"],
            "estreias": h["estreias"],
            "anbima": anbima["total"][mes],
            "diferenca": h["fluxo"] - anbima["total"][mes],
            "pl_monitor": h["pl"],
            "pl_anbima": pl_anbima,
            "cobertura": (pl_anbima / h["pl"]) if pl_anbima and h["pl"] else None,
            "renda_fixa": {"monitor": rf["fluxo"], "anbima": anbima.get("renda_fixa", {}).get(mes),
                           "pl_monitor": rf["pl"], "pl_anbima": anexo["pl"].get("renda_fixa", {}).get(mes)},
            "renda_variavel": {"monitor": rv["fluxo"], "anbima": anbima.get("renda_variavel", {}).get(mes),
                               "pl_monitor": rv["pl"], "pl_anbima": anexo["pl"].get("renda_variavel", {}).get(mes)},
        })

    def soma(meses: list[str]) -> dict:
        por_mes = {l["mes"]: l for l in linhas}
        return {
            "de": meses[0], "ate": meses[-1],
            "monitor": sum(por_mes[m]["monitor"] for m in meses),
            "anbima": sum(por_mes[m]["anbima"] for m in meses),
            "estreias": sum(por_mes[m]["estreias"] for m in meses),
        }

    doze = comuns[-12:]
    ano = [m for m in doze if m[:4] == comuns[-1][:4]]
    ultimo = linhas[-1]
    return {
        "arquivo": anexo["arquivo"],
        "ultimo_mes": comuns[-1],
        "linhas": linhas,
        "doze_meses": soma(doze),
        "no_ano": soma(ano) if ano else None,
        "fundos_anbima": anexo.get("fundos", {}).get(comuns[-1]),
        "etfs_monitor": historico[comuns[-1]]["etfs_com_dado"],
        "cobertura": ultimo["cobertura"],
        "destaques": [
            {
                "mes": l["mes"], "diferenca": l["diferenca"], "estreias": l["estreias"],
                "etfs_em_estreia": _estreias_do_mes(bloco, l["mes"]),
            }
            for l in linhas if abs(l["diferenca"]) >= DIFERENCA_RELEVANTE
        ],
    }


def _estreias_do_mes(bloco: dict, mes: str) -> list[dict]:
    """Os ETFs que estrearam no mes, do maior PL ao menor (so meses da visao)."""
    estreias = [
        {"ticker": e.get("ticker"), "nome": e["nome"], "gestora": e["gestora"], "valor": e["fluxo"].get(mes)}
        for e in bloco.get("etfs") or []
        if (e.get("origem") or {}).get(mes, "").startswith("estreia") and e["fluxo"].get(mes)
    ]
    return sorted(estreias, key=lambda x: -x["valor"])[:5]


def leitura_da_conferencia(conferencia: dict) -> str:
    """A frase que resume a conferencia, com a janela e os numeros dela."""
    from shelf_monitor.market.etf_captacao_mensal import reais_com_sinal
    from shelf_monitor.market.etf_competicao import pct
    from shelf_monitor.texto import competencia_br

    doze = conferencia["doze_meses"]
    frase = (
        f"Conferência com o boletim da ANBIMA: de {competencia_br(doze['de'])} a "
        f"{competencia_br(doze['ate'])}, o monitor soma {reais_com_sinal(doze['monitor'])} e a "
        f"ANBIMA {reais_com_sinal(doze['anbima'])}"
    )
    ano = conferencia.get("no_ano")
    if ano:
        frase += f"; no ano, {reais_com_sinal(ano['monitor'])} e {reais_com_sinal(ano['anbima'])}"
    frase += (
        f". A base da ANBIMA tem {pct(conferencia['cobertura'])} do PL de ETF da CVM, e o monitor "
        f"conta como captação o PL de estreia ({reais_com_sinal(doze['estreias'])} na janela), que "
        "a ANBIMA nem sempre conta"
    )
    destaque = max(conferencia["destaques"], key=lambda d: abs(d["diferenca"]), default=None)
    if destaque:
        frase += (
            f" - em {competencia_br(destaque['mes'])}, a diferença foi de "
            f"{reais_com_sinal(destaque['diferenca'])}, com {reais_com_sinal(destaque['estreias'])} "
            "de estreias no mês"
        )
    return frase + "."


def anexar_conferencia(bloco: dict, anexo: dict | None) -> dict:
    """Poe a conferencia com a ANBIMA no bloco mensal (e a frase na leitura)."""
    if not bloco.get("available") or not anexo:
        return bloco
    conferencia = conciliar(bloco, anexo)
    if conferencia:
        bloco["anbima"] = conferencia
        bloco["leitura"] = [*bloco["leitura"], leitura_da_conferencia(conferencia)]
    return bloco
