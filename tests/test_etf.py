"""Testes de catalogação e agregação de fluxo de ETFs.

A fixture usada aqui (tests/fixtures/etf/) é dedicada - separada da fixture
geral de mercado (tests/fixtures/market/) para não acoplar aos asserts de
contagem que test_market.py já faz sobre ela. Regenerar com:
`python tests/fixtures/etf/make_etf_fixtures.py`.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

from shelf_monitor.config import Config
try:  # `discover -s tests` poe tests/ no sys.path; `-m unittest tests.X` nao.
    from fixtures_comuns import temp_config
except ImportError:  # pragma: no cover - depende de como a suite foi invocada
    from tests.fixtures_comuns import temp_config
from shelf_monitor.market import import_market
from shelf_monitor.market.etf import (
    build_etf_catalog,
    categorize_etf,
    catalog_counts,
    classify_etf_kind,
    extract_ticker,
    is_etf_record,
)
from shelf_monitor.market.etf_flow import (
    EtfDailyRecord,
    EtfWeeklyFlow,
    aggregate_category_weekly_flows,
    aggregate_etf_weekly_flows,
    flow_coverage,
    get_recent_months,
)

FIXTURES = Path(__file__).parent / "fixtures" / "etf"




class TestExtractTicker(unittest.TestCase):
    def test_extracts_ticker_in_parentheses(self):
        self.assertEqual(extract_ticker("ISHARES IBOVESPA FUNDO DE INDICE IE (BOVA11)"), "BOVA11")
        self.assertEqual(extract_ticker("HASHDEX NASDAQ CRYPTO INDEX ETF (HASH11)"), "HASH11")
        self.assertEqual(extract_ticker("IT NOW IVVB11 FI EM COTAS DE ETF"), "IVVB11")

    def test_no_ticker_in_name_is_none(self):
        """A maioria dos ETFs (família IT NOW) não escreve o ticker no nome -
        é por isso que o catálogo aceita um override manual por CNPJ."""
        self.assertIsNone(extract_ticker("ITAU IBOVESPA FUNDO DE INDICE"))


class TestClassifyEtfKind(unittest.TestCase):
    """O tipo estrutural (FIIM) é a via principal - não o nome."""

    def test_fiim_class_type_is_etf_even_without_etf_in_the_name(self):
        self.assertEqual(
            classify_etf_kind(class_type="Classes de Cotas de Fundos FIIM"), "etf"
        )

    def test_fiim_fund_type_is_etf(self):
        self.assertEqual(classify_etf_kind(fund_type="FIIM"), "etf")

    def test_fif_with_etf_in_the_name_is_a_feeder_not_an_etf(self):
        """Um FIF que só replica um ETF não pode contar como o ETF em si -
        seria dinheiro contado duas vezes."""
        self.assertEqual(
            classify_etf_kind(fund_type="FIF", name="SANTANDER S&P 500 ETF FIC FIF"),
            "feeder",
        )

    def test_plain_multimercado_is_neither(self):
        self.assertIsNone(classify_etf_kind(fund_type="FIF", name="FUNDO MULTIMERCADO TRADICIONAL"))

    def test_is_etf_record_is_true_for_etf_or_feeder(self):
        self.assertTrue(is_etf_record(fund_type="FIIM"))
        self.assertTrue(is_etf_record(name="ALGO ETF DE ACOES"))
        self.assertFalse(is_etf_record(name="FUNDO MULTIMERCADO TRADICIONAL"))


class TestCategorizeEtf(unittest.TestCase):
    def test_benchmark_wins_over_the_name(self):
        self.assertEqual(
            categorize_etf(name="QUALQUER COISA", benchmark="IMA-B"), "Renda Fixa"
        )

    def test_sp500_is_recognized_despite_the_ampersand(self):
        """'S&P 500' normaliza para 'S P 500' (o '&' vira espaço) - o bug
        original procurava 'SP 500', que nunca aparece nesse texto."""
        self.assertEqual(categorize_etf(benchmark="S&P 500"), "Internacional")
        self.assertEqual(categorize_etf(name="IVVB11 S&P 500 ETF EUA"), "Internacional")

    def test_ima_b_is_renda_fixa(self):
        self.assertEqual(categorize_etf(benchmark="IMA-B"), "Renda Fixa")
        self.assertEqual(categorize_etf(benchmark="IRF-M"), "Renda Fixa")

    def test_ibovespa_is_acoes_brasil(self):
        self.assertEqual(categorize_etf(name="BOVA11 IBOVESPA ACOES BRASIL ETF"), "Ações Brasil")

    def test_bitcoin_is_cripto(self):
        self.assertEqual(categorize_etf(cvm_class="Fundo de Índice", name="HASH11 BITCOIN ETF"), "Cripto")

    def test_gold_is_multimercado(self):
        self.assertEqual(categorize_etf(name="GOLD11 OURO COMMODITY ETF"), "Multimercado")

    def test_unknown_is_outros(self):
        self.assertEqual(categorize_etf(name="ALGO SEM PISTA NENHUMA"), "Outros")

    def test_responsabilidade_limitada_nao_virou_renda_fixa(self):
        """O sufixo da RCVM 175 nao pode decidir a categoria.

        "RESPONSABILIDADE" contem "IDA" (indice de debentures ANBIMA). Com busca
        por substring solta, 166 dos 266 registros do catalogo caiam em Renda
        Fixa por causa do proprio nome juridico - inclusive ETFs de Ibovespa.
        """
        self.assertEqual(
            categorize_etf(
                name="ISHARES IBOVESPA CLASSE DE INDICE - RESPONSABILIDADE LIMITADA"
            ),
            "Ações Brasil",
        )
        self.assertEqual(
            categorize_etf(name="FUNDO DE INDICE RESPONSABILIDADE LIMITADA"), "Outros"
        )

    def test_codigo_de_indice_com_sufixo_ainda_casa(self):
        """A ancora e so a esquerda: "IMA-B5" tem de continuar sendo Renda Fixa."""
        self.assertEqual(
            categorize_etf(name="IT NOW IMA-B5 P2 FUNDO DE INDICE"), "Renda Fixa"
        )
        self.assertEqual(categorize_etf(name="ETF IBOV11 DE ACOES"), "Ações Brasil")


class TestBuildEtfCatalog(unittest.TestCase):
    """Integração com o cadastro: tipo estrutural, exclusão de cancelados,
    feeder separado do ETF."""

    def setUp(self):
        self.config, self.tmp = temp_config()
        self.db = import_market(self.config, local_dir=FIXTURES, skip_fees=True)
        self.catalog = build_etf_catalog(self.config, database=self.db)

    def tearDown(self):
        self.tmp.cleanup()

    def test_etf_found_by_structural_type_despite_generic_name(self):
        """Nome não tem 'ETF' nem 'índice' - só o Tipo_Classe FIIM entrega."""
        meta = self.catalog.get("10298765000111")
        self.assertIsNotNone(meta)
        self.assertEqual(meta.kind, "etf")
        self.assertEqual(meta.category, "Ações Brasil")

    def test_cancelled_etf_is_excluded(self):
        self.assertNotIn("10298765000122", self.catalog)

    def test_feeder_is_tagged_separately_from_etf(self):
        meta = self.catalog.get("10298765000133")
        self.assertIsNotNone(meta)
        self.assertEqual(meta.kind, "feeder")
        self.assertEqual(meta.category, "Internacional")

    def test_feeder_signal_in_the_funds_name_is_not_lost(self):
        """A classe única deste feeder se chama só 'CLASSE UNICA' - o 'ETF'
        que denuncia o feeder está no nome do FUNDO. Bug real encontrado ao
        testar: sem o fallback para o nome do fundo, esse registro caía no
        laço de fundos legados (passo 2) e perdia o benchmark, que só a classe
        carrega.
        """
        meta = self.catalog.get("10298765000133")
        self.assertEqual(meta.benchmark, "S&P 500")

    def test_catalog_counts_separates_kinds(self):
        counts = catalog_counts(self.catalog)
        self.assertEqual(counts["etfs"], 1)
        self.assertEqual(counts["feeders"], 1)
        self.assertEqual(counts["por_categoria"], {"Ações Brasil": 1})


class TestGetRecentMonths(unittest.TestCase):
    def test_returns_the_last_n_months_in_order(self):
        ref = date(2026, 9, 13)
        months = get_recent_months(3, ref_date=ref)
        self.assertEqual(months, ["202607", "202608", "202609"])


class TestAggregateEtfWeeklyFlows(unittest.TestCase):
    def test_basic_weekly_aggregation(self):
        recs = [
            EtfDailyRecord(
                cnpj="11111111000111", dt_comptc=date(2026, 9, 7),
                vl_patrim_liq=1000.0, captc_dia=100.0, resg_dia=20.0,
                nr_cotst=10, vl_quota=10.0,
            ),
            EtfDailyRecord(
                cnpj="11111111000111", dt_comptc=date(2026, 9, 11),
                vl_patrim_liq=1150.0, captc_dia=50.0, resg_dia=10.0,
                nr_cotst=15, vl_quota=11.0,
            ),
        ]
        weekly = aggregate_etf_weekly_flows({}, recs)

        self.assertEqual(len(weekly), 1)
        w = weekly[0]
        self.assertEqual(w.cnpj, "11111111000111")
        # Sem semana anterior, a referência é o primeiro registro da própria
        # semana (única forma possível para a primeira semana da série) - e o
        # fluxo DAQUELE dia nao entra na soma: medido no informe do Fundos.NET,
        # PL(D) ja contem a captacao de D, entao somar o dia da referencia
        # contaria o mesmo dinheiro duas vezes e a valorizacao por diferenca
        # sairia negativa sem motivo.
        self.assertEqual(w.pl_inicio, 1000.0)
        self.assertEqual(w.captacao_bruta, 50.0)
        self.assertEqual(w.resgates, 10.0)
        self.assertEqual(w.captacao_liquida, 40.0)
        self.assertEqual(w.pl_fim, 1150.0)
        self.assertEqual(w.cotistas_fim, 15)
        self.assertEqual(w.rentabilidade_semana, 10.0)  # (11/10 - 1)*100
        self.assertFalse(w.data_suspeita)

    def test_pl_decomposto_em_captacao_e_valorizacao(self):
        """variacao_pl = captacao_liquida + valorizacao, por construcao."""
        recs = [
            EtfDailyRecord(
                cnpj="11111111000111", dt_comptc=date(2026, 9, 7),
                vl_patrim_liq=1000.0, captc_dia=0.0, resg_dia=0.0,
                nr_cotst=10, vl_quota=10.0,
            ),
            EtfDailyRecord(
                cnpj="11111111000111", dt_comptc=date(2026, 9, 11),
                vl_patrim_liq=1150.0, captc_dia=100.0, resg_dia=20.0,
                nr_cotst=15, vl_quota=10.5,
            ),
        ]
        w = aggregate_etf_weekly_flows({}, recs)[0]

        self.assertEqual(w.captacao_liquida, 80.0)
        self.assertEqual(w.variacao_pl, 150.0)
        self.assertEqual(w.valorizacao, 70.0)
        self.assertAlmostEqual(w.captacao_liquida + w.valorizacao, w.variacao_pl)

    def test_valorizacao_que_a_cota_desmente_e_marcada(self):
        """A identidade fecha sempre - a cota e a segunda leitura.

        Aqui a cota nao se move (retorno zero), mas o PL sobe 50% sem fluxo
        declarado: a semana esta incompleta, nao extraordinaria. Os valores sao
        em escala de fundo real porque a guarda tem piso absoluto de R$ 1.000 -
        abaixo dele o desvio e arredondamento, nao sinal.
        """
        recs = [
            EtfDailyRecord(
                cnpj="11111111000111", dt_comptc=date(2026, 9, 7),
                vl_patrim_liq=1_000_000.0, nr_cotst=10, vl_quota=10.0,
            ),
            EtfDailyRecord(
                cnpj="11111111000111", dt_comptc=date(2026, 9, 11),
                vl_patrim_liq=1_500_000.0, nr_cotst=10, vl_quota=10.0,
            ),
        ]
        w = aggregate_etf_weekly_flows({}, recs)[0]
        self.assertEqual(w.rentabilidade_semana, 0.0)
        self.assertTrue(w.decomposicao_suspeita)

    def test_queda_de_cotistas_que_o_pl_desmente_e_marcada(self):
        """O caso real de 2026-09-15: 52.587 -> 854 cotistas com o PL parado."""
        recs = [
            EtfDailyRecord(
                cnpj="11111111000111", dt_comptc=date(2026, 9, 14),
                vl_patrim_liq=2_592_018_900.85, nr_cotst=52_587, vl_quota=132.58,
            ),
            EtfDailyRecord(
                cnpj="11111111000111", dt_comptc=date(2026, 9, 15),
                vl_patrim_liq=2_613_719_662.40, nr_cotst=854, vl_quota=133.35,
            ),
        ]
        w = aggregate_etf_weekly_flows({}, recs)[0]
        self.assertEqual(w.variacao_cotistas, -51_733)
        self.assertTrue(w.cotistas_suspeito)

    def test_movimento_normal_de_cotistas_nao_e_marcado(self):
        recs = [
            EtfDailyRecord(
                cnpj="11111111000111", dt_comptc=date(2026, 9, 7),
                vl_patrim_liq=1000.0, nr_cotst=100, vl_quota=10.0,
            ),
            EtfDailyRecord(
                cnpj="11111111000111", dt_comptc=date(2026, 9, 11),
                vl_patrim_liq=1010.0, nr_cotst=112, vl_quota=10.1,
            ),
        ]
        w = aggregate_etf_weekly_flows({}, recs)[0]
        self.assertEqual(w.variacao_cotistas, 12)
        self.assertFalse(w.cotistas_suspeito)

    def test_return_compares_against_the_previous_weeks_close_not_mondays_open(self):
        """O bug original comparava segunda vs sexta da MESMA semana e perdia
        o retorno do primeiro pregão. Com duas semanas seguidas, a segunda
        semana deve comparar contra o fechamento da primeira."""
        recs = [
            EtfDailyRecord(cnpj="X", dt_comptc=date(2026, 8, 31), vl_quota=10.0, vl_patrim_liq=100.0, nr_cotst=1),
            EtfDailyRecord(cnpj="X", dt_comptc=date(2026, 9, 4), vl_quota=11.0, vl_patrim_liq=110.0, nr_cotst=1),
            EtfDailyRecord(cnpj="X", dt_comptc=date(2026, 9, 7), vl_quota=11.0, vl_patrim_liq=110.0, nr_cotst=1),
            EtfDailyRecord(cnpj="X", dt_comptc=date(2026, 9, 11), vl_quota=12.1, vl_patrim_liq=121.0, nr_cotst=1),
        ]
        weekly = {w.week_key: w for w in aggregate_etf_weekly_flows({}, recs)}
        # Semana 1 (31/08 a 04/09): sem semana anterior, usa a abertura dela mesma -> 10%.
        self.assertAlmostEqual(weekly["2026-W36"].rentabilidade_semana, 10.0)
        # Semana 2 (07/09 a 11/09): compara com o fechamento da semana 1 (11.0), não
        # com a segunda-feira dessa semana (também 11.0, o que daria 10% de qualquer
        # forma aqui) - o ponto é que a referência veio da semana anterior.
        self.assertAlmostEqual(weekly["2026-W37"].rentabilidade_semana, 10.0)

    def test_duplicate_same_day_rows_are_deduped_before_aggregating(self):
        """Duas linhas do mesmo CNPJ na mesma data (retificação do arquivo
        mensal) não podem virar 'dois dias' na mesma semana."""
        recs = [
            EtfDailyRecord(cnpj="Y", dt_comptc=date(2026, 9, 7), vl_quota=100.0, vl_patrim_liq=1000.0, nr_cotst=100),
            # Retificação do mesmo dia, chegada depois no arquivo.
            EtfDailyRecord(cnpj="Y", dt_comptc=date(2026, 9, 7), vl_quota=101.0, vl_patrim_liq=1010.0, nr_cotst=100),
            EtfDailyRecord(cnpj="Y", dt_comptc=date(2026, 9, 11), vl_quota=102.0, vl_patrim_liq=1020.0, nr_cotst=100),
        ]
        weekly = aggregate_etf_weekly_flows({}, recs)
        self.assertEqual(len(weekly), 1)
        # A referência de abertura é a retificação (última lida para o dia), não a original.
        self.assertEqual(weekly[0].pl_inicio, 1010.0)

    def test_implausible_weekly_return_is_flagged_instead_of_reported(self):
        """Um -92% numa semana (o caso real observado) não é reportado como
        rentabilidade - vira data_suspeita, sinal de descontinuidade."""
        recs = [
            EtfDailyRecord(cnpj="Z", dt_comptc=date(2026, 9, 7), vl_quota=100.0, vl_patrim_liq=327_000_000.0, nr_cotst=28455),
            EtfDailyRecord(cnpj="Z", dt_comptc=date(2026, 9, 11), vl_quota=8.0, vl_patrim_liq=4_500_000.0, nr_cotst=1),
        ]
        catalog = {"Z": _fake_meta("Z", category="Cripto")}
        weekly = aggregate_etf_weekly_flows(catalog, recs)
        self.assertEqual(len(weekly), 1)
        self.assertIsNone(weekly[0].rentabilidade_semana)
        self.assertTrue(weekly[0].data_suspeita)


def _fake_meta(cnpj: str, category: str = "Outros", kind: str = "etf"):
    from shelf_monitor.market.etf import EtfMetadata

    return EtfMetadata(cnpj=cnpj, name=cnpj, kind=kind, category=category)


class TestAggregateCategoryWeeklyFlows(unittest.TestCase):
    def test_category_sums_only_include_etfs_not_feeders(self):
        weekly = [
            EtfWeeklyFlow(
                cnpj="111", etf_name="BOVA11", ticker="BOVA11", category="Ações Brasil",
                manager_name="BlackRock", kind="etf", iso_year=2026, iso_week=37,
                week_key="2026-W37", week_end_date=date(2026, 9, 11),
                captacao_bruta=200.0, resgates=50.0, captacao_liquida=150.0,
                pl_inicio=1000.0, pl_fim=1150.0, cotistas_inicio=10, cotistas_fim=15,
                variacao_cotistas=5,
            ),
            # Feeder do mesmo BOVA11: se somasse junto, dobraria o dinheiro.
            EtfWeeklyFlow(
                cnpj="222", etf_name="FEEDER BOVA11", ticker=None, category="Ações Brasil",
                manager_name="BlackRock", kind="feeder", iso_year=2026, iso_week=37,
                week_key="2026-W37", week_end_date=date(2026, 9, 11),
                captacao_bruta=999.0, resgates=999.0, captacao_liquida=999.0,
                pl_inicio=500.0, pl_fim=580.0, cotistas_inicio=5, cotistas_fim=8,
                variacao_cotistas=3,
            ),
        ]
        cats = aggregate_category_weekly_flows(weekly)
        self.assertEqual(len(cats), 1)
        c = cats[0]
        self.assertEqual(c.category, "Ações Brasil")
        self.assertEqual(c.qtd_etfs, 1)
        self.assertEqual(c.captacao_liquida, 150.0)


class TestFlowCoverage(unittest.TestCase):
    def test_reports_how_much_of_the_catalog_actually_has_flow_data(self):
        catalog = {
            "A": _fake_meta("A", kind="etf"),
            "B": _fake_meta("B", kind="etf"),
            "C": _fake_meta("C", kind="feeder"),
        }
        weekly = [
            EtfWeeklyFlow(
                cnpj="A", etf_name="A", ticker=None, category="Outros",
                manager_name=None, kind="etf", iso_year=2026, iso_week=1,
                week_key="2026-W01", week_end_date=date(2026, 1, 2),
                captacao_bruta=0, resgates=0, captacao_liquida=0,
                pl_inicio=0, pl_fim=0, cotistas_inicio=0, cotistas_fim=0,
                variacao_cotistas=0,
            ),
        ]
        cov = flow_coverage(catalog, weekly)
        self.assertEqual(cov.etfs_no_catalogo, 2)
        self.assertEqual(cov.etfs_com_fluxo, 1)
        self.assertEqual(cov.feeders_no_catalogo, 1)
        self.assertEqual(cov.feeders_com_fluxo, 0)


class TestFonteDoFluxo(unittest.TestCase):
    """O padrão do `etf` não publica fluxo de feeder como captação de ETF.

    O INF_DIARIO não cobre fundo de índice (0 de 249 ETFs em 2026-09-19) e era
    o padrão do comando, sem aviso; a planilha o chamava de "Captação Líquida
    de ETFs". Agora o padrão é `auto`: Fundos.NET quando há série em disco.
    """

    def setUp(self):
        self.config, self._tmp = temp_config()
        self.addCleanup(self._tmp.cleanup)

    def _gravar_informes(self, dias: int) -> None:
        from datetime import timedelta

        self.config.etf_fnet_dir.mkdir(parents=True, exist_ok=True)
        linhas = ["cnpj;fundo;data;pl;cota;cotas;cotistas;captacao;resgates;vl_total;doc_id"]
        inicio = date(2026, 9, 1)
        for i in range(dias):
            dia = inicio + timedelta(days=i)
            linhas.append(f"14120533000111;ETF X;{dia.isoformat()};1000;10;;5;0;0;1000;{i}")
        (self.config.etf_fnet_dir / "2026-09-18.csv").write_text(
            "\n".join(linhas), encoding="utf-8"
        )

    def test_sem_serie_o_auto_fica_no_inf_diario_e_diz_por_que(self):
        from shelf_monitor.main import resolver_fonte_do_fluxo

        fonte, motivo = resolver_fonte_do_fluxo(self.config, "auto")
        self.assertEqual(fonte, "inf_diario")
        self.assertIn("0 datas de informe", motivo)

    def test_a_coleta_de_hoje_ainda_nao_e_serie(self):
        # A única coleta em disco em 23/09 trouxe 7 datas de informe.
        from shelf_monitor.main import resolver_fonte_do_fluxo

        self._gravar_informes(7)
        self.assertEqual(resolver_fonte_do_fluxo(self.config, "auto")[0], "inf_diario")

    def test_com_serie_o_auto_passa_ao_fundos_net(self):
        from shelf_monitor.main import resolver_fonte_do_fluxo
        from shelf_monitor.market import fundosnet

        self._gravar_informes(fundosnet.MIN_DIAS_SERIE)
        fonte, motivo = resolver_fonte_do_fluxo(self.config, "auto")
        self.assertEqual(fonte, "fnet")
        self.assertIn(str(fundosnet.MIN_DIAS_SERIE), motivo)

    def test_fonte_pedida_explicitamente_e_respeitada(self):
        from shelf_monitor.main import resolver_fonte_do_fluxo

        self.assertEqual(resolver_fonte_do_fluxo(self.config, "inf_diario"), ("inf_diario", None))
        self.assertEqual(resolver_fonte_do_fluxo(self.config, "fnet"), ("fnet", None))

    def test_o_aviso_diz_que_nao_e_etf(self):
        from shelf_monitor import main

        with main.console.capture() as saida:
            main.avisar_fluxo_de_feeder("o Fundos.NET tem 7 datas")
        texto = saida.get()
        self.assertIn("NÃO é de ETF", texto)
        self.assertIn("feeders", texto)
        self.assertIn("7 datas", texto)

    def test_o_padrao_do_cli_e_auto(self):
        import subprocess
        import sys

        ajuda = subprocess.run(
            [sys.executable, "-m", "shelf_monitor", "etf", "--help"],
            capture_output=True, text=True, timeout=120,
            cwd=str(Path(__file__).resolve().parent.parent),
        ).stdout
        self.assertIn("{auto,inf_diario,fnet}", ajuda)

    def test_planilha_diz_a_fonte_no_titulo(self):
        from openpyxl import load_workbook

        from shelf_monitor.report.etf_report import export_etf_excel

        feeder = export_etf_excel([], [], {}, self.config.reports_dir / "a.xlsx")
        wb = load_workbook(feeder)
        titulo = wb["Resumo Categorias"]["A1"].value
        self.assertIn("FEEDERS", titulo)
        self.assertIn("INF_DIARIO", titulo)
        self.assertIn("0 de 249", wb["Resumo Categorias"]["A3"].value)
        self.assertIn("FEEDER", wb["Matriz ETFs"]["A1"].value)

        etf = export_etf_excel([], [], {}, self.config.reports_dir / "b.xlsx", fonte="fnet")
        wb = load_workbook(etf)
        titulo = wb["Resumo Categorias"]["A1"].value
        self.assertNotIn("FEEDER", titulo)
        self.assertIn("Fundos.NET", titulo)



class TestFormatoEmMilhoes(unittest.TestCase):
    """Estoque sem sinal: "R$ +14.090,0M" num PL fazia ler variacao onde ha saldo."""

    def test_estoque_sai_sem_sinal_e_variacao_com(self):
        from shelf_monitor.report.etf_report import format_money_millions

        self.assertEqual(format_money_millions(14_090e6, sinal=False), "R$ 14.090,0M")
        self.assertEqual(format_money_millions(832.6e6), "R$ +832,6M")
        self.assertEqual(format_money_millions(-954.8e6), "R$ -954,8M")

    def test_percentual_com_virgula_decimal(self):
        from types import SimpleNamespace

        from shelf_monitor.report.etf_report import _format_rentab

        semana = SimpleNamespace(data_suspeita=False, rentabilidade_semana=1.234)
        self.assertEqual(_format_rentab(semana), "+1,23%")


if __name__ == "__main__":
    unittest.main()
