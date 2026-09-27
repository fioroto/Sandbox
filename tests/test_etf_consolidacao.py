"""Testes da consolidacao: uma linha por ETF, tudo numa planilha so.

Monta em disco um mundo pequeno no formato que o `etf-bruto` grava - um ano
de pregoes, informes do FNET, fechamentos mensais, cadastro e carteira da
CVM, serie do Ibovespa e as planilhas manuais - e confere cada familia de
coluna. Sem rede.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import unittest
from datetime import date, timedelta

from openpyxl import load_workbook

try:  # `discover -s tests` poe tests/ no sys.path; `-m unittest tests.X` nao.
    from fixtures_comuns import temp_config
    import etf_bruto_apoio as apoio
except ImportError:  # pragma: no cover - depende de como a suite foi invocada
    from tests.fixtures_comuns import temp_config
    from tests import etf_bruto_apoio as apoio

from shelf_monitor.market.etf_bruto import armazem, b3_arquivos, consolidacao, manual, universo
from shelf_monitor.market.etf_bruto.composicao import pasta_cda
from shelf_monitor.market.etf_bruto.fnet import pasta_informe
from shelf_monitor.market.etf_bruto.indices_b3 import pasta as pasta_indices
from shelf_monitor.report.etf_consolidado import escrever_csv, escrever_planilha

HOJE = date(2026, 9, 25)
BOVA, NOVO, RENF = "10406511000161", "22222222000122", "33333333000133"


def _dias():
    return sorted(b3_arquivos.dias_uteis(date(2025, 8, 18), date(2026, 9, 24)))


def _preco_bova(i: int) -> float:
    return round(100 * (1 + 0.001 * i) * (1.01 if i % 2 else 0.99), 2)


class _Mundo(unittest.TestCase):
    """BOVA11 (acoes, antigo, com data de inicio reiniciada pela B3), NOVO11
    (indice estrangeiro, estreou no fim da serie), RENF11 (renda fixa) e um
    cancelado que so existe no cadastro da B3."""

    @classmethod
    def setUpClass(cls):
        # O mundo e montado uma vez por classe: sao ~280 arquivos de pregao.
        cls.config, cls.tmp = temp_config()
        cls.addClassCleanup(cls.tmp.cleanup)
        cls.dias = _dias()
        cls.estreia = cls.dias[-20]
        self = cls
        armazem.gravar_csv(self.config.etf_universo_path, universo.COLUNAS_UNIVERSO, [
            {"chave": BOVA, "cnpj": BOVA, "ticker": "BOVA11", "isin": "BRBOVACTF003", "nome": "ISHARES IBOVESPA FUNDO DE INDICE",
             "tipo_b3": "ETF", "categoria_b3": "ETF EQUITIES", "inicio_negociacao": "2026-07-16", "liquidacao_dias": "2",
             "lote": "1", "listado": "sim", "negocia": "sim", "na_cvm": "sim", "gestora": "BLACKROCK BRASIL"},
            {"chave": NOVO, "cnpj": NOVO, "ticker": "NOVO11", "nome": "TREND ETF NOVO S&P 500 FUNDO DE INDICE",
             "tipo_b3": "ETF", "categoria_b3": "ETF FOREIGN INDEX", "inicio_negociacao": self.estreia.isoformat(),
             "liquidacao_dias": "2", "listado": "sim", "negocia": "sim", "na_cvm": "sim"},
            {"chave": RENF, "cnpj": RENF, "ticker": "RENF11", "nome": "IT NOW IMA-B FUNDO DE INDICE", "tipo_b3": "ETF-RF",
             "categoria_b3": "FIXED INCOME TRADABLE INSTRUMENT T1", "inicio_negociacao": "2020-06-29",
             "liquidacao_dias": "1", "listado": "sim", "negocia": "sim", "na_cvm": "sim"},
            {"chave": "ticker:XBIO11", "ticker": "XBIO11", "nome": "TREND ETF BIOTECHNOLOGY", "categoria_b3": "ETF FOREIGN INDEX",
             "inicio_negociacao": "9999-12-31", "listado": "não", "negocia": "não", "na_cvm": "não"},
        ])
        self._negocios()
        self._informes()
        self._mensal()
        self._cvm()
        self._carteira()
        self._ibovespa()

    @classmethod
    def _negocios(self):
        cabecalho = apoio.CABECALHO_NEGOCIOS + list(b3_arquivos.COLUNAS_EXTRAS)
        for i, dia in enumerate(self.dias):
            linhas = [
                {"RptDt": dia.isoformat(), "TckrSymb": "BOVA11", "SgmtNm": "CASH", "LastPric": f"{_preco_bova(i)}".replace(".", ","),
                 "TradQty": "100", "NtlFinVol": f"{1000 + i}", "x_tipo_linha": "padrao", "x_ticker_base": "BOVA11"},
                # Lote de bloco gigante: nao pode mexer na mediana.
                {"RptDt": dia.isoformat(), "TckrSymb": "BOVA11Q", "SgmtNm": "Equity Block Trading Lot", "LastPric": "1",
                 "TradQty": "1", "NtlFinVol": "99999999", "x_tipo_linha": "bloco", "x_ticker_base": "BOVA11"},
                {"RptDt": dia.isoformat(), "TckrSymb": "RENF11", "SgmtNm": "FORWARD", "LastPric": "110,00",
                 "TradQty": "10", "NtlFinVol": "500", "x_tipo_linha": "padrao", "x_ticker_base": "RENF11"},
            ]
            # NOVO11 negocia so em pregao par a partir da estreia.
            if dia >= self.estreia and i % 2 == 0:
                linhas.append({"RptDt": dia.isoformat(), "TckrSymb": "NOVO11", "SgmtNm": "CASH", "LastPric": "50,00",
                               "TradQty": "4", "NtlFinVol": "200", "x_tipo_linha": "padrao", "x_ticker_base": "NOVO11"})
            armazem.gravar_csv(b3_arquivos.pasta_negocios(self.config) / f"{dia.isoformat()}.csv", cabecalho, linhas)

    @classmethod
    def _informes(self):
        cabecalho = ("doc_id", "CNPJ_FDO", "DT_COMPT", "VL_QUOTA", "PATRIM_LIQ", "NR_COTST", "CAPTC_DIA", "RESG_DIA")
        for n, dia in enumerate(self.dias[-4:-1]):
            i = len(self.dias) - 4 + n
            armazem.gravar_csv(pasta_informe(self.config) / f"{dia.isoformat()}.csv", cabecalho, [
                {"doc_id": str(100 + n), "CNPJ_FDO": BOVA, "DT_COMPT": dia.strftime("%d/%m/%Y"),
                 "VL_QUOTA": str(_preco_bova(i) * 1.01).replace(".", ","), "PATRIM_LIQ": "1000000,00", "NR_COTST": "5000",
                 "CAPTC_DIA": "300,00", "RESG_DIA": "100,00"},
            ])
        # O dia mais recente sem cota: vale o anterior para o premio.
        armazem.gravar_csv(pasta_informe(self.config) / "2026-09-24.csv", cabecalho, [
            {"doc_id": "200", "CNPJ_FDO": BOVA, "DT_COMPT": "24/09/2026", "VL_QUOTA": "", "PATRIM_LIQ": "1000500,00",
             "NR_COTST": "5001"}])

    @classmethod
    def _mensal(self):
        meses = [f"{2025 + (m - 1) // 12}{(m - 1) % 12 + 1:02d}" for m in range(8, 21)]  # 202508 a 202608
        for k, mes in enumerate(meses):
            ano, numero = int(mes[:4]), int(mes[4:])
            fim = (date(ano + numero // 12, numero % 12 + 1, 1) - timedelta(days=1)).isoformat()
            armazem.gravar_csv(self.config.etf_fnet_mensal_dir / f"{mes}.csv", ("cnpj", "data", "cota", "pl"), [
                {"cnpj": BOVA, "data": fim, "cota": str(100 + k), "pl": "1000"}])

    @classmethod
    def _cvm(self):
        pasta = universo.pasta_cvm_cadastro(self.config) / HOJE.isoformat()
        armazem.gravar_csv(pasta / "registro_classe_etf.csv",
                           ("ID_Registro_Fundo", "CNPJ_Classe", "Data_Constituicao", "Custodiante", "Auditor"), [
                               {"ID_Registro_Fundo": "1", "CNPJ_Classe": BOVA, "Data_Constituicao": "2025-04-23",
                                "Custodiante": "BANCO X", "Auditor": "AUDITOR Y"},
                               {"ID_Registro_Fundo": "2", "CNPJ_Classe": NOVO, "Data_Constituicao": "2026-06-01"}])
        armazem.gravar_csv(pasta / "registro_fundo_etf.csv",
                           ("ID_Registro_Fundo", "CNPJ_Fundo", "Data_Constituicao", "Gestor", "Administrador"), [
                               {"ID_Registro_Fundo": "1", "CNPJ_Fundo": BOVA, "Data_Constituicao": "2008-11-19",
                                "Gestor": "BLACKROCK BRASIL GESTORA", "Administrador": "BANCO X"},
                               {"ID_Registro_Fundo": "2", "CNPJ_Fundo": NOVO, "Data_Constituicao": "2026-06-01"}])

    @classmethod
    def _carteira(self):
        cabecalho = ("TP_FUNDO_CLASSE", "CNPJ_FUNDO_CLASSE", "VL_PATRIM_LIQ", "TP_APLIC", "CD_ATIVO", "DS_ATIVO",
                     "EMISSOR", "CD_SELIC", "DT_VENC", "VL_MERC_POS_FINAL", "PAIS")
        linhas = [
            ("10.406.511/0001-61", "Ações", "VALE3", "", "", "", "", "300"),
            ("10.406.511/0001-61", "Ações e outros TVM cedidos em empréstimo", "VALE3", "", "", "", "", "100"),
            ("10.406.511/0001-61", "Ações", "ITUB4", "", "", "", "", "200"),
            ("10.406.511/0001-61", "Disponibilidades", "", "", "", "", "", "50"),
            ("10.406.511/0001-61", "Valores a pagar", "", "", "", "", "", "-10"),
            ("22.222.222/0001-22", "Investimento no Exterior", "9000032", "", "ISHARES CORE SP500", "", "", "1000"),
            ("33.333.333/0001-33", "Títulos Públicos", "", "", "", "760199", "2035-05-15", "600"),
        ]
        registros = []
        for cnpj, aplic, codigo, desc, emissor, selic, venc, valor in linhas:
            pl = {"10.406.511/0001-61": "1000", "22.222.222/0001-22": "1000", "33.333.333/0001-33": "1000"}[cnpj]
            registros.append({"TP_FUNDO_CLASSE": "CLASSES FIIM", "CNPJ_FUNDO_CLASSE": cnpj, "VL_PATRIM_LIQ": pl,
                              "TP_APLIC": aplic, "CD_ATIVO": codigo, "DS_ATIVO": desc, "EMISSOR": emissor, "CD_SELIC": selic,
                              "DT_VENC": venc, "VL_MERC_POS_FINAL": valor,
                              "PAIS": "ESTADOS UNIDOS DA AMÉRICA" if cnpj.startswith("22") else ""})
        armazem.gravar_csv(pasta_cda(self.config) / "cda_fie_etf_202608.csv", cabecalho, registros)

    @classmethod
    def _ibovespa(self):
        # Nivel = preco do BOVA11 * 10: beta 1 contra o proprio indice.
        por_ano: dict[int, list[dict]] = {}
        for i, dia in enumerate(self.dias):
            linhas = por_ano.setdefault(dia.year, [{"day": d} for d in range(1, 32)])
            linhas[dia.day - 1][f"rateValue{dia.month}"] = f"{_preco_bova(i) * 10:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        for ano, linhas in por_ano.items():
            armazem.gravar_json(pasta_indices(self.config) / "serie" / "IBOV" / f"{ano}.json",
                                {"indice": "IBOV", "ano": ano, "resposta": {"results": linhas}})

    def _linhas(self):
        c = consolidacao.consolidar(self.config, HOJE)
        return c, {l["ticker"]: l for l in c.linhas}


class TestConsolidacao(_Mundo):
    def test_uma_linha_por_etf_na_ordem_da_situacao(self):
        c, _ = self._linhas()
        self.assertEqual([l["ticker"] for l in c.linhas], ["BOVA11", "NOVO11", "RENF11", "XBIO11"])
        self.assertEqual(c.linhas[-1]["situacao"], "cancelado (só no cadastro da B3)")
        self.assertEqual(len({col.chave for col in c.colunas}), len(c.colunas))

    def test_pdf_fica_pendente_ate_alguem_conferir(self):
        _, linhas = self._linhas()
        bova = linhas["BOVA11"]
        for coluna in ("taxa_administracao", "taxa_maxima", "taxa_fundo_investido", "custos_extras", "taxa_faixa_pl",
                       "custo_total", "taxa_performance", "receita_aluguel", "indice_referencia", "publico_alvo",
                       "politica_distribuicao"):
            self.assertEqual(bova[coluna], consolidacao.PENDENTE_PDF, coluna)
        # Renda fixa: PMRC, duration e a aliquota dependem do regulamento.
        renf = linhas["RENF11"]
        self.assertEqual((renf["pmrc"], renf["duration"], renf["tributacao"]), (consolidacao.PENDENTE_PDF,) * 3)
        self.assertEqual((bova["pmrc"], bova["tributacao"]), (consolidacao.NAO_SE_APLICA, "15%"))
        # Uma pessoa confere a taxa no regulamento: a planilha passa a mostrar o valor.
        caminho = manual.caminho(self.config, "regulamento.csv")
        self.addCleanup(caminho.unlink)  # o mundo e da classe: nao pode vazar para os outros testes
        armazem.gravar_csv(caminho,
                           ("chave", "campo", "valor_conferido", "conferido_por"),
                           [{"chave": BOVA, "campo": "taxa_administracao", "valor_conferido": "0,10%", "conferido_por": "Ana"},
                            {"chave": BOVA, "campo": "taxa_maxima", "valor_conferido": "0,25%", "conferido_por": ""}])
        _, linhas = self._linhas()
        self.assertEqual(linhas["BOVA11"]["taxa_administracao"], "0,10%")
        self.assertEqual(linhas["BOVA11"]["taxa_maxima"], consolidacao.PENDENTE_PDF)  # sem quem conferiu, nao vale

    def test_liquidez_de_63_pregoes_sem_lote_de_bloco(self):
        _, linhas = self._linhas()
        bova = linhas["BOVA11"]
        janela = list(range(len(self.dias) - 63, len(self.dias)))
        self.assertEqual(bova["pregoes_janela"], 63)
        self.assertAlmostEqual(bova["volume_mediano"], sorted(1000 + i for i in janela)[31])
        self.assertEqual((bova["presenca"], bova["elegivel_presenca"]), (1.0, "sim"))
        # NOVO11 estreou 20 pregoes antes do fim: a janela e desde a estreia e o dia sem negocio conta zero.
        novo = linhas["NOVO11"]
        self.assertEqual(novo["pregoes_janela"], 20)
        self.assertAlmostEqual(novo["presenca"], 10 / 20)
        self.assertEqual(novo["elegivel_presenca"], "não")
        self.assertIn(novo["volume_mediano"], (0.0, 100.0, 200.0))

    def test_rentabilidade_volatilidade_queda_e_beta(self):
        c, linhas = self._linhas()
        bova = linhas["BOVA11"]
        precos = {d: _preco_bova(i) for i, d in enumerate(self.dias)}
        inicio, fim = c.janelas.doze_meses
        self.assertAlmostEqual(bova["rent_12m"], precos[fim] / precos[inicio] - 1)
        base_do_ano = max(d for d in self.dias if d.year == 2025)
        self.assertAlmostEqual(bova["rent_ano"], precos[fim] / precos[base_do_ano] - 1)
        self.assertGreater(bova["volatilidade_12m"], 0)
        self.assertLess(bova["drawdown_12m"], 0)
        self.assertAlmostEqual(bova["beta_12m"], 1.0, places=6)  # o IBOV e o proprio BOVA11 * 10
        # ETF novo: sem 12 meses, com rentabilidade desde a estreia; S&P 500 nao tem beta contra o IBOV.
        novo = linhas["NOVO11"]
        self.assertIsNone(novo.get("rent_12m"))
        self.assertEqual(novo["inicio_serie"], min(d for i, d in enumerate(self.dias) if d >= self.estreia and i % 2 == 0))
        self.assertAlmostEqual(novo["rent_inicio"], 0.0)
        self.assertIsNone(novo.get("beta_12m"))

    def test_pregao_sem_arquivo_completo_nao_vira_retorno_de_dois_dias(self):
        # Consolidado da B3 sem o mercado a vista: o dia fica fora do disco, mas
        # segue sendo pregao - o retorno que atravessa ele nao e diario.
        c, _ = self._linhas()
        inicio, fim = c.janelas.doze_meses
        i = len(self.dias) - 100
        arquivo = b3_arquivos.pasta_negocios(self.config) / f"{self.dias[i].isoformat()}.csv"
        guardado = arquivo.with_suffix(".guardado")
        arquivo.rename(guardado)
        self.addCleanup(guardado.rename, arquivo)
        incompletos = b3_arquivos.caminho_incompletos(self.config)
        armazem.gravar_json(incompletos, {self.dias[i].isoformat(): "sem mercado a vista"})
        self.addCleanup(incompletos.unlink)
        _, linhas = self._linhas()
        precos = [(d, _preco_bova(k)) for k, d in enumerate(self.dias) if inicio <= d <= fim and k != i]
        retornos = [p2 / p1 - 1 for (d1, p1), (_, p2) in zip(precos, precos[1:]) if d1 != self.dias[i - 1]]
        self.assertAlmostEqual(linhas["BOVA11"]["volatilidade_12m"], consolidacao.volatilidade_anual(retornos))

    def test_cota_indice_e_diferenca_na_mesma_janela(self):
        c, linhas = self._linhas()
        bova = linhas["BOVA11"]
        self.assertEqual(c.janelas.cota_12m, ("202508", "202608"))
        self.assertAlmostEqual(bova["rent_cota_12m"], 112 / 100 - 1)
        self.assertEqual(bova["codigo_indice"], "IBOV")
        self.assertIsNotNone(bova["retorno_indice_12m"])
        self.assertAlmostEqual(bova["tracking_difference"], (bova["rent_cota_12m"] - bova["retorno_indice_12m"]) * 100)

    def test_premio_no_dia_mais_recente_com_preco_e_cota(self):
        _, linhas = self._linhas()
        bova = linhas["BOVA11"]
        # 24/09 tem PL mas nao cota: o premio e o de 23/09, e a cota exibida tambem.
        self.assertEqual(bova["data_premio"], date(2026, 9, 23))
        self.assertAlmostEqual(bova["premio"], 1 / 1.01 - 1, places=4)
        self.assertEqual(bova["data_fechamento"], date(2026, 9, 24))
        self.assertEqual(bova["captacao_liquida"], 600.0)  # 3 informes de +200 nos ultimos 30 dias

    def test_carteira_emprestimo_soma_e_caixa_fica_de_fora(self):
        _, linhas = self._linhas()
        bova = linhas["BOVA11"]
        self.assertEqual(bova["n_ativos"], 2)
        self.assertEqual(bova["maiores_posicoes"], "VALE3 40,0%; ITUB4 20,0%")
        self.assertAlmostEqual(bova["top10"], 0.6)
        self.assertEqual(linhas["NOVO11"]["maiores_posicoes"], "ISHARES CORE SP500 100,0%")
        self.assertEqual(linhas["NOVO11"]["paises"], "ESTADOS UNIDOS DA AMÉRICA 100%")
        self.assertEqual(linhas["RENF11"]["maiores_posicoes"], "NTN-B 15/05/2035 60,0%")

    def test_idade_pela_data_mais_antiga_e_fundo_novo(self):
        _, linhas = self._linhas()
        # A B3 reiniciou o inicio do BOVA11 (16/07/2026) e a CVM a constituicao da classe: vale a do fundo, 2008.
        bova = linhas["BOVA11"]
        self.assertEqual(bova["constituicao"], date(2008, 11, 19))
        self.assertEqual(bova["fundo_novo"], "não")
        self.assertEqual(linhas["NOVO11"]["fundo_novo"], "sim")
        self.assertEqual((linhas["XBIO11"]["situacao"], linhas["XBIO11"].get("inicio_negociacao")),
                         ("cancelado (só no cadastro da B3)", None))


class TestEstreia(unittest.TestCase):
    def _insumos(self, pregoes, constituicao=None):
        class Falso:
            pass
        ins = Falso()
        ins.pregoes = pregoes
        ins.cvm = {"1": {"constituicao": constituicao}}
        return ins

    def test_data_reiniciada_nao_e_estreia(self):
        pregoes = [date(2026, 1, 5), date(2026, 7, 16), date(2026, 7, 17)]
        etf = {"cnpj": "1", "inicio_negociacao": "2026-07-16"}
        # Negociava antes da data: a B3 reiniciou.
        self.assertIsNone(consolidacao._estreia(self._insumos(pregoes), etf, [(date(2026, 1, 5), 10.0)]))
        # Constituicao anos antes da data: tambem reiniciada.
        self.assertIsNone(consolidacao._estreia(self._insumos(pregoes, date(2019, 4, 10)), etf, [(date(2026, 7, 16), 10.0)]))
        # Estreia de verdade.
        self.assertEqual(consolidacao._estreia(self._insumos(pregoes, date(2026, 5, 1)), etf, [(date(2026, 7, 16), 10.0)]),
                         date(2026, 7, 16))


class TestPlanilha(_Mundo):
    def test_uma_aba_de_dados_com_pdf_destacado_e_csv_igual(self):
        c = consolidacao.consolidar(self.config, HOJE)
        caminho = escrever_planilha(c, self.config.reports_dir / "etf-consolidado.xlsx")
        livro = load_workbook(caminho)
        self.assertEqual(livro.sheetnames, ["ETFs", "Dicionário", "Fontes"])
        aba = livro["ETFs"]
        cabecalho = [celula.value for celula in aba[6]]
        self.assertEqual(cabecalho, [col.titulo for col in c.colunas])
        self.assertEqual(aba.max_row, 6 + len(c.linhas))
        coluna_taxa = cabecalho.index("Taxa de administração") + 1
        celula = aba.cell(row=7, column=coluna_taxa)
        self.assertEqual(celula.value, consolidacao.PENDENTE_PDF)
        self.assertEqual(celula.fill.fgColor.rgb[-6:], "FFF2CC")
        coluna_rent = next(i for i, t in enumerate(cabecalho, start=1) if t.startswith("Rentabilidade 12 meses"))
        self.assertEqual(aba.cell(row=7, column=coluna_rent).number_format, "0.00%")
        self.assertIn("24/09/2026", cabecalho[coluna_rent - 1])  # a janela vai no cabecalho
        csv = armazem.ler_csv(escrever_csv(c, self.config.reports_dir / "etf_consolidado.csv"))
        self.assertEqual(len(csv), len(c.linhas))
        self.assertEqual(csv[0]["Taxa de administração"], consolidacao.PENDENTE_PDF)
        self.assertEqual(csv[0]["Data do fechamento"], "2026-09-24")


if __name__ == "__main__":
    unittest.main()


class TestDesdobramento(unittest.TestCase):
    def test_salto_perto_de_razao_simples_e_evento_e_movimento_forte_nao(self):
        d = [date(2026, 1, 20), date(2026, 1, 21), date(2026, 1, 22), date(2026, 1, 23), date(2026, 1, 26)]
        precos = [(d[0], 400.0), (d[1], 400.49), (d[2], 49.77), (d[3], 35.0), (d[4], 350.0)]
        eventos = consolidacao.eventos_de_cota(precos)
        # 400,49 -> 49,77 e 1:8 (SPXI11 em 22/01/2026); 49,77 -> 35 (-30%, cripto) nao e evento; 35 -> 350 e 10:1.
        self.assertEqual([(e[0], e[2]) for e in eventos], [(d[2], "desdobramento 1:8"), (d[4], "grupamento 10:1")])
        self.assertAlmostEqual(consolidacao.fator_ate(eventos, d[1]), 1 / 8 * 10)
        self.assertAlmostEqual(consolidacao.fator_ate(eventos, d[3]), 10)
        self.assertEqual(consolidacao.fator_ate(eventos, d[4]), 1.0)


def _planilha(caminho, cabecalho, linhas):
    from openpyxl import Workbook

    livro = Workbook()
    aba = livro.active
    aba.title = "comdinheiro"
    aba.append(cabecalho)
    for linha in linhas:
        aba.append(linha)
    livro.save(caminho)
    return caminho


class TestComparacao(_Mundo):
    def _exportacoes(self):
        from pathlib import Path

        from shelf_monitor.market.etf_bruto import comparacao

        pasta = Path(self.tmp.name)
        cadastral = _planilha(pasta / "cad.xlsx", [
            "Identificador", "Razão Social do Fundo", "CNPJ do Fundo", "Data da Constituição do Fundo", "ISIN",
            "Razão Social do Administrador", "Razão Social do Auditor", "Razão Social do Custodiante",
            "Razão Social do Gestor", "Público Alvo", "É para investidores qualificados?", "Taxa de Administracao",
            "Taxa de Administracao Máxima", "Taxa de Performance", "Taxa de Administração Financeira", "DY(12m,,)",
        ], [
            ["BOVA11", "ISHARES IBOVESPA", "10.406.511/0001-61", "19/11/2008", "BRBOVACTF003", "BANCO X", "AUDITOR Y",
             "BANCO X", "BLACKROCK BRASIL GESTORA", "Investidores em geral.", "NÃO", 0.1, None, 0, None, 0],
            # Ticker antigo do NOVO11: mesmo CNPJ, outra linha, sem preco.
            ["OLDX11", "TREND ETF NOVO", "22.222.222/0001-22", "01/06/2026", "BROLDXCTF001", "", "", "", "", "", "", None,
             None, None, None, 0],
        ])
        analise = self.dias[-2]
        i = len(self.dias) - 2
        delta = 250
        mercado = _planilha(pasta / "mer.xlsx", [
            "Empresas", "TICKER", "Data da Análise", "Preço de Fechamento", "Volume Negociado", "NEGOCIOS_DIA",
            "Retorno 12 meses", "Delta du 12m Data Análise", "Delta du 03m Data Análise", "Retorno 3 meses",
            "Delta du 03m Data Análise", "Retorno IBOV 12 meses", "Volatilidade Anualizada 12 meses", "Beta 12 meses",
            "DY_12M", "VAR95",
        ], [
            ["BOVA11", "BOVA11", analise.strftime("%d/%m/%Y"), _preco_bova(i), (1000 + i) / 1e9, 100,
             (_preco_bova(i) / _preco_bova(i - delta) - 1) * 100, delta, 64,
             (_preco_bova(i) / _preco_bova(i - 64) - 1) * 100, 1254,
             (_preco_bova(i) / _preco_bova(i - delta) - 1) * 100, 12.0, 1.0, 0, 18.8],
        ])
        return comparacao, comparacao.ler_exportacoes([cadastral, mercado])

    def test_mesma_fonte_bate_e_ticker_antigo_e_explicado(self):
        armazem.gravar_csv(manual.caminho(self.config, "regulamento.csv"), ("chave", "campo", "candidato"),
                           [{"chave": BOVA, "campo": "taxa_administracao", "candidato": "0,10%"}])
        self.addCleanup(manual.caminho(self.config, "regulamento.csv").unlink)
        comparacao, exportacao = self._exportacoes()
        resultado = comparacao.comparar(self.config, exportacao, HOJE)
        por_campo = {d.campo: d for d in resultado.diferencas if d.ticker == "BOVA11"}
        for campo in ("CNPJ", "ISIN", "Gestor", "Administrador", "Custodiante", "Auditor", "Constituição",
                      "Preço de fechamento", "Volume do dia", "Negócios do dia", "Retorno 12 meses",
                      "Retorno 3 meses", "Retorno do Ibovespa 12 meses", "Beta 12 meses", "Taxa de administração"):
            self.assertEqual(por_campo[campo].situacao, comparacao.IGUAL, campo)
        # A coluna "Delta du 03m" repetida (a segunda e a de 60 meses): vale a primeira.
        self.assertAlmostEqual(por_campo["Retorno 3 meses"].diferenca, 0.0, places=6)
        self.assertEqual(por_campo["Taxa de administração"].observacao, "sugestão do PDF")
        self.assertEqual(por_campo["DY 12 meses"].situacao, comparacao.SO_COMDINHEIRO)
        self.assertEqual(por_campo["Público-alvo"].situacao, comparacao.SO_COMDINHEIRO)
        self.assertEqual(resultado.so_comdinheiro,
                         [("OLDX11", "ticker antigo: o mesmo CNPJ negocia hoje como NOVO11")])
        self.assertIn("RENF11", [t for t, _ in resultado.so_captura])

    def test_calendario_anbima_conta_24_e_31_de_dezembro(self):
        from shelf_monitor.market.etf_bruto.comparacao import calendario_anbima

        pregoes = [date(2025, 12, 22), date(2025, 12, 23), date(2025, 12, 26), date(2025, 12, 29),
                   date(2025, 12, 30), date(2026, 1, 2)]
        calendario = calendario_anbima(pregoes)
        self.assertIn(date(2025, 12, 24), calendario)
        self.assertIn(date(2025, 12, 31), calendario)
        self.assertNotIn(date(2025, 12, 25), calendario)

    def test_planilha_da_comparacao(self):
        from shelf_monitor.report import etf_comparacao

        comparacao, exportacao = self._exportacoes()
        resultado = comparacao.comparar(self.config, exportacao, HOJE)
        caminho = etf_comparacao.escrever(resultado, self.config.reports_dir / "comparacao.xlsx")
        livro = load_workbook(caminho)
        self.assertEqual(livro.sheetnames, ["Resumo", "Detalhe", "Cobertura", "Fontes"])
        self.assertEqual(livro["Cobertura"]["A2"].value, "OLDX11")
