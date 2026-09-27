"""Captacao de ETF pelo informe diario do Fundos.NET: captacao, resgates e
diferenca de PL.

Os testes travam:

- a leitura da pagina do informe (`exibirDocumento`), que traz o mes inteiro
  do fundo - com os numeros reais do BUENA VISTA DEX VETTAFI NEOS BOOSTED
  BITCOIN HIGH INCOME em 09/2026;
- a soma do mes, dia a dia: o declarado quando fecha com a variacao de cotas;
  a diferenca de PL quando nao fecha (o IT NOW Ibovespa de 14/09/2026, que
  declarou zero com 600 mil cotas a menos; resgate declarado com as cotas
  paradas); o dia de desdobramento com o declarado; a estreia contra o PL;
- o bloco com o informe: as colunas de cada mes, a identidade liquida =
  captacao - resgates + ajuste, o ETF sem a pagina do mes e as superficies;
- o coletor: uma pagina por ETF por mes, so o que falta, e a parada quando o
  site deixa de responder o informe.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import csv
import os
import unittest
from datetime import date, datetime
from unittest import mock

from shelf_monitor.market import etf_captacao_mensal as mensal
from shelf_monitor.market import etf_sources
from shelf_monitor.market.etf_captacao_mensal import (
    build_captacao_mensal,
    com_informe,
    fluxo_do_mes,
    fluxo_pelo_informe,
)
from shelf_monitor.market.fundosnet import (
    InformeDiarioEtf,
    parse_informe_do_mes_html,
    read_informes_csv,
    write_informes_csv,
)

from fixtures_comuns import temp_config
from test_etf_captacao_mensal import MI, _cenario, _informe

# A pagina do informe como o Fundos.NET a serve (os tres primeiros dias de
# 09/2026 do documento 1324334), com uma linha cancelada e um dia repetido.
PAGINA = """<html><body>
<h2 align="center">Informe Di&aacute;rio</h2>
<table align="center" width="70%">
 <tr><td><span class="titulo-dado">Compet&ecirc;ncia: </span></td>
     <td colspan="3"><span class="dado-cabecalho">01/09/2026</span></td></tr>
 <tr><td><span class="titulo-dado">Nome do Fundo/Classe: </span></td>
     <td><span class="dado-cabecalho">BUENA VISTA DEX VETTAFI NEOS BOOSTED BITCOIN HIGH INCOME FUNDO DE
     &Iacute;NDICE DE RESPONSABILIDADE LIMITAD</span></td>
     <td><span class="titulo-dado">CNPJ: </span></td><td><span class="dado-cabecalho">65.602.949/0001-07</span></td></tr>
 <tr><td><span class="titulo-dado">Administrador: </span></td>
     <td><span class="dado-cabecalho">HEDGE INVESTMENTS DTVM LTDA.</span></td>
     <td><span class="titulo-dado">CNPJ: </span></td><td><span class="dado-cabecalho">07.253.654/0001-76</span></td></tr>
</table>
<table align="center" width="70%">
 <tr><td><span class="titulo-dado">Dia</span></td><td><span class="titulo-dado">Quota (R$)</span></td>
     <td><span class="titulo-dado">Capta&ccedil;&atilde;o no dia (R$)</span></td>
     <td><span class="titulo-dado">Resgate no dia (R$)</span></td>
     <td><span class="titulo-dado">Patrim&ocirc;nio Liquido (R$)</span></td>
     <td><span class="titulo-dado">Total de Carteira (R$)</span></td>
     <td><span class="titulo-dado">N.&ordm; Total de Cotistas</span></td>
     <td><span class="titulo-dado">Data da pr&oacute;xima informa&ccedil;&atilde;o do PL</span></td></tr>
 <tr><td>01</td><td>93,502570070000</td><td>0,00</td><td>0,00</td><td>50491387,84</td>
     <td>50543947,62</td><td>1882</td><td>02/09/2026</td></tr>
 <tr><td>02</td><td>1,000000000000</td><td>0,00</td><td>0,00</td><td>1,00</td>
     <td>1,00</td><td>1</td><td>03/09/2026</td></tr>
 <tr><td>02</td><td>89,690502850000</td><td>0,00</td><td>0,00</td><td>48432871,54</td>
     <td>48489139,57</td><td>1894</td><td>03/09/2026</td></tr>
 <tr><td>03</td><td>95,871273070000</td><td>3834850,92</td><td>0,00</td><td>55605338,39</td>
     <td>55763330,24</td><td>1915</td><td>04/09/2026</td></tr>
 <tr><td>04 C</td><td>99,000000000000</td><td>9999999,99</td><td>0,00</td><td>60000000,00</td>
     <td>60000000,00</td><td>1920</td><td>05/09/2026</td></tr>
</table>
<p>C = Cancelado.</p>
</body></html>"""


def _dia(dia: date, cotas: float, cota: float, captacao: float = 0.0, resgates: float = 0.0) -> InformeDiarioEtf:
    """Um dia do informe com PL = cotas x cota."""
    return InformeDiarioEtf(cnpj="X", data=dia, pl=cotas * cota, cota=cota, captacao=captacao, resgates=resgates)


class TestPaginaDoMes(unittest.TestCase):
    def test_le_o_mes_inteiro_do_fundo(self):
        dias = parse_informe_do_mes_html(PAGINA, doc_id=1324334)
        self.assertEqual([d.data for d in dias], [date(2026, 9, 1), date(2026, 9, 2), date(2026, 9, 3)])
        terceiro = dias[2]
        self.assertEqual(terceiro.cnpj, "65602949000107")
        self.assertEqual(
            (terceiro.cota, terceiro.captacao, terceiro.resgates, terceiro.pl, terceiro.vl_total, terceiro.cotistas),
            (95.87127307, 3834850.92, 0.0, 55605338.39, 55763330.24, 1915),
        )
        self.assertEqual(terceiro.doc_id, 1324334)
        self.assertTrue(terceiro.fundo.startswith("BUENA VISTA DEX VETTAFI"))
        # Dia repetido fica com a ultima linha; a linha cancelada fica de fora.
        self.assertEqual(dias[1].pl, 48432871.54)

    def test_pagina_que_nao_e_o_informe_da_lista_vazia(self):
        self.assertEqual(parse_informe_do_mes_html("<html><body>Just a moment...</body></html>"), [])
        self.assertEqual(parse_informe_do_mes_html(""), [])

    def test_o_declarado_da_pagina_fecha_com_a_variacao_de_cotas(self):
        """03/09/2026: 40 mil cotas novas a R$ 95,87 - o que o administrador declarou."""
        dias = parse_informe_do_mes_html(PAGINA)
        soma = fluxo_pelo_informe(dias[1:], vespera=dias[0])
        self.assertAlmostEqual(soma.liquida, 3834850.92, places=0)
        self.assertEqual((soma.captacao, soma.resgates, soma.dias, soma.dias_pelo_pl), (3834850.92, 0.0, 2, 0))


class TestFluxoPeloInforme(unittest.TestCase):
    def test_declarado_zero_com_as_cotas_mudando_vale_a_diferenca_de_pl(self):
        """IT NOW Ibovespa, 14/09/2026: 39.350.264 -> 38.750.264 cotas, nada declarado."""
        vespera = _dia(date(2026, 9, 11), 39_350_264, 193.312707)
        soma = fluxo_pelo_informe([_dia(date(2026, 9, 14), 38_750_264, 191.554891)], vespera)
        self.assertAlmostEqual(soma.liquida / MI, -114.93, places=2)
        self.assertEqual((soma.captacao, soma.resgates, soma.dias_pelo_pl), (0.0, 0.0, 1))

    def test_resgate_declarado_com_as_cotas_paradas_nao_e_saida(self):
        """O padrao da Buena Vista em 16/09/2026: resgate declarado, cotas iguais."""
        vespera = _dia(date(2026, 9, 15), 1_000_000, 10.0)
        soma = fluxo_pelo_informe([_dia(date(2026, 9, 16), 1_000_000, 9.8, resgates=200_000.0)], vespera)
        self.assertAlmostEqual(soma.liquida, 0.0)
        self.assertEqual((soma.resgates, soma.dias_pelo_pl), (200_000.0, 1))

    def test_folga_de_cinco_por_cento_fica_com_o_declarado(self):
        """IT NOW Bitcoin, 15/09/2026: R$ 438 mil declarados, R$ 421 mil pelas cotas."""
        vespera = _dia(date(2026, 9, 14), 100_000, 10.0)
        soma = fluxo_pelo_informe([_dia(date(2026, 9, 15), 142_100, 10.0, captacao=437_548.64)], vespera)
        self.assertEqual((soma.liquida, soma.dias_pelo_pl), (437_548.64, 0))

    def test_dia_de_desdobramento_fica_com_o_declarado(self):
        vespera = _dia(date(2026, 1, 14), 3_420_000, 411.29)
        # Cota dividida por 8 e 8 vezes mais cotas, com R$ 5 mi de captacao no dia.
        soma = fluxo_pelo_informe([_dia(date(2026, 1, 15), 27_460_000, 51.41, captacao=5 * MI)], vespera)
        self.assertEqual((soma.liquida, soma.dias_pelo_pl), (5 * MI, 0))

    def test_estreia_confere_o_primeiro_dia_contra_o_pl(self):
        """B-Index LFT, 17/10/2025: estreou com R$ 355 mi e nada declarado."""
        sem_declarar = fluxo_pelo_informe([_dia(date(2025, 10, 17), 3_550_000, 100.0)], None, estreia=True)
        self.assertEqual((sem_declarar.liquida, sem_declarar.dias_pelo_pl), (355 * MI, 1))
        declarando = fluxo_pelo_informe(
            [_dia(date(2025, 10, 17), 3_550_000, 100.0, captacao=355 * MI)], None, estreia=True
        )
        self.assertEqual((declarando.liquida, declarando.dias_pelo_pl), (355 * MI, 0))
        # ETF que ja existia, sem o fechamento anterior: o primeiro dia fica com o declarado.
        sem_vespera = fluxo_pelo_informe([_dia(date(2026, 1, 5), 3_550_000, 100.0, captacao=10 * MI)], None)
        self.assertEqual(sem_vespera.liquida, 10 * MI)

    def test_dia_sem_pl_entra_no_trecho_do_dia_seguinte(self):
        """O dinheiro do dia sem PL aparece nas cotas do dia seguinte: conta uma vez."""
        vespera = _dia(date(2026, 3, 2), 1_000_000, 10.0)
        buraco = InformeDiarioEtf(cnpj="X", data=date(2026, 3, 3), captacao=1 * MI)
        soma = fluxo_pelo_informe([buraco, _dia(date(2026, 3, 4), 1_100_000, 10.0)], vespera)
        self.assertEqual((soma.liquida, soma.captacao, soma.dias_pelo_pl), (1 * MI, 1 * MI, 0))
        # Sem nada declarado no trecho, vale a variacao de cotas desde o dia 2.
        sem_declarar = fluxo_pelo_informe(
            [InformeDiarioEtf(cnpj="X", data=date(2026, 3, 3)), _dia(date(2026, 3, 4), 1_100_000, 10.0)],
            vespera,
        )
        self.assertEqual((sem_declarar.liquida, sem_declarar.dias_pelo_pl), (1 * MI, 1))

    def test_com_informe_guarda_a_estreia_e_o_resto_vira_informe(self):
        estreia = fluxo_do_mes("202601", None, _informe("E", "202601", 600 * MI, 10.0), False)
        dias = [_dia(date(2026, 1, 20), 50 * MI, 10.0, captacao=500 * MI),
                _dia(date(2026, 1, 31), 60 * MI, 10.0, captacao=100 * MI)]
        f = com_informe(estreia, dias, None, estreia=True)
        self.assertEqual((f.origem, f.fonte, f.fluxo, f.captacao), ("estreia", "informe", 600 * MI, 600 * MI))
        self.assertEqual(f.diferenca_pl, 600 * MI)
        sem_informe = fluxo_do_mes("202601", None, _informe("F", "202601", 130 * MI, 10.0), True)
        self.assertEqual(com_informe(sem_informe, [], None, estreia=False), sem_informe)
        resolvido = com_informe(sem_informe, [_dia(date(2026, 1, 31), 13 * MI, 10.0, captacao=5 * MI)], None, False)
        self.assertEqual((resolvido.origem, resolvido.fluxo), ("informe", 5 * MI))


def _informes_do_cenario() -> dict:
    """Informe do mes de tres ETFs do cenario, fechando com os fechamentos."""
    return {
        # XP Ouro (A) em 12/2025: 110 -> 130 mi de cotas a R$ 10, declarado inteiro.
        "202512": {"A": [
            _dia(date(2025, 12, 15), 125 * MI, 10.0, captacao=150 * MI),
            _dia(date(2025, 12, 31), 130 * MI, 10.0, captacao=50 * MI),
        ]},
        "202601": {
            # Itau (D): 230 -> 260 mi de cotas, nada declarado.
            "D": [_dia(date(2026, 1, 15), 245 * MI, 10.0), _dia(date(2026, 1, 31), 260 * MI, 10.0)],
            # Estreia (E), declarada desde o primeiro dia.
            "E": [_dia(date(2026, 1, 20), 50 * MI, 10.0, captacao=500 * MI),
                  _dia(date(2026, 1, 31), 60 * MI, 10.0, captacao=100 * MI)],
        },
    }


class TestBlocoComInforme(unittest.TestCase):
    def setUp(self):
        self.catalogo, fechamentos, balancete = _cenario()
        with mock.patch.object(mensal, "MESES_DA_VISAO", 3):
            self.bloco = build_captacao_mensal(
                self.catalogo, fechamentos, balancete, informes=_informes_do_cenario()
            )
        self.etfs = {e["cnpj"]: e for e in self.bloco["etfs"]}
        self.config, self.tmp = temp_config()

    def tearDown(self):
        self.tmp.cleanup()

    def test_colunas_do_mes_e_a_identidade_da_liquida(self):
        b = self.bloco
        self.assertEqual(b["metodo"], "informe")
        detalhe = b["mercado"]["detalhe"]
        # Mesma liquida do cenario: o informe fecha com os fechamentos.
        self.assertEqual({m: round(v / MI) for m, v in b["mercado"]["fluxo"].items()},
                         {"202512": -130, "202601": 930, "202602": 65})
        self.assertAlmostEqual(detalhe["captacao"]["202601"], 600 * MI)
        self.assertAlmostEqual(detalhe["ajuste"]["202601"], 300 * MI)  # o Itau, que nao declarou
        # Os ETFs sem a pagina do mes (A, B, C, F em 01/2026): a liquida deles pelos fechamentos.
        self.assertAlmostEqual(detalhe["sem_informe"]["202601"], 30 * MI)
        for m in b["meses"]:
            partes = [detalhe[c][m] or 0.0 for c in ("captacao", "ajuste", "sem_informe")]
            self.assertAlmostEqual(b["mercado"]["fluxo"][m], sum(partes) - (detalhe["resgates"][m] or 0.0))
        # Mes em que nenhum ETF tem o informe: sem captacao declarada, nao zero.
        self.assertIsNone(detalhe["captacao"]["202602"])
        self.assertEqual(detalhe["com_informe"], {"202512": 1, "202601": 2, "202602": 0})
        self.assertAlmostEqual(b["mercado"]["historico"]["202601"]["estreias"], 600 * MI)

    def test_cada_etf_com_captacao_resgates_e_diferenca_de_pl(self):
        ouro, itau = self.etfs["A"], self.etfs["D"]
        self.assertAlmostEqual(ouro["detalhe"]["captacao"]["202512"], 200 * MI)
        self.assertAlmostEqual(ouro["detalhe"]["diferenca_pl"]["202512"], 200 * MI)
        self.assertEqual(ouro["mes_a_mes"]["202512"]["fonte"], "informe")
        self.assertEqual(itau["mes_a_mes"]["202601"]["dias_pelo_pl"], 2)
        self.assertAlmostEqual(itau["detalhe"]["ajuste"]["202601"], 300 * MI)
        # Mes com informe em disco sem a pagina deste ETF: vale o fechamento, e isso aparece.
        self.assertEqual(self.etfs["C"]["origem"]["202601"], mensal.ORIGENS["fechamentos"])
        cobertura = self.bloco["cobertura"]
        self.assertEqual(cobertura["dias_pelo_pl"], 2)
        self.assertEqual([a["gestora"] for a in cobertura["ajuste_por_gestora"]][:1], ["Itaú Asset (It Now)"])

    def test_leitura_prazos_e_relatorio(self):
        from shelf_monitor.report.etf_captacao import render_captacao_html, render_captacao_markdown

        self.assertTrue(any(f.startswith("No informe diário do Fundos.NET, em 12 meses: captação de")
                            for f in self.bloco["leitura"]))
        self.assertTrue(mensal.prazos_captacao(self.bloco)[0].startswith("captação: informe diário do Fundos.NET"))
        md = render_captacao_markdown(self.bloco, datetime(2026, 9, 25, 10, 0))
        # A coluna dos ETFs sem a pagina do mes so aparece porque ha algum.
        self.assertIn(
            "| Mês | Captação | Resgates | Ajuste pela diferença de PL | Sem o informe do mês | "
            "Captação líquida | Diferença de PL | Das quais estreias | Estado | ETFs com informe |", md
        )
        self.assertIn("| jan/2026 | R$ 600 mi | R$ 0 mi | +R$ 300 mi | +R$ 30 mi | +R$ 930 mi |", md)
        self.assertIn("Onde o ajuste pesa em 12 meses: Itaú Asset (It Now), +R$ 300 mi", md)
        html = render_captacao_html(self.bloco, datetime(2026, 9, 25, 10, 0))
        self.assertIn("Captação declarada", html)
        self.assertIn("Ajuste pela diferença de PL", html)
        self.assertIn("Sem o informe do mês (fechamentos)", html)

    def test_planilha_e_csv(self):
        from openpyxl import load_workbook

        from shelf_monitor.report.etf_report import export_captacao_csv, export_etf_excel

        caminho = export_etf_excel([], [], self.catalogo, self.config.reports_dir / "etf.xlsx", captacao=self.bloco)
        livro = load_workbook(caminho)
        self.assertEqual(livro.sheetnames[4], "Captação - informe mês a mês")
        linhas = [[c.value for c in linha] for linha in livro["Captação - informe mês a mês"].iter_rows()]
        do_itau = next(l for l in linhas if l[4] == "IT NOW IBOVESPA FUNDO DE ÍNDICE" and l[5] == "jan/26")
        # Captacao, resgates, ajuste, liquida, PL inicio e fim, diferenca e valorizacao.
        self.assertEqual(do_itau[6:14], [0, 0, 300 * MI, 300 * MI, 2300 * MI, 2600 * MI, 300 * MI, 0])
        rotulos = [l[1] for l in (
            [c.value for c in linha] for linha in livro["Captação - Competição"].iter_rows()
        )]
        self.assertIn("Ajuste pela diferença de PL", rotulos)

        csv_path = export_captacao_csv(self.bloco, self.config.reports_dir / "captacao.csv")
        with csv_path.open(encoding="utf-8-sig", newline="") as handle:
            registros = list(csv.DictReader(handle, delimiter=";"))
        itau = next(r for r in registros if r["cnpj"] == "D" and r["competencia"] == "202601")
        self.assertEqual((itau["origem"], itau["captacao_declarada"], itau["ajuste_diferenca_pl"]),
                         ("informe", "0.0", "300000000.0"))
        self.assertEqual(itau["dias_pela_diferenca_pl"], "2")


class TestFluxoSemanalPeloFundosNet(unittest.TestCase):
    """O fluxo semanal pela mesma conta: declarado, conferido contra a diferenca de PL."""

    def _serie_do_itau(self) -> list[InformeDiarioEtf]:
        return [
            InformeDiarioEtf(cnpj="D", data=date(2026, 9, 11), pl=39_350_264 * 193.312707,
                             cota=193.312707, captacao=0.0, resgates=0.0, cotistas=100),
            InformeDiarioEtf(cnpj="D", data=date(2026, 9, 14), pl=38_750_264 * 191.554891,
                             cota=191.554891, captacao=0.0, resgates=0.0, cotistas=100),
            InformeDiarioEtf(cnpj="D", data=date(2026, 9, 15), pl=38_750_264 * 192.587708,
                             cota=192.587708, captacao=0.0, resgates=0.0, cotistas=100),
        ]

    def test_registro_do_dia_leva_o_ajuste(self):
        from shelf_monitor.market.fundosnet import records_from_informes

        registros = records_from_informes(self._serie_do_itau())
        catorze = next(r for r in registros if r.dt_comptc == date(2026, 9, 14))
        self.assertEqual((catorze.captc_dia, catorze.resg_dia, catorze.flow_source), (0.0, 0.0, "diferenca_pl"))
        self.assertAlmostEqual(catorze.ajuste_dia / MI, -114.93, places=2)
        self.assertAlmostEqual(catorze.net_inflow, catorze.ajuste_dia)
        # O dia seguinte, sem mudanca de cotas, fica com o declarado (zero).
        self.assertEqual(registros[-1].ajuste_dia, 0.0)

    def test_semana_soma_captacao_resgates_e_ajuste(self):
        from shelf_monitor.market.etf_flow import aggregate_category_weekly_flows, aggregate_etf_weekly_flows
        from shelf_monitor.market.fundosnet import records_from_informes

        semanas = aggregate_etf_weekly_flows({}, records_from_informes(self._serie_do_itau()))
        semana = next(s for s in semanas if s.week_key == "2026-W38")
        self.assertAlmostEqual(semana.ajuste_pl / MI, -114.93, places=2)
        self.assertAlmostEqual(semana.captacao_liquida, semana.captacao_bruta - semana.resgates + semana.ajuste_pl)
        # A valorizacao sai do PL menos a liquida: a queda de cotas nao vira valorizacao.
        self.assertAlmostEqual(semana.valorizacao, semana.variacao_pl - semana.captacao_liquida)
        categoria = aggregate_category_weekly_flows(semanas)
        self.assertAlmostEqual(sum(c.ajuste_pl for c in categoria if c.week_key == "2026-W38"), semana.ajuste_pl)

    def test_competicao_le_o_fluxo_conferido(self):
        from shelf_monitor.market.etf_competicao import _cotistas_e_fluxo

        _cotistas, fluxo, de, ate = _cotistas_e_fluxo(self._serie_do_itau())
        self.assertAlmostEqual(fluxo["D"] / MI, -114.93, places=2)
        self.assertEqual((de, ate), (date(2026, 9, 11), date(2026, 9, 15)))


class TestLeituraDoFundosNetEmDisco(unittest.TestCase):
    def setUp(self):
        self.config, self.tmp = temp_config()

    def tearDown(self):
        self.tmp.cleanup()

    def test_pagina_do_mes_fechado_vence_a_parcial_e_o_xml(self):
        from shelf_monitor.market.fundosnet import informes_recentes, read_fnet_csvs

        dia, cnpj = date(2026, 8, 31), "11111111000111"
        write_informes_csv(self.config.etf_fnet_dir / "2026-09-01.csv", [
            InformeDiarioEtf(cnpj=cnpj, data=dia, pl=1.0, cota=1.0),
            InformeDiarioEtf(cnpj=cnpj, data=date(2026, 9, 1), pl=5.0, cota=1.0),
        ])
        write_informes_csv(self.config.etf_fnet_diario_dir / "202608-parcial.csv",
                           [InformeDiarioEtf(cnpj=cnpj, data=dia, pl=2.0, cota=1.0)])
        write_informes_csv(self.config.etf_fnet_diario_dir / "202608.csv",
                           [InformeDiarioEtf(cnpj=cnpj, data=dia, pl=3.0, cota=1.0)])
        lidos = [(i.cnpj, i.data, i.pl) for i in read_fnet_csvs(self.config)]
        self.assertEqual(sorted(lidos), [(cnpj, dia, 3.0), (cnpj, date(2026, 9, 1), 5.0)])
        self.assertEqual([i.data for i in informes_recentes(self.config, datas=1)], [date(2026, 9, 1)])


class _ClienteDePaginas:
    def __init__(self, paginas: dict[int, str], docs: list | None = None):
        self.paginas = paginas
        self.docs = docs or []
        self.pedidas: list[int] = []

    def fund_types(self):
        return {"ETF": 3}

    def search(self, *_args, **_kwargs):
        return iter(self.docs)

    def exibir(self, doc_id: int) -> str:
        self.pedidas.append(doc_id)
        return self.paginas.get(doc_id, "<html>fora do ar</html>")


class TestColetaDoInformeDoMes(unittest.TestCase):
    def setUp(self):
        self.config, self.tmp = temp_config()
        write_informes_csv(self.config.etf_fnet_mensal_dir / "202609.csv", [
            InformeDiarioEtf(cnpj="65602949000107", data=date(2026, 9, 30), pl=1.0, cota=1.0, doc_id=1324334),
            InformeDiarioEtf(cnpj="11111111000111", data=date(2026, 9, 30), pl=1.0, cota=1.0, doc_id=777),
        ])

    def tearDown(self):
        self.tmp.cleanup()

    def _coletar(self, cliente, **kwargs):
        with mock.patch("shelf_monitor.market.etf_sources.fundosnet.FundosNetClient", return_value=cliente), \
                mock.patch.object(etf_sources, "_PAUSA_ENTRE_PAGINAS", 0):
            return etf_sources.collect_informes_do_mes(self.config, {}, meses=["202609"], **kwargs)

    def test_uma_pagina_por_etf_e_depois_so_o_que_faltou(self):
        cliente = _ClienteDePaginas({1324334: PAGINA})
        resumo = self._coletar(cliente)
        self.assertEqual((resumo["paginas"], resumo["falhas"]), (2, 1))
        self.assertEqual(resumo["meses"], {"202609": "1 de 2 ETFs"})
        gravados = read_informes_csv(self.config.etf_fnet_diario_dir / "202609.csv")
        self.assertEqual({(i.cnpj, i.data.day) for i in gravados},
                         {("65602949000107", 1), ("65602949000107", 2), ("65602949000107", 3)})
        # Na proxima vez, so o ETF cuja pagina falhou.
        cliente.pedidas.clear()
        self._coletar(cliente)
        self.assertEqual(cliente.pedidas, [777])

    def test_para_depois_de_falhas_seguidas(self):
        cliente = _ClienteDePaginas({})
        with mock.patch.object(etf_sources, "_FALHAS_SEGUIDAS", 1):
            resumo = self._coletar(cliente)
        self.assertEqual(cliente.pedidas, [1324334])
        self.assertIn("a coleta parou", resumo["interrompida"])
        self.assertFalse((self.config.etf_fnet_diario_dir / "202609.csv").exists())


class _DocDaBusca:
    def __init__(self, doc_id: int, ref: date, fundo: str):
        self.id, self.data_referencia, self.fundo = doc_id, ref, fundo
        self.cnpj_fundo = None


class TestMesEmAberto(unittest.TestCase):
    def setUp(self):
        self.config, self.tmp = temp_config()

    def tearDown(self):
        self.tmp.cleanup()

    def test_meses_depois_do_ultimo_fechado(self):
        self.assertEqual(etf_sources.meses_abertos(date(2026, 9, 25)), ["202609"])
        # No comeco do mes, o anterior ainda esta sendo entregue.
        self.assertEqual(etf_sources.meses_abertos(date(2026, 10, 3)), ["202609", "202610"])

    def test_pagina_do_documento_mais_recente_de_cada_etf_uma_vez_por_dia(self):
        from shelf_monitor.market.etf import EtfMetadata

        catalogo = {"65602949000107": EtfMetadata(cnpj="65602949000107", name="BUENA VISTA DEX", kind="etf")}
        nome = "BUENA VISTA DEX VETTAFI NEOS BOOSTED BITCOIN HIGH INCOME"
        cliente = _ClienteDePaginas({1324334: PAGINA}, docs=[
            _DocDaBusca(1324300, date(2026, 9, 17), nome),
            _DocDaBusca(1324334, date(2026, 9, 18), nome),   # o mais recente: a pagina dele
            _DocDaBusca(1300000, date(2026, 8, 31), nome),   # mes fechado: fora
        ])
        with mock.patch("shelf_monitor.market.etf_sources.fundosnet.FundosNetClient", return_value=cliente), \
                mock.patch.object(etf_sources, "_PAUSA_ENTRE_PAGINAS", 0):
            resumo = etf_sources.collect_informes_do_mes_aberto(self.config, catalogo, day=date(2026, 9, 25))
            self.assertEqual(cliente.pedidas, [1324334])
            self.assertEqual(resumo["meses"], {"202609": "1 ETF, até 03/09/2026"})
            parcial = self.config.etf_fnet_diario_dir / "202609-parcial.csv"
            gravados = read_informes_csv(parcial)
            self.assertEqual({i.cnpj for i in gravados}, {"65602949000107"})
            # Gravado no mesmo dia: nao busca de novo.
            momento = datetime(2026, 9, 25, 12, 0).timestamp()
            os.utime(parcial, (momento, momento))
            de_novo = etf_sources.collect_informes_do_mes_aberto(self.config, catalogo, day=date(2026, 9, 25))
        self.assertEqual(cliente.pedidas, [1324334])
        self.assertEqual(de_novo["meses"], {"202609": "em disco (hoje)"})


if __name__ == "__main__":
    unittest.main()
