"""Testes das listagens publicas: extracao do DOM, paginacao e diagnostico.

Os sites reais estao bloqueados neste ambiente (e duas das tres listas barram
cliente HTTP mesmo fora dele), entao aqui se exercita o que e determinavel sem
rede: o HTML renderizado vira fundos, a paginacao nao trunca em silencio, e uma
falha de TLS e classificada em vez de virar "falha de TLS" e mais nada.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import unittest

from shelf_monitor.collectors import open_finance as of
from shelf_monitor.collectors.browser import (
    extract_from_dom,
    looks_like_fund_name,
    to_offer,
)
from shelf_monitor.collectors.inter import _SITEMAP_URLS
from shelf_monitor.collectors.xp import XpCollector
from shelf_monitor.config import PLATFORMS, Config
from shelf_monitor.diagnose import _tls_cause
from shelf_monitor.models import AssetClass, CollectorStrategy, SnapshotStatus

# --- HTML de exemplo ---------------------------------------------------------

# Formato da lista de rentabilidade do Itau: tabela com colunas de retorno.
TABELA_RENTABILIDADE = """
<html><body>
<table>
  <thead><tr><th>Fundo</th><th>Classe</th><th>Aplicação</th><th>Taxa</th><th>12 meses</th></tr></thead>
  <tbody>
    <tr>
      <td>Itau Renda Fixa Referenciado DI FIC</td><td>Renda Fixa</td>
      <td>R$ 1,00</td><td>0,30%</td><td>10,45%</td>
    </tr>
    <tr>
      <td>Kinea Chronos Multimercado FIC FIM</td><td>Multimercado</td>
      <td>R$ 5.000,00</td><td>1,90%</td><td>13,20%</td>
    </tr>
    <tr>
      <td>Itau Ações Dividendos FIA</td><td>Ações</td>
      <td>R$ 500,00</td><td>2,00%</td><td>-4,10%</td>
    </tr>
  </tbody>
</table>
</body></html>
"""

# Formato de vitrine de produtos: cartoes repetidos com a mesma classe CSS.
CARTOES = """
<html><body>
<nav><a href="/fundos">Fundos de investimento</a><a href="/simule">Simule agora</a></nav>
<div class="shelf">
  <div class="fund-card">
    <h3 class="fund-card__name">BTG Pactual Digital Renda Fixa FI</h3>
    <span class="fund-card__category">Renda Fixa</span>
    <p>Aplicação mínima R$ 500,00</p>
    <a href="/fundos-de-investimento/btg-rf">Ver detalhes</a>
  </div>
  <div class="fund-card">
    <h3 class="fund-card__name">BTG Pactual Absoluto Institucional FIC FIA</h3>
    <span class="fund-card__category">Ações</span>
    <p>Aplicação mínima R$ 25.000,00</p>
    <a href="/fundos-de-investimento/btg-absoluto">Ver detalhes</a>
  </div>
  <div class="fund-card">
    <h3 class="fund-card__name">BTG Pactual Cambial Dólar FIC FIM</h3>
    <span class="fund-card__category">Cambial</span>
    <p>Aplicação mínima R$ 1.000,00</p>
    <a href="/fundos-de-investimento/btg-cambial">Ver detalhes</a>
  </div>
</div>
</body></html>
"""

# Nem tabela nem cartao: o CNPJ no texto e a unica ancora.
CNPJ_SOLTO = """
<html><body>
<section>
  <article>
    <h2>Inter Conservador Renda Fixa FI</h2>
    <p>CNPJ: 12.345.678/0001-90</p>
  </article>
  <article>
    <h2>Inter Selection Multimercado FIC</h2>
    <p>CNPJ: 98.765.432/0001-10</p>
  </article>
</section>
</body></html>
"""

# Uma mencao isolada num rodape nao e prateleira.
RODAPE = """
<html><body>
<main><h1>Fundos de investimento</h1><p>Conheca nossas opções.</p></main>
<footer>
  <p>Inter Conservador Renda Fixa FI - CNPJ: 12.345.678/0001-90</p>
</footer>
</body></html>
"""


class TestFundNameHeuristic(unittest.TestCase):
    def test_accepts_real_fund_names(self):
        for name in (
            "Itau Renda Fixa Referenciado DI FIC",
            "BTG Pactual Absoluto Institucional FIC FIA",
            "XP Selection Multimercado FIC FIM",
        ):
            self.assertTrue(looks_like_fund_name(name), name)

    def test_rejects_navigation_labels(self):
        """Rotulo de menu tem token de fundo mas nao e fundo."""
        for label in (
            "Fundos de investimento",
            "Renda Fixa",
            "Simule agora",
            "Ver todos os fundos",
            "Fundos",
        ):
            self.assertFalse(looks_like_fund_name(label), label)

    def test_cnpj_alongside_relaxes_the_criterion(self):
        """Com CNPJ na mesma caixa, o nome nao precisa trazer token conhecido."""
        self.assertFalse(looks_like_fund_name("Nu Reserva Imediata"))
        self.assertTrue(looks_like_fund_name("Nu Reserva Imediata", has_cnpj=True))

    def test_a_cnpj_is_not_a_name(self):
        """Sem isso o proprio numero viraria o nome do fundo."""
        self.assertFalse(looks_like_fund_name("CNPJ: 12.345.678/0001-90", has_cnpj=True))


class TestJsonExtraction(unittest.TestCase):
    def test_reads_the_cvm_code_however_the_feed_spells_it(self):
        """O codigo costuma vir como numero, e a grafia varia por casa."""
        for key, raw in (
            ("codigoCVM", 555001),
            ("cd_cvm", "0555001"),
            ("codCvm", "555001"),
        ):
            item = {"nomeFundo": "Fundo A", "cnpj": "12.345.678/0001-91", key: raw}
            fund = to_offer("itau", item)
            self.assertEqual(fund.cvm_code, "555001", key)

    def test_a_fund_without_cvm_code_keeps_it_empty(self):
        fund = to_offer("itau", {"nomeFundo": "Fundo A", "cnpj": "12345678000191"})
        self.assertIsNone(fund.cvm_code)


class TestDomExtraction(unittest.TestCase):
    def test_reads_a_rentability_table(self):
        offers = extract_from_dom("itau", TABELA_RENTABILIDADE, "https://www.itau.com.br/x")
        names = [o.name for o in offers]
        self.assertEqual(len(offers), 3, names)
        self.assertIn("Kinea Chronos Multimercado FIC FIM", names)

        kinea = next(o for o in offers if o.name.startswith("Kinea"))
        self.assertEqual(kinea.raw_category, "Multimercado")
        self.assertEqual(kinea.asset_class, AssetClass.MULTIMERCADO)
        self.assertEqual(kinea.min_initial, 5000.0)

    def test_header_row_is_not_a_fund(self):
        offers = extract_from_dom("itau", TABELA_RENTABILIDADE)
        self.assertNotIn("Fundo", [o.name for o in offers])

    def test_percent_columns_are_not_read_as_fees(self):
        """Numa lista de rentabilidade quase todo "%" e retorno, nao taxa.

        Confundir os dois e pior que nao ter o dado: o dashboard passaria a
        ranquear custo por uma coluna de desempenho.
        """
        offers = extract_from_dom("itau", TABELA_RENTABILIDADE)
        self.assertTrue(all(o.admin_fee_pct is None for o in offers))

    def test_reads_repeated_cards(self):
        offers = extract_from_dom("btg", CARTOES, "https://investimentos.btgpactual.com/x")
        names = [o.name for o in offers]
        self.assertEqual(len(offers), 3, names)

        absoluto = next(o for o in offers if "Absoluto" in o.name)
        self.assertEqual(absoluto.raw_category, "Ações")
        self.assertEqual(absoluto.min_initial, 25000.0)
        self.assertEqual(
            absoluto.detail_url,
            "https://investimentos.btgpactual.com/fundos-de-investimento/btg-absoluto",
        )

    def test_navigation_links_do_not_become_funds(self):
        names = [o.name for o in extract_from_dom("btg", CARTOES)]
        self.assertNotIn("Fundos de investimento", names)
        self.assertNotIn("Simule agora", names)

    def test_reads_from_cnpj_in_the_text(self):
        """Sem tabela e sem cartao repetido, o CNPJ e a ancora."""
        offers = extract_from_dom("inter", CNPJ_SOLTO)
        self.assertEqual(len(offers), 2, [o.name for o in offers])
        self.assertEqual(
            {o.cnpj for o in offers}, {"12345678000190", "98765432000110"}
        )
        self.assertEqual(
            {o.name for o in offers},
            {"Inter Conservador Renda Fixa FI", "Inter Selection Multimercado FIC"},
        )

    def test_a_lone_mention_is_not_a_shelf(self):
        """Um fundo citado no rodape nao pode virar a prateleira do dia."""
        self.assertEqual(extract_from_dom("inter", RODAPE), [])

    def test_empty_page_yields_nothing(self):
        self.assertEqual(extract_from_dom("btg", "<html><body></body></html>"), [])


# Cartoes com CNPJ: cada fundo tem de sair uma vez, nao duas.
CARTOES_COM_CNPJ = """
<html><body>
<div class="shelf">
  <div class="fund-card">
    <h3 class="fund-card__name">BTG Pactual Digital Renda Fixa FI</h3>
    <p>CNPJ: 11.111.111/0001-11 - aplicação mínima R$ 500,00</p>
  </div>
  <div class="fund-card">
    <h3 class="fund-card__name">BTG Pactual Absoluto FIC FIA</h3>
    <p>CNPJ: 22.222.222/0001-22 - aplicação mínima R$ 25.000,00</p>
  </div>
  <div class="fund-card">
    <h3 class="fund-card__name">BTG Pactual Cambial Dólar FIC FIM</h3>
    <p>CNPJ: 33.333.333/0001-33 - aplicação mínima R$ 1.000,00</p>
  </div>
</div>
</body></html>
"""

# Master e feeder: nomes quase iguais, CNPJs diferentes. Sao dois fundos.
MASTER_E_FEEDER = """
<html><body><table><tbody>
  <tr><td>Kinea Chronos Master FIM</td><td>Multimercado</td><td>11.111.111/0001-11</td></tr>
  <tr><td>Kinea Chronos FIC FIM</td><td>Multimercado</td><td>22.222.222/0001-22</td></tr>
  <tr><td>Kinea Atlas Master FIA</td><td>Ações</td><td>33.333.333/0001-33</td></tr>
</tbody></table></body></html>
"""


# Duas mencoes soltas: passam do minimo, entao servem de caso para a limpeza
# do nome (uma so seria descartada antes de chegar la).
RODAPE_DUPLO = """
<html><body><footer>
  <p>Inter Conservador Renda Fixa FI - CNPJ: 12.345.678/0001-90</p>
  <p>Inter Selection Multimercado FIC - CNPJ: 98.765.432/0001-10</p>
</footer></body></html>
"""

class TestNoDoubleCounting(unittest.TestCase):
    """Um fundo lido por duas caixas da pagina continua sendo um fundo.

    A pagina renderizada expoe o mesmo dado em niveis aninhados (o cartao e o
    titulo dentro dele) e por vias diferentes (tabela e cartao). Sem consolidar,
    a prateleira do BTG sairia com o dobro do tamanho - e o diff do dia seguinte
    leria a correcao como uma saida em massa.
    """

    def test_a_card_and_its_heading_are_one_fund(self):
        offers = extract_from_dom("btg", CARTOES_COM_CNPJ)
        self.assertEqual(len(offers), 3, [(o.name, o.cnpj) for o in offers])
        self.assertEqual(
            {o.cnpj for o in offers},
            {"11111111000111", "22222222000122", "33333333000133"},
        )

    def test_the_richer_reading_wins(self):
        """Entre a leitura com CNPJ e minimo e a sem nada, fica a com dados."""
        offers = extract_from_dom("btg", CARTOES_COM_CNPJ)
        digital = next(o for o in offers if "Digital" in o.name)
        self.assertEqual(digital.cnpj, "11111111000111")
        self.assertEqual(digital.min_initial, 500.0)

    def test_master_and_feeder_are_not_merged(self):
        """Nomes quase iguais com CNPJs distintos sao fundos distintos."""
        offers = extract_from_dom("itau", MASTER_E_FEEDER)
        self.assertEqual(len(offers), 3, [(o.name, o.cnpj) for o in offers])
        chronos = [o for o in offers if "Chronos" in o.name]
        self.assertEqual(len(chronos), 2)
        self.assertEqual(
            {o.cnpj for o in chronos}, {"11111111000111", "22222222000122"}
        )

    def test_the_cnpj_does_not_leak_into_the_name(self):
        """Numa caixa sem marcacao o nome vem grudado no numero."""
        offers = extract_from_dom("inter", RODAPE_DUPLO)
        self.assertEqual(
            {o.name for o in offers},
            {"Inter Conservador Renda Fixa FI", "Inter Selection Multimercado FIC"},
        )



class TestOpenFinancePagination(unittest.TestCase):
    """A paginacao nao pode cortar a prateleira sem dizer."""

    def setUp(self):
        self.config = Config(data_dir="/tmp/nao-usado-nestes-testes")
        self.original_get_json = of.get_json
        self.original_max_pages = of._MAX_PAGES

    def tearDown(self):
        of.get_json = self.original_get_json
        of._MAX_PAGES = self.original_max_pages

    def _serve(self, total_pages: int, per_page: int, total_records: int):
        """Finge uma instituicao que devolve `per_page` itens por pagina."""
        self.requested: list[dict] = []

        def fake_get_json(session, url, config, params=None):
            self.requested.append(params or {})
            page = (params or {}).get("page", 1)
            return {
                "data": [
                    {
                        "name": f"Fundo {page}-{i}",
                        "cnpjNumber": f"{page:07d}{i:07d}",
                    }
                    for i in range(per_page)
                ],
                "meta": {"totalPages": total_pages, "totalRecords": total_records},
            }

        of.get_json = fake_get_json

    def _client(self) -> of.OpenFinanceFunds:
        platform = self.config.platform("btg")
        platform.open_finance_base_url = "https://exemplo.invalido"
        return of.OpenFinanceFunds(self.config, platform, session=None)

    def test_records_the_effective_page_size(self):
        """O servidor que ignora o page-size pedido tem de ficar registrado."""
        self._serve(total_pages=3, per_page=10, total_records=30)
        client = self._client()
        offers = client.fetch()

        self.assertEqual(self.requested[0]["page-size"], of._PAGE_SIZE)
        self.assertEqual(client.effective_page_size, 10)
        self.assertEqual(client.total_records, 30)
        self.assertEqual(client.pages_read, 3)
        self.assertEqual(len(offers), 30)
        self.assertFalse(client.truncated)
        self.assertEqual(client.warnings, [])

    def test_truncation_is_reported_not_silent(self):
        of._MAX_PAGES = 2
        self._serve(total_pages=20, per_page=10, total_records=200)
        client = self._client()
        offers = client.fetch()

        self.assertEqual(len(offers), 20, "parou no teto, como esperado")
        self.assertTrue(client.truncated)
        self.assertEqual(len(client.warnings), 1)
        self.assertIn("truncada", client.warnings[0])
        self.assertIn("200", client.warnings[0])

    def test_a_truncated_collection_is_partial_not_ok(self):
        """Meia prateleira nao pode virar a baseline do diff de amanha."""
        of._MAX_PAGES = 2
        self._serve(total_pages=20, per_page=10, total_records=200)

        collector = XpCollector(self.config, self._client().platform)
        collector._collect_playwright = lambda: []
        snapshot = collector.collect()

        self.assertEqual(snapshot.status, SnapshotStatus.PARTIAL)
        self.assertIn("truncada", snapshot.error)

    def test_missing_meta_midway_does_not_end_pagination(self):
        """A instituicao que omite `totalPages` numa pagina do meio nao pode
        encerrar o laco como se a paginacao tivesse acabado."""
        def fake_get_json(session, url, config, params=None):
            page = (params or {}).get("page", 1)
            meta = (
                {"totalPages": 4, "totalRecords": 40} if page == 1 else {}
            )
            return {
                "data": [
                    {"name": f"Fundo {page}-{i}", "cnpjNumber": f"{page:07d}{i:07d}"}
                    for i in range(10)
                ],
                "meta": meta,
            }

        of.get_json = fake_get_json
        client = self._client()
        offers = client.fetch()

        self.assertEqual(client.pages_read, 4)
        self.assertEqual(len(offers), 40)
        self.assertFalse(client.truncated)
        self.assertEqual(client.warnings, [])

    def test_pages_exhausted_before_declared_total_is_reported(self):
        """Itens acabaram antes do total anunciado: a prateleira nao e a inteira."""
        def fake_get_json(session, url, config, params=None):
            page = (params or {}).get("page", 1)
            data = (
                [
                    {"name": f"Fundo {page}-{i}", "cnpjNumber": f"{page:07d}{i:07d}"}
                    for i in range(10)
                ]
                if page < 4
                else []
            )
            return {"data": data, "meta": {"totalPages": 10, "totalRecords": 100}}

        of.get_json = fake_get_json
        client = self._client()
        offers = client.fetch()

        self.assertEqual(len(offers), 30)
        self.assertTrue(client.truncated)
        self.assertEqual(len(client.warnings), 1)
        self.assertIn("de 10", client.warnings[0])


class TestPlatformSources(unittest.TestCase):
    """As URLs que o usuario indicou, e a estrategia que as alcanca."""

    def test_btg_points_at_the_product_listing(self):
        btg = PLATFORMS["btg"]
        self.assertEqual(
            btg.shelf_url,
            "https://investimentos.btgpactual.com/fundos-de-investimento/produtos",
        )
        self.assertIn(CollectorStrategy.PLAYWRIGHT, btg.strategies)

    def test_itau_points_at_the_public_rentability_list(self):
        itau = PLATFORMS["itau"]
        self.assertEqual(
            itau.shelf_url, "https://www.itau.com.br/investimentos/fundos/rentabilidade"
        )
        self.assertNotIn("personnalite", itau.shelf_url)
        self.assertIn(CollectorStrategy.PLAYWRIGHT, itau.strategies)

    def test_pages_that_block_plain_http_do_not_try_html_scrape(self):
        """XP, BTG e Itau devolvem 403 ou casca vazia a cliente HTTP simples."""
        for platform_id in ("xp", "btg", "itau"):
            self.assertNotIn(
                CollectorStrategy.HTML_SCRAPE,
                PLATFORMS[platform_id].strategies,
                platform_id,
            )

    def test_open_finance_comes_first_everywhere(self):
        """E a fonte de maior confianca; o site acrescenta, nao substitui."""
        for platform_id, platform in PLATFORMS.items():
            self.assertIn(
                CollectorStrategy.OPEN_FINANCE, platform.strategies, platform_id
            )

    def test_inter_only_keeps_the_sitemap_that_exists(self):
        """/sitemap-0.xml devolve 404 - so poluia o log a cada coleta."""
        self.assertEqual(_SITEMAP_URLS, ["https://inter.co/sitemap.xml"])


class TestTlsDiagnosis(unittest.TestCase):
    """"Falha de TLS" nao diz o que corrigir; a causa diz."""

    def test_incomplete_chain(self):
        cause = _tls_cause(
            "certificate verify failed: unable to get local issuer certificate",
            {"certificados_apresentados": 1},
            "api.openbanking.bancointer.com.br",
        )
        self.assertIn("cadeia incompleta", cause)

    def test_unknown_issuer_with_full_chain_is_not_the_same_thing(self):
        cause = _tls_cause(
            "certificate verify failed: unable to get local issuer certificate",
            {"certificados_apresentados": 3},
            "exemplo.com.br",
        )
        self.assertIn("emissor desconhecido", cause)

    def test_expired_certificate(self):
        cause = _tls_cause(
            "certificate verify failed: certificate has expired",
            {"valido_ate": "2024-01-01T00:00:00+00:00"},
            "exemplo.com.br",
        )
        self.assertIn("vencido", cause)

    def test_wrong_host(self):
        cause = _tls_cause(
            "certificate verify failed: Hostname mismatch",
            {"nomes": ["outro.com.br"]},
            "exemplo.com.br",
        )
        self.assertIn("outro nome", cause)

    def test_intercepting_proxy(self):
        cause = _tls_cause(
            "certificate verify failed: self signed certificate in certificate chain",
            {"auto_assinado": True},
            "exemplo.com.br",
        )
        self.assertIn("auto-assinado", cause)


if __name__ == "__main__":
    unittest.main()


class TestBtgApi(unittest.TestCase):
    """O funds/list/v2 mapeado contra a resposta real capturada pelo diagnose.

    A fixture e um recorte literal do payload que a pagina /produtos recebeu na
    maquina do usuario (2026-08-31) - nao um exemplo inventado.
    """

    @classmethod
    def setUpClass(cls):
        import json
        from pathlib import Path

        fixture = Path(__file__).parent / "fixtures" / "btg" / "funds_list_v2.json"
        cls.payload = json.loads(fixture.read_text(encoding="utf-8"))

    def test_maps_the_rich_fields(self):
        from shelf_monitor.collectors.btg import offer_from_btg_item

        offer = offer_from_btg_item(self.payload["items"][0])
        self.assertEqual(offer.name, "A1 D30 FICFIRF LP CrPr RL")
        self.assertEqual(offer.cnpj, "63446494000152")
        self.assertEqual(offer.raw_category, "Renda Fixa")
        self.assertEqual(offer.asset_class, AssetClass.RENDA_FIXA)
        self.assertEqual(offer.min_initial, 1000.0)
        self.assertEqual(offer.admin_fee_pct, 0.9)
        self.assertEqual(offer.performance_fee.amount_pct, 20.0)
        self.assertEqual(offer.performance_fee.benchmark, "CDI")
        self.assertEqual(offer.redemption_quotation_days, 30)   # rescueQuota D+30
        self.assertEqual(offer.redemption_payment_days, 31)     # D+31
        self.assertEqual(offer.manager_name, "ASSET1 CREDITO INVESTIMENTOS LTDA")
        self.assertEqual(offer.taxation, "Longo Prazo")
        self.assertIs(offer.closed_for_investment, False)       # applyOpenPortal=True
        self.assertEqual(offer.extra["classificacao_anbima"], "RF duração livre crédito livre")

    def test_dplus_parsing(self):
        from shelf_monitor.collectors.btg import parse_dplus

        self.assertEqual(parse_dplus("D+30"), 30)
        self.assertEqual(parse_dplus("D+0"), 0)
        self.assertEqual(parse_dplus("d + 1"), 1)
        self.assertIsNone(parse_dplus(None))
        self.assertIsNone(parse_dplus("no ato"))

    def test_fee_in_reais_is_not_reported_as_percent(self):
        """administrationFeeType e o discriminador; fora de PERCENTAGE, nada."""
        import copy

        from shelf_monitor.collectors.btg import offer_from_btg_item

        item = copy.deepcopy(self.payload["items"][0])
        item["Detail"]["detail"]["administrationFeeType"] = "NOMINAL"
        item["Detail"]["detail"]["administrationFee"] = 5000
        self.assertIsNone(offer_from_btg_item(item).admin_fee_pct)

    def test_paginates_until_total_pages(self):
        from shelf_monitor.collectors.btg import paginate_btg_api

        pages: list[str] = []

        def fetch(url):
            pages.append(url)
            page = int(url.split("page=")[1].split("&")[0])
            items = [
                {"Name": f"Fundo {page}-{i} FIC FIM", "CNPJ": f"{page:07d}{i:07d}"}
                for i in range(3)
            ]
            return {"items": items, "total": 9, "page": page, "total_pages": 3}

        offers, urls = paginate_btg_api(fetch, "https://exemplo.invalido")
        self.assertEqual(len(pages), 3)
        self.assertEqual(len(offers), 9)
        self.assertEqual(len(urls), 1)
        self.assertIn("funds/list/v2", urls[0])

    def test_unexpected_envelope_fails_loud(self):
        from shelf_monitor.collectors.btg import paginate_btg_api
        from shelf_monitor.http import SourceUnavailable

        with self.assertRaises(SourceUnavailable):
            paginate_btg_api(
                lambda url: {"error": "maintenance"}, "https://exemplo.invalido"
            )

    def test_page_cap_overflow_is_reported(self):
        """Prateleira cortada pelo teto nao pode sair como coleta inteira."""
        from unittest import mock

        from shelf_monitor.collectors import btg as btg_mod
        from shelf_monitor.collectors.btg import paginate_btg_api

        def fetch(url):
            page = int(url.split("page=")[1].split("&")[0])
            return {
                "items": [{"Name": f"Fundo {page} FIC FIM", "CNPJ": f"{page:014d}"}],
                "total": 5,
                "page": page,
                "total_pages": 5,
            }

        warnings: list[str] = []
        with mock.patch.object(btg_mod, "_MAX_PAGES", 2):
            offers, _ = paginate_btg_api(
                fetch, "https://exemplo.invalido", warnings
            )

        self.assertEqual(len(offers), 2)
        self.assertEqual(len(warnings), 1)
        self.assertIn("truncada", warnings[0])

    def test_pagination_without_warnings_list_still_works(self):
        """O parametro e opcional: quem nao passa a lista nao quebra."""
        from shelf_monitor.collectors.btg import paginate_btg_api

        def fetch(url):
            return {"items": [{"Name": "Fundo A FIC FIM", "CNPJ": "1" * 14}],
                    "total": 1, "page": 1, "total_pages": 1}

        offers, _ = paginate_btg_api(fetch, "https://exemplo.invalido")
        self.assertEqual(len(offers), 1)


class TestInterCollector(unittest.TestCase):
    def test_cap_overflow_becomes_a_warning_not_silence(self):
        """574 fundos no sitemap com teto menor: avisa e o snapshot sai PARTIAL."""
        from unittest import mock

        from shelf_monitor.collectors import inter as inter_mod
        from shelf_monitor.collectors.inter import InterCollector

        config = Config(data_dir="/tmp/nao-usado")
        collector = InterCollector(config, config.platform("inter"))
        urls = [f"https://inter.co/pra-voce/investimentos/fundos-de-investimento/f{i}/" for i in range(5)]
        collector._discover_fund_urls = lambda: urls
        collector._parse_fund_page = lambda url: offer_for(url)
        collector._collect_open_finance = lambda: []

        def offer_for(url):
            from shelf_monitor.models import FundOffer
            return FundOffer(platform="inter", name=f"Fundo {url[-3:]} FIC FIM", detail_url=url)

        with mock.patch.object(inter_mod, "_MAX_FUNDS", 3):
            snapshot = collector.collect()

        self.assertEqual(len(snapshot.funds), 3)
        self.assertEqual(snapshot.status, SnapshotStatus.PARTIAL)
        self.assertIn("teto", snapshot.error)

    def _collector_over(self, url_count: int, failing: set[int],
                        page_data_fails: bool = False):
        """Coletor do Inter sobre `url_count` paginas, com `failing` caindo.

        Stubba o `get` (e nao `_parse_fund_page`) para os contadores rodarem de
        verdade - e o caminho que perdia fundo em silencio.
        """
        from unittest import mock

        from shelf_monitor.collectors import inter as inter_mod
        from shelf_monitor.collectors.inter import InterCollector
        from shelf_monitor.http import SourceUnavailable

        prefix = "https://inter.co/pra-voce/investimentos/fundos-de-investimento/"
        urls = [f"{prefix}f{i}/" for i in range(url_count)]

        class FakeResponse:
            def __init__(self, text):
                self.text = text

            def json(self):
                raise ValueError("sem json")

        def fake_get(session, url, config, **kwargs):
            if "/page-data/" in url:
                if page_data_fails:
                    raise SourceUnavailable("page-data fora do ar")
                return FakeResponse("{}")
            index = int(url.rstrip("/").rsplit("f", 1)[1])
            if index in failing:
                raise SourceUnavailable("HTTP 503")
            return FakeResponse(
                f"<html><h1>Inter Fundo {index} FIC FIM</h1>"
                f"<p>11.111.{index:03d}/0001-11</p></html>"
            )

        config = Config(data_dir="/tmp/nao-usado")
        collector = InterCollector(config, config.platform("inter"))
        collector._discover_fund_urls = lambda: urls
        collector._collect_open_finance = lambda: []
        return collector, mock.patch.object(inter_mod, "get", fake_get)

    def test_mass_page_failures_downgrade_to_partial(self):
        """Pagina que nao abre e fundo perdido: sem aviso, o diff de amanha le
        a perda como saida da prateleira."""
        collector, patched = self._collector_over(20, failing=set(range(10)))
        with patched:
            snapshot = collector.collect()

        self.assertEqual(len(snapshot.funds), 10)
        self.assertEqual(snapshot.status, SnapshotStatus.PARTIAL)
        self.assertIn("páginas de fundo falharam", snapshot.error)

    def test_single_page_failure_does_not_downgrade(self):
        """Uma falha isolada e ruido de rede, nao prateleira incompleta."""
        collector, patched = self._collector_over(20, failing={3})
        with patched:
            snapshot = collector.collect()

        self.assertEqual(len(snapshot.funds), 19)
        self.assertEqual(snapshot.status, SnapshotStatus.OK)
        self.assertEqual(collector.warnings, [])

    def test_page_data_mass_failure_is_a_note_not_partial(self):
        """A prateleira esta inteira; o que faltou foram os fatos."""
        collector, patched = self._collector_over(
            20, failing=set(), page_data_fails=True
        )
        with patched:
            snapshot = collector.collect()

        self.assertEqual(len(snapshot.funds), 20)
        self.assertEqual(snapshot.status, SnapshotStatus.OK)
        self.assertIn("page-data.json falhou", snapshot.error)

    def test_fill_missing_keeps_html_identity(self):
        """page-data completa o que falta; nome e CNPJ do HTML prevalecem."""
        from shelf_monitor.collectors.inter import _fill_missing
        from shelf_monitor.models import FundOffer

        html = FundOffer(platform="inter", name="Inter Conservador FIRF", cnpj="11111111000111")
        json_offer = FundOffer(
            platform="inter", name="outro nome", cnpj="99999999000199",
            min_initial=100.0, admin_fee_pct=0.5, manager_name="Inter Asset",
            asset_class=AssetClass.RENDA_FIXA,
        )
        merged = _fill_missing(html, json_offer)
        self.assertEqual(merged.name, "Inter Conservador FIRF")
        self.assertEqual(merged.cnpj, "11111111000111")
        self.assertEqual(merged.min_initial, 100.0)
        self.assertEqual(merged.admin_fee_pct, 0.5)
        self.assertEqual(merged.manager_name, "Inter Asset")
        self.assertEqual(merged.asset_class, AssetClass.RENDA_FIXA)


class TestTrustStore(unittest.TestCase):
    def test_context_carries_cas_and_keeps_verification(self):
        """O contexto soma stores; nunca desliga hostname nem cadeia."""
        import ssl

        from shelf_monitor.http import _trust_context

        context = _trust_context()
        if context is None:
            self.skipTest("ambiente sem CA nenhuma")
        self.assertTrue(context.cert_store_stats().get("x509_ca"))
        self.assertTrue(context.check_hostname)
        self.assertEqual(context.verify_mode, ssl.CERT_REQUIRED)


# Cabecalho e linhas copiados literalmente da tabela renderizada na maquina do
# usuario (data/cache/debug/itau/pagina.html, 2026-08-31) - incluindo o espaco
# duro (\xa0) do "R$ 1,00".
ITAU_HEADER = [
    "nome", "risco", "aplicação inicial",
    "12 meses referente ao período de 01/08/2025 a 31/07/2026",
    "no ano referente ao período de 01/01/2026 a 31/07/2026",
    "no mês referente ao período de 01/08/2026 a 28/08/2026",
    "taxa máxima do produto Referente a taxa máxima praticada",
    "resgate", "+ detalhes",
]
ITAU_PAGE_1 = [
    ITAU_HEADER,
    ["Diferenciado Crédito Privado Renda Fixa", "baixo risco", "R$\xa01,00",
     "14,46%", "9,04%", "1,01% (até 28/08)", "0,45%",
     "crédito em conta em 1 dia útil", "mais info"],
    ["Privilege DI", "baixo risco", "R$\xa01,00", "14,60%", "9,23%",
     "1,04% (até 28/08)", "0,30%", "crédito em conta no mesmo dia", "mais info"],
]
ITAU_PAGE_2 = [
    ITAU_HEADER,
    ["Bahia AM Mutá Seleção Multimercado", "médio risco", "R$\xa0500,00",
     "13,20%", "8,00%", "0,90% (até 28/08)", "1,90%",
     "crédito em conta em 30 dias úteis", "mais info"],
]


class TestItauRentabilidade(unittest.TestCase):
    """Parser da tabela de rentabilidade, contra as linhas reais."""

    def test_maps_by_header_not_position(self):
        from shelf_monitor.collectors.itau import offers_from_rentability_rows

        offers = offers_from_rentability_rows([ITAU_PAGE_1, ITAU_PAGE_2])
        self.assertEqual(len(offers), 3, [o.name for o in offers])

        diferenciado = offers[0]
        self.assertEqual(diferenciado.name, "Diferenciado Crédito Privado Renda Fixa")
        self.assertEqual(diferenciado.min_initial, 1.0)
        self.assertEqual(diferenciado.admin_fee_pct, 0.45)
        self.assertEqual(diferenciado.redemption_payment_days, 1)
        self.assertEqual(diferenciado.asset_class, AssetClass.RENDA_FIXA)
        self.assertEqual(diferenciado.extra["risco"], "baixo risco")

    def test_returns_do_not_leak_into_the_fee(self):
        """As colunas de 14,46%/9,04% sao retorno - a taxa e so a da coluna taxa."""
        from shelf_monitor.collectors.itau import offers_from_rentability_rows

        offers = offers_from_rentability_rows([ITAU_PAGE_1])
        self.assertEqual([o.admin_fee_pct for o in offers], [0.45, 0.30])

    def test_same_day_redemption_is_zero(self):
        from shelf_monitor.collectors.itau import offers_from_rentability_rows

        privilege = offers_from_rentability_rows([ITAU_PAGE_1])[1]
        self.assertEqual(privilege.name, "Privilege DI")
        self.assertEqual(privilege.redemption_payment_days, 0)

    def test_thirty_business_days(self):
        from shelf_monitor.collectors.itau import parse_redemption_days

        self.assertEqual(parse_redemption_days("crédito em conta em 30 dias úteis"), 30)
        self.assertIsNone(parse_redemption_days("consulte o regulamento"))

    def test_repeated_pages_do_not_duplicate(self):
        """Se o clique nao trocar a pagina, a mesma tabela nao vira fundos 2x."""
        from shelf_monitor.collectors.itau import offers_from_rentability_rows

        offers = offers_from_rentability_rows([ITAU_PAGE_1, ITAU_PAGE_1])
        self.assertEqual(len(offers), 2)

    def test_synthetic_payload_reaches_the_mapper(self):
        import tempfile
        from pathlib import Path

        from shelf_monitor.collectors.itau import LeituraDaTabela

        with tempfile.TemporaryDirectory() as sem_base:
            leitura = LeituraDaTabela(Path(sem_base))
            offers = leitura.mapear(
                "https://www.itau.com.br/x (tabela)", {"itau_rentabilidade": [ITAU_PAGE_1]}
            )
            self.assertEqual(len(offers), 2)
            # Payload que nao e a tabela segue para o extrator generico.
            self.assertIsNone(leitura.mapear("https://x", {"data": []}))


class TestMergeReconciliation(unittest.TestCase):
    """Fundo do site sem CNPJ nao pode duplicar o gemeo do Open Finance."""

    def test_site_fund_merges_into_open_finance_twin_by_slug(self):
        from shelf_monitor.merge import merge_offers

        groups = {
            "open-finance": [offer(
                "itau", "Privilege DI FIC FIRF", "11111111000111",
                manager_name="Itau Asset",
            )],
            "playwright": [offer(
                "itau", "Privilege DI", None, admin_fee_pct=0.30,
            )],
        }
        merged = merge_offers(groups)
        # Slugs diferentes ("privilege-di" so casa se o slug bater exato):
        # aqui "Privilege DI FIC FIRF" tem slug identico porque FIC/FIRF sao
        # tokens genericos descartados - e o caso real do Itau.
        self.assertEqual(len(merged), 1, [(o.name, o.cnpj) for o in merged])
        fund = merged[0]
        self.assertEqual(fund.cnpj, "11111111000111")
        self.assertEqual(fund.sources, ["open-finance", "playwright"])
        self.assertEqual(fund.admin_fee_pct, 0.30)
        self.assertIn("Privilege DI", fund.extra["name_variants"])

    def test_distinct_slugs_stay_apart(self):
        from shelf_monitor.merge import merge_offers

        groups = {
            "open-finance": [offer("itau", "Kinea Chronos FIC FIM", "11111111000111")],
            "playwright": [offer("itau", "Bahia AM Mutá Seleção", None)],
        }
        merged = merge_offers(groups)
        self.assertEqual(len(merged), 2)


def offer(platform: str, name: str, cnpj: str | None = None, **kwargs):
    from shelf_monitor.models import FundOffer

    return FundOffer(platform=platform, name=name, cnpj=cnpj, **kwargs)


class TestEnrichment(unittest.TestCase):
    """Join das prateleiras com a base oficial - a taxonomia unica."""

    def setUp(self):
        import shutil
        import tempfile
        from pathlib import Path

        self.tmp = tempfile.TemporaryDirectory()
        self.config = Config(data_dir=self.tmp.name)
        self.config.market_dir.mkdir(parents=True, exist_ok=True)
        fixtures = Path(__file__).parent / "fixtures" / "enrich"
        for name in ("mercado.csv", "subclasses.csv"):
            shutil.copy(fixtures / name, self.config.market_dir / name)

    def tearDown(self):
        self.tmp.cleanup()

    def _snapshot(self, funds):
        from datetime import datetime

        from shelf_monitor.models import CollectorStrategy, ShelfSnapshot, SnapshotStatus

        return ShelfSnapshot(
            platform="btg", collected_at=datetime(2026, 8, 31),
            strategy=CollectorStrategy.OPEN_FINANCE, status=SnapshotStatus.OK,
            funds=funds,
        )

    def test_official_classification_lands_on_the_offer(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="btg", name="A1 D30 FICFIRF", cnpj="63446494000152")
        stats = enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertEqual(fund.extra["classificacao_anbima"], "RF duração livre crédito livre")
        self.assertEqual(fund.extra["gestor_oficial"], "ASSET1 CREDITO INVESTIMENTOS LTDA")
        self.assertEqual(fund.extra["situacao_cvm"], "EM FUNCIONAMENTO NORMAL")
        self.assertEqual(fund.extra["pl_cvm"], 11584330.36)
        self.assertEqual(fund.manager_name, "ASSET1 CREDITO INVESTIMENTOS LTDA")
        self.assertEqual(stats.matched, 1)
        self.assertEqual(stats.anbima_filled, 1)

    def test_house_value_survives_divergence(self):
        """A classificacao que a casa publica nao se perde - divergencia e dado."""
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(
            platform="btg", name="A1 D30", cnpj="63446494000152",
            extra={"classificacao_anbima": "Renda Fixa Duração Livre"},
        )
        enrich_snapshots(self.config, [self._snapshot([fund])])
        self.assertEqual(fund.extra["classificacao_anbima"], "RF duração livre crédito livre")
        self.assertEqual(fund.extra["classificacao_anbima_casa"], "Renda Fixa Duração Livre")

    def test_legacy_fund_matches_and_fixes_asset_class(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import AssetClass, FundOffer

        fund = FundOffer(platform="btg", name="Fundo Legado", cnpj="99999999000199")
        self.assertEqual(fund.asset_class, AssetClass.OUTROS)
        enrich_snapshots(self.config, [self._snapshot([fund])])
        self.assertEqual(fund.asset_class, AssetClass.ACOES)
        self.assertEqual(fund.extra["base_cvm"], "nao_adaptado")

    def test_no_market_base_degrades_to_none(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        (self.config.market_dir / "mercado.csv").unlink()
        fund = FundOffer(platform="btg", name="Qualquer", cnpj="63446494000152")
        self.assertIsNone(enrich_snapshots(self.config, [self._snapshot([fund])]))
        self.assertNotIn("classificacao_anbima", fund.extra)

    def test_unmatched_cnpj_is_counted_not_invented(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="btg", name="Desconhecido", cnpj="55555555000155")
        stats = enrich_snapshots(self.config, [self._snapshot([fund])])
        self.assertEqual(stats.with_cnpj, 1)
        self.assertEqual(stats.matched, 0)
        self.assertNotIn("classificacao_anbima", fund.extra)

    def test_official_manager_wins_over_what_the_house_publishes(self):
        """Sem isso, a mesma gestora aparece com tres nomes e o total infla."""
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(
            platform="btg", name="A1 D30", cnpj="63446494000152",
            manager_name="Asset1", manager_cnpj="99999999000199",
        )
        stats = enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertEqual(fund.manager_name, "ASSET1 CREDITO INVESTIMENTOS LTDA")
        self.assertEqual(fund.manager_cnpj, "41287689000164")
        self.assertEqual(fund.extra["gestor_casa"], "Asset1")
        self.assertEqual(fund.extra["cnpj_gestor_casa"], "99999999000199")
        self.assertEqual(stats.managers_from_cvm, 1)

    def test_same_manager_written_differently_is_not_a_divergence(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(
            platform="btg", name="A1 D30", cnpj="63446494000152",
            manager_name="Asset1 Crédito Investimentos Ltda.",
        )
        enrich_snapshots(self.config, [self._snapshot([fund])])
        self.assertEqual(fund.manager_name, "ASSET1 CREDITO INVESTIMENTOS LTDA")
        self.assertNotIn("gestor_casa", fund.extra)

    def test_a_shelf_without_managers_is_filled_from_the_registry(self):
        """O sintoma que motivou o lookup: prateleira inteira sem gestora."""
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        funds = [
            FundOffer(platform="inter", name="A1 D30", cnpj="63446494000152"),
            FundOffer(platform="inter", name="Privilege DI", cnpj="11111111000111"),
            FundOffer(platform="inter", name="Legado", cnpj="99999999000199"),
        ]
        stats = enrich_snapshots(self.config, [self._snapshot(funds)])

        self.assertTrue(all(f.manager_name for f in funds))
        self.assertEqual(stats.managers_from_cvm, 3)


class TestMarketBaseBootstrap(TestEnrichment):
    """A base ausente era falha silenciosa: sem ela nao ha gestora nenhuma."""

    def test_existing_base_is_not_downloaded_again(self):
        from unittest import mock

        from shelf_monitor import main

        with mock.patch.object(main, "run_import") as run_import:
            main.ensure_market_base(self.config)
        run_import.assert_not_called()

    def test_missing_base_triggers_the_import(self):
        from unittest import mock

        from datetime import datetime

        from shelf_monitor import main
        from shelf_monitor.market.models import MarketDatabase

        (self.config.market_dir / "mercado.csv").unlink()
        empty = (MarketDatabase(imported_at=datetime.now()), {})
        with mock.patch.object(main, "run_import", return_value=empty) as run_import:
            main.ensure_market_base(self.config)
        run_import.assert_called_once()

    def test_a_failed_import_does_not_bring_the_run_down(self):
        """A coleta do dia vale mais que a base: o relatorio sai com o aviso."""
        from unittest import mock

        from shelf_monitor import main

        (self.config.market_dir / "mercado.csv").unlink()
        with mock.patch.object(main, "run_import", side_effect=OSError("sem rede")):
            main.ensure_market_base(self.config)  # nao pode levantar


class TestCvmCodeLookup(TestEnrichment):
    """Casamento pelo codigo CVM - a chave de quem nao publica CNPJ."""

    def test_matches_by_class_code(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="itau", name="A1 D30", cvm_code="C100")
        stats = enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertEqual(fund.manager_name, "ASSET1 CREDITO INVESTIMENTOS LTDA")
        self.assertEqual(fund.extra["match_cvm"], "codigo_classe")
        self.assertEqual(stats.matched_by_cvm_code, 1)
        self.assertEqual(stats.with_cvm_code, 1)
        self.assertEqual(stats.with_cnpj, 0)

    def test_matches_by_subclass_code_through_the_parent_class(self):
        """A subclasse partilha CNPJ e carteira: a resposta e a da classe-mae."""
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="itau", name="A1 D30 Varejo", cvm_code="S777")
        enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertEqual(fund.manager_name, "ASSET1 CREDITO INVESTIMENTOS LTDA")
        self.assertEqual(fund.extra["classificacao_anbima"], "RF duração livre crédito livre")
        self.assertEqual(fund.extra["match_cvm"], "codigo_subclasse")

    def test_matches_by_fund_code_of_a_multi_class_envelope(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="itau", name="Fundo Multi", cvm_code="400400")
        enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertEqual(fund.manager_name, "GESTORA MULTI LTDA")
        self.assertEqual(fund.extra["match_cvm"], "codigo_fundo")
        # As duas classes concordam, entao a classificacao desce ao envelope.
        self.assertEqual(fund.extra["classificacao_anbima"], "Multimercados Livre")

    def test_envelope_code_does_not_invent_a_classification(self):
        """Classes que discordam: gestora desce, classificacao nao."""
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="itau", name="Fundo Diverge", cvm_code="600600")
        enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertEqual(fund.manager_name, "GESTORA DIVERGE")
        self.assertNotIn("classificacao_anbima", fund.extra)

    def test_legacy_fund_code_also_matches(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="itau", name="Legado", cvm_code="999099")
        enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertEqual(fund.manager_name, "GESTORA LEGADA LTDA")
        self.assertEqual(fund.extra["match_cvm"], "codigo_fundo")

    def test_cnpj_is_tried_before_the_code(self):
        """Com as duas chaves, vale a mais confiavel."""
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(
            platform="itau", name="A1 D30", cnpj="63446494000152", cvm_code="C200",
        )
        stats = enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertEqual(fund.manager_name, "ASSET1 CREDITO INVESTIMENTOS LTDA")
        self.assertEqual(fund.extra["match_cvm"], "cnpj")
        self.assertEqual(stats.matched_by_cvm_code, 0)

    def test_the_most_specific_level_wins_a_code_collision(self):
        """COL1 e codigo de uma subclasse de C4 e de outra classe."""
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="itau", name="Colidido", cvm_code="COL1")
        enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertEqual(fund.manager_name, "GESTORA MULTI LTDA")
        self.assertEqual(fund.extra["match_cvm"], "codigo_subclasse")

    def test_code_normalization_crosses_the_join(self):
        """Zero a esquerda e numero cru sao o mesmo codigo dos dois lados."""
        from shelf_monitor.enrich import MarketLookup

        lookup = MarketLookup.load(self.config)
        for value in ("400400", "0400400", 400400, 400400.0):
            hit = lookup.get_by_cvm_code(value)
            self.assertIsNotNone(hit, repr(value))
            self.assertEqual(hit[0].gestor, "GESTORA MULTI LTDA")

    def test_unknown_code_is_not_forced_into_a_match(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="itau", name="Fora da base", cvm_code="Z999")
        stats = enrich_snapshots(self.config, [self._snapshot([fund])])
        self.assertEqual(stats.matched, 0)
        self.assertIsNone(fund.manager_name)

    def test_a_base_without_subclasses_still_loads(self):
        """Base importada antes desta versao nao tem o arquivo - nao pode quebrar."""
        from shelf_monitor.enrich import MarketLookup

        (self.config.market_dir / "subclasses.csv").unlink()
        lookup = MarketLookup.load(self.config)

        self.assertIsNone(lookup.get_by_cvm_code("S777"))
        self.assertIsNotNone(lookup.get_by_cvm_code("C100"))
        self.assertIsNotNone(lookup.get("63446494000152"))

    def test_a_base_without_the_code_columns_still_loads(self):
        """mercado.csv antigo: sem as colunas, o indice fica vazio e so."""
        import csv

        from shelf_monitor.enrich import MarketLookup

        path = self.config.market_dir / "mercado.csv"
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle, delimiter=";"))
        headers = [h for h in rows[0] if not h.startswith("codigo_cvm")]
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(
                handle, fieldnames=headers, delimiter=";", extrasaction="ignore"
            )
            writer.writeheader()
            writer.writerows(rows)

        lookup = MarketLookup.load(self.config)
        self.assertIsNone(lookup.get_by_cvm_code("C100"))
        self.assertIsNotNone(lookup.get("63446494000152"))


class TestListingBlock(unittest.TestCase):
    """A visao de vitrine nao pode se misturar com o catalogo do Open Finance."""

    def _report(self):
        from datetime import datetime

        from shelf_monitor.models import (
            CollectorStrategy, FundOffer, RunReport, ShelfSnapshot, SnapshotStatus,
        )

        def fund(platform, name, cnpj, sources, manager=None, fee=None):
            f = FundOffer(platform=platform, name=name, cnpj=cnpj,
                          manager_name=manager, admin_fee_pct=fee)
            f.sources = sources
            return f

        btg = ShelfSnapshot(
            platform="btg", collected_at=datetime(2026, 8, 31),
            strategy=CollectorStrategy.OPEN_FINANCE, status=SnapshotStatus.OK,
            funds=[
                fund("btg", "BTG Vitrine A", "11111111000111",
                     ["open-finance", "api-json"], manager="BTG PACTUAL ASSET"),
                fund("btg", "BTG Vitrine B", "22222222000122",
                     ["api-json"], manager="KAPITALO GESTAO"),
                fund("btg", "BTG Só Catálogo", "33333333000133", ["open-finance"]),
            ],
        )
        itau = ShelfSnapshot(
            platform="itau", collected_at=datetime(2026, 8, 31),
            strategy=CollectorStrategy.OPEN_FINANCE, status=SnapshotStatus.OK,
            funds=[
                fund("itau", "Compartilhado", "22222222000122", ["playwright"]),
            ],
        )
        return RunReport(generated_at=datetime(2026, 8, 31), snapshots=[btg, itau])

    def test_catalog_only_funds_stay_out_of_the_listing(self):
        from shelf_monitor.report.aggregate import build_listing_block

        block = build_listing_block(self._report())
        btg = next(p for p in block["platforms"] if p["platform_id"] == "btg")
        self.assertEqual(btg["total"], 2)
        self.assertEqual(btg["catalog_total"], 3)

    def test_own_vs_third_party_uses_the_watchlist(self):
        from shelf_monitor.report.aggregate import build_listing_block

        btg = next(
            p for p in build_listing_block(self._report())["platforms"]
            if p["platform_id"] == "btg"
        )
        self.assertEqual(btg["own_product"], 1)     # BTG PACTUAL ASSET
        self.assertEqual(btg["third_party"], 1)     # KAPITALO

    def test_exclusives_count_only_listing_presence(self):
        from shelf_monitor.report.aggregate import build_listing_block

        block = build_listing_block(self._report())
        btg = next(p for p in block["platforms"] if p["platform_id"] == "btg")
        itau = next(p for p in block["platforms"] if p["platform_id"] == "itau")
        # "BTG Vitrine A" so esta na vitrine do BTG; "Compartilhado" esta nas duas.
        self.assertEqual(btg["exclusive"], 1)
        self.assertEqual(itau["exclusive"], 0)
        self.assertEqual(block["shared_all"], 1)


class TestGapBlock(unittest.TestCase):
    """Gaps da XP: vitrine x vitrine, com supressao em vez de numero falso."""

    def _fund(self, platform, name, cnpj, sources, manager=None, extra=None,
              min_initial=None, fee=None):
        from shelf_monitor.models import FundOffer

        f = FundOffer(platform=platform, name=name, cnpj=cnpj,
                      manager_name=manager, min_initial=min_initial,
                      admin_fee_pct=fee)
        f.sources = list(sources)
        if extra:
            f.extra.update(extra)
        return f

    def _snap(self, platform, funds, status=None, strategy=None, error=None):
        from datetime import datetime

        from shelf_monitor.models import (
            CollectorStrategy, ShelfSnapshot, SnapshotStatus,
        )

        return ShelfSnapshot(
            platform=platform, collected_at=datetime(2026, 8, 31),
            strategy=strategy or CollectorStrategy.PLAYWRIGHT,
            status=status or SnapshotStatus.OK,
            funds=funds, error=error,
        )

    def _xp(self):
        return self._snap("xp", [
            self._fund("xp", "XP Macro FIC FIM", "10000000000101",
                       ["playwright"], manager="XP GESTAO",
                       extra={"classificacao_anbima": "Multimercado Macro"}),
            self._fund("xp", "Compartilhado RF", "20000000000202",
                       ["playwright", "open-finance"],
                       extra={"classificacao_anbima": "Renda Fixa Duração Baixa"}),
            self._fund("xp", "XP Só Catálogo", "30000000000303",
                       ["open-finance"]),
        ])

    def _btg(self):
        return self._snap("btg", [
            self._fund("btg", "Compartilhado RF", "20000000000202", ["api-json"],
                       extra={"classificacao_anbima": "Renda Fixa Duração Baixa"}),
            self._fund("btg", "Kapitalo Zeta FIC FIM", "40000000000404",
                       ["api-json"], manager="KAPITALO GESTAO",
                       extra={"classificacao_anbima": "Multimercado Macro"},
                       min_initial=5000.0, fee=2.0),
            self._fund("btg", "XP Só Catálogo", "30000000000303", ["api-json"]),
            self._fund("btg", "BTG Restrito", "50000000000505", ["api-json"],
                       extra={"exclusivo_cvm": "S"}),
            self._fund("btg", "BTG Sem CNPJ", None, ["api-json"]),
        ])

    def _itau(self):
        return self._snap("itau", [
            self._fund("itau", "Kapitalo Zeta FIC FIM", "40000000000404",
                       ["playwright"],
                       extra={"classificacao_anbima": "Multimercado Macro"},
                       min_initial=1000.0, fee=1.5),
        ])

    def test_competitor_listing_fund_missing_from_xp_listing_is_a_gap(self):
        from shelf_monitor.report.aggregate import build_gap_block

        block = build_gap_block([self._xp(), self._btg(), self._itau()])
        self.assertTrue(block["available"])
        cnpjs = [f["cnpj"] for f in block["funds"]]
        # Kapitalo (2 vitrines) e "XP So Catalogo" (1 vitrine); o compartilhado
        # esta na vitrine da XP e fica fora.
        self.assertEqual(block["total_gaps"], 2)
        self.assertEqual(block["multi_house_gaps"], 1)
        self.assertNotIn("20000000000202", cnpjs)
        kapitalo = block["funds"][0]
        self.assertEqual(kapitalo["cnpj"], "40000000000404")
        self.assertEqual(kapitalo["houses_count"], 2)
        self.assertEqual(kapitalo["houses_label"], "BTG Pactual, Itaú")
        # Melhores condicoes entre as casas que vitrinam.
        self.assertEqual(kapitalo["min_initial"], 1000.0)
        self.assertEqual(kapitalo["admin_fee_pct"], 1.5)

    def test_fund_only_in_xp_catalog_still_counts_as_gap(self):
        from shelf_monitor.report.aggregate import build_gap_block

        block = build_gap_block([self._xp(), self._btg()])
        cnpjs = [f["cnpj"] for f in block["funds"]]
        # A XP tem o fundo, mas so via Open Finance: a base e a vitrine.
        self.assertIn("30000000000303", cnpjs)

    def test_xp_without_listing_suppresses_block(self):
        from shelf_monitor.report.aggregate import build_gap_block

        xp = self._snap("xp", [
            self._fund("xp", "XP Só Catálogo", "30000000000303",
                       ["open-finance"]),
        ])
        block = build_gap_block([xp, self._btg()])
        self.assertFalse(block["available"])
        self.assertIn("vitrine", block["reason"])

    def test_missing_or_unavailable_xp_suppresses_the_block(self):
        from shelf_monitor.models import SnapshotStatus
        from shelf_monitor.report.aggregate import build_gap_block

        without_xp = build_gap_block([self._btg(), self._itau()])
        self.assertFalse(without_xp["available"])

        xp = self._snap("xp", [], status=SnapshotStatus.UNAVAILABLE,
                        error="timeout")
        unavailable = build_gap_block([xp, self._btg()])
        self.assertFalse(unavailable["available"])
        self.assertIn("indisponível", unavailable["reason"])

    def test_partial_xp_keeps_block_with_warning(self):
        from shelf_monitor.models import SnapshotStatus
        from shelf_monitor.report.aggregate import build_gap_block

        xp = self._xp()
        xp.status = SnapshotStatus.PARTIAL
        xp.error = "playwright caiu na página 3"
        block = build_gap_block([xp, self._btg()])
        self.assertTrue(block["available"])
        self.assertIn("superestimados", block["baseline_warning"])

    def test_unavailable_competitor_leaves_the_universe(self):
        from shelf_monitor.models import SnapshotStatus
        from shelf_monitor.report.aggregate import build_gap_block

        itau = self._snap("itau", [], status=SnapshotStatus.UNAVAILABLE)
        block = build_gap_block([self._xp(), self._btg(), itau])
        self.assertIn("itau", [e["platform_id"] for e in block["excluded"]])
        for fund in block["funds"]:
            self.assertNotIn("itau", fund["houses"])

    def test_catalog_only_competitor_is_excluded_with_reason(self):
        from shelf_monitor.models import CollectorStrategy
        from shelf_monitor.report.aggregate import build_gap_block

        nu = self._snap("nu", [
            self._fund("nu", "Nu Reserva", "60000000000606", ["open-finance"]),
        ], strategy=CollectorStrategy.OPEN_FINANCE)
        block = build_gap_block([self._xp(), nu])
        excluded = next(e for e in block["excluded"] if e["platform_id"] == "nu")
        self.assertIn("Open Finance", excluded["reason"])
        self.assertEqual(block["total_gaps"], 0)

    def test_restricted_and_cnpjless_funds_stay_out(self):
        from shelf_monitor.report.aggregate import build_gap_block

        block = build_gap_block([self._xp(), self._btg()])
        cnpjs = [f["cnpj"] for f in block["funds"]]
        names = [f["name"] for f in block["funds"]]
        self.assertNotIn("50000000000505", cnpjs)
        self.assertNotIn("BTG Sem CNPJ", names)

    def test_category_block_compares_xp_to_best_competitor(self):
        from shelf_monitor.report.aggregate import build_gap_block

        block = build_gap_block([self._xp(), self._btg(), self._itau()])
        macro = next(
            c for c in block["categories"] if c["category"] == "Multimercado Macro"
        )
        self.assertEqual(macro["xp"], 1)            # XP Macro FIC FIM
        self.assertEqual(macro["gap_count"], 1)     # Kapitalo
        self.assertEqual(macro["best_competitor"]["count"], 1)
        self.assertEqual(macro["by_platform"], {"btg": 1, "itau": 1})

    def test_manager_and_slug_heuristics_never_remove(self):
        from shelf_monitor.report.aggregate import build_gap_block

        btg = self._btg()
        btg.funds.append(self._fund(
            "btg", "XP Macro Fundo de Investimento em Cotas Multimercado",
            "70000000000707", ["api-json"], manager="XP GESTAO",
        ))
        block = build_gap_block([self._xp(), btg])
        feeder = next(f for f in block["funds"] if f["cnpj"] == "70000000000707")
        # O sinal aparece, mas o fundo continua na lista.
        self.assertEqual(feeder["possible_xp_equivalent"], "XP Macro FIC FIM")
        self.assertTrue(feeder["manager_on_xp"])
        kapitalo = next(f for f in block["funds"] if f["cnpj"] == "40000000000404")
        self.assertFalse(kapitalo["manager_on_xp"])


class TestDeck(unittest.TestCase):
    """O deck abre e tem os slides que os dados pedem - nem mais, nem menos."""

    def _report(self, with_listing=True):
        from datetime import datetime

        from shelf_monitor.models import (
            CollectorStrategy, FundOffer, RunReport, ShelfSnapshot, SnapshotStatus,
        )

        funds = []
        if with_listing:
            for i in range(6):
                f = FundOffer(
                    platform="btg", name=f"Fundo {i} FIC FIM", cnpj=f"{i:014d}",
                    admin_fee_pct=1.0 + i * 0.1, min_initial=1000.0,
                    manager_name="BTG PACTUAL ASSET" if i % 2 else "TERCEIRA GESTORA",
                )
                f.sources = ["api-json"]
                funds.append(f)
        snapshot = ShelfSnapshot(
            platform="btg", collected_at=datetime(2026, 8, 31),
            strategy=CollectorStrategy.API_JSON, status=SnapshotStatus.OK,
            funds=funds,
        )
        return RunReport(generated_at=datetime(2026, 8, 31), snapshots=[snapshot])

    def test_deck_renders_and_reopens(self):
        import tempfile

        from pptx import Presentation

        from shelf_monitor.report.deck import render_deck

        with tempfile.TemporaryDirectory() as tmp:
            config = Config(data_dir=tmp)
            path = render_deck(config, self._report())
            deck = Presentation(path)
            # capa, numeros, arquitetura, custo, movimentos, metodologia, fontes
            # (sem matriz - nao ha base ANBIMA; sem radar - nao ha CVM)
            self.assertEqual(len(deck.slides), 7)

    def test_deck_survives_an_empty_run(self):
        import tempfile

        from pptx import Presentation

        from shelf_monitor.report.deck import render_deck

        with tempfile.TemporaryDirectory() as tmp:
            config = Config(data_dir=tmp)
            path = render_deck(config, self._report(with_listing=False))
            self.assertGreaterEqual(len(Presentation(path).slides), 4)

    def test_gap_slide_appears_with_xp_baseline(self):
        import tempfile
        from datetime import datetime

        from pptx import Presentation

        from shelf_monitor.models import FundOffer, ShelfSnapshot
        from shelf_monitor.models import CollectorStrategy, SnapshotStatus
        from shelf_monitor.report.deck import render_deck

        report = self._report()
        xp_fund = FundOffer(platform="xp", name="XP Único FIC FIM",
                            cnpj="99000000000199")
        xp_fund.sources = ["playwright"]
        report.snapshots.insert(0, ShelfSnapshot(
            platform="xp", collected_at=datetime(2026, 8, 31),
            strategy=CollectorStrategy.PLAYWRIGHT, status=SnapshotStatus.OK,
            funds=[xp_fund],
        ))

        with tempfile.TemporaryDirectory() as tmp:
            config = Config(data_dir=tmp)
            path = render_deck(config, report)
            # capa, numeros, gaps (os 6 fundos BTG estao fora da vitrine XP),
            # arquitetura, custo, movimentos, metodologia, fontes
            self.assertEqual(len(Presentation(path).slides), 8)


class TestFundEnvelopeJoin(TestEnrichment):
    """Prateleira anuncia o CNPJ do envelope; o registro tem os das classes.

    Foi o sintoma da primeira rodada real: fundo multi-classes anunciado pelo
    CNPJ do fundo saia "sem gestora e sem classificacao" porque o indice so
    conhecia CNPJ de classe.
    """

    def test_envelope_cnpj_reaches_the_fund_record(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="btg", name="Fundo Multi", cnpj="66666666000166")
        stats = enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertEqual(stats.matched, 1)
        self.assertEqual(fund.extra["gestor_oficial"], "GESTORA MULTI LTDA")
        # As duas classes concordam em Multimercados Livre - o envelope herda.
        self.assertEqual(fund.extra["classificacao_anbima"], "Multimercados Livre")

    def test_divergent_classes_do_not_invent_a_classification(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="btg", name="Fundo Diverge", cnpj="20202020000120")
        enrich_snapshots(self.config, [self._snapshot([fund])])

        # RF numa classe, Acoes na outra: atribuir qualquer uma seria inventar.
        self.assertNotIn("classificacao_anbima", fund.extra)
        self.assertNotIn("classificacao_cvm", fund.extra)
        # Mas o gestor e do envelope - vale sempre.
        self.assertEqual(fund.extra["gestor_oficial"], "GESTORA DIVERGE")

    def test_class_cnpj_still_wins_over_fund_index(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="btg", name="Multi A", cnpj="77777777000177")
        enrich_snapshots(self.config, [self._snapshot([fund])])
        self.assertEqual(fund.extra["publico_alvo_cvm"], "Investidores em geral")
        self.assertEqual(fund.extra["pl_cvm"], 5000000.0)


class TestCategoryFallback(unittest.TestCase):
    """Classe em tres camadas: ANBIMA -> CVM -> heuristica -> sem classe."""

    def test_tiers(self):
        from shelf_monitor.models import AssetClass, FundOffer
        from shelf_monitor.report.aggregate import anbima_category

        with_anbima = FundOffer(platform="x", name="A", extra={
            "classificacao_anbima": "RF referenciado DI", "classificacao_cvm": "Renda Fixa",
        })
        self.assertEqual(anbima_category(with_anbima), "RF referenciado DI")

        cvm_only = FundOffer(platform="x", name="B", extra={"classificacao_cvm": "Renda Fixa"})
        self.assertEqual(anbima_category(cvm_only), "Renda Fixa")

        heuristic = FundOffer(platform="x", name="C", asset_class=AssetClass.MULTIMERCADO)
        self.assertEqual(anbima_category(heuristic), "Multimercado")

        nothing = FundOffer(platform="x", name="D")
        self.assertEqual(anbima_category(nothing), "Sem classificação")


class TestPlanilhaComArquivoPreso(unittest.TestCase):
    """O Excel trava o arquivo aberto para escrita.

    Uma execucao de dezenas de minutos nao pode terminar sem planilha porque
    alguem estava com a anterior aberta - aconteceu em 2026-09-20, com o `run`
    completo devolvendo "Permission denied" no ultimo passo.
    """

    def test_cai_para_um_nome_vizinho_quando_o_alvo_esta_travado(self):
        import tempfile
        from pathlib import Path as _Path
        from openpyxl import Workbook

        from shelf_monitor.report.spreadsheet import _salvar

        with tempfile.TemporaryDirectory() as tmp:
            alvo = _Path(tmp) / "prateleiras-2026-09-20.xlsx"
            original = Workbook.save

            def travado(self, caminho):
                if _Path(caminho).name == alvo.name:
                    raise PermissionError(13, "Permission denied")
                return original(self, caminho)

            Workbook.save = travado
            try:
                destino = _salvar(Workbook(), alvo)
            finally:
                Workbook.save = original

        self.assertEqual(destino.name, "prateleiras-2026-09-20-2.xlsx")


class TestRetailProfile(unittest.TestCase):
    """Exclusivos/restritos ficam fora da analise principal, com motivo."""

    def _fund(self, name="Fundo Comum FIC FIM", **extra):
        from shelf_monitor.models import FundOffer
        return FundOffer(platform="btg", name=name, cnpj="11111111000111", extra=extra)

    def test_reasons_in_order_of_evidence(self):
        from shelf_monitor.profile import restriction_reason

        self.assertIn("registro CVM", restriction_reason(self._fund(exclusivo_cvm="S")))
        self.assertIn("profissional", restriction_reason(
            self._fund(publico_alvo_cvm="Investidor Profissional")))
        self.assertIn("própria casa", restriction_reason(self._fund(restrito_casa=True)))
        self.assertIn("nome", restriction_reason(
            self._fund(name="Família Silva Exclusivo FIC FIM")))

    def test_retail_funds_pass_untouched(self):
        from shelf_monitor.profile import restriction_reason

        self.assertIsNone(restriction_reason(self._fund()))
        # Qualificado NAO e motivo de corte - varejo de alta renda compra.
        self.assertIsNone(restriction_reason(
            self._fund(publico_alvo_cvm="Investidor Qualificado")))
        # "Reserva" e "Selection" sao nomes de produto de varejo, nao restricao.
        self.assertIsNone(restriction_reason(self._fund(name="Nu Reserva Imediata FIC")))
        self.assertIsNone(restriction_reason(self._fund(name="XP Selection FIC FIM")))

    def test_main_analyses_exclude_and_auxiliary_reports(self):
        from datetime import datetime

        from shelf_monitor.models import (
            CollectorStrategy, FundOffer, RunReport, ShelfSnapshot, SnapshotStatus,
        )
        from shelf_monitor.report.aggregate import (
            build_listing_block, build_restricted_block, summarize_platform,
        )

        retail = FundOffer(platform="btg", name="BTG Varejo FIC FIM", cnpj="11111111000111")
        retail.sources = ["api-json"]
        restricted = FundOffer(
            platform="btg", name="BTG Cliente Único FIC FIM", cnpj="22222222000122",
            extra={"exclusivo_cvm": "S"},
        )
        restricted.sources = ["api-json"]
        snapshot = ShelfSnapshot(
            platform="btg", collected_at=datetime(2026, 8, 31),
            strategy=CollectorStrategy.API_JSON, status=SnapshotStatus.OK,
            funds=[retail, restricted],
        )
        report = RunReport(generated_at=datetime(2026, 8, 31), snapshots=[snapshot])

        summary = summarize_platform(snapshot, set())
        self.assertEqual(summary.listing_funds, 1, "restrito não conta na análise")
        self.assertEqual(summary.catalog_funds, 1, "catálogo de varejo também exclui o restrito")
        self.assertEqual(summary.collected_funds, 2, "a coleta trouxe os dois")

        listing = build_listing_block(report)
        btg = next(p for p in listing["platforms"] if p["platform_id"] == "btg")
        self.assertEqual(btg["total"], 1)

        aux = build_restricted_block(report)
        self.assertEqual(aux["total"], 1)
        entry = aux["platforms"][0]
        self.assertEqual(entry["restricted"], 1)
        self.assertIn("exclusivo no registro CVM", entry["by_reason"])
        self.assertEqual(entry["sample"][0]["name"], "BTG Cliente Único FIC FIM")


class TestManagerInference(unittest.TestCase):
    """Gestora pelo prefixo do nome - inequivoca ou nada."""

    MANAGERS = [
        "KINEA INVESTIMENTOS LTDA",
        "BAHIA AM RENDA VARIAVEL LTDA",
        "ITAU UNIBANCO ASSET MANAGEMENT LTDA",
        "ITAU DTVM S.A.",           # segunda "ITAU": ambiguidade proposital
        "INTRAG DTVM LTDA",
        "BRASIL CAPITAL GESTAO DE RECURSOS LTDA",
    ]

    def test_unambiguous_brands_fill_the_manager(self):
        from shelf_monitor.enrich import build_brand_index, infer_manager
        from shelf_monitor.models import FundOffer

        index = build_brand_index(self.MANAGERS)
        fund = FundOffer(platform="itau", name="Bahia AM Muta Seleção Multimercado")
        self.assertTrue(infer_manager(fund, index))
        self.assertEqual(fund.manager_name, "BAHIA AM RENDA VARIAVEL LTDA")
        self.assertIn("gestor_inferido", fund.extra)

        kinea = FundOffer(platform="itau", name="Kinea Chronos FIC FIM")
        infer_manager(kinea, index)
        self.assertEqual(kinea.manager_name, "KINEA INVESTIMENTOS LTDA")

    def test_ambiguous_brand_fills_nothing(self):
        """"ITAU" aponta para Itau Asset E Itau DTVM: melhor vazio que errado."""
        from shelf_monitor.enrich import build_brand_index, infer_manager
        from shelf_monitor.models import FundOffer

        index = build_brand_index(self.MANAGERS)
        self.assertNotIn("ITAU", index)
        fund = FundOffer(platform="itau", name="Itau Carteira Livre FIC")
        self.assertFalse(infer_manager(fund, index))
        self.assertIsNone(fund.manager_name)

    def test_same_manager_spelled_differently_is_not_ambiguity(self):
        from shelf_monitor.enrich import build_brand_index

        index = build_brand_index(
            ["KINEA INVESTIMENTOS LTDA", "KINEA INVESTIMENTOS"]
        )
        self.assertEqual(index["KINEA"], "KINEA INVESTIMENTOS")

    def test_existing_manager_is_never_overwritten(self):
        from shelf_monitor.enrich import build_brand_index, infer_manager
        from shelf_monitor.models import FundOffer

        index = build_brand_index(self.MANAGERS)
        fund = FundOffer(platform="itau", name="Kinea Chronos", manager_name="OFICIAL")
        self.assertFalse(infer_manager(fund, index))
        self.assertEqual(fund.manager_name, "OFICIAL")


class TestFuzzyReconciliation(unittest.TestCase):
    """Fuzzy conservador: casa o obvio, nunca o master com o feeder."""

    def _merge(self, of_name, site_name):
        from shelf_monitor.merge import merge_offers

        return merge_offers({
            "open-finance": [offer("itau", of_name, "11111111000111")],
            "playwright": [offer("itau", site_name, None, admin_fee_pct=0.5)],
        })

    def test_typo_level_difference_merges(self):
        merged = self._merge(
            "Itau Global Dinâmico Renda Fixa Longo Prazo",
            "Itau Global Dinâmico Renda Fixa Longo Praz",  # truncado pela tabela
        )
        self.assertEqual(len(merged), 1, [(o.name, o.cnpj) for o in merged])
        self.assertEqual(merged[0].cnpj, "11111111000111")
        self.assertIn("nome aproximado", merged[0].extra["reconciliacao"])

    def test_master_and_feeder_never_merge_by_similarity(self):
        merged = self._merge(
            "Kinea Chronos Master Multimercado",
            "Kinea Chronos Multimercado",
        )
        self.assertEqual(len(merged), 2, "MASTER e token distintivo - não funde")

    def test_numbered_siblings_never_merge(self):
        merged = self._merge(
            "A1 D30 Renda Fixa Crédito Privado",
            "A1 D60 Renda Fixa Crédito Privado",
        )
        self.assertEqual(len(merged), 2)

    def test_ambiguous_candidates_do_not_merge(self):
        from shelf_monitor.merge import merge_offers

        merged = merge_offers({
            "open-finance": [
                offer("itau", "Alpha Global Equities USDX FIC", "11111111000111"),
                offer("itau", "Alpha Global Equities USDY FIC", "22222222000122"),
            ],
            "playwright": [offer("itau", "Alpha Global Equities USD FIC", None)],
        })
        # Dois candidatos igualmente parecidos: ambiguo, o do site fica separado.
        self.assertEqual(len(merged), 3, [(o.name, o.cnpj) for o in merged])
