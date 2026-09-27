"""Conferencia da captacao de ETF com o boletim de fundos da ANBIMA.

O anexo e montado aqui, na forma do arquivo que a ANBIMA publica: cabecalho
"Tipos ANBIMA", meses como data e o ultimo como texto ("fev-26"), valores em
R$ milhoes, e as linhas "ETF", "ETF Renda Fixa" e "ETF Renda Variavel". O
bloco mensal e o cenario de `test_etf_captacao_mensal`, com numeros redondos.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import os
import time
import unittest
from datetime import datetime
from unittest import mock

from openpyxl import Workbook

from shelf_monitor.market import anbima_boletim
from shelf_monitor.market import etf_captacao_mensal as mensal
from shelf_monitor.market.anbima_boletim import anexar_conferencia, conciliar, ler_anexo, localizar_anexo
from shelf_monitor.market.etf_captacao_mensal import build_captacao_mensal

from fixtures_comuns import temp_config
from test_etf_captacao_mensal import MI, _cenario

MESES = [datetime(2025, 11, 1), datetime(2025, 12, 1), datetime(2026, 1, 1), "fev-26"]
CAPTACAO = {  # R$ milhoes, como o anexo
    "ETF": [200, -100, 400, 80],
    "ETF Renda Fixa": [0, 0, 600, 0],
    "ETF Renda Variável": [200, -100, -200, 80],
}
PL = {
    "ETF": [13000, 13500, 14000, 14200],
    "ETF Renda Fixa": [0, 0, 600, 750],
    "ETF Renda Variável": [13000, 13500, 13400, 13450],
}


def _anexo(caminho, fundos=5):
    livro = Workbook()
    livro.remove(livro.active)
    for aba, valores in ((anbima_boletim.ABA_PL, PL), (anbima_boletim.ABA_CAPTACAO, CAPTACAO)):
        ws = livro.create_sheet(aba)
        ws.append([None, "ANBIMA » Fundos de Investimento | Relatórios"])
        ws.append([None, "Dados em R$ milhões"])
        ws.append([None, "Tipos ANBIMA", *MESES, "Total Captação Líquida - 12 meses"])
        ws.append([None, "Renda fixa", 1, 2, 3, 4, 10])
        for tipo, serie in valores.items():
            ws.append([None, tipo, *serie, sum(serie)])
    ws = livro.create_sheet(anbima_boletim.ABA_FUNDOS)
    ws.append(["Período", "Renda Fixa (3)", "Ações", "ETF"])
    ws.append([datetime(2026, 2, 1), 10, 20, fundos])
    livro.save(caminho)
    return caminho


class TestConferencia(unittest.TestCase):
    def setUp(self):
        self.config, self.tmp = temp_config()
        catalogo, fechamentos, balancete = _cenario()
        with mock.patch.object(mensal, "MESES_DA_VISAO", 3):
            self.bloco = build_captacao_mensal(catalogo, fechamentos, balancete)
        pasta = self.config.etf_dir / "anbima"
        pasta.mkdir(parents=True)
        self.caminho = _anexo(pasta / "Anexo_boletim_fundos_investimento_fevereiro.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def test_le_o_anexo(self):
        anexo = ler_anexo(self.caminho)
        self.assertEqual(anexo["captacao"]["total"],
                         {"202511": 200 * MI, "202512": -100 * MI, "202601": 400 * MI, "202602": 80 * MI})
        self.assertEqual(anexo["pl"]["renda_fixa"]["202602"], 750 * MI)
        self.assertEqual(anexo["fundos"], {"202602": 5})

    def test_localiza_o_anexo_mais_recente_e_recusa_o_que_nao_e_anexo(self):
        self.assertEqual(localizar_anexo(self.config), self.caminho)
        novo = _anexo(self.caminho.with_name("Anexo_boletim_fundos_investimento_marco.xlsx"))
        os.utime(novo, (time.time() + 10, time.time() + 10))
        self.assertEqual(localizar_anexo(self.config, self.caminho.parent), novo)
        self.assertEqual(localizar_anexo(self.config, self.caminho), self.caminho)
        outro = self.caminho.with_name("qualquer.xlsx")
        Workbook().save(outro)
        self.assertIsNone(ler_anexo(outro))

    def test_mes_a_mes_na_janela_dos_dois_com_as_estreias(self):
        conferencia = conciliar(self.bloco, ler_anexo(self.caminho))
        # O mes-base da visao (11/2025) entra: o boletim cobre ele.
        self.assertEqual([l["mes"] for l in conferencia["linhas"]], ["202511", "202512", "202601", "202602"])
        janeiro = conferencia["linhas"][2]
        self.assertEqual((janeiro["monitor"], janeiro["estreias"], janeiro["anbima"]), (930 * MI, 600 * MI, 400 * MI))
        self.assertAlmostEqual(janeiro["diferenca"], 530 * MI)
        self.assertAlmostEqual(janeiro["renda_fixa"]["monitor"], 600 * MI)
        self.assertAlmostEqual(janeiro["renda_variavel"]["monitor"], 330 * MI)
        self.assertAlmostEqual(conferencia["doze_meses"]["monitor"], 1075 * MI)
        self.assertAlmostEqual(conferencia["doze_meses"]["anbima"], 580 * MI)
        self.assertEqual((conferencia["no_ano"]["de"], conferencia["no_ano"]["ate"]), ("202601", "202602"))
        self.assertEqual(conferencia["fundos_anbima"], 5)
        self.assertEqual(conferencia["destaques"], [])

    def test_destaque_traz_os_etfs_que_estrearam_no_mes(self):
        with mock.patch.object(anbima_boletim, "DIFERENCA_RELEVANTE", 500 * MI):
            conferencia = conciliar(self.bloco, ler_anexo(self.caminho))
        self.assertEqual(len(conferencia["destaques"]), 1)
        destaque = conferencia["destaques"][0]
        self.assertEqual(destaque["mes"], "202601")
        self.assertEqual([e["nome"] for e in destaque["etfs_em_estreia"]], ["IT NOW IMA-B5 P2 FUNDO DE INDICE"])

    def test_leitura_e_relatorio(self):
        from shelf_monitor.report.etf_captacao import render_captacao_markdown

        anexar_conferencia(self.bloco, ler_anexo(self.caminho))
        self.assertIn(
            "Conferência com o boletim da ANBIMA: de nov/2025 a fev/2026, o monitor soma +R$ 1,1 bi "
            "e a ANBIMA +R$ 580 mi; no ano, +R$ 995 mi e +R$ 480 mi.",
            self.bloco["leitura"][-1],
        )
        md = render_captacao_markdown(self.bloco, datetime(2026, 9, 24, 10, 0))
        self.assertIn("## Conferência com a ANBIMA", md)
        self.assertIn("| jan/2026 | +R$ 930 mi | +R$ 600 mi | +R$ 400 mi | +R$ 530 mi |", md)
        self.assertIn("| **12 meses** | **+R$ 1,1 bi** | +R$ 600 mi | **+R$ 580 mi** | +R$ 495 mi | |", md)
        self.assertIn("[Boletim de Fundos de Investimento da ANBIMA]", md)

    def test_planilha_ganha_a_aba_da_conferencia(self):
        from openpyxl import load_workbook

        from shelf_monitor.report.etf_report import export_etf_excel

        anexar_conferencia(self.bloco, ler_anexo(self.caminho))
        caminho = export_etf_excel([], [], {}, self.config.reports_dir / "etf.xlsx", captacao=self.bloco)
        livro = load_workbook(caminho)
        self.assertEqual(livro.sheetnames[4], "Captação - ANBIMA")
        valores = [c.value for c in next(
            linha for linha in livro["Captação - ANBIMA"].iter_rows() if linha[0].value == "jan/26"
        )]
        self.assertEqual(valores[1:5], [930 * MI, 600 * MI, 400 * MI, 530 * MI])

    def test_sem_anexo_o_bloco_fica_como_esta(self):
        antes = list(self.bloco["leitura"])
        anexar_conferencia(self.bloco, None)
        self.assertNotIn("anbima", self.bloco)
        self.assertEqual(self.bloco["leitura"], antes)


if __name__ == "__main__":
    unittest.main()
