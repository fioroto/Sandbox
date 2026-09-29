"""Bloco da industria (Fase 1C): janelas, fluxos, decomposicao, planilha.

Um painel sintetico de 13 competencias, com numeros redondos calculados a mao:

- A: fundo direto do universo, PL de 1.000 em 08/2025 que rende 1% ao mes e
  capta 10 por mes; 20% do PL e de outros fundos o tempo todo;
- B: exclusivo (fora do universo);
- C: lancado em 03/2026 (cadastro de 02/2026), no universo;
- D: estava no universo e sumiu do informe em 06/2026;
- E: evento na cota em 04/2026 (cota x1000): o retorno sai, o fluxo cai no rateio.

Trava as identidades: variacao do PL = fluxo + rendimento + residuo; as
macroclasses somam o universo; o fluxo do investidor mais o dos fundos e a
captacao declarada; o retorno absurdo fica fora da mediana.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import csv
import json
import unittest
from datetime import date

from openpyxl import load_workbook

from shelf_monitor.industria.adjacentes import camada_etf
from shelf_monitor.industria.base import BaseIndustria
from shelf_monitor.industria.bloco import montar_bloco
from shelf_monitor.industria.cadastro import GERAL, Cadastro, Produto
from shelf_monitor.industria.cotistas import CHAVES, PerfilClasse
from shelf_monitor.industria.curadoria import Curadoria
from shelf_monitor.industria.fluxos import fluxos_por, lancamentos_e_encerramentos, total
from shelf_monitor.industria.janelas import janelas
from shelf_monitor.industria.painel import linhas_do_mes
from shelf_monitor.report.industria_planilha import gravar_planilha
from shelf_monitor.report.rentabilidade import MesSubclasse, SerieMensal

from fixtures_comuns import temp_config

MESES = ["202508", "202509", "202510", "202511", "202512", "202601", "202602",
         "202603", "202604", "202605", "202606", "202607", "202608"]


def _perfil(fundos: float, conta_e_ordem: float = 0.0) -> PerfilClasse:
    fatias = []
    for c in CHAVES:
        if c == "fundos":
            fatias.append(fundos)
        elif c == "conta_e_ordem":
            fatias.append(conta_e_ordem)
        elif c == "pf_varejo":
            fatias.append(1 - fundos - conta_e_ordem)
        else:
            fatias.append(0.0)
    return PerfilClasse(cnpj="x", versao=4, fatias=tuple(fatias), cotistas=tuple(10 for _ in CHAVES),
                        previdencia=0.0, maior_cotista=None, soma_informada=1.0)


def _produto(cnpj: str, **campos) -> Produto:
    padrao = dict(
        classe=cnpj, subclasse="", nivel="classe", nome=f"FUNDO {cnpj}", gestor="GESTORA",
        gestor_cnpj="99999999000199", tipo="Classes de Cotas de Fundos FIF",
        situacao="Em Funcionamento Normal", classificacao_anbima="Renda Fixa Duração Baixa Grau de Invest.",
        classificacao_cvm="Renda Fixa", publico=GERAL, exclusivo=False, inicio=date(2015, 1, 1),
    )
    padrao.update(campos)
    return Produto(**padrao)


def _dia(mes: str, dia: int) -> date:
    return date(int(mes[:4]), int(mes[4:]), dia)


def _serie() -> SerieMensal:
    dados: dict = {("A", ""): {}, ("B", ""): {}, ("C", ""): {}, ("D", ""): {}, ("E", ""): {}}
    pl, cota = 1000.0, 1.0
    for i, mes in enumerate(MESES):
        # A: PL(t) = PL(t-1) x 1,01 + 10 -> a decomposicao fecha sem residuo.
        captacao = 10.0 if i else 0.0
        if i:
            pl, cota = pl * 1.01 + captacao, cota * 1.01
        dados[("A", "")][mes] = MesSubclasse(
            cota_ini=cota / 1.01 if i else cota, cota_fim=cota, pl_fim=pl, cotistas_fim=100,
            captacao=captacao, resgates=0.0, pregoes=21, dia_ini=_dia(mes, 1), dia_fim=_dia(mes, 28),
        )
        dados[("B", "")][mes] = MesSubclasse(cota_ini=1, cota_fim=1, pl_fim=500.0, captacao=50.0,
                                            pregoes=21, dia_ini=_dia(mes, 1), dia_fim=_dia(mes, 28))
        if mes >= "202603":
            dados[("C", "")][mes] = MesSubclasse(cota_ini=1, cota_fim=1, pl_fim=100.0,
                                                captacao=100.0 if mes == "202603" else 0.0,
                                                pregoes=21, dia_ini=_dia(mes, 1), dia_fim=_dia(mes, 28))
        if mes < "202606":
            dados[("D", "")][mes] = MesSubclasse(cota_ini=1, cota_fim=1, pl_fim=300.0, pregoes=21,
                                                dia_ini=_dia(mes, 1), dia_fim=_dia(mes, 28))
        cota_e = 1000.0 if mes >= "202604" else 1.0
        dados[("E", "")][mes] = MesSubclasse(cota_ini=cota_e, cota_fim=cota_e, pl_fim=200.0,
                                            captacao=4.0, pregoes=21, dia_ini=_dia(mes, 1), dia_fim=_dia(mes, 28))
    return SerieMensal(list(MESES), dados)


def _base() -> BaseIndustria:
    cadastro = Cadastro("202609", {
        ("A", ""): _produto("A"),
        ("B", ""): _produto("B", exclusivo=True),
        ("C", ""): _produto("C", inicio=date(2026, 2, 20), classificacao_anbima="Ações Livre",
                           classificacao_cvm="Ações"),
        ("D", ""): _produto("D"),
        ("E", ""): _produto("E"),
    })
    perfis = {mes: {c: _perfil(0.2 if c == "A" else 0.0, conta_e_ordem=0.1 if mes == "202608" else 0.05)
                    for c in "ABCDE"} for mes in MESES}
    return BaseIndustria(
        competencia="202608", meses=list(MESES), cadastro=cadastro, serie=_serie(), perfis=perfis,
        curadoria=Curadoria({}, [], [], [], {}),
    )


class TestJanelas(unittest.TestCase):
    def test_doze_meses_com_datas_e_base(self):
        jan = janelas(_base())
        doze = jan["12m"]
        self.assertEqual(doze.competencias, tuple(MESES[1:]))
        self.assertEqual(doze.base, "202508")
        self.assertEqual((doze.de, doze.ate), (date(2025, 9, 1), date(2026, 8, 28)))
        self.assertEqual(doze.data_base, date(2025, 8, 28))
        self.assertEqual(jan["ano"].competencias[0], "202601")
        self.assertEqual(jan["1m"].competencias, ("202608",))


class TestFluxosEDecomposicao(unittest.TestCase):
    def setUp(self):
        self.base = _base()
        self.jan = janelas(self.base)

    def test_decomposicao_fecha_sem_residuo_quando_o_fluxo_e_consistente(self):
        a = fluxos_por(self.base, self.jan["12m"], lambda l, d: l.classe)["A"]
        # 12 meses de +10 declarados; 20% do PL e de fundos, constante.
        self.assertAlmostEqual(a.captacao_liquida, 120.0)
        self.assertAlmostEqual(a.variacao_pl, a.fluxo_investidor + a.rendimento_investidor + a.residuo)
        self.assertAlmostEqual(a.residuo, 0.0, places=6)
        self.assertAlmostEqual(a.fluxo_investidor, 120.0 * 0.8, places=6)

    def test_grupos_somam_o_total_e_exclusivo_nao_entra(self):
        doze = self.jan["12m"]
        por_n1 = fluxos_por(self.base, doze, lambda l, d: d.n1)
        soma = sum(f.fluxo_investidor for f in por_n1.values())
        self.assertAlmostEqual(soma, total(self.base, doze).fluxo_investidor)
        classes = {l.classe for l in linhas_do_mes(self.base, "202608") if l.motivo == "universo"}
        self.assertNotIn("B", classes)

    def test_fluxo_do_investidor_mais_o_dos_fundos_e_o_declarado(self):
        for linha in linhas_do_mes(self.base, "202608"):
            self.assertAlmostEqual(linha.fluxo_investidor + linha.fluxo_fundos, linha.captacao_liquida)

    def test_evento_na_cota_sai_do_retorno_e_cai_no_rateio(self):
        linha = next(l for l in linhas_do_mes(self.base, "202604") if l.classe == "E")
        self.assertTrue(linha.evento_na_cota)
        self.assertIsNone(linha.retorno)
        self.assertAlmostEqual(linha.fluxo_fundos, 4.0 * linha.fatia_fundos)

    def test_lancamento_e_encerramento(self):
        movimentos = lancamentos_e_encerramentos(self.base, self.jan["12m"])
        self.assertEqual([c for c, _, _ in movimentos["entrou"]], ["C"])
        self.assertEqual([(c, m) for c, _, m in movimentos["saiu"]], [("D", "202605")])


class TestBloco(unittest.TestCase):
    def setUp(self):
        self.config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        self.bloco = montar_bloco(_base(), self.config)

    def test_contrato_do_bloco(self):
        bloco = self.bloco
        self.assertTrue(bloco["available"])
        for chave in ("janelas", "prazos", "funil", "camadas", "segmentos", "mercado", "adjacentes",
                      "publico", "estrutura", "estrategias", "familias", "gestoras",
                      "plataformizacao", "lancamentos", "qualidade", "leitura", "fontes"):
            self.assertIn(chave, bloco, chave)
        self.assertNotIn("foco", bloco)
        self.assertTrue(bloco["leitura"])
        self.assertEqual(len(bloco["prazos"]), 4)
        publico = {k: v for k, v in bloco.items() if not k.startswith("_")}
        json.dumps(publico, default=str)  # serializavel

    def test_macroclasses_somam_o_universo(self):
        n1 = self.bloco["estrategias"]["n1"]
        doze = self.bloco["mercado"]["janelas"]["12m"]
        self.assertAlmostEqual(sum(i["janelas"]["12m"]["fluxo_investidor"] for i in n1),
                               doze["fluxo_investidor"], places=1)
        self.assertAlmostEqual(sum(i["pl_investidor"] for i in n1), self.bloco["camadas"]["pl_investidor"], places=1)

    def test_retorno_absurdo_fica_fora_do_peer_group(self):
        e = next(l for l in self.bloco["_explorador"] if l["nome"] == "FUNDO E")
        self.assertIsNone(e["retorno_12m"])
        a = next(l for l in self.bloco["_explorador"] if l["nome"] == "FUNDO A")
        self.assertAlmostEqual(a["retorno_12m"], 1.01 ** 12 - 1, places=4)

    def test_planilha_tem_as_abas_na_ordem_e_fontes_com_link(self):
        caminho = gravar_planilha(self.config, self.bloco)
        self.assertEqual(caminho.name, "industria-2026-08.xlsx")
        wb = load_workbook(caminho)
        self.assertEqual(wb.sheetnames[:3], ["Leitura", "Funil", "Estratégias"])
        self.assertIn("Período de cada aba", wb.sheetnames)
        fontes = wb["Fontes"]
        links = [c.hyperlink for row in fontes.iter_rows() for c in row if c.hyperlink]
        self.assertTrue(links)
        explorador = wb["Explorador"]
        self.assertIsNotNone(explorador.auto_filter.ref)


class TestCamadaEtf(unittest.TestCase):
    def test_sem_arquivo_diz_como_gerar(self):
        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        camada = camada_etf(config, janelas(_base()))
        self.assertFalse(camada["available"])
        self.assertIn("etf --export", camada["reason"])

    def test_soma_a_captacao_das_competencias_da_janela(self):
        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        caminho = config.reports_dir / "etf_captacao_mensal.csv"
        with caminho.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle, delimiter=";")
            writer.writerow(["competencia", "cnpj", "captacao", "pl_fim"])
            for mes in MESES:
                writer.writerow([mes, "1", "10", "1000"])
        camada = camada_etf(config, janelas(_base()))
        self.assertTrue(camada["available"])
        self.assertEqual(camada["janelas"]["12m"]["fluxo_investidor"], 120.0)
        self.assertEqual(camada["janelas"]["1m"]["fluxo_investidor"], 10.0)


if __name__ == "__main__":
    unittest.main()
