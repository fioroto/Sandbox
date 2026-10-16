"""Testes do custo relativo a categoria.

O caso central e o que motivou o modulo: duas casas que cobram exatamente o
mesmo preco em cada categoria nao podem aparecer com custos diferentes so
porque uma tem mais renda fixa que a outra.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

from shelf_monitor.distribuivel import (
    DISTRIBUI,
    PODE_DISTRIBUIR,
    ArquivoLido,
    BaseDistribuivel,
)
from shelf_monitor.models import (
    AssetClass,
    CollectorStrategy,
    FundOffer,
    ShelfSnapshot,
    SnapshotStatus,
)
from shelf_monitor.report.custo import (
    _MAX_IDADE_MESES,
    _MIN_AMOSTRA_CATEGORIA,
    FONTE_BASE_INTERNA,
    FONTE_VITRINE,
    CustoCasa,
    CustoClasse,
    IndiceCusto,
    build_custo,
    classes_com_indice,
    nota_das_classes,
)
from shelf_monitor.report.shared import MIN_COMPARAVEIS_CLASSE_CUSTO

HOJE = date.today()


def recente(meses_atras: int = 1) -> date:
    ano, mes = HOJE.year, HOJE.month - meses_atras
    while mes <= 0:
        mes += 12
        ano -= 1
    return date(ano, mes, 1)


def classe(
    cnpj: str,
    taxa: float | None = 1.0,
    categoria: str | None = "Renda Fixa Simples",
    percentual: bool = True,
    data_base: date | None = None,
    **kwargs,
) -> CustoClasse:
    return CustoClasse(
        cnpj=cnpj,
        categoria=categoria,
        taxa_adm=taxa,
        taxa_percentual=percentual,
        data_base=data_base if data_base is not None else recente(),
        **kwargs,
    )


def indice_com(*classes: CustoClasse) -> IndiceCusto:
    return IndiceCusto({c.cnpj: c for c in classes})


def universo(categoria: str, taxa: float, quantos: int = _MIN_AMOSTRA_CATEGORIA):
    """Classes bastantes para a categoria virar referencia, todas na mesma taxa."""
    return [
        classe(f"{categoria[:3]}{i:011d}", taxa=taxa, categoria=categoria)
        for i in range(quantos)
    ]


def oferta(cnpj: str, nome: str = "Fundo") -> FundOffer:
    return FundOffer(
        platform="xp",
        name=nome,
        cnpj=cnpj,
        sources=[CollectorStrategy.PLAYWRIGHT.value],
    )


class TestElegibilidade(unittest.TestCase):
    def test_taxa_em_reais_nao_entra_no_indice(self):
        """TAXA_ADM carrega percentual e reais no mesmo campo. Um fundo de
        "R$ 30.000 por ano" no indice lideraria como se cobrasse 30.000%."""
        idx = indice_com(
            classe("1", taxa=30000.0, percentual=False), *universo("Renda Fixa Simples", 1.0)
        )
        avaliado = idx.avaliar(oferta("1"))
        self.assertIsNone(avaliado.indice)
        self.assertIn("reais", avaliado.motivo_fora)

    def test_entrega_antiga_sai_do_indice(self):
        """O extrato guarda uma linha por data-base desde 2015; uma taxa de
        2015 nao e a taxa de hoje."""
        antiga = classe("1", data_base=date(2015, 10, 2))
        idx = indice_com(antiga, *universo("Renda Fixa Simples", 1.0))
        avaliado = idx.avaliar(oferta("1"))
        self.assertIsNone(avaliado.indice)
        self.assertIn("meses", avaliado.motivo_fora)
        self.assertNotIn(str(avaliado.idade_meses), avaliado.motivo_fora)
        self.assertGreater(avaliado.idade_meses, _MAX_IDADE_MESES)

    def test_categoria_com_amostra_pequena_nao_vira_referencia(self):
        idx = indice_com(*universo("Nicho", 1.0, quantos=_MIN_AMOSTRA_CATEGORIA - 1))
        avaliado = idx.avaliar(oferta(f"Nic{0:011d}"))
        self.assertIsNone(avaliado.indice)
        self.assertIn("sem referência", avaliado.motivo_fora)
        # O motivo e agrupavel: sem a contagem no texto, cada fundo viraria uma
        # chave propria e o resumo "por motivo" deixaria de resumir.
        self.assertNotIn("9", avaliado.motivo_fora)

    def test_fundo_fora_da_base_diz_que_nao_casou(self):
        idx = indice_com(*universo("Renda Fixa Simples", 1.0))
        avaliado = idx.avaliar(oferta("99999999999999"))
        self.assertIn("não encontrado na base", avaliado.motivo_fora)


class TestIndiceRemoveComposicao(unittest.TestCase):
    """O motivo de o modulo existir."""

    def _casa(self, platform_id: str, rf: int, acoes: int) -> CustoCasa:
        # Todas as casas cobram 1,0% em renda fixa e 2,0% em acoes - exatamente
        # a mediana de cada categoria. So o MIX difere.
        idx = indice_com(
            *universo("Renda Fixa Simples", 1.0),
            *universo("Ações Livre", 2.0),
            *[classe(f"{platform_id}rf{i:09d}", taxa=1.0) for i in range(rf)],
            *[
                classe(f"{platform_id}ac{i:09d}", taxa=2.0, categoria="Ações Livre")
                for i in range(acoes)
            ],
        )
        ofertas = [oferta(f"{platform_id}rf{i:09d}") for i in range(rf)]
        ofertas += [oferta(f"{platform_id}ac{i:09d}") for i in range(acoes)]
        return CustoCasa(
            platform_id=platform_id,
            label=platform_id.upper(),
            vitrine=len(ofertas),
            fundos=[idx.avaliar(o) for o in ofertas],
        )

    def test_mix_diferente_com_preco_igual_da_o_mesmo_indice(self):
        so_rf = self._casa("aa", rf=40, acoes=2)
        so_acoes = self._casa("bb", rf=2, acoes=40)
        # A taxa mediana bruta separaria as duas (1,0% contra 2,0%); o indice
        # mediano nao, porque cada fundo ja foi comparado com a propria
        # categoria antes de a casa virar um numero.
        self.assertAlmostEqual(so_rf.indice_mediano, 1.0)
        self.assertAlmostEqual(so_acoes.indice_mediano, 1.0)

    def test_cobrar_acima_da_categoria_aparece_no_indice(self):
        idx = indice_com(
            *universo("Renda Fixa Simples", 1.0),
            classe("caro", taxa=2.0),
        )
        avaliado = idx.avaliar(oferta("caro"))
        self.assertAlmostEqual(avaliado.indice, 2.0)
        self.assertEqual(avaliado.faixa, "caro para a categoria")


class TestReferenciaDeVarejo(unittest.TestCase):
    """A mediana da categoria sai do universo de VAREJO, nao do universo todo.

    Medido em 2026-09-20: das 13.093 classes com taxa percentual e entrega
    recente, 47,2% sao exclusivas e 21,1% cobram zero. Com as duas dentro, a
    mediana de "Multimercados Livre" caia para 0,080% e a primeira execucao
    devolveu indice mediano de 6,79x com 82% da vitrine do BTG marcada como
    cara.
    """

    def test_classe_exclusiva_nao_entra_na_mediana(self):
        idx = indice_com(
            *[classe(f"ex{i:012d}", taxa=0.05, exclusivo=True) for i in range(40)],
            *universo("Renda Fixa Simples", 1.0),
            classe("alvo", taxa=1.0),
        )
        # So as 20 de varejo contam: mediana 1,0%, indice 1,00x.
        self.assertAlmostEqual(idx.avaliar(oferta("alvo")).indice, 1.0)

    def test_taxa_zero_nao_entra_na_mediana(self):
        idx = indice_com(
            *[classe(f"ze{i:012d}", taxa=0.0) for i in range(40)],
            *universo("Renda Fixa Simples", 1.0),
            classe("alvo", taxa=1.0),
        )
        self.assertAlmostEqual(idx.avaliar(oferta("alvo")).indice, 1.0)

    def test_publico_profissional_nao_entra_na_mediana(self):
        idx = indice_com(
            *[classe(f"pr{i:012d}", taxa=0.05, publico_alvo="PROFISSIONAL")
              for i in range(40)],
            *universo("Renda Fixa Simples", 1.0),
            classe("alvo", taxa=1.0),
        )
        self.assertAlmostEqual(idx.avaliar(oferta("alvo")).indice, 1.0)

    def test_qualificado_continua_na_mediana(self):
        """`profile.py:12-13` decidiu que qualificado nao e motivo de corte.

        Cortar aqui seria desfazer essa decisao por outra porta - e mudaria o
        numero: no universo real, cortar qualificado leva Multimercados Livre de
        0,800% para 1,000%.
        """
        idx = indice_com(
            *[classe(f"qu{i:012d}", taxa=2.0, publico_alvo="QUALIFICADO")
              for i in range(_MIN_AMOSTRA_CATEGORIA)],
            classe("alvo", taxa=2.0),
        )
        # Os qualificados sao a categoria inteira: a mediana e 2,0% e o alvo 1,00x.
        self.assertAlmostEqual(idx.avaliar(oferta("alvo")).indice, 1.0)

    def test_fundo_da_vitrine_com_taxa_zero_sai_com_motivo(self):
        """Feeder que cobra na outra ponta nao e "infinitamente barato"."""
        idx = indice_com(classe("feeder", taxa=0.0), *universo("Renda Fixa Simples", 1.0))
        avaliado = idx.avaliar(oferta("feeder"))
        self.assertIsNone(avaliado.indice)
        self.assertIn("outra ponta", avaliado.motivo_fora)


class TestPortaDeEntrada(unittest.TestCase):
    def test_classe_mais_barata_do_mesmo_fundo_e_sinalizada(self):
        idx = indice_com(
            classe("cara", minimo=50_000.0, id_fundo="F1"),
            classe("barata", minimo=100.0, id_fundo="F1"),
            *universo("Renda Fixa Simples", 1.0),
        )
        avaliado = idx.avaliar(oferta("cara"))
        self.assertEqual(avaliado.porta_menor, 100.0)

    def test_sem_classe_mais_barata_nao_inventa_porta(self):
        idx = indice_com(
            classe("unica", minimo=100.0, id_fundo="F1"),
            *universo("Renda Fixa Simples", 1.0),
        )
        self.assertIsNone(idx.avaliar(oferta("unica")).porta_menor)


class TestLiquidezEIsencao(unittest.TestCase):
    def test_resgate_ate_d1_conta_como_liquidez_imediata(self):
        idx = indice_com(
            classe("rapido", dias_cotizacao=1.0),
            classe("lento", dias_cotizacao=30.0),
            *universo("Renda Fixa Simples", 1.0),
        )
        self.assertTrue(idx.avaliar(oferta("rapido")).liquidez_imediata)
        self.assertFalse(idx.avaliar(oferta("lento")).liquidez_imediata)

    def test_sem_prazo_publicado_nao_e_falso(self):
        idx = indice_com(classe("x"), *universo("Renda Fixa Simples", 1.0))
        self.assertIsNone(idx.avaliar(oferta("x")).liquidez_imediata)

    def test_share_sem_entrada_ou_saida_ignora_quem_nao_respondeu(self):
        idx = indice_com(
            classe("livre", tem_taxa_entrada="N", tem_taxa_saida="N"),
            classe("cobra", tem_taxa_entrada="S", tem_taxa_saida="N"),
            classe("mudo"),
            *universo("Renda Fixa Simples", 1.0),
        )
        casa = CustoCasa(
            platform_id="xp", label="XP", vitrine=3,
            fundos=[idx.avaliar(oferta(c)) for c in ("livre", "cobra", "mudo")],
        )
        # 1 de 2 respondidos, nao 1 de 3.
        self.assertAlmostEqual(casa.share_sem_entrada_ou_saida, 0.5)


class TestCustoPorClasse(unittest.TestCase):
    """A quebra por classe de fundo (pedido de 16/10/2026): a conta da casa,
    dentro de cada classe - o nivel 1 da ANBIMA."""

    def test_a_classe_e_o_nivel_1_da_categoria(self):
        idx = indice_com(
            classe("rf"), classe("ac", taxa=2.0, categoria="Ações Livre"),
            *universo("Renda Fixa Simples", 1.0), *universo("Ações Livre", 2.0),
        )
        self.assertEqual(idx.avaliar(oferta("rf")).classe, "Renda Fixa")
        self.assertEqual(idx.avaliar(oferta("ac")).classe, "Ações")
        # Fora do cadastro e sem categoria: o balde de sempre, nao um palpite.
        self.assertEqual(idx.avaliar(oferta("fora")).classe, "Sem classificação")

    def test_estruturado_sem_categoria_pelo_tipo_do_cadastro(self):
        """O fundo da base interna chega so com o CNPJ - sem o nome comercial
        ("XP Malls FII") que o classificador da vitrine le."""
        idx = indice_com(
            classe("fidc", categoria=None, tipo="Classes de Cotas de Fundos FIDC"),
            classe("fif", categoria=None, tipo="Classes de Cotas de Fundos FIF"),
        )
        fidc = idx.avaliar(oferta("fidc", "CREDITO XYZ RESPONSABILIDADE LIMITADA"))
        self.assertEqual(fidc.classe, "Estruturado")
        # So o rotulo: o motivo de ficar fora do indice e o de sempre.
        self.assertEqual(fidc.motivo_fora, "fundo sem categoria oficial")
        self.assertEqual(idx.avaliar(oferta("fif")).classe, "Sem classificação")

    def test_indice_da_classe_com_amostra_minima(self):
        idx = indice_com(
            *universo("Renda Fixa Simples", 1.0), *universo("Ações Livre", 2.0),
            *[classe(f"rf{i}", taxa=2.0) for i in range(MIN_COMPARAVEIS_CLASSE_CUSTO)],
            *[classe(f"ac{i}", taxa=2.0, categoria="Ações Livre") for i in range(2)],
        )
        ofertas = [oferta(f"ac{i}") for i in range(2)]
        ofertas += [oferta(f"rf{i}") for i in range(MIN_COMPARAVEIS_CLASSE_CUSTO)]
        casa = CustoCasa("xp", "XP", len(ofertas), [idx.avaliar(o) for o in ofertas])
        por_classe = {c["classe"]: c for c in casa.por_classe()}
        # A ordem e a do eixo, nao a de chegada.
        self.assertEqual(list(por_classe), ["Renda Fixa", "Ações"])
        renda_fixa = por_classe["Renda Fixa"]
        self.assertAlmostEqual(renda_fixa["indice_mediano"], 2.0)
        self.assertAlmostEqual(renda_fixa["taxa_mediana"], 2.0)
        self.assertEqual(renda_fixa["share_caro"], 1.0)
        # Dois comparaveis nao fazem indice de classe; a contagem fica.
        self.assertIsNone(por_classe["Ações"]["indice_mediano"])
        self.assertEqual(por_classe["Ações"]["comparaveis"], 2)
        self.assertEqual(casa.to_dict()["por_classe"], casa.por_classe())

    def test_colunas_e_nota_saem_das_casas_do_recorte(self):
        casas = [
            {"por_classe": [{"classe": "Ações", "indice_mediano": 1.1},
                            {"classe": "Cambial", "indice_mediano": None}]},
            {"por_classe": [{"classe": "Renda Fixa", "indice_mediano": 0.9},
                            {"classe": "Estruturado", "indice_mediano": None}]},
        ]
        self.assertEqual(classes_com_indice(casas), ["Renda Fixa", "Ações"])
        self.assertIn("fora da tabela: Cambial e Estruturado.", nota_das_classes(casas))
        # So a primeira casa (o recorte do digest): renda fixa sai das colunas.
        self.assertEqual(classes_com_indice(casas[:1]), ["Ações"])


def _vitrine(platform: str, cnpjs: list[str]) -> ShelfSnapshot:
    return ShelfSnapshot(
        platform=platform, collected_at=datetime(2026, 10, 15, 8, 0),
        strategy=CollectorStrategy.PLAYWRIGHT, status=SnapshotStatus.OK,
        funds=[
            FundOffer(platform=platform, name=f"Fundo {c}", cnpj=c,
                      sources=[CollectorStrategy.PLAYWRIGHT.value])
            for c in cnpjs
        ],
    )


def _indice_fixo(idx: IndiceCusto) -> type[IndiceCusto]:
    """Um `classe_indice` que devolve o universo pronto, sem mercado.csv."""

    class Fixo(IndiceCusto):
        @classmethod
        def load(cls, market_dir):
            return idx

    return Fixo


class TestPrateleiraDaXpPelaBaseInterna(unittest.TestCase):
    """O custo da XP sai da lista distribuivel; o das demais casas, da vitrine
    (pedido de 16/10/2026)."""

    N = 30
    # Categoria que nao e de caixa: "Renda Fixa Simples" o `caixa.sem_caixa` tira.
    CATEGORIA = "Multimercados Livre"
    # O BTGP Tesouro Selic esta como "Duracao Livre Grau de Investimento" no
    # cadastro: e o nome que o marca como caixa.
    SELIC = "Renda Fixa Duração Livre Grau de Investimento"

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.pasta = Path(self.tmp.name) / "distribuivel" / "entrada"
        cat = self.CATEGORIA
        # Vitrines na mediana (1,0x); a base interna da XP cobra o dobro.
        self.idx = indice_com(
            *universo(cat, 1.0),
            *[classe(f"xpv{i:011d}", categoria=cat) for i in range(self.N)],
            *[classe(f"btg{i:011d}", categoria=cat) for i in range(self.N)],
            *[
                classe(f"int{i:011d}", taxa=2.0, categoria=cat, nome=f"Interno {i}")
                for i in range(self.N)
            ],
            classe("exclusiva00001", taxa=2.0, categoria=cat, exclusivo=True),
            classe("profissional1", taxa=2.0, categoria=cat, publico_alvo="PROFISSIONAL"),
            classe("selicinterno1", taxa=0.2, categoria=self.SELIC,
                   nome="XP TESOURO SELIC FIRF"),
        )
        btg = _vitrine("btg", [f"btg{i:011d}" for i in range(self.N)])
        btg.funds.append(FundOffer(
            platform="btg", name="BTGP TESOURO SELIC FI RF REF", cnpj="selicbtg00001",
            asset_class=AssetClass.RENDA_FIXA, sources=[CollectorStrategy.PLAYWRIGHT.value],
        ))
        self.snapshots = [_vitrine("xp", [f"xpv{i:011d}" for i in range(self.N)]), btg]

    def _lista(self) -> BaseDistribuivel:
        situacao = {f"int{i:011d}": DISTRIBUI for i in range(self.N)}
        situacao |= {"exclusiva00001": DISTRIBUI, "profissional1": DISTRIBUI}
        situacao |= {"selicinterno1": DISTRIBUI}
        situacao |= {"pode0000000001": PODE_DISTRIBUIR, "pode0000000002": PODE_DISTRIBUIR}
        return BaseDistribuivel(
            pasta=self.pasta, situacao=situacao,
            arquivos=[ArquivoLido(nome="lista.csv", linhas=len(situacao), validas=len(situacao),
                                  data_referencia=date(2026, 10, 16))],
        )

    def _custo(self, distribuivel, snapshots=None) -> dict:
        return build_custo(
            SimpleNamespace(market_dir=Path(self.tmp.name)), snapshots or self.snapshots,
            classe_indice=_indice_fixo(self.idx), distribuivel=distribuivel,
        )

    def test_xp_pela_lista_e_as_demais_pela_vitrine(self):
        custo = self._custo(self._lista())
        casas = {c["platform_id"]: c for c in custo["casas"]}
        self.assertEqual(casas["xp"]["fonte"], FONTE_BASE_INTERNA)
        self.assertAlmostEqual(casas["xp"]["indice_mediano"], 2.0)
        self.assertEqual(casas["btg"]["fonte"], FONTE_VITRINE)
        self.assertAlmostEqual(casas["btg"]["indice_mediano"], 1.0)
        # Publico profissional fica, como fica nas vitrines; exclusiva e
        # "pode distribuir" nao estao na prateleira.
        self.assertEqual(casas["xp"]["vitrine"], self.N + 1)
        base = custo["prateleira_da_base"]
        self.assertEqual(
            (base["fundos"], base["exclusivas"], base["caixa"], base["so_capacidade"]),
            (self.N + 1, 1, 1, 2),
        )
        self.assertEqual(
            custo["nota_prateleira"],
            "XP Investimentos pela base interna (lista distribuível de 16/10/2026): "
            "entram 31 fundos que a casa distribui. Ficam fora 1 classe exclusiva "
            "(veículo de um cotista, não produto de prateleira), 1 fundo caixa/simples "
            "e 2 fundos só com capacidade de distribuir. As demais casas, pela vitrine "
            "do site.",
        )
        # O nome vem do cadastro: a lista chega so com o CNPJ.
        nomes = {f["nome"] for f in custo["mais_caros"]}
        self.assertIn("Interno 0", nomes)

    def test_caixa_sai_da_lista_da_xp_como_sai_das_vitrines(self):
        """O Tesouro Selic da lista leva a categoria do cadastro na oferta - sem
        ela, `caixa.oferta_caixa` o deixava passar, e so as vitrines perdiam o
        caixa: a XP saia mais barata do que e."""
        custo = self._custo(self._lista())
        casas = {c["platform_id"]: c for c in custo["casas"]}
        self.assertEqual(custo["caixa_fora"], 2)
        self.assertEqual(casas["btg"]["vitrine"], self.N)
        self.assertNotIn("XP TESOURO SELIC FIRF", {f["nome"] for f in custo["mais_caros"]})

    def test_sem_lista_na_pasta_a_xp_volta_a_vitrine_e_a_nota_diz(self):
        custo = self._custo(BaseDistribuivel(pasta=self.pasta))
        xp = next(c for c in custo["casas"] if c["platform_id"] == "xp")
        self.assertEqual(xp["fonte"], FONTE_VITRINE)
        self.assertAlmostEqual(xp["indice_mediano"], 1.0)
        self.assertTrue(custo["nota_prateleira"].startswith(
            "XP Investimentos pela vitrine do site, como as demais casas: a base "
            "interna não entrou - nenhum CSV na pasta distribuivel/entrada"
        ))

    def test_sem_lista_lida_nada_muda_e_nao_ha_nota(self):
        """A previdencia nao le a lista: o custo segue so de vitrine, sem nota."""
        custo = self._custo(None)
        self.assertIsNone(custo["prateleira_da_base"])
        self.assertEqual(custo["nota_prateleira"], "")
        self.assertTrue(all(c["fonte"] == FONTE_VITRINE for c in custo["casas"]))

    def test_xp_sem_vitrine_entra_pela_lista(self):
        custo = self._custo(self._lista(), snapshots=self.snapshots[1:])
        self.assertEqual([c["platform_id"] for c in custo["casas"]], ["xp", "btg"])
        self.assertFalse(any(e["platform_id"] == "xp" for e in custo["excluded"]))


class TestCustoPorClasseNasSuperficies(unittest.TestCase):
    """A tabela casa x classe e a nota da base no digest, na planilha e no deck."""

    def _custo(self) -> dict:
        teste = TestPrateleiraDaXpPelaBaseInterna()
        teste.setUp()
        try:
            return teste._custo(teste._lista())
        finally:
            teste.tmp.cleanup()

    def test_digest_markdown_e_html(self):
        from shelf_monitor.report.digest import _html_custo_por_classe, _md_custo_por_classe

        custo = self._custo()
        md = "\n".join(_md_custo_por_classe(custo))
        self.assertIn("| Casa | Multimercados |", md)
        self.assertIn("| XP Investimentos | 2,00x (31) |", md)
        self.assertIn("| BTG Pactual | 1,00x (30) |", md)
        self.assertIn("_XP Investimentos pela base interna", md)
        html = _html_custo_por_classe(custo)
        self.assertIn("Custo por classe de fundo", html)
        self.assertIn("31 fundos", html)
        self.assertIn("pela base interna", html)

    def test_planilha_tem_a_aba_por_classe_e_a_base_de_cada_casa(self):
        from openpyxl import load_workbook

        from shelf_monitor.report.spreadsheet import render_spreadsheet
        from fixtures_comuns import temp_config

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        teste = TestPrateleiraDaXpPelaBaseInterna()
        teste.setUp()
        self.addCleanup(teste.tmp.cleanup)
        caminho = render_spreadsheet(
            config, teste.snapshots, custo=teste._custo(teste._lista())
        )
        livro = load_workbook(caminho)
        linhas = list(livro["Custo por classe"].iter_rows(values_only=True))
        topo = next(i for i, linha in enumerate(linhas) if linha[0] == "Casa")
        dados = [linha[:9] for linha in linhas[topo + 1:] if linha[0]]
        self.assertIn(
            ("XP Investimentos", "Multimercados", 31, 31, 2.0, 2.0, 1.0, 0.0,
             "base interna (lista distribuível)"),
            dados,
        )
        resumo = list(livro["Custo e acesso"].iter_rows(values_only=True))
        self.assertTrue(str(resumo[4][0]).startswith("XP Investimentos pela base interna"))
        cabecalho = next(linha for linha in resumo if linha[0] == "Casa")
        self.assertEqual(cabecalho[9], "Base da prateleira")

    def test_deck_ganha_a_tabela_por_classe(self):
        from pptx import Presentation

        from shelf_monitor.report.deck import _costs

        prs = Presentation()
        _costs(prs, [], self._custo())
        textos = [
            celula.text for forma in prs.slides[0].shapes if forma.has_table
            for linha in forma.table.rows for celula in linha.cells
        ]
        self.assertIn("Por classe", textos)
        self.assertIn("2,00x (31)", textos)


if __name__ == "__main__":
    unittest.main()
