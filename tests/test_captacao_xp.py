"""Captacao na XP ao lado das maiores captacoes e resgates (pedido de 20/10/2026).

Os testes travam:

- a leitura das bases internas (`captacao_interna`): CSV e Excel, colunas por
  nome, unidade no cabecalho, liquida ou bruta menos resgate, mes em varios
  formatos, linhas do mesmo dia somando e o dia repetido em outro arquivo
  valendo uma vez, linha ruim listada com o motivo;
- as janelas (`report.captacao_na_xp`): o mes fechado pela base mensal, o
  corrente e a semana pela diaria, ate o ultimo dia completo do informe; a
  janela que as bases nao cobrem inteira fica sem numero, com o motivo;
- o cruzamento: o fundo que a XP distribui, ou os dela com o mesmo master; o
  que ela nao distribui fica sem numero (e "pode distribuir" nao conta);
- o digest (coluna Na XP, frase na lista e nota) e a planilha (colunas da aba
  Captacao por fundo), e os dois sem coluna quando nao ha base.

O cenario e o do digest v2 (`test_digest_v2._relatorio`): 3 meses de 03/08 a
09/10/2026 e a semana de 05/10 a 09/10/2026.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

from shelf_monitor.captacao_interna import BaseCaptacaoInterna
from shelf_monitor.distribuivel import DISTRIBUI, PODE_DISTRIBUIR, BaseDistribuivel
from shelf_monitor.familias import MapaDeFamilias
from shelf_monitor.report.captacao_na_xp import (
    LIGACAO_MESMO_FUNDO,
    LIGACAO_MESMO_MASTER,
    captacao_na_xp,
)

try:
    from carteira_apoio import Cda, classe
    from fixtures_comuns import temp_config
    from test_digest_v2 import FERIADO, MI, A, B, C, E, _lista_da_casa, _relatorio, _texto
except ImportError:  # rodando da raiz, como pacote
    from tests.carteira_apoio import Cda, classe
    from tests.fixtures_comuns import temp_config
    from tests.test_digest_v2 import (
        FERIADO, MI, A, B, C, E, _lista_da_casa, _relatorio, _texto,
    )

# O espelho que a XP distribui do mesmo master de E (que so o Itau exibe).
X = "16000000000202"
MASTER_E = "16000000000999"
GESTORA = ("GAMA GESTORA LTDA", "77777777000191")


def _uteis(de: date, ate: date) -> list[date]:
    dias, dia = [], de
    while dia <= ate:
        if dia.weekday() < 5 and dia != FERIADO:
            dias.append(dia)
        dia += timedelta(days=1)
    return dias


def _escrever(pasta: Path, tipo: str, nome: str, linhas: list[str]) -> None:
    (pasta / tipo).mkdir(parents=True, exist_ok=True)
    (pasta / tipo / nome).write_text("\n".join(linhas), encoding="utf-8")


def _diaria(pasta: Path, de: date, ate: date, por_dia: dict[str, str]) -> None:
    """Uma linha por fundo e dia util, com o valor em texto brasileiro."""
    linhas = ["data;cnpj;captacao_liquida"]
    for dia in _uteis(de, ate):
        linhas += [f"{dia:%d/%m/%Y};{cnpj};{valor}" for cnpj, valor in por_dia.items()]
    _escrever(pasta, "diaria", "diaria.csv", linhas)


def _mensal(pasta: Path, linhas: list[str]) -> None:
    _escrever(pasta, "mensal", "mensal.csv", ["mes;cnpj;captacao_liquida", *linhas])


def _familias() -> MapaDeFamilias:
    """E e X investem no mesmo master: X e o espelho da XP."""
    cda = Cda("202608")
    cda.pl(MASTER_E, 1000.0).titulo(MASTER_E, 1000.0)
    cda.pl(E, 100.0).cota(E, MASTER_E, 100.0)
    cda.pl(X, 100.0).cota(X, MASTER_E, 100.0)
    return MapaDeFamilias.montar([
        classe(MASTER_E, "GAMA MASTER FIF", *GESTORA),
        classe(E, "FUNDO E FIC FIM", *GESTORA),
        classe(X, "GAMA ADVISORY XP FIC", *GESTORA),
    ], [cda.resumo()])


# --- leitura das bases ------------------------------------------------------------


class TestLeitura(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.pasta = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_unidade_do_cabecalho_vale(self):
        """"(R$ mil)" multiplica por mil: o numero sai na ordem de grandeza da CVM."""
        _escrever(self.pasta, "diaria", "a.csv", [
            "Data;CNPJ;Captação Líquida (R$ mil)",
            "05/10/2026;13.000.000/0001-01;1.500,5",
        ])
        base = BaseCaptacaoInterna.ler(self.pasta)
        self.assertAlmostEqual(base.dias[date(2026, 10, 5)][C], 1_500_500.0)

    def test_bruta_menos_resgate_com_o_sinal_que_vier(self):
        _escrever(self.pasta, "diaria", "a.csv", [
            "dt_movimento;cnpj_fundo;aplicacao;resgate",
            f"2026-10-05;{C};1000;-300",
            f"2026-10-06;{C};1000;300",
        ])
        base = BaseCaptacaoInterna.ler(self.pasta)
        self.assertEqual(base.dias[date(2026, 10, 5)][C], 700.0)
        self.assertEqual(base.dias[date(2026, 10, 6)][C], 700.0)

    def test_mesmo_dia_soma_no_arquivo_e_vale_uma_vez_entre_arquivos(self):
        """Linhas do mesmo fundo e dia somam (canal, subclasse); o dia repetido
        em outro arquivo vale o do ultimo, para reexportar nao contar duas vezes."""
        _escrever(self.pasta, "diaria", "a.csv", [
            "data;cnpj;captacao_liquida", f"05/10/2026;{C};100", f"05/10/2026;{C};50",
        ])
        _escrever(self.pasta, "diaria", "b.csv", [
            "data;cnpj;captacao_liquida", f"05/10/2026;{C};400",
        ])
        base = BaseCaptacaoInterna.ler(self.pasta)
        self.assertEqual(base.dias[date(2026, 10, 5)][C], 400.0)
        self.assertIn("diaria/b.csv repete 1 dia de arquivo anterior - valem os dele",
                      base.avisos)
        _escrever(self.pasta, "diaria", "b.csv", ["data;cnpj;captacao_liquida"])
        self.assertEqual(BaseCaptacaoInterna.ler(self.pasta).dias[date(2026, 10, 5)][C], 150.0)

    def test_linha_ruim_listada_com_o_motivo(self):
        _escrever(self.pasta, "diaria", "a.csv", [
            "data;cnpj;captacao_liquida",
            "05/10/2026;123;100",
            f"32/13/2026;{C};100",
            f"05/10/2026;{C};",
            f"06/10/2026;{C};100",
        ])
        _escrever(self.pasta, "mensal", "sem_cnpj.csv", ["mes;fundo;captacao_liquida", "ago/2026;X;1"])
        base = BaseCaptacaoInterna.ler(self.pasta)
        motivos = [(i.arquivo, i.linha, i.motivo) for i in base.invalidas]
        self.assertEqual(motivos, [
            ("diaria/a.csv", 2, "CNPJ inválido '123'"),
            ("diaria/a.csv", 3, "data ilegível '32/13/2026'"),
            ("diaria/a.csv", 4, "sem valor de captação"),
            ("mensal/sem_cnpj.csv", 0,
             "sem cabeçalho com coluna de CNPJ (cnpj, cod_cnpj, codigo_cnpj ou codigo_cvm) "
             "nas 20 primeiras linhas"),
        ])
        self.assertEqual(base.dias, {date(2026, 10, 6): {C: 100.0}})

    def test_mes_em_varios_formatos(self):
        _mensal(self.pasta, [
            f"ago/2026;{C};1", f"2026-09;{C};2", f"10/2026;{C};3", f"01/07/2026;{C};4",
            f"agosto de 2025;{C};5",
        ])
        base = BaseCaptacaoInterna.ler(self.pasta)
        self.assertEqual(
            sorted(base.meses), ["202508", "202607", "202608", "202609", "202610"]
        )
        self.assertEqual(base.no_mes([C], "202609"), 2.0)

    def test_excel_com_titulo_antes_do_cabecalho(self):
        """A exportacao do BI abre com titulo; o Excel guarda CNPJ como numero
        (sem os zeros a esquerda) e data como data."""
        pasta = Workbook()
        aba = pasta.active
        aba.append(["Captação diária - fundos"])
        aba.append([])
        aba.append(["Data", "CNPJ", "Captação Líquida (R$)"])
        aba.append([datetime(2026, 10, 5), 1300000000101, 250.0])
        aba.append([datetime(2026, 10, 5), float(C), 10.0])
        (self.pasta / "diaria").mkdir()
        pasta.save(self.pasta / "diaria" / "bi.xlsx")
        base = BaseCaptacaoInterna.ler(self.pasta)
        self.assertEqual(
            base.dias, {date(2026, 10, 5): {"01300000000101": 250.0, C: 10.0}}
        )

    def test_layout_das_bases_da_xp(self):
        """As bases da XP (20/10/2026): `cod_cnpj` em numero na coluna N e
        `captação` na coluna S. A data, sem nome conhecido, sai da primeira
        coluna com datas - e o aviso diz qual."""
        pasta = Workbook()
        aba = pasta.active
        cabecalho = [f"campo {letra}" for letra in "ABCDEFGHIJKLMNOPQRS"]
        cabecalho[1] = "Dt Posição"
        cabecalho[2] = "Assessor"
        cabecalho[13] = "cod_cnpj"
        cabecalho[18] = "Captação"
        aba.append(cabecalho)
        for dia, cnpj, valor in (
            (datetime(2026, 10, 5), int(C), 1500.25),
            (datetime(2026, 10, 5), int(C), -500.25),
            (datetime(2026, 10, 6), 1300000000101, 2000.0),
        ):
            linha = ["x"] * 19
            linha[1], linha[13], linha[18] = dia, cnpj, valor
            aba.append(linha)
        (self.pasta / "diaria").mkdir()
        pasta.save(self.pasta / "diaria" / "captacao.xlsx")
        base = BaseCaptacaoInterna.ler(self.pasta)
        self.assertEqual(aba["N1"].value, "cod_cnpj")
        self.assertEqual(aba["S1"].value, "Captação")
        self.assertEqual(base.dias, {
            date(2026, 10, 5): {C: 1000.0},
            date(2026, 10, 6): {"01300000000101": 2000.0},
        })
        self.assertIn(
            "diaria/captacao.xlsx: sem coluna com nome de data - usada a coluna "
            "'dt_posicao', a primeira com datas",
            base.avisos,
        )

    def test_captacao_sozinha_e_a_liquida_mesmo_com_resgate_ao_lado(self):
        _escrever(self.pasta, "mensal", "a.csv", [
            "mes;cod_cnpj;resgate;captacao", f"ago/2026;{C};300;1000",
        ])
        self.assertEqual(BaseCaptacaoInterna.ler(self.pasta).no_mes([C], "202608"), 1000.0)

    def test_cnpj_em_numero_no_csv(self):
        """O CSV salvo do Excel: "11000000000101,0" vale; "1,1E+13" perdeu os
        digitos e e recusado dizendo por que."""
        _escrever(self.pasta, "diaria", "a.csv", [
            "data;cod_cnpj;captacao",
            "05/10/2026;13000000000101,0;10",
            "05/10/2026;1,3E+13;10",
        ])
        base = BaseCaptacaoInterna.ler(self.pasta)
        self.assertEqual(base.dias[date(2026, 10, 5)], {C: 10.0})
        self.assertEqual(
            base.invalidas[0].motivo,
            "CNPJ em notação científica '1,3E+13' - os dígitos se perderam; use o .xlsx "
            "ou a coluna como texto",
        )

    def test_sem_arquivo_diz_onde_por(self):
        base = BaseCaptacaoInterna.ler(self.pasta / "captacao_xp")
        self.assertIn("captacao_xp/diaria", base.motivo_da_supressao())
        self.assertIn("modelos em", base.motivo_da_supressao())


def _planilha_da_xp(caminho: Path, coluna_de_tempo: str, linhas: list[tuple],
                    dimensao_a1: bool = False) -> None:
    """O layout das bases da XP (21/10/2026): o tempo na coluna A, `cod_cnpj` em
    numero na N, `Captação` liquida com sinal na S. `dimensao_a1` grava a aba como
    a exportacao do BI, com a dimensao "A1"."""
    import re
    import zipfile

    pasta = Workbook()
    aba = pasta.active
    cabecalho = [f"campo {letra}" for letra in "ABCDEFGHIJKLMNOPQRS"]
    cabecalho[0], cabecalho[13], cabecalho[18] = coluna_de_tempo, "cod_cnpj", "Captação"
    aba.append(cabecalho)
    for quando, valor in linhas:
        linha = ["x"] * 19
        linha[0], linha[13], linha[18] = quando, int(C), valor
        aba.append(linha)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    pasta.save(caminho)
    if dimensao_a1:
        copia = caminho.with_suffix(".tmp")
        with zipfile.ZipFile(caminho) as origem, zipfile.ZipFile(copia, "w") as destino:
            for item in origem.infolist():
                dados = origem.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    dados = re.sub(rb'<dimension ref="[^"]+"/>', b'<dimension ref="A1"/>', dados)
                destino.writestr(item, dados)
        copia.replace(caminho)


class TestFormatosDaXp(unittest.TestCase):
    """As bases da XP nao foram lidas na primeira execucao (21/10/2026): diaria
    com a data, `cod_cnpj` em numero e a captacao liquida com sinal; mensal com o
    mes 202607. Cada caso abaixo era recusado - quase sempre em silencio."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.pasta = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _dias(self) -> dict:
        return {d: v for d, v in BaseCaptacaoInterna.ler(self.pasta).dias.items()}

    def test_data_em_numero(self):
        """O numero serial do Excel (46311, a celula sem formato de data) e o
        aaaammdd (20261016) sao 16/10/2026."""
        _planilha_da_xp(self.pasta / "diaria" / "a.xlsx", "Data", [(46311, 1500.5)])
        _planilha_da_xp(self.pasta / "diaria" / "b.xlsx", "Data", [(20261015, -200.0)])
        self.assertEqual(self._dias(), {
            date(2026, 10, 16): {C: 1500.5}, date(2026, 10, 15): {C: -200.0},
        })

    def test_exportacao_do_bi_com_a_dimensao_a1(self):
        """O openpyxl no modo de leitura confiava na dimensao gravada e lia so A1:
        "sem linhas de dados"."""
        _planilha_da_xp(self.pasta / "diaria" / "bi.xlsx", "Data",
                        [(datetime(2026, 10, 16), 1500.5)], dimensao_a1=True)
        self.assertEqual(self._dias(), {date(2026, 10, 16): {C: 1500.5}})

    def test_qualquer_pasta_dentro_de_captacao_xp(self):
        """A pasta com acento ("Diária") e o arquivo solto na raiz eram ignorados."""
        _planilha_da_xp(self.pasta / "Diária" / "a.xlsx", "Data",
                        [(datetime(2026, 10, 16), 1500.5)])
        _planilha_da_xp(self.pasta / "mensal_xp.xlsx", "Mês", [(202607, -3000.0)])
        base = BaseCaptacaoInterna.ler(self.pasta)
        self.assertEqual(base.dias, {date(2026, 10, 16): {C: 1500.5}})
        self.assertEqual(base.meses, {"202607": {C: -3000.0}})
        self.assertEqual({a.nome: a.tipo for a in base.arquivos},
                         {"Diária/a.xlsx": "diaria", "mensal_xp.xlsx": "mensal"})

    def test_csv_com_titulo_e_negativo_contabil(self):
        _escrever(self.pasta, "diaria", "a.csv", [
            "Captação diária - fundos;;",
            ";;",
            "Data;cod_cnpj;Captação",
            f"16/10/2026;{C};(1.234,56)",
            f"15/10/2026;{C};2.000,00-",
        ])
        self.assertEqual(self._dias(), {
            date(2026, 10, 16): {C: -1234.56}, date(2026, 10, 15): {C: -2000.0},
        })

    def test_formato_que_nao_e_lido_diz_o_que_fazer(self):
        (self.pasta / "diaria").mkdir()
        (self.pasta / "diaria" / "antiga.xls").write_bytes(b"\xd0\xcf")
        base = BaseCaptacaoInterna.ler(self.pasta)
        self.assertEqual(base.invalidas[0].motivo,
                         "formato .xls não é lido - salve como .xlsx ou .csv")

    def test_diagnostico_mostra_como_cada_arquivo_foi_lido(self):
        _planilha_da_xp(self.pasta / "diaria" / "a.xlsx", "Data", [
            (datetime(2026, 10, 13), 10.0), (datetime(2026, 10, 16), 20.0),
            ("32/10/2026", 5.0),
        ])
        texto = "\n".join(BaseCaptacaoInterna.ler(self.pasta).diagnostico())
        self.assertIn("- diaria/a.xlsx: base diária, 3 linhas, 2 válidas, 1 recusada", texto)
        self.assertIn("colunas: CNPJ 'cod_cnpj', data ou mês 'data', captação 'captacao'", texto)
        self.assertIn("1 recusada - data ilegível (a primeira, linha 4: data ilegível "
                      "'32/10/2026')", texto)
        # 14 e 15/10 nao vieram: o diagnostico diz, porque sem eles a semana fica
        # sem numero.
        self.assertIn("sem 2 dias úteis no meio (feriado também aparece aqui): "
                      "14/10/2026, 15/10/2026", texto)


class TestColunaDoCnpj(unittest.TestCase):
    """O report de 20/10/2026 recusou a base diaria inteira - 297.354 linhas, "CNPJ
    nao encontrado no cadastro da CVM" -, com os dados colados no modelo: a
    coluna "cnpj" do modelo recebeu outro campo da exportacao."""

    CLIENTE = "17905206000163"  # o CNPJ do print, fora do cadastro

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.pasta = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    @staticmethod
    def _resolver(cnpj, codigo):
        """O cadastro so conhece o fundo C."""
        return [cnpj] if cnpj == C else []

    def test_entre_duas_colunas_de_cnpj_vale_a_do_cadastro(self):
        """"CNPJ" do cliente antes do `cod_cnpj` do fundo: o nome sozinho
        escolhia a do cliente, e a base inteira caia."""
        _escrever(self.pasta, "diaria", "a.csv", [
            "Data;CNPJ;cod_cnpj;Captação",
            f"16/10/2026;{self.CLIENTE};{C};1000",
            f"15/10/2026;{self.CLIENTE};{C};-200",
        ])
        base = BaseCaptacaoInterna.ler(self.pasta, self._resolver)
        self.assertEqual(base.dias, {
            date(2026, 10, 16): {C: 1000.0}, date(2026, 10, 15): {C: -200.0},
        })
        self.assertEqual(base.arquivos[0].colunas["cnpj"], "cod_cnpj")

    def test_base_recusada_inteira_diz_que_a_coluna_nao_e_a_do_fundo(self):
        _escrever(self.pasta, "diaria", "modelo_diaria.csv", [
            "data;cnpj;nome;captacao_bruta;resgate;captacao_liquida",
            f"16/10/2026;{self.CLIENTE};Cliente A;x;y;1000",
            f"16/10/2026;{int(self.CLIENTE) + 1};Cliente B;x;y;2000",
        ])
        base = BaseCaptacaoInterna.ler(self.pasta, self._resolver)
        self.assertEqual(base.dias, {})
        self.assertEqual(
            base.recusas_por_motivo("diaria/modelo_diaria.csv")[0][:2],
            ("CNPJ não encontrado no cadastro da CVM", 2),
        )
        self.assertIn(
            "diaria/modelo_diaria.csv: nenhum valor da coluna 'cnpj' está no cadastro da "
            "CVM - ela não parece trazer o CNPJ do fundo", " ".join(base.avisos),
        )
        texto = "\n".join(base.diagnostico())
        self.assertIn("  2 recusadas - CNPJ não encontrado no cadastro da CVM (a primeira, "
                      f"linha 2: CNPJ {self.CLIENTE} não encontrado no cadastro da CVM)", texto)

    def test_console_do_report_resume_as_recusas_por_motivo(self):
        import io
        from unittest import mock

        from rich.console import Console

        from shelf_monitor import main

        _escrever(self.pasta, "diaria", "modelo_diaria.csv", [
            "data;cnpj;captacao_liquida",
            *(f"16/10/2026;{int(self.CLIENTE) + i};1000" for i in range(5)),
        ])
        base = BaseCaptacaoInterna.ler(self.pasta, self._resolver)
        saida = io.StringIO()
        report = _relatorio()
        with mock.patch.object(main, "console", Console(file=saida, width=300)), \
                mock.patch.object(main, "carregar_captacao_interna", return_value=base):
            main.attach_captacao_na_xp(None, report)
        texto = saida.getvalue()
        self.assertIn(
            "Base da XP diaria/modelo_diaria.csv: 5 linhas recusadas de 5 - 5 CNPJ não "
            "encontrado no cadastro da CVM (a primeira, linha 2: CNPJ 17905206000163 não "
            "encontrado no cadastro da CVM)", texto,
        )
        self.assertIn("ela não parece trazer o CNPJ do fundo", texto)


# --- janelas e cruzamento -------------------------------------------------------


class _ComCenario(unittest.TestCase):
    """A no lista da casa, C na vitrine da XP, E so no Itau com o espelho X na
    lista, B (produto proprio do BTG) fora da XP."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.pasta = Path(self._tmp.name)
        self.lista = _lista_da_casa(**{A: "Distribuição", X: "Distribuição"})
        self.report = _relatorio(lista=self.lista, com_itau=True)
        self.familias = _familias()

    def tearDown(self):
        self._tmp.cleanup()

    def _bases_completas(self) -> None:
        # Outubro dia a dia (7 dias uteis ate 09/10), agosto e setembro no mes.
        _diaria(self.pasta, date(2026, 10, 1), date(2026, 10, 9), {
            A: "1.000.000,00", C: "100.000,00", X: "200.000,00",
        })
        _mensal(self.pasta, [
            f"ago/2026;{A};5000000", f"set/2026;{A};6000000",
            f"ago/2026;{C};1000000", f"set/2026;{C};2000000",
            f"ago/2026;{X};300000",
        ])

    def _bloco(self, familias="padrao", lista=None) -> dict:
        return captacao_na_xp(
            self.report.retail_flow, self.report.snapshots, lista or self.lista,
            self.familias if familias == "padrao" else familias,
            BaseCaptacaoInterna.ler(self.pasta),
        )


class TestJanelas(_ComCenario):
    def test_mes_fechado_pela_mensal_e_o_corrente_pela_diaria(self):
        self._bases_completas()
        bloco = self._bloco()
        self.assertEqual(
            [(p["mes"], p["fonte"]) for p in bloco["tres_meses"]["pedacos"]],
            [("202608", "mensal"), ("202609", "mensal"), ("202610", "diaria")],
        )
        a = bloco["por_cnpj"][A]
        self.assertEqual(a["xp_ligacao"], LIGACAO_MESMO_FUNDO)
        self.assertAlmostEqual(a["captacao_xp_3m"], (5 + 6 + 7) * MI)
        self.assertAlmostEqual(a["captacao_xp_semana"], 5 * MI)

    def test_base_diaria_curta_deixa_a_janela_sem_numero(self):
        """Dois meses e meio da XP ao lado de tres da CVM pareceriam uma fatia
        menor: a janela incompleta fica sem numero, com o motivo."""
        _diaria(self.pasta, date(2026, 10, 1), date(2026, 10, 7), {A: "1000000"})
        _mensal(self.pasta, [f"ago/2026;{A};5000000", f"set/2026;{A};6000000"])
        bloco = self._bloco()
        for janela in ("tres_meses", "semana"):
            self.assertFalse(bloco[janela]["medida"])
            self.assertEqual(bloco[janela]["motivo"], "a base diária da XP vai até 07/10/2026")
        self.assertIsNone(bloco["por_cnpj"][A]["captacao_xp_3m"])
        self.assertIsNone(bloco["por_cnpj"][A]["captacao_xp_semana"])

    def test_mes_fechado_sem_a_mensal_sai_da_diaria_inteira(self):
        _diaria(self.pasta, date(2026, 8, 1), date(2026, 10, 9), {A: "1000000"})
        bloco = self._bloco()
        self.assertEqual(
            {p["fonte"] for p in bloco["tres_meses"]["pedacos"]}, {"diaria"}
        )
        dias = len(_uteis(date(2026, 8, 3), date(2026, 10, 9)))
        self.assertAlmostEqual(bloco["por_cnpj"][A]["captacao_xp_3m"], dias * MI)

    def test_mes_fechado_sem_mensal_nem_diaria(self):
        _diaria(self.pasta, date(2026, 10, 1), date(2026, 10, 9), {A: "1000000"})
        _mensal(self.pasta, [f"ago/2026;{A};5000000"])
        bloco = self._bloco()
        self.assertEqual(
            bloco["tres_meses"]["motivo"],
            "sem set/2026 na base mensal, e a base diária da XP começa em 01/10/2026",
        )
        # A semana segue medida: a diaria a cobre.
        self.assertAlmostEqual(bloco["por_cnpj"][A]["captacao_xp_semana"], 5 * MI)


class TestCruzamento(_ComCenario):
    def setUp(self):
        super().setUp()
        self._bases_completas()

    def test_espelho_da_xp_mede_o_fundo_que_ela_nao_tem(self):
        """E so no Itau; a XP distribui X, do mesmo master: a captacao e a de X."""
        e = self._bloco()["por_cnpj"][E]
        self.assertEqual(e["xp_ligacao"], LIGACAO_MESMO_MASTER)
        self.assertEqual([v["cnpj"] for v in e["xp_fundos"]], [X])
        self.assertEqual(e["xp_fundos"][0]["nome"], "GAMA ADVISORY XP FIC")
        # Agosto pela mensal; setembro sem linha de X, zero; outubro 7 x 0,2.
        self.assertAlmostEqual(e["captacao_xp_3m"], (0.3 + 1.4) * MI)
        self.assertAlmostEqual(e["captacao_xp_semana"], 1.0 * MI)

    def test_sem_o_mapa_de_familias_so_o_proprio_fundo(self):
        bloco = self._bloco(familias=None)
        self.assertNotIn(E, bloco["por_cnpj"])
        self.assertIn(A, bloco["por_cnpj"])

    def test_fundo_que_a_xp_nao_distribui_fica_sem_numero(self):
        """Nem zero: B nao esta na XP. "Pode distribuir" e capacidade, nao fundo
        que o cliente compra - nao conta."""
        lista = BaseDistribuivel(
            pasta=Path("entrada"), situacao={A: DISTRIBUI, B: PODE_DISTRIBUIR},
        )
        bloco = self._bloco(lista=lista)
        self.assertNotIn(B, bloco["por_cnpj"])
        self.assertIn(A, bloco["por_cnpj"])
        self.assertIn(C, bloco["por_cnpj"])  # pela vitrine da XP

    def test_distribuido_sem_linha_na_base_captou_zero(self):
        c = BaseCaptacaoInterna.ler(self.pasta)
        self.assertEqual(c.no_mes(["31000000000000"], "202608"), 0.0)
        self.assertEqual(self._bloco()["por_cnpj"]["31000000000000"]["captacao_xp_3m"], 0.0)


# --- digest e planilha ------------------------------------------------------------


class TestSuperficies(_ComCenario):
    def setUp(self):
        super().setUp()
        self._bases_completas()
        self.report.captacao_xp = self._bloco()

    def test_digest_traz_a_captacao_na_xp_nas_listas(self):
        from shelf_monitor.report.digest import render_html, render_markdown

        md = render_markdown(self.report, None, "semanal")
        self.assertIn("- XP, BTG, Itaú - na XP +18,0 MM em 3 meses, +5,0 MM na semana\n", md)
        self.assertIn("- BTG - na XP: -\n", md)
        self.assertIn("- Itaú - na XP +1,7 MM em 3 meses, +1,0 MM na semana, pelo espelho\n", md)
        self.assertIn("_Na XP: a captação líquida do fundo na XP, pelas bases internas", md)
        texto = _texto(render_html(self.report, None, "semanal"))
        self.assertIn("Fundo Semana 3 meses Na XP Cotistas Rent. 12m Casas", texto)
        self.assertIn("+42,0 MM +1,7 MM +1,0 MM na semana · pelo espelho", texto)
        self.assertIn("+245,0 MM - +49", texto)

    def test_sem_base_o_digest_fica_como_era(self):
        from shelf_monitor.report.digest import render_html, render_markdown

        self.report.captacao_xp = {"available": False, "reason": "nenhum arquivo"}
        md = render_markdown(self.report, None, "semanal")
        self.assertNotIn("na XP", md)
        self.assertIn("Fundo Semana 3 meses Cotistas Rent. 12m Casas",
                      _texto(render_html(self.report, None, "semanal")))

    def test_planilha_traz_as_colunas_da_xp(self):
        from shelf_monitor.report.spreadsheet import render_spreadsheet

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        caminho = render_spreadsheet(
            config, self.report.snapshots, retail_flow=self.report.retail_flow,
            familias=self.familias, distribuivel=self.lista,
            captacao_xp=self.report.captacao_xp,
        )
        aba = load_workbook(caminho)["Captação por fundo"]
        cabecalho = [c.value for c in aba[7]]
        colunas = ["Fundo", "Captação líquida (R$)", "Captação líquida na XP (R$)",
                   "Captação na semana na XP (R$)", "Cruzamento com a XP",
                   "Fundo da XP no cruzamento"]
        linhas = {
            linha[0]: [linha[cabecalho.index(c)] for c in colunas[1:]]
            for linha in aba.iter_rows(min_row=8, values_only=True) if linha[0]
        }
        self.assertEqual(linhas["Fundo A FIC FIM"], [590 * MI, 18 * MI, 5 * MI, "mesmo fundo", None])
        self.assertEqual(linhas["BTG Proprio B FIC FIM"],
                         [245 * MI, None, None, "a XP não distribui", None])
        self.assertEqual(
            linhas["Fundo E FIC FIM"],
            [42 * MI, 1.7 * MI, 1 * MI, "pelo espelho", "Gama Advisory XP FIC (16.000.000/0002-02)"],
        )
        self.assertIn("Na XP: a captação líquida do fundo na XP", aba["A5"].value)
        # A grade nao passa pelo cruzamento: sem as colunas da XP.
        grade = render_spreadsheet(config, self.report.snapshots,
                                   retail_flow=self.report.retail_flow)
        cabecalho = [c.value for c in load_workbook(grade)["Captação por fundo"][7]]
        self.assertNotIn("Captação líquida na XP (R$)", cabecalho)


class TestDiagnosticoDoCruzamento(_ComCenario):
    def test_diz_quantos_fundos_das_bases_a_xp_distribui(self):
        from shelf_monitor.report.captacao_na_xp import diagnostico_do_cruzamento

        # A na lista da casa; B, produto proprio do BTG, fora da XP.
        _mensal(self.pasta, [f"ago/2026;{A};1", f"ago/2026;{B};1"])
        base = BaseCaptacaoInterna.ler(self.pasta)
        linhas = diagnostico_do_cruzamento(base, self.report.snapshots, self.lista)
        self.assertTrue(linhas[0].startswith("Cruzamento: 1 fundo das bases, de 2, na vitrine"))
        self.assertEqual(linhas[1], "  Das bases e fora da XP, por exemplo: 12.000.000/0001-01.")
        # Nenhum batendo: os dois lados, para a pessoa comparar.
        _mensal(self.pasta, [f"ago/2026;{B};1"])
        linhas = diagnostico_do_cruzamento(
            BaseCaptacaoInterna.ler(self.pasta), self.report.snapshots, self.lista
        )
        self.assertTrue(linhas[-1].startswith("  Da XP, por exemplo: "))


class TestNaExecucao(unittest.TestCase):
    def test_attach_le_a_pasta_e_anexa_o_bloco(self):
        from shelf_monitor.main import attach_captacao_na_xp

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        report = _relatorio(com_itau=True)
        attach_captacao_na_xp(config, report)
        self.assertFalse(report.captacao_xp["available"])
        self.assertIn("captacao_xp/diaria", report.captacao_xp["reason"])
        _diaria(config.captacao_xp_dir, date(2026, 8, 1), date(2026, 10, 9), {C: "10"})
        attach_captacao_na_xp(config, report)
        self.assertTrue(report.captacao_xp["available"])
        self.assertIn(C, report.captacao_xp["por_cnpj"])


if __name__ == "__main__":
    unittest.main()
