"""Testes da captacao (informe diario da CVM).

Tudo offline, contra as fixtures de tests/fixtures/market/ - os zips mensais
reproduzem o formato real (latin-1, ';', decimal com ponto) nos dois schemas,
o pre e o pos RCVM 175.

Rode com: python -m unittest tests.test_flows -v
"""
from __future__ import annotations

import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from shelf_monitor.config import Config
from shelf_monitor.market.export import write_database
from shelf_monitor.market.flows import (
    import_flows,
    load_summaries,
    month_window,
    parse_informe,
    run_flows_import,
    summarize_watchlist_managers,
    write_flows,
)
from shelf_monitor.market.importer import run_import
from shelf_monitor.market.models import MarketDatabase

FIXTURES = Path(__file__).parent / "fixtures" / "market"
TODAY = date(2026, 8, 15)  # data congelada: 2026-08 e o mes parcial

BTG = "12345678000191"
KINEA = "42345678000194"
VERDE = "70345678000170"


def temp_config() -> tuple[Config, tempfile.TemporaryDirectory]:
    tmp = tempfile.TemporaryDirectory()
    config = Config(data_dir=tmp.name)
    config.ensure_dirs()
    return config, tmp


class TestMonthWindow(unittest.TestCase):
    def test_window_ends_at_the_current_month(self):
        self.assertEqual(
            month_window(date(2026, 1, 15), 3), ["2025-11", "2025-12", "2026-01"]
        )

    def test_default_window_is_13_closed_months_plus_the_current(self):
        """O 13o mes fechado e a base das janelas de 12m (cota e cotistas)."""
        window = month_window(TODAY)
        self.assertEqual(len(window), 14)
        self.assertEqual(window[0], "2025-07")
        self.assertEqual(window[-1], "2026-08")


class TestParseInforme(unittest.TestCase):
    def setUp(self):
        self.flows, self.unknown = parse_informe(
            FIXTURES / "inf_diario_fi_202607.zip", "2026-07"
        )

    def test_masked_cnpj_matches_the_unmasked_registry(self):
        """O informe pode vir com mascara - mesma causa n.1 de join vazio."""
        self.assertIn(BTG, self.flows)

    def test_inflow_and_outflow_sum_over_the_month(self):
        flow = self.flows[BTG]
        self.assertAlmostEqual(flow.inflow, 3_500_000.0)
        self.assertAlmostEqual(flow.outflow, 750_000.0)
        self.assertAlmostEqual(flow.net, 2_750_000.0)

    def test_net_worth_is_the_last_date_not_the_file_order(self):
        """As linhas de 31/07 vem antes da de 30/07 na fixture de proposito."""
        flow = self.flows[BTG]
        self.assertEqual(flow.last_date, date(2026, 7, 31))
        self.assertAlmostEqual(flow.net_worth, 1_000_000_000.0)

    def test_subclasses_of_the_same_day_sum_their_net_worth(self):
        """Subclasses partilham o CNPJ da classe; o PL do dia e a soma delas."""
        self.assertEqual(self.flows[BTG].shareholders, 11_000)

    def test_days_count_distinct_dates(self):
        self.assertEqual(self.flows[BTG].days, 2)
        self.assertEqual(self.flows[KINEA].days, 1)

    def test_redemption_only_class_goes_negative(self):
        self.assertAlmostEqual(self.flows[KINEA].net, -3_000_000.0)

    def test_unknown_column_is_reported_not_swallowed(self):
        self.assertEqual(self.unknown, ["ID_INFORME"])

    def test_old_schema_still_parses(self):
        """Uma janela de 14 meses pode atravessar a virada da RCVM 175."""
        flows, unknown = parse_informe(
            FIXTURES / "inf_diario_fi_202412.zip", "2024-12"
        )
        self.assertAlmostEqual(flows[BTG].inflow, 100_000.0)
        self.assertAlmostEqual(flows[VERDE].net, 3_000_000.0)
        self.assertEqual(unknown, [])
        # Sem ID_SUBCLASSE a serie e a propria classe.
        self.assertAlmostEqual(flows[BTG].quota_end, 1.4)
        self.assertEqual(flows[BTG].quota_series, "")

    def test_quota_comes_from_the_representative_series_not_a_sum(self):
        """Subclasses tem cotas distintas: a classe recebe a de maior PL."""
        flow = self.flows[BTG]
        self.assertAlmostEqual(flow.quota_end, 1.523456)
        self.assertEqual(flow.quota_series, "3001")

    def test_daily_log_returns_are_accumulated_per_series(self):
        import math

        flow = self.flows[BTG]
        self.assertEqual(flow.ret_n, 1)
        self.assertAlmostEqual(flow.ret_sum, math.log(1.523456 / 1.522))
        self.assertAlmostEqual(flow.ret_sumsq, math.log(1.523456 / 1.522) ** 2)
        # Um unico dia informado nao forma par.
        self.assertEqual(self.flows[KINEA].ret_n, 0)
        self.assertAlmostEqual(self.flows[KINEA].quota_end, 2.1)

    def test_dirty_quotas_leave_a_gap_instead_of_a_fake_return(self):
        """Cota zero/vazia sai da serie; salto absurdo e descartado."""
        import tempfile
        import zipfile

        header = "TP_FUNDO_CLASSE;CNPJ_FUNDO_CLASSE;ID_SUBCLASSE;DT_COMPTC;VL_TOTAL;VL_QUOTA;VL_PATRIM_LIQ;CAPTC_DIA;RESG_DIA;NR_COTST"
        rows = [
            # Serie com cota suja no ultimo dia: quota_end fica None.
            "CLASSES - FIF;11111111000111;;2026-05-27;10;1.500000;10;0;0;1",
            "CLASSES - FIF;11111111000111;;2026-05-28;10;0;10;0;0;1",
            "CLASSES - FIF;11111111000111;;2026-05-29;10;;10;0;0;1",
            # Serie com grupamento de cota: o salto nao vira volatilidade.
            "CLASSES - FIF;22222222000122;;2026-05-28;10;1.000000;10;0;0;1",
            "CLASSES - FIF;22222222000122;;2026-05-29;10;3.000000;10;0;0;1",
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "inf_diario_fi_202605.zip"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr(
                    "inf_diario_fi_202605.csv", "\n".join([header, *rows])
                )
            flows, _ = parse_informe(path, "2026-05")

        dirty = flows["11111111000111"]
        self.assertIsNone(dirty.quota_end)
        self.assertEqual(dirty.ret_n, 0)
        self.assertEqual(dirty.days, 3)

        jump = flows["22222222000122"]
        self.assertAlmostEqual(jump.quota_end, 3.0)
        self.assertEqual(jump.ret_n, 0)


class TestImportFlows(unittest.TestCase):
    def setUp(self):
        self.config, self._tmp = temp_config()
        self.addCleanup(self._tmp.cleanup)
        self.result = import_flows(
            self.config, months=4, local_dir=FIXTURES, today=TODAY
        )

    def test_missing_month_is_a_note_not_an_exception(self):
        """O caso tipico: o zip do mes ainda nao publicado pela CVM."""
        self.assertEqual(
            self.result.window, ["2026-05", "2026-06", "2026-07", "2026-08"]
        )
        self.assertEqual(self.result.missing_months, ["2026-05"])

    def test_current_month_is_partial_and_closed_months_are_not(self):
        self.assertFalse(self.result.flows[(BTG, "2026-06")].partial)
        self.assertFalse(self.result.flows[(BTG, "2026-07")].partial)
        self.assertTrue(self.result.flows[(BTG, "2026-08")].partial)

    def test_counts_reflect_what_actually_came(self):
        counts = self.result.counts()
        self.assertEqual(counts["meses"], 3)
        self.assertEqual(counts["classes_com_captacao"], 2)


class TestWriteFlows(unittest.TestCase):
    def setUp(self):
        self.config, self._tmp = temp_config()
        self.addCleanup(self._tmp.cleanup)
        self.result = import_flows(
            self.config, months=4, local_dir=FIXTURES, today=TODAY
        )
        self.path = write_flows(self.config, self.result)

    def test_rerun_writes_identical_bytes(self):
        """Sem timestamp no CSV, um rerun sem mudanca nao gera diff."""
        first = self.path.read_bytes()
        write_flows(self.config, self.result)
        self.assertEqual(self.path.read_bytes(), first)

    def test_file_has_bom_and_semicolon_like_the_other_tables(self):
        raw = self.path.read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        self.assertIn(b"cnpj_classe;mes;captacao", raw)
        self.assertIn(b"cota_fim_mes;cota_serie;ret_diario_soma", raw)

    def test_quota_keeps_its_precision_in_the_csv(self):
        """Duas casas decimais zerariam um retorno diario inteiro."""
        text = self.path.read_text(encoding="utf-8-sig")
        july = next(line for line in text.splitlines() if line.startswith("12.345.678/0001-91;2026-07"))
        self.assertIn(";1.523456;3001;", july)

    def test_manifest_merge_preserves_the_registry_section(self):
        """Dois produtores, um manifesto: nenhum apaga a secao do outro."""
        manifest_path = self.config.market_dir / "_manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        self.assertIn("captacao", manifest)
        self.assertEqual(manifest["captacao"]["missing_months"], ["2026-05"])

        # A reimportacao do cadastro regrava o manifesto inteiro - e nao pode
        # perder a secao de captacao no processo.
        write_database(self.config, MarketDatabase(imported_at=TODAY))
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        self.assertIn("captacao", manifest)
        self.assertIn("imported_at", manifest)

    def test_run_flows_import_returns_none_when_nothing_came(self):
        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        empty = tempfile.TemporaryDirectory()
        self.addCleanup(empty.cleanup)
        result, path = run_flows_import(
            config, months=2, local_dir=Path(empty.name)
        )
        self.assertIsNone(path)
        self.assertEqual(len(result.missing_months), 2)


class TestLoadSummaries(unittest.TestCase):
    def setUp(self):
        self.config, self._tmp = temp_config()
        self.addCleanup(self._tmp.cleanup)
        result = import_flows(self.config, months=3, local_dir=FIXTURES, today=TODAY)
        write_flows(self.config, result)
        self.summaries = load_summaries(self.config)

    def test_windows_cover_only_closed_months(self):
        """Mes de 3 dias uteis nao entra na mesma conta que meses de 21."""
        btg = self.summaries[BTG]
        self.assertAlmostEqual(btg.net_1m, 2_750_000.0)
        # Junho (+50 mil) e julho (+2,75 mi) fechados; agosto parcial fora.
        self.assertAlmostEqual(btg.net_12m, 2_800_000.0)
        self.assertEqual(btg.months_12m, 2)

    def test_partial_month_is_exposed_separately(self):
        self.assertAlmostEqual(self.summaries[BTG].current_partial, 300_000.0)
        self.assertAlmostEqual(self.summaries[KINEA].current_partial, -1_000_000.0)

    def test_net_worth_comes_from_the_freshest_informe(self):
        """O PL de agosto (parcial) e mais fresco que o de julho fechado."""
        btg = self.summaries[BTG]
        self.assertAlmostEqual(btg.net_worth, 1_001_000_000.0)
        self.assertEqual(btg.as_of, date(2026, 8, 3))
        self.assertEqual(btg.shareholders, 11_050)

    def test_windows_are_anchored_on_the_last_closed_month(self):
        self.assertEqual(self.summaries[BTG].window_end, "2026-07")
        # net_3m cobre junho e julho (k=1 e k=0).
        self.assertAlmostEqual(self.summaries[BTG].net_3m, 2_800_000.0)

    def test_return_needs_the_base_quota_exactly_k_months_before(self):
        """Estrito: sem a cota-base, None - nunca um '12m' que e 2 meses."""
        btg = self.summaries[BTG]
        self.assertAlmostEqual(btg.return_1m, 1.523456 / 1.51 - 1)
        self.assertIsNone(btg.return_3m)
        self.assertIsNone(btg.return_12m)
        self.assertEqual(btg.quota_months, 2)

    def test_volatility_needs_enough_daily_pairs(self):
        btg = self.summaries[BTG]
        self.assertEqual(btg.vol_obs, 2)  # 1 par em junho + 1 em julho
        self.assertIsNone(btg.vol_12m)

    def test_ticket_and_shareholder_deltas(self):
        btg = self.summaries[BTG]
        self.assertAlmostEqual(btg.ticket, 1_001_000_000.0 / 11_050)
        self.assertEqual(btg.shareholders_delta_1m, 11_000 - 10_810)
        self.assertIsNone(btg.shareholders_delta_12m)
        self.assertEqual(self.summaries[KINEA].shareholders_delta_1m, -5)

    def test_missing_base_returns_none(self):
        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        self.assertIsNone(load_summaries(config))


class TestSummaryWindows(unittest.TestCase):
    """Formulas e bordas da leitura, com CSV inline (sem manifesto)."""

    HEADER = (
        "cnpj_classe;mes;captacao;resgates;captacao_liquida;pl_fim_mes;"
        "cotistas_fim_mes;data_ultimo_informe;dias_informados;mes_parcial;"
        "cota_fim_mes;cota_serie;ret_diario_soma;ret_diario_soma2;ret_diario_n"
    )

    def _load(self, rows: list[str], header: str | None = None):
        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        (config.market_dir / "captacao.csv").write_text(
            "\n".join([header or self.HEADER, *rows]) + "\n", encoding="utf-8-sig"
        )
        return load_summaries(config)

    def test_volatility_is_the_pooled_daily_stdev_annualized(self):
        import math

        rows = [
            f"11111111000111;2026-0{m};0;0;0;100;10;2026-0{m}-28;20;N;1.0;;0.1;0.002;100"
            for m in (5, 6, 7)
        ]
        summary = self._load(rows)["11111111000111"]
        n, s1, s2 = 300, 0.3, 0.006
        expected = math.sqrt((s2 - s1 * s1 / n) / (n - 1)) * math.sqrt(252)
        self.assertEqual(summary.vol_obs, 300)
        self.assertAlmostEqual(summary.vol_12m, expected)

    def test_too_few_pairs_give_no_volatility(self):
        rows = ["11111111000111;2026-07;0;0;0;100;10;2026-07-31;20;N;1.0;;0.1;0.002;59"]
        self.assertIsNone(self._load(rows)["11111111000111"].vol_12m)

    def test_return_12m_uses_the_13th_closed_month_as_base(self):
        rows = [
            "11111111000111;2025-07;0;0;0;100;10;2025-07-31;20;N;1.00;;0;0;0",
            "11111111000111;2026-04;0;0;0;100;10;2026-04-30;20;N;1.05;;0;0;0",
            "11111111000111;2026-06;0;0;0;100;10;2026-06-30;20;N;1.08;;0;0;0",
            "11111111000111;2026-07;0;0;0;100;12;2026-07-31;20;N;1.10;;0;0;0",
        ]
        summary = self._load(rows)["11111111000111"]
        self.assertAlmostEqual(summary.return_12m, 0.10)
        self.assertAlmostEqual(summary.return_3m, 1.10 / 1.05 - 1)
        self.assertAlmostEqual(summary.return_1m, 1.10 / 1.08 - 1)
        self.assertEqual(summary.shareholders_delta_12m, 2)
        self.assertEqual(summary.shareholders_start_12m, 10)

    def test_a_different_series_never_produces_a_return(self):
        """A representativa mudou de subclasse: cotas incomparaveis."""
        rows = [
            "11111111000111;2026-06;0;0;0;100;10;2026-06-30;20;N;1.00;3001;0;0;0",
            "11111111000111;2026-07;0;0;0;100;10;2026-07-31;20;N;1.10;3002;0;0;0",
        ]
        self.assertIsNone(self._load(rows)["11111111000111"].return_1m)

    def test_class_without_the_anchor_month_has_no_1m_window(self):
        """Janela ancorada na base, nao no ultimo mes de cada classe."""
        rows = [
            "11111111000111;2026-07;0;0;0;100;10;2026-07-31;20;N;1.10;;0;0;0",
            "22222222000122;2026-06;100;0;100;100;10;2026-06-30;20;N;1.00;;0;0;0",
        ]
        summaries = self._load(rows)
        stale = summaries["22222222000122"]
        self.assertEqual(stale.window_end, "2026-07")
        self.assertIsNone(stale.net_1m)
        self.assertAlmostEqual(stale.net_12m, 100.0)
        self.assertIsNone(stale.return_1m)

    def test_old_csv_without_quota_columns_still_loads(self):
        """Base importada antes da cota: captacao e ticket seguem, retorno nao."""
        header = (
            "cnpj_classe;mes;captacao;resgates;captacao_liquida;pl_fim_mes;"
            "cotistas_fim_mes;data_ultimo_informe;dias_informados;mes_parcial"
        )
        rows = ["11111111000111;2026-07;500;0;500;1000;10;2026-07-31;20;N"]
        summary = self._load(rows, header)["11111111000111"]
        self.assertAlmostEqual(summary.net_1m, 500.0)
        self.assertAlmostEqual(summary.ticket, 100.0)
        self.assertIsNone(summary.return_1m)
        self.assertEqual(summary.quota_months, 0)


class TestCdi(unittest.TestCase):
    """A serie do CDI como benchmark, do JSON do SGS ao acumulado por janela."""

    def setUp(self):
        from shelf_monitor.market.benchmarks import run_cdi_import

        self.config, self._tmp = temp_config()
        self.addCleanup(self._tmp.cleanup)
        self.path, self.notes = run_cdi_import(
            self.config, months=4, local_dir=FIXTURES, today=TODAY
        )

    def test_import_writes_the_series_and_its_manifest_section(self):
        from shelf_monitor.market.benchmarks import load_cdi

        self.assertTrue(self.path.is_file())
        rates = load_cdi(self.config)
        self.assertEqual(len(rates), 22 + 23)  # dias uteis de jun + jul/2026
        self.assertAlmostEqual(rates[date(2026, 7, 31)], 0.05)
        manifest = json.loads(
            (self.config.market_dir / "_manifest.json").read_text(encoding="utf-8")
        )
        self.assertEqual(manifest["cdi"]["count"], 45)
        self.assertEqual(manifest["cdi"]["last"], "2026-07-31")

    def test_window_accumulates_only_complete_months(self):
        from shelf_monitor.market.benchmarks import cdi_block

        block = cdi_block(self.config, "2026-07")
        self.assertAlmostEqual(block["windows"]["1m"], 1.0005**23 - 1)
        # Maio nao existe na fixture: 3m e 12m nao saem pela metade.
        self.assertIsNone(block["windows"]["3m"])
        self.assertIsNone(block["windows"]["12m"])
        self.assertEqual(block["window_end"], "2026-07")
        self.assertEqual(block["last_rate"], "2026-07-31")

    def test_no_series_or_no_window_end_means_no_benchmark(self):
        from shelf_monitor.market.benchmarks import cdi_block

        self.assertIsNone(cdi_block(self.config, None))
        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        self.assertIsNone(cdi_block(config, "2026-07"))

    def test_unreachable_source_keeps_the_previous_series(self):
        from shelf_monitor.market.benchmarks import load_cdi, run_cdi_import

        empty = tempfile.TemporaryDirectory()
        self.addCleanup(empty.cleanup)
        path, notes = run_cdi_import(
            self.config, months=4, local_dir=Path(empty.name),
            today=date(2026, 9, 15),
        )
        self.assertEqual(path, self.path)
        self.assertEqual(len(load_cdi(self.config)), 45)
        self.assertTrue(any("mantido" in note for note in notes))

    def test_nothing_on_disk_and_no_source_returns_none(self):
        from shelf_monitor.market.benchmarks import run_cdi_import

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        empty = tempfile.TemporaryDirectory()
        self.addCleanup(empty.cleanup)
        path, notes = run_cdi_import(config, months=4, local_dir=Path(empty.name))
        self.assertIsNone(path)
        self.assertTrue(notes)

    def test_series_up_to_date_skips_the_fetch(self):
        from shelf_monitor.market.benchmarks import run_cdi_import

        _, notes = run_cdi_import(
            self.config, months=4, local_dir=FIXTURES, today=date(2026, 7, 31)
        )
        self.assertTrue(any("em dia" in note for note in notes))


class TestFlowsBaseStatus(unittest.TestCase):
    def test_csv_without_quota_columns_is_incomplete(self):
        from shelf_monitor.main import flows_base_status

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        self.assertEqual(flows_base_status(config)[0], "ausente")
        (config.market_dir / "captacao.csv").write_text(
            "cnpj_classe;mes;captacao\n", encoding="utf-8-sig"
        )
        state, detail = flows_base_status(config)
        self.assertEqual(state, "incompleta")
        self.assertIn("market --skip-fees", detail)


class TestWatchlistManagers(unittest.TestCase):
    def setUp(self):
        self.config, self._tmp = temp_config()
        self.addCleanup(self._tmp.cleanup)
        run_import(self.config, local_dir=FIXTURES)
        result = import_flows(self.config, months=3, local_dir=FIXTURES, today=TODAY)
        write_flows(self.config, result)
        self.summary = summarize_watchlist_managers(self.config)

    def test_houses_aggregate_their_managers_flows(self):
        houses = {h["house"]: h for h in self.summary["houses"]}
        self.assertAlmostEqual(houses["btg"]["net_12m"], 2_800_000.0)
        self.assertAlmostEqual(houses["itau"]["net_12m"], -3_000_000.0)
        self.assertEqual(
            houses["itau"]["managers"][0]["name"], "KINEA INVESTIMENTOS LTDA"
        )

    def test_market_totals_cover_every_class_with_flows(self):
        market = self.summary["market"]
        self.assertEqual(market["classes"], 2)
        self.assertAlmostEqual(market["net_12m"], -200_000.0)

    def test_coverage_is_explicit(self):
        coverage = self.summary["coverage"]
        self.assertEqual(coverage["classes_total"], 2)
        self.assertEqual(coverage["classes_com_gestora"], 2)

    def test_as_of_labels_the_freshest_informe(self):
        self.assertEqual(self.summary["as_of"], "2026-08-03")
        self.assertEqual(self.summary["window_months"], 2)
        self.assertEqual(self.summary["window_end"], "2026-07")

    def test_houses_carry_their_shareholders(self):
        houses = {h["house"]: h for h in self.summary["houses"]}
        self.assertEqual(houses["btg"]["shareholders"], 11_050)
        self.assertEqual(houses["itau"]["shareholders"], 208)
        self.assertEqual(self.summary["market"]["shareholders"], 11_258)
        # Sem a base de 12 meses, o delta nao fala por ninguem.
        self.assertEqual(houses["btg"]["classes_with_delta"], 0)

    def test_missing_either_base_returns_none(self):
        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        self.assertIsNone(summarize_watchlist_managers(config))


class TestFlowsEnrichment(unittest.TestCase):
    """A captacao chegando as ofertas da prateleira, via enrich."""

    CAPTACAO = "\n".join(
        [
            "cnpj_classe;mes;captacao;resgates;captacao_liquida;pl_fim_mes;"
            "cotistas_fim_mes;data_ultimo_informe;dias_informados;mes_parcial;"
            "cota_fim_mes;cota_serie;ret_diario_soma;ret_diario_soma2;ret_diario_n",
            "63.446.494/0001-52;2026-06;0.00;0.00;0.00;"
            "11000000.00;1150;2026-06-30;21;N;1.00;;0.05;0.0005;20",
            "63.446.494/0001-52;2026-07;1000000.00;250000.00;750000.00;"
            "11584330.36;1200;2026-07-31;22;N;1.05;;0.05;0.0005;21",
            "63.446.494/0001-52;2026-08;100000.00;0.00;100000.00;"
            "11684330.36;1210;2026-08-03;2;S;1.06;;0.01;0.0001;1",
            "77.777.777/0001-77;2026-07;500000.00;0.00;500000.00;"
            "5000000.00;100;2026-07-31;22;N;2.0;;0;0;0",
            "88.888.888/0001-88;2026-07;0.00;200000.00;-200000.00;"
            "7000000.00;50;2026-07-30;21;N;3.0;;0;0;0",
        ]
    )

    def setUp(self):
        import shutil

        self.config, self._tmp = temp_config()
        self.addCleanup(self._tmp.cleanup)
        enrich_fixtures = Path(__file__).parent / "fixtures" / "enrich"
        for name in ("mercado.csv", "subclasses.csv"):
            shutil.copy(enrich_fixtures / name, self.config.market_dir / name)
        (self.config.market_dir / "captacao.csv").write_text(
            self.CAPTACAO + "\n", encoding="utf-8-sig"
        )

    def _snapshot(self, funds):
        from datetime import datetime

        from shelf_monitor.models import (
            CollectorStrategy,
            ShelfSnapshot,
            SnapshotStatus,
        )

        return ShelfSnapshot(
            platform="btg", collected_at=datetime(2026, 8, 31),
            strategy=CollectorStrategy.OPEN_FINANCE, status=SnapshotStatus.OK,
            funds=funds,
        )

    def test_class_match_stamps_the_windows(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="btg", name="A1 D30", cnpj="63446494000152")
        stats = enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertAlmostEqual(fund.extra["captacao_liquida_1m"], 750_000.0)
        self.assertAlmostEqual(fund.extra["captacao_liquida_12m"], 750_000.0)
        self.assertEqual(fund.extra["captacao_janela_meses"], 2)
        self.assertEqual(fund.extra["captacao_data_base"], "2026-08-03")
        self.assertEqual(stats.flows_matched, 1)
        self.assertEqual(stats.by_platform["btg"]["flows"], 1)

    def test_class_match_stamps_returns_ticket_and_shareholders(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="btg", name="A1 D30", cnpj="63446494000152")
        stats = enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertAlmostEqual(fund.extra["rentabilidade_1m"], 0.05)
        self.assertNotIn("rentabilidade_12m", fund.extra)  # sem cota-base
        self.assertEqual(fund.extra["rentabilidade_janela_meses"], 2)
        self.assertNotIn("volatilidade_12m", fund.extra)  # 41 pares < 60
        self.assertAlmostEqual(fund.extra["ticket_medio"], 11_684_330.36 / 1210)
        self.assertEqual(fund.extra["cotistas_var_1m"], 50)
        self.assertNotIn("cotistas_var_12m", fund.extra)
        self.assertEqual(stats.returns_matched, 1)
        self.assertEqual(stats.by_platform["btg"]["returns"], 1)

    def test_informe_net_worth_is_fresher_than_the_registry_one(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="btg", name="A1 D30", cnpj="63446494000152")
        enrich_snapshots(self.config, [self._snapshot([fund])])

        # pl_cvm (cadastro) e pl_informe (informe) convivem: o segundo e o
        # mais fresco, o primeiro fica como referencia do registro.
        self.assertAlmostEqual(fund.extra["pl_informe"], 11_684_330.36)
        self.assertAlmostEqual(fund.extra["pl_cvm"], 11_584_330.36)
        self.assertEqual(fund.extra["cotistas_informe"], 1210)

    def test_envelope_offer_sums_its_classes(self):
        """Oferta que so publicou o CNPJ do fundo: captacao = soma das classes."""
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        fund = FundOffer(platform="btg", name="Fundo Multi", cnpj="66666666000166")
        stats = enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertAlmostEqual(fund.extra["captacao_liquida_12m"], 300_000.0)
        self.assertAlmostEqual(fund.extra["pl_informe"], 12_000_000.0)
        # O minimo das datas: a serie combinada so esta completa ate onde a
        # classe mais atrasada informou.
        self.assertEqual(fund.extra["captacao_data_base"], "2026-07-30")
        self.assertEqual(stats.flows_matched, 1)
        # Cotas de classes diferentes nao se combinam: sem retorno no envelope;
        # cotistas somam e o ticket deriva do PL/cotistas combinados.
        self.assertNotIn("rentabilidade_1m", fund.extra)
        self.assertNotIn("rentabilidade_janela_meses", fund.extra)
        self.assertEqual(fund.extra["cotistas_informe"], 150)
        self.assertAlmostEqual(fund.extra["ticket_medio"], 12_000_000.0 / 150)
        self.assertEqual(stats.returns_matched, 0)

    def test_without_captacao_base_nothing_is_stamped(self):
        from shelf_monitor.enrich import enrich_snapshots
        from shelf_monitor.models import FundOffer

        (self.config.market_dir / "captacao.csv").unlink()
        fund = FundOffer(platform="btg", name="A1 D30", cnpj="63446494000152")
        stats = enrich_snapshots(self.config, [self._snapshot([fund])])

        self.assertNotIn("captacao_liquida_12m", fund.extra)
        self.assertEqual(stats.flows_matched, 0)
        self.assertEqual(stats.matched, 1)  # o join do cadastro segue normal


class TestFlowsBlock(unittest.TestCase):
    """O bloco 'flows' do payload - fonte unica das quatro superficies."""

    def _snapshot(self, platform, funds):
        from datetime import datetime

        from shelf_monitor.models import (
            CollectorStrategy,
            ShelfSnapshot,
            SnapshotStatus,
        )

        return ShelfSnapshot(
            platform=platform, collected_at=datetime(2026, 8, 31),
            strategy=CollectorStrategy.OPEN_FINANCE, status=SnapshotStatus.OK,
            funds=funds,
        )

    def _offer(self, platform, name, cnpj, net_12m, net_3m=None, pl=None):
        from shelf_monitor.models import FundOffer

        return FundOffer(
            platform=platform, name=name, cnpj=cnpj,
            extra={
                "captacao_liquida_12m": net_12m,
                "captacao_liquida_3m": net_3m if net_3m is not None else net_12m,
                "captacao_janela_meses": 12,
                "captacao_data_base": "2026-08-29",
                "pl_informe": pl,
            },
        )

    def _report(self, snapshots, flows=None):
        from datetime import datetime

        from shelf_monitor.models import RunReport

        return RunReport(
            generated_at=datetime(2026, 8, 31), snapshots=snapshots, flows=flows
        )

    def test_platform_sums_carry_their_coverage(self):
        from shelf_monitor.models import FundOffer
        from shelf_monitor.report.aggregate import build_flows_block

        snapshot = self._snapshot(
            "btg",
            [
                self._offer("btg", "Alpha", "11111111000111", 1_000_000.0),
                self._offer("btg", "Beta", "22222222000122", -400_000.0),
                FundOffer(platform="btg", name="Sem informe", cnpj="33333333000133"),
            ],
        )
        block = build_flows_block(self._report([snapshot]))

        platform = block["platforms"][0]
        self.assertAlmostEqual(platform["net_12m"], 600_000.0)
        self.assertEqual((platform["matched"], platform["total"]), (2, 3))
        self.assertTrue(block["available"])
        self.assertEqual(block["as_of"], "2026-08-29")
        self.assertEqual(block["window_months"], 12)

    def test_top_movers_dedupe_the_same_fund_across_shelves(self):
        """Fundo em duas prateleiras e UM fluxo no mercado, nao dois."""
        from shelf_monitor.report.aggregate import build_flows_block

        block = build_flows_block(
            self._report(
                [
                    self._snapshot(
                        "btg",
                        [self._offer("btg", "Alpha", "11111111000111", 1_000_000.0)],
                    ),
                    self._snapshot(
                        "inter",
                        [
                            self._offer(
                                "inter", "Alpha", "11111111000111", 1_000_000.0
                            ),
                            self._offer(
                                "inter", "Gamma", "44444444000144", -2_000_000.0
                            ),
                        ],
                    ),
                ]
            )
        )

        self.assertEqual(len(block["top_inflow"]), 1)
        self.assertEqual(
            sorted(block["top_inflow"][0]["platforms"]), ["BTG Pactual", "Banco Inter"]
        )
        self.assertEqual(len(block["top_outflow"]), 1)
        self.assertEqual(block["top_outflow"][0]["cnpj"], "44444444000144")

    def test_without_any_flow_data_the_block_says_unavailable(self):
        from shelf_monitor.models import FundOffer
        from shelf_monitor.report.aggregate import build_flows_block

        snapshot = self._snapshot(
            "btg", [FundOffer(platform="btg", name="Sem informe")]
        )
        block = build_flows_block(self._report([snapshot]))
        self.assertFalse(block["available"])
        self.assertIsNone(block["managers"])

    def test_managers_view_rides_along_when_present(self):
        from shelf_monitor.report.aggregate import build_flows_block

        managers = {"houses": [], "market": {}, "as_of": "2026-08-29"}
        block = build_flows_block(self._report([], flows=managers))
        self.assertTrue(block["available"])
        self.assertIs(block["managers"], managers)


class TestPerformanceBlock(unittest.TestCase):
    """Rentabilidade e risco por casa e por classe, no payload unico."""

    def _snapshot(self, platform, funds):
        from datetime import datetime

        from shelf_monitor.models import (
            CollectorStrategy,
            ShelfSnapshot,
            SnapshotStatus,
        )

        return ShelfSnapshot(
            platform=platform, collected_at=datetime(2026, 8, 31),
            strategy=CollectorStrategy.OPEN_FINANCE, status=SnapshotStatus.OK,
            funds=funds,
        )

    def _offer(self, platform, name, cnpj, ret, vol=None, anbima="Multimercados Macro"):
        from shelf_monitor.models import FundOffer

        extra = {
            "rentabilidade_12m": ret, "rentabilidade_janela_meses": 13,
            "captacao_data_base": "2026-07-31", "classificacao_anbima": anbima,
        }
        if vol is not None:
            extra["volatilidade_12m"] = vol
        return FundOffer(platform=platform, name=name, cnpj=cnpj, extra=extra)

    def _report(self, snapshots, benchmarks=None):
        from datetime import datetime

        from shelf_monitor.models import RunReport

        return RunReport(
            generated_at=datetime(2026, 8, 31), snapshots=snapshots,
            benchmarks=benchmarks,
        )

    def test_platform_medians_and_cdi_share(self):
        from shelf_monitor.report.aggregate import build_performance_block

        funds = [
            self._offer("btg", "A", "11111111000111", 0.15, vol=0.10),
            self._offer("btg", "B", "22222222000122", 0.08, vol=0.05, anbima="RF DI"),
            self._offer("btg", "C", "33333333000133", 0.12),
        ]
        cdi = {"window_end": "2026-07", "windows": {"12m": 0.10, "3m": None, "1m": 0.008}}
        block = build_performance_block(self._report([self._snapshot("btg", funds)], cdi))

        platform = block["platforms"][0]
        self.assertAlmostEqual(platform["return_median"], 0.12)
        self.assertAlmostEqual(platform["vol_median"], 0.075)
        self.assertEqual(platform["above_cdi"], 2)
        self.assertAlmostEqual(platform["above_cdi_share"], 2 / 3)
        self.assertAlmostEqual(platform["median_pct_cdi"], 1.2)
        self.assertEqual(block["window_end"], "2026-07")
        self.assertEqual(block["window_months"], 13)
        self.assertTrue(block["available"])

    def test_without_benchmark_nothing_claims_to_beat_the_cdi(self):
        from shelf_monitor.report.aggregate import build_performance_block

        funds = [self._offer("btg", "A", "11111111000111", 0.15)]
        block = build_performance_block(self._report([self._snapshot("btg", funds)]))
        platform = block["platforms"][0]
        self.assertIsNone(platform["above_cdi"])
        self.assertIsNone(platform["median_pct_cdi"])
        self.assertIsNone(block["top"][0]["pct_cdi"])

    def test_matrix_carries_median_return_per_category_and_house(self):
        from shelf_monitor.report.aggregate import build_performance_block

        block = build_performance_block(
            self._report(
                [
                    self._snapshot("btg", [
                        self._offer("btg", "A", "11111111000111", 0.10, anbima="Acoes Livre"),
                        self._offer("btg", "B", "22222222000122", 0.30, anbima="Acoes Livre"),
                    ]),
                    self._snapshot("xp", [
                        self._offer("xp", "C", "33333333000133", -0.05, anbima="Acoes Livre"),
                    ]),
                ]
            )
        )
        row = next(r for r in block["matrix"] if r["category"] == "Acoes Livre")
        self.assertAlmostEqual(row["medians"]["btg"], 0.20)
        self.assertAlmostEqual(row["medians"]["xp"], -0.05)
        self.assertEqual(row["counts"], {"btg": 2, "xp": 1})

    def test_top_and_bottom_dedupe_across_shelves(self):
        from shelf_monitor.report.aggregate import build_performance_block

        block = build_performance_block(
            self._report(
                [
                    self._snapshot("btg", [self._offer("btg", "A", "11111111000111", 0.2)]),
                    self._snapshot("xp", [
                        self._offer("xp", "A", "11111111000111", 0.2),
                        self._offer("xp", "Z", "99999999000199", -0.3),
                    ]),
                ]
            )
        )
        self.assertEqual(len(block["top"]), 2)
        self.assertEqual(block["top"][0]["cnpj"], "11111111000111")
        self.assertEqual(sorted(block["top"][0]["platforms"]), ["BTG Pactual", "XP Investimentos"])
        self.assertEqual(block["bottom"][0]["cnpj"], "99999999000199")

    def test_without_returns_the_block_is_unavailable(self):
        from shelf_monitor.models import FundOffer
        from shelf_monitor.report.aggregate import build_performance_block

        snapshot = self._snapshot("btg", [FundOffer(platform="btg", name="Sem cota")])
        block = build_performance_block(self._report([snapshot]))
        self.assertFalse(block["available"])
        self.assertEqual(block["matrix"], [])


class TestInvestorsBlock(unittest.TestCase):
    """Cotistas e ticket medio por casa, no payload unico."""

    def _snapshot(self, platform, funds):
        from datetime import datetime

        from shelf_monitor.models import (
            CollectorStrategy,
            ShelfSnapshot,
            SnapshotStatus,
        )

        return ShelfSnapshot(
            platform=platform, collected_at=datetime(2026, 8, 31),
            strategy=CollectorStrategy.OPEN_FINANCE, status=SnapshotStatus.OK,
            funds=funds,
        )

    def _offer(self, platform, name, cnpj, shareholders, ticket=None, delta=None):
        from shelf_monitor.models import FundOffer

        extra = {"cotistas_informe": shareholders, "captacao_janela_meses": 12}
        if ticket is not None:
            extra["ticket_medio"] = ticket
        if delta is not None:
            extra["cotistas_var_12m"] = delta
            extra["cotistas_var_pct_12m"] = delta / (shareholders - delta)
        return FundOffer(platform=platform, name=name, cnpj=cnpj, extra=extra)

    def _report(self, snapshots, flows=None):
        from datetime import datetime

        from shelf_monitor.models import RunReport

        return RunReport(generated_at=datetime(2026, 8, 31), snapshots=snapshots, flows=flows)

    def test_platform_profile(self):
        from shelf_monitor.report.aggregate import build_investors_block

        funds = [
            self._offer("btg", "A", "11111111000111", 1000, ticket=5_000.0, delta=100),
            self._offer("btg", "B", "22222222000122", 200, ticket=50_000.0, delta=-50),
            self._offer("btg", "C", "33333333000133", 300, ticket=8_000.0),
        ]
        block = build_investors_block(self._report([self._snapshot("btg", funds)]))
        platform = block["platforms"][0]
        self.assertEqual(platform["shareholders"], 1500)
        self.assertAlmostEqual(platform["ticket_median"], 8_000.0)
        self.assertAlmostEqual(platform["small_ticket_share"], 2 / 3)
        self.assertEqual(platform["delta_12m"], 50)
        self.assertEqual(platform["delta_12m_matched"], 2)
        # base = (1000-100) + (200+50) = 1150
        self.assertAlmostEqual(platform["delta_12m_pct"], 50 / 1150)
        self.assertEqual(block["top_gain"][0]["cnpj"], "11111111000111")
        self.assertEqual(block["top_loss"][0]["cnpj"], "22222222000122")
        self.assertTrue(block["available"])

    def test_gainers_dedupe_across_shelves(self):
        from shelf_monitor.report.aggregate import build_investors_block

        block = build_investors_block(
            self._report(
                [
                    self._snapshot("btg", [self._offer("btg", "A", "11111111000111", 1000, delta=10)]),
                    self._snapshot("xp", [self._offer("xp", "A", "11111111000111", 1000, delta=10)]),
                ]
            )
        )
        self.assertEqual(len(block["top_gain"]), 1)
        # Por prateleira os cotistas contam nas duas casas - e o perfil de cada vitrine.
        self.assertEqual([p["shareholders"] for p in block["platforms"]], [1000, 1000])

    def test_managers_view_alone_keeps_the_block_available(self):
        from shelf_monitor.report.aggregate import build_investors_block

        managers = {"market": {"shareholders": 5_000_000}, "houses": []}
        block = build_investors_block(self._report([], flows=managers))
        self.assertTrue(block["available"])
        block = build_investors_block(self._report([]))
        self.assertFalse(block["available"])


class TestFlowsRendering(unittest.TestCase):
    """A secao de captacao nas quatro superficies."""

    def _snapshot(self, platform, funds):
        from datetime import datetime

        from shelf_monitor.models import (
            CollectorStrategy,
            ShelfSnapshot,
            SnapshotStatus,
        )

        return ShelfSnapshot(
            platform=platform, collected_at=datetime(2026, 8, 31),
            strategy=CollectorStrategy.OPEN_FINANCE, status=SnapshotStatus.OK,
            funds=funds,
        )

    def _report(self):
        from datetime import datetime

        from shelf_monitor.diff import diff_shelf
        from shelf_monitor.models import FundOffer, RunReport

        funds = [
            FundOffer(
                platform="btg", name="Alpha Captador", cnpj="11111111000111",
                manager_name="Gestora Alpha",
                extra={
                    "captacao_liquida_12m": 1_000_000.0,
                    "captacao_liquida_3m": 400_000.0,
                    "captacao_janela_meses": 12,
                    "captacao_data_base": "2026-08-29",
                    "pl_informe": 50_000_000.0,
                },
            ),
            FundOffer(
                platform="btg", name="Beta Resgatado", cnpj="22222222000122",
                extra={
                    "captacao_liquida_12m": -2_500_000.0,
                    "captacao_liquida_3m": -900_000.0,
                    "captacao_janela_meses": 12,
                    "captacao_data_base": "2026-08-29",
                },
            ),
        ]
        snapshot = self._snapshot("btg", funds)
        managers = {
            "as_of": "2026-08-29",
            "window_months": 12,
            "market": {"net_3m": -1e9, "net_12m": 5e9, "pl": 9.3e12, "classes": 70000},
            "houses": [
                {
                    "house": "btg", "label": "BTG Pactual", "net_3m": 2e8,
                    "net_12m": 1.5e9, "pl": 3e11, "classes": 900, "managers": [],
                }
            ],
            "coverage": {"classes_com_gestora": 68000, "classes_total": 70000},
        }
        return RunReport(
            generated_at=datetime(2026, 8, 31), snapshots=[snapshot],
            shelf_diffs=[diff_shelf(snapshot, None)], flows=managers,
        )

    def test_markdown_carries_the_flows_section(self):
        from shelf_monitor.report.digest import render_markdown

        markdown = render_markdown(self._report())
        self.assertIn("## Captacao (CVM)", markdown)
        self.assertIn("Alpha Captador", markdown)
        self.assertIn("Beta Resgatado", markdown)
        self.assertIn("BTG Pactual", markdown)
        # A janela rotulada acompanha o numero - nunca um total sem contexto.
        self.assertIn("meses fechados", markdown)

    def test_markdown_without_flows_says_how_to_get_them(self):
        from datetime import datetime

        from shelf_monitor.diff import diff_shelf
        from shelf_monitor.models import FundOffer, RunReport
        from shelf_monitor.report.digest import render_markdown

        snapshot = self._snapshot(
            "btg", [FundOffer(platform="btg", name="Sem informe")]
        )
        report = RunReport(
            generated_at=datetime(2026, 8, 31), snapshots=[snapshot],
            shelf_diffs=[diff_shelf(snapshot, None)],
        )
        markdown = render_markdown(report)
        self.assertIn("Captacao indisponivel", markdown)
        self.assertIn("shelf_monitor market", markdown)

    def test_digest_html_flows_section_has_no_script(self):
        from shelf_monitor.report.digest import render_html

        html = render_html(self._report())
        self.assertIn("Captacao (CVM)", html)
        self.assertNotIn("<script", html)

    def test_dashboard_embeds_the_flows_block(self):
        import json

        from shelf_monitor.report.dashboard import render_dashboard

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        content = render_dashboard(config, self._report()).read_text(encoding="utf-8")
        self.assertIn('id="flows-section"', content)
        start = content.index('id="dashboard-data">') + len('id="dashboard-data">')
        data = json.loads(content[start:content.index("</script>", start)])
        self.assertTrue(data["flows"]["available"])
        self.assertEqual(data["flows"]["managers"]["houses"][0]["label"], "BTG Pactual")

    def test_spreadsheet_gains_the_flows_tab_and_columns(self):
        from openpyxl import load_workbook

        from shelf_monitor.report.spreadsheet import render_spreadsheet

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        report = self._report()
        path = render_spreadsheet(config, report.snapshots)
        workbook = load_workbook(path)

        self.assertIn("Captacao", workbook.sheetnames)
        flows_sheet = workbook["Captacao"]
        # Ordenada pela captacao liquida de 12m: captador no topo.
        self.assertEqual(flows_sheet.cell(row=2, column=2).value, "Alpha Captador")
        self.assertEqual(flows_sheet.cell(row=2, column=7).value, 1_000_000.0)

        consolidated = workbook["Consolidado"]
        headers = [cell.value for cell in consolidated[1]]
        self.assertIn("Captacao liq. 12m (R$)", headers)

    def test_deck_renders_the_flows_slide(self):
        try:
            from shelf_monitor.report.deck import render_deck
            from pptx import Presentation
        except ImportError:
            self.skipTest("python-pptx nao instalado")

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        path = render_deck(config, self._report())
        texts = []
        for slide in Presentation(path).slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
        self.assertTrue(any("Captacao (CVM)" in t for t in texts))


if __name__ == "__main__":
    unittest.main()
