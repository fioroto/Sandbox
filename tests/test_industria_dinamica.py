"""Fase 2A e 2B do monitor da industria: CDI e macro do SGS, giro,
concentracao, entrantes, fluxo x retorno e captacao por tipo de cotista.

Sobre o painel sintetico de `test_industria_bloco` (A rende 1% ao mes e capta
10 por mes, com 20% do PL em maos de fundos), mais uma serie do SGS montada a
mao: CDI de 0,05% em todo dia do calendario e IPCA de 0,5% em todo mes.

Trava:
- o CDI da janela vai da base (exclusive) ao fim (inclusive), e sai None
  quando a serie nao cobre a janela;
- o cache do SGS nao volta a rede para competencia fechada, e sem rede vale o
  cache, com aviso;
- os segmentos do investidor mais o residuo dao o fluxo do investidor, e o
  segmento "fundos" e o fluxo dos fundos do painel;
- os quartis de retorno somam a captacao do universo.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import dataclasses
import unittest
from datetime import date, timedelta
from unittest import mock

from openpyxl import load_workbook

from shelf_monitor.http import SourceUnavailable
from shelf_monitor.industria.agregados import dimensoes
from shelf_monitor.industria.bloco import leitura_com_topicos, montar_bloco
from shelf_monitor.industria.desempenho import rentabilidade_por_peer_group
from shelf_monitor.industria.dinamica import (
    captacao_por_segmento,
    concentracao_por,
    fatia_dos_maiores,
    fluxo_por_quartil,
    gestoras_que_entraram_e_sairam,
    giro_por,
    hhi,
)
from shelf_monitor.industria.fluxos import total
from shelf_monitor.industria.janelas import janelas
from shelf_monitor.industria.macro import Macro
from shelf_monitor.industria.painel import linhas_do_mes
from shelf_monitor.market import sgs
from shelf_monitor.report.industria_email import render_email_html, render_email_markdown
from shelf_monitor.report.industria_painel import dados_do_painel, linha_do_macro
from shelf_monitor.report.industria_planilha import gravar_planilha
from shelf_monitor.report.shared.fontes import BCB_SGS

from fixtures_comuns import temp_config
from test_industria_bloco import MESES, _base

CDI_DIARIO = 0.05  # % ao dia
IPCA_MENSAL = 0.5  # % no mes


def _macro() -> Macro:
    dias = [date(2025, 7, 1) + timedelta(days=i) for i in range(500)]
    meses = [date(2025 + (7 + i - 1) // 12, (7 + i - 1) % 12 + 1, 1) for i in range(16)]
    return Macro({
        "cdi": [(d, CDI_DIARIO) for d in dias],
        "selic_meta": [(d, 15.0 if d < date(2026, 3, 1) else 14.0) for d in dias],
        "ipca": [(m, IPCA_MENSAL) for m in meses],
        "dolar": [(d, 5.0 if d < date(2026, 1, 1) else 5.5) for d in dias],
    })


def _base_com_macro():
    base = _base()
    base.macro = _macro()
    return base


class TestMacro(unittest.TestCase):
    def test_cdi_da_base_exclusive_ao_fim_inclusive(self):
        macro = _macro()
        base, fim = date(2025, 8, 28), date(2026, 8, 28)
        dias = (fim - base).days  # todo dia do calendario tem taxa nesta serie
        self.assertAlmostEqual(macro.cdi_entre(base, fim), (1 + CDI_DIARIO / 100) ** dias - 1, places=12)

    def test_cdi_sem_cobertura_sai_none(self):
        macro = _macro()
        self.assertIsNone(macro.cdi_entre(date(2025, 8, 28), date(2027, 12, 31)))
        self.assertIsNone(macro.cdi_entre(date(2024, 1, 1), date(2025, 8, 28)))
        self.assertIsNone(macro.cdi_entre(None, date(2026, 8, 28)))

    def test_ipca_composto_e_valor_no_dia(self):
        macro = _macro()
        self.assertAlmostEqual(macro.ipca_acumulado(MESES[1:]), (1 + IPCA_MENSAL / 100) ** 12 - 1, places=12)
        self.assertIsNone(macro.ipca_acumulado(["209901"]))
        self.assertEqual(macro.valor_em("selic_meta", date(2026, 8, 28)), 14.0)
        self.assertEqual(macro.valor_em("dolar", date(2025, 12, 31)), 5.0)
        self.assertIsNone(macro.valor_em("dolar", date(2020, 1, 1)))


class TestSgs(unittest.TestCase):
    def setUp(self):
        self.config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)

    def test_le_a_resposta_da_api(self):
        lidas = sgs.ler_resposta([
            {"data": "02/01/2026", "valor": "0.050788"},
            {"data": "01/01/2026", "valor": "0,05"},
            {"data": "xx", "valor": "1"},
            {"data": "05/01/2026", "valor": ""},
        ])
        self.assertEqual(lidas, [(date(2026, 1, 1), 0.05), (date(2026, 1, 2), 0.050788)])

    def test_competencia_fechada_nao_volta_a_rede(self):
        observacoes = [(date(2025, 8, 1) + timedelta(days=i), 0.05) for i in range(420)]
        with mock.patch.object(sgs, "_baixar", return_value=observacoes) as baixar:
            primeira, avisos = sgs.carregar(self.config, ("cdi",), date(2025, 9, 1), date(2026, 8, 31),
                                            hoje=date(2026, 10, 10))
            segunda, _ = sgs.carregar(self.config, ("cdi",), date(2025, 9, 1), date(2026, 8, 31),
                                      hoje=date(2026, 10, 11))
        self.assertEqual(baixar.call_count, 1)
        self.assertEqual(avisos, [])
        self.assertEqual(primeira, segunda)
        self.assertEqual(primeira["cdi"][-1][0], date(2026, 8, 31))

    def test_sem_rede_vale_o_cache_com_aviso(self):
        observacoes = [(date(2025, 8, 1) + timedelta(days=i), 0.05) for i in range(100)]
        with mock.patch.object(sgs, "_baixar", return_value=observacoes):
            sgs.carregar(self.config, ("cdi",), date(2025, 9, 1), date(2025, 10, 31), hoje=date(2025, 11, 1))
        falha = SourceUnavailable("timeout")
        with mock.patch.object(sgs, "_baixar", side_effect=falha):
            series, avisos = sgs.carregar(self.config, ("cdi",), date(2025, 9, 1), date(2026, 8, 31),
                                          hoje=date(2026, 10, 10))
        self.assertTrue(series["cdi"])
        self.assertTrue(any("indisponível" in a for a in avisos))

    def test_offline_sem_cache_avisa(self):
        series, avisos = sgs.carregar(self.config, ("ipca",), date(2025, 9, 1), date(2026, 8, 31), baixar=False)
        self.assertEqual(series["ipca"], [])
        self.assertTrue(avisos)


class TestRetornoContraOCdi(unittest.TestCase):
    def test_mediana_em_porcentagem_do_cdi_e_quem_passou(self):
        base = _base()
        doze = janelas(base)["12m"]
        resultado = rentabilidade_por_peer_group(base, doze, cdi=0.10)
        a = resultado["por_classe"]["A"]
        self.assertAlmostEqual(a["sobre_cdi"], a["retorno"] / 0.10, places=3)
        linha_a = next(l for l in linhas_do_mes(base, "202608") if l.classe == "A")
        dim = dimensoes(base, linha_a)
        grupo = next(g for g in resultado["n2"] if (g["n1"], g["n2"]) == (dim.n1, dim.n2))
        # A e D estao no mesmo peer group; so A tem retorno (D saiu): 12,7% contra 10%.
        self.assertEqual(grupo["acima_do_cdi"], 1.0)
        self.assertAlmostEqual(grupo["mediana_sobre_cdi"], round(grupo["mediana"] / 0.10, 4))
        sem = rentabilidade_por_peer_group(base, doze)
        self.assertIsNone(sem["por_classe"]["A"]["sobre_cdi"])
        self.assertTrue(all(g["acima_do_cdi"] is None for g in sem["n2"]))


class TestGiroEConcentracao(unittest.TestCase):
    def setUp(self):
        self.base = _base()
        self.doze = janelas(self.base)["12m"]

    def test_giro_e_entrada_sobre_o_pl_medio(self):
        a = giro_por(self.base, self.doze, lambda l, d: l.classe)["A"]
        pls = [next(l.pl for l in linhas_do_mes(self.base, m) if l.classe == "A") for m in MESES[1:]]
        self.assertAlmostEqual(a["captacao"], 120.0, places=2)
        self.assertEqual(a["resgates"], 0.0)
        self.assertAlmostEqual(a["pl_medio"], sum(pls) / 12, places=2)
        self.assertAlmostEqual(a["entrada"], round(120.0 / (sum(pls) / 12), 4))
        self.assertEqual(a["giro"], 0.0)

    def test_hhi_e_fatia_dos_maiores(self):
        self.assertEqual(hhi([50, 50]), 5000.0)
        self.assertEqual(hhi([3, 1]), 6250.0)
        self.assertEqual(hhi([10, 0, -5]), 10000.0)
        self.assertIsNone(hhi([0, 0]))
        self.assertEqual(fatia_dos_maiores([1] * 10, 5), 0.5)

    def test_uma_gestora_so_e_concentracao_total(self):
        conc = concentracao_por(self.base, self.doze, lambda l, d: "total")["total"]
        self.assertEqual(conc["gestoras"], 1)
        self.assertEqual(conc["hhi"], 10000.0)
        self.assertEqual(conc["top5"], 1.0)
        self.assertEqual(conc["variacao_hhi"], 0.0)


class TestEntrantes(unittest.TestCase):
    def test_grupo_que_entrou_e_grupo_que_saiu(self):
        base = _base()
        produtos = dict(base.cadastro.produtos)
        produtos[("C", "")] = dataclasses.replace(produtos[("C", "")], gestor="NOVA GESTORA",
                                                  gestor_cnpj="11111111000111")
        produtos[("D", "")] = dataclasses.replace(produtos[("D", "")], gestor="GESTORA QUE SAIU",
                                                  gestor_cnpj="22222222000122")
        base.cadastro = type(base.cadastro)(base.cadastro.competencia, produtos)
        nova, saiu = base.estrutura.grupo_da_classe("C").nome, base.estrutura.grupo_da_classe("D").nome
        movimento = gestoras_que_entraram_e_sairam(base, janelas(base)["12m"])
        self.assertEqual([i["grupo"] for i in movimento["entraram"]["lista"]], [nova])
        self.assertEqual(movimento["entraram"]["lista"][0]["primeira"], "202603")
        self.assertEqual([(i["grupo"], i["ultima"]) for i in movimento["sairam"]["lista"]], [(saiu, "202605")])
        self.assertEqual(movimento["grupos_na_base"], 2)
        self.assertEqual(movimento["grupos_no_fim"], 2)


class TestFluxoXRetornoESegmentos(unittest.TestCase):
    def setUp(self):
        self.base = _base()
        self.doze = janelas(self.base)["12m"]

    def test_quartis_somam_a_captacao_do_universo(self):
        por_classe = rentabilidade_por_peer_group(self.base, self.doze)["por_classe"]
        quartis = fluxo_por_quartil(self.base, self.doze, por_classe)
        soma = sum(q["fluxo_investidor"] for q in quartis)
        self.assertAlmostEqual(soma, total(self.base, self.doze).fluxo_investidor, places=1)

    def test_segmentos_mais_residuo_dao_o_fluxo_do_investidor(self):
        for chave in (None, lambda l, d: d.n1, lambda l, d: l.classe):
            for grupo in captacao_por_segmento(self.base, self.doze, chave).values():
                soma = sum(s["captacao"] for s in grupo["segmentos"]) + grupo["residuo"]
                self.assertAlmostEqual(soma, grupo["fluxo_investidor"], places=1)

    def test_fundos_e_o_fluxo_dos_fundos_do_painel(self):
        segmentos = captacao_por_segmento(self.base, self.doze)["total"]
        self.assertAlmostEqual(segmentos["fundos"], total(self.base, self.doze).fluxo_fundos, places=1)

    def test_conta_e_ordem_de_a_calculada_a_mao(self):
        # Fatias constantes ate 07/2026 (5% conta e ordem, 75% PF varejo): o
        # segmento leva a fatia dos 10 captados. Em 08/2026 a conta e ordem
        # vai a 10%: PL(t) x 0,10 - PL(t-1) x 1,01 x 0,05.
        a = captacao_por_segmento(self.base, self.doze, lambda l, d: l.classe)["A"]
        valores = {s["chave"]: s["captacao"] for s in a["segmentos"]}
        pl_ago = next(l.pl for l in linhas_do_mes(self.base, "202608") if l.classe == "A")
        self.assertAlmostEqual(valores["conta_e_ordem"], 11 * 0.5 + 0.05 * pl_ago + 0.5, places=1)
        self.assertAlmostEqual(a["fundos"], 0.2 * 10 * 12, places=1)
        self.assertAlmostEqual(a["residuo"], 0.0, places=1)
        self.assertGreater(a["ruido"], 0.0)
        pj = next(s for s in a["segmentos"] if s["chave"] == "pj_private")
        self.assertTrue(pj["no_ruido"])  # zero e sempre "dentro do ruido"


class TestBlocoESuperficies(unittest.TestCase):
    def setUp(self):
        self.config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        self.bloco = montar_bloco(_base_com_macro(), self.config)

    def test_bloco_traz_macro_desempenho_e_dinamica(self):
        bloco = self.bloco
        self.assertTrue(bloco["macro"]["available"])
        doze = bloco["janelas"]["12m"]
        dias = (date.fromisoformat(doze["ate"]) - date.fromisoformat(doze["data_base"])).days
        self.assertAlmostEqual(bloco["macro"]["janelas"]["12m"]["cdi"], round((1.0005 ** dias) - 1, 6))
        self.assertEqual(bloco["desempenho"]["12m"]["cdi"], bloco["macro"]["janelas"]["12m"]["cdi"])
        self.assertEqual(len(bloco["macro"]["meses"]), 12)
        for chave in ("giro", "concentracao", "gestoras", "fluxo_x_retorno", "captacao_por_segmento"):
            self.assertIn(chave, bloco["dinamica"])
        com_retorno = [i for i in bloco["estrategias"]["n2"] if (i["rentabilidade"]["12m"] or {}).get("mediana")]
        self.assertTrue(any(i["rentabilidade"]["12m"].get("mediana_sobre_cdi") for i in com_retorno))
        self.assertIn(BCB_SGS.nome, bloco["fontes"])

    def test_leitura_do_fluxo_x_retorno_com_o_cdi(self):
        bloco = dict(self.bloco)
        bloco["dinamica"] = {**bloco["dinamica"], "fluxo_x_retorno": [
            {"quartil": 1, "sobre_pl": 0.031}, {"quartil": 4, "sobre_pl": -0.024},
        ]}
        frases = dict(leitura_com_topicos(bloco))
        self.assertIn("**+3,1%**", frases["retorno"])
        self.assertIn("**-2,4%**", frases["retorno"])
        self.assertIn("CDI de", frases["retorno"])

    def test_sem_sgs_o_bloco_sai_sem_macro_e_sem_a_fonte(self):
        bloco = montar_bloco(_base(), self.config)
        self.assertFalse(bloco["macro"]["available"])
        self.assertIn("SGS", bloco["macro"]["reason"])
        self.assertNotIn(BCB_SGS.nome, bloco["fontes"])
        self.assertIsNone(bloco["desempenho"]["12m"]["cdi"])
        self.assertEqual(linha_do_macro(bloco), "")

    def test_painel_leva_o_cdi_e_deixa_os_segmentos_por_peer_group_na_planilha(self):
        dados = dados_do_painel(self.bloco)
        self.assertEqual(dados["meta"]["cdi"]["12m"], self.bloco["desempenho"]["12m"]["cdi"])
        self.assertIn("CDI em 12 meses", dados["meta"]["macro"])
        self.assertNotIn("n2", dados["dinamica"]["captacao_por_segmento"])
        self.assertIn("total", dados["dinamica"]["captacao_por_segmento"])
        self.assertIn(BCB_SGS.url, [f["url"] for f in dados["fontes"]])

    def test_email_traz_o_retorno_contra_o_cdi_nos_dois_formatos(self):
        html = render_email_html(self.bloco, self.config)
        markdown = render_email_markdown(self.bloco, self.config)
        for texto in (html, markdown):
            self.assertIn("Os fundos entregam?", texto)
            self.assertIn("Macro: CDI em 12 meses", texto)
            self.assertIn("Quartil de retorno", texto)
        for proibido in ("<svg", "<style", "<script", "<img"):
            self.assertNotIn(proibido, html)
        self.assertLess(len(html.encode("utf-8")), 100_000)

    def test_planilha_ganha_dinamica_e_macro(self):
        wb = load_workbook(gravar_planilha(self.config, self.bloco))
        self.assertIn("Dinâmica", wb.sheetnames)
        self.assertIn("Macro", wb.sheetnames)
        self.assertLess(wb.sheetnames.index("Gestoras"), wb.sheetnames.index("Dinâmica"))
        textos = [c.value for row in wb["Estratégias"].iter_rows() for c in row if isinstance(c.value, str)]
        self.assertIn("Mediana 12m em % do CDI", textos)


if __name__ == "__main__":
    unittest.main()
