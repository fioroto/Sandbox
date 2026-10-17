"""Digest v2 (14/10/2026): semana de referencia, oportunidade de grade e recorte.

O digest passou a ser semanal de fato e mais curto. Os testes travam:

- a regra da semana (`report/shared/semanas.py`): o dia que ainda esta
  chegando sai, o feriado com linha avulsa sai, a semana e a do ultimo dia
  completo (com ao menos 3 dias) e a comparacao sao as 12 anteriores;
- a captacao semanal por fundo e o ritmo dela, pela regra do ritmo mensal, e a
  janela de 3 meses terminando no mesmo ultimo dia completo;
- o mes anterior revalidado na CVM (a semana pode cruzar a virada do mes);
- a oportunidade de grade: gaps reais que captam, sem o produto proprio do
  concorrente nem o que a lista da casa ja tem; desde 15/10/2026 sem a frase
  de abertura, sem a frase do funil e sem a evolucao semanal da soma;
- o recorte as casas do digest (XP, BTG e Itau): Nu e Inter fora, e sem o radar
  da CVM desde 17/10/2026 - nem no corpo, nem no assunto;
- a legenda do ritmo no fim da secao de captacao (17/10/2026);
- o servico da lista da casa (Distribuicao ou FaaS) ao lado do fundo da XP;
- a edicao mensal no primeiro digest do mes, com rentabilidade e custo.

O informe diario e substituido por registros sinteticos de 16 semanas, com
numeros redondos.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import re
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest import mock

from shelf_monitor.classify import KEY_CATEGORIA
from shelf_monitor.config import Config
from shelf_monitor.diff import diff_shelf
from shelf_monitor.distribuivel import ArquivoLido, BaseDistribuivel
from shelf_monitor.models import (
    CollectorStrategy,
    CvmDiff,
    CvmFund,
    CvmSnapshot,
    FundOffer,
    RunReport,
    ShelfSnapshot,
    SnapshotStatus,
)
from shelf_monitor.report.retail_flow import (
    DailyFlow,
    SemanaDoFundo,
    build_retail_flow,
    ritmo_da_semana,
    weekly_by_fund,
)
from shelf_monitor.report.shared.semanas import (
    Semana,
    dias_completos,
    prazos_da_semana,
    recorte_semanal,
)

from fixtures_comuns import temp_config

MI = 1_000_000.0
PL = 1_000 * MI
CATEGORIA = "Multimercados Livre"
INDEPENDENTE = "GESTORA INDEPENDENTE LTDA"
BTG_ASSET = "BTG PACTUAL ASSET MANAGEMENT S/A DTVM"

# Os fundos do caso. A capta 10 por dia e acelera para 30 na ultima semana; B
# (produto proprio do BTG) capta 5 por dia, estavel; C (na vitrine das duas
# casas, e na lista da casa como FaaS) capta 2 por dia; D so o Nu exibe.
A = "11000000000101"
B = "12000000000101"
C = "13000000000101"
D = "14000000000101"
# So a vitrine do Itau exibe. Perde 4 por dia ate julho, fica parado e capta 4
# por dia na ultima semana: virou entrada, e capta nos 3 meses.
E = "16000000000101"
# Produto proprio do BTG que acelera (de 12 para 40 por dia) e capta mais que A
# em 3 meses, abaixo do proprio PL: so entra na vitrine com `com_proprio`.
F = "17000000000101"
_ENCHIMENTO = 32  # cada vitrine com 30 fundos ou mais (MIN_VITRINE)

PRIMEIRO_DIA = date(2026, 6, 1)
ULTIMO_COMPLETO = date(2026, 10, 9)   # sexta
DIA_PARCIAL = date(2026, 10, 12)      # so A chegou
FERIADO = date(2026, 9, 7)            # so C, com uma linha avulsa
INICIO_DA_SEMANA = date(2026, 10, 5)


def _dias_uteis() -> list[date]:
    dias, dia = [], PRIMEIRO_DIA
    while dia <= ULTIMO_COMPLETO:
        if dia.weekday() < 5:
            dias.append(dia)
        dia += timedelta(days=1)
    return dias


def _registros() -> list[DailyFlow]:
    registros: list[DailyFlow] = []
    enchimento = [f"{casa}{i:012d}" for casa in (31, 32) for i in range(_ENCHIMENTO)]
    for n, dia in enumerate(_dias_uteis()):
        if dia == FERIADO:
            registros.append(DailyFlow(cnpj=C, dia=dia, pl=PL, cota=10.0, cotistas=1000 + n))
            continue
        na_semana = dia >= INICIO_DA_SEMANA
        a = 30 * MI if na_semana else 10 * MI
        e = 4 * MI if na_semana else -4 * MI if dia < date(2026, 8, 1) else 0.0
        registros += [
            DailyFlow(cnpj=A, dia=dia, captacao=a, pl=PL, cota=10.0, cotistas=1000 + n),
            DailyFlow(cnpj=B, dia=dia, captacao=5 * MI, pl=PL, cota=10.0, cotistas=1000 + n),
            DailyFlow(cnpj=C, dia=dia, captacao=2 * MI, pl=PL, cota=10.0, cotistas=1000 + n),
            DailyFlow(cnpj=D, dia=dia, pl=PL, cota=10.0, cotistas=1000),
            DailyFlow(cnpj=E, dia=dia, captacao=max(e, 0.0), resgates=max(-e, 0.0), pl=PL,
                      cota=10.0, cotistas=1000 + n),
            DailyFlow(cnpj=F, dia=dia, captacao=(40 if na_semana else 12) * MI, pl=PL,
                      cota=10.0, cotistas=1000 + n),
        ]
        registros += [
            DailyFlow(cnpj=cnpj, dia=dia, pl=PL, cota=10.0, cotistas=1000)
            for cnpj in enchimento
        ]
    registros.append(DailyFlow(cnpj=A, dia=DIA_PARCIAL, captacao=30 * MI, pl=PL, cota=10.0,
                               cotistas=5000))
    return registros


def _oferta(platform: str, cnpj: str, nome: str, gestor: str = INDEPENDENTE) -> FundOffer:
    fundo = FundOffer(
        platform=platform, name=nome, cnpj=cnpj, manager_name=gestor,
        sources=[CollectorStrategy.API_JSON.value],
    )
    fundo.extra[KEY_CATEGORIA] = CATEGORIA
    return fundo


def _snapshot(platform: str, fundos: list[FundOffer], quando: datetime) -> ShelfSnapshot:
    return ShelfSnapshot(
        platform=platform, collected_at=quando, status=SnapshotStatus.OK,
        strategy=CollectorStrategy.API_JSON, funds=fundos,
    )


def _snapshots(
    quando: datetime = datetime(2026, 10, 14, 8, 0), com_itau: bool = False,
    com_proprio: bool = False,
) -> list[ShelfSnapshot]:
    def enchimento(casa: int, platform: str) -> list[FundOffer]:
        return [
            _oferta(platform, f"{casa}{i:012d}", f"{platform} fundo {i}")
            for i in range(_ENCHIMENTO)
        ]

    xp = [_oferta("xp", C, "Fundo C FIC FIM")] + enchimento(31, "xp")
    btg = [
        _oferta("btg", A, "Fundo A FIC FIM"),
        _oferta("btg", B, "BTG Proprio B FIC FIM", BTG_ASSET),
        _oferta("btg", C, "Fundo C FIC FIM"),
    ] + enchimento(32, "btg")
    if com_proprio:
        btg.append(_oferta("btg", F, "BTG Proprio F FIC FIM", BTG_ASSET))
    nu = [_oferta("nu", D, "Fundo D FIC FIM")] + enchimento(32, "nu")
    casas = [_snapshot("xp", xp, quando), _snapshot("btg", btg, quando)]
    if com_itau:
        # A nas duas vitrines concorrentes; E so na do Itau.
        itau = [_oferta("itau", A, "Fundo A FIC FIM"), _oferta("itau", E, "Fundo E FIC FIM")]
        casas.append(_snapshot("itau", itau + enchimento(32, "itau"), quando))
    return casas + [_snapshot("nu", nu, quando)]


def _fluxo(snapshots: list[ShelfSnapshot]) -> tuple[dict, mock.MagicMock, mock.MagicMock]:
    registros = _registros()

    def ler(config, yyyymm, alvo, **_):
        return [r for r in registros if r.dia.strftime("%Y%m") == yyyymm and r.cnpj in alvo]

    with mock.patch("shelf_monitor.report.retail_flow.ler_filtrado",
                    side_effect=ler) as principal, \
            mock.patch("shelf_monitor.report.retail_flow.ler_filtrado_sem_guardar",
                       side_effect=ler) as anterior, \
            mock.patch("shelf_monitor.report.retail_flow.months_back",
                       return_value=["202608", "202609", "202610"]):
        flow = build_retail_flow(Config(), snapshots)
    return flow, principal, anterior


def _lista_da_casa(**servicos: str) -> BaseDistribuivel:
    """A lista da casa ja lida: {cnpj: servico}."""
    return BaseDistribuivel(
        pasta=Path("entrada"),
        situacao={cnpj: "distribui" for cnpj in servicos},
        servico=dict(servicos),
        arquivos=[ArquivoLido(nome="lista.csv", linhas=len(servicos), validas=len(servicos))],
    )


def _relatorio(
    lista: BaseDistribuivel | None = None, com_itau: bool = False, com_proprio: bool = False
) -> RunReport:
    snapshots = _snapshots(com_itau=com_itau, com_proprio=com_proprio)
    flow, _, _ = _fluxo(snapshots)
    antes = [_snapshot(s.platform, list(s.funds), datetime(2026, 10, 7)) for s in snapshots]
    novo = _oferta("btg", "15000000000101", "BTG Novo Credito FIC FIM")
    snapshots[1].funds.append(novo)
    cvm = CvmSnapshot(
        collected_at=datetime(2026, 10, 14, 7, 0), status=SnapshotStatus.OK,
        funds=[
            CvmFund(cnpj="70000000000101", name="XP NOVO FIC", house="xp", matched_by="name",
                    manager_name="XP ALLOCATION ASSET MANAGEMENT LTDA",
                    situation="Fase Pré-Operacional"),
            CvmFund(cnpj="71000000000101", name="BTG NOVO FIC", house="btg", matched_by="name",
                    manager_name="EMPIRICUS ASSET MANAGEMENT LTDA.",
                    situation="Fase Pré-Operacional"),
            CvmFund(cnpj="72000000000101", name="KINEA NOVO FIC", house="itau",
                    matched_by="name", manager_name="KINEA INVESTIMENTOS LTDA.",
                    situation="Em Funcionamento Normal"),
            CvmFund(cnpj="73000000000101", name="NU NOVO FIC", house="nu", matched_by="name",
                    manager_name="NU ASSET MANAGEMENT LTDA", situation="Fase Pré-Operacional"),
        ],
    )
    cvm_diff = CvmDiff(prev_date=date(2026, 10, 7), curr_date=date(2026, 10, 14),
                       new_registrations=list(cvm.funds))
    return RunReport(
        generated_at=datetime(2026, 10, 14, 8, 0), snapshots=snapshots,
        shelf_diffs=[diff_shelf(s, a) for s, a in zip(snapshots, antes)],
        cvm_snapshot=cvm, cvm_diff=cvm_diff, retail_flow=flow, distribuivel=lista,
    )


def _texto(html: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html))


# --- a semana de referencia ----------------------------------------------------


class TestSemanas(unittest.TestCase):
    def test_dia_que_chega_sai_do_fim_e_feriado_sai_do_meio(self):
        contagem = {date(2026, 9, d): 100 for d in (1, 2, 3, 4, 8, 9, 10, 11)}
        contagem[date(2026, 9, 7)] = 2     # feriado com linha avulsa
        contagem[date(2026, 9, 14)] = 40   # ainda chegando
        dias = dias_completos(contagem)
        self.assertEqual(dias.incompletos, ((date(2026, 9, 14), 0.4),))
        self.assertEqual(dias.descartados, (date(2026, 9, 7),))
        self.assertNotIn(date(2026, 9, 7), dias.completos)
        self.assertEqual(dias.completos[-1], date(2026, 9, 11))

    def test_semana_com_menos_de_tres_dias_cede_a_anterior(self):
        dias = [date(2026, 9, d) for d in range(1, 31) if date(2026, 9, d).weekday() < 5]
        ate_terca = {d: 10 for d in dias if d <= date(2026, 9, 22)}
        recorte = recorte_semanal(ate_terca)
        self.assertEqual(recorte.semana.segunda, date(2026, 9, 14))
        ate_quarta = {d: 10 for d in dias if d <= date(2026, 9, 23)}
        self.assertEqual(recorte_semanal(ate_quarta).semana.segunda, date(2026, 9, 21))

    def test_doze_semanas_de_comparacao_e_os_prazos(self):
        contagem = {d: 10 for d in _dias_uteis()}
        recorte = recorte_semanal(contagem)
        self.assertEqual(recorte.semana, Semana(INICIO_DA_SEMANA, tuple(
            INICIO_DA_SEMANA + timedelta(days=i) for i in range(5))))
        self.assertEqual(recorte.inicio_da_comparacao, date(2026, 7, 13))
        self.assertEqual(len(recorte.semanas_anteriores()), 12)
        self.assertEqual(recorte.comparacao[0], date(2026, 7, 13))
        self.assertEqual(recorte.comparacao[-1], date(2026, 10, 2))
        self.assertEqual(
            prazos_da_semana(recorte.semana.to_dict(), recorte.comparacao_dict()),
            ["semana: de 05/10/2026 a 09/10/2026 (5 dias úteis com informe)",
             "comparação: 12 semanas anteriores, de 13/07/2026 a 02/10/2026 (60 dias úteis)"],
        )


# --- a captacao semanal ----------------------------------------------------------


class TestCaptacaoSemanal(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.flow, cls.principal, cls.anterior = _fluxo(_snapshots())
        cls.fundos = {f["cnpj"]: f for f in cls.flow["fundos"]}

    def test_semana_comparacao_e_o_que_ficou_fora(self):
        recorte = self.flow["recorte_semanal"]
        self.assertEqual(recorte["semana"]["de"], "2026-10-05")
        self.assertEqual(recorte["semana"]["ate"], "2026-10-09")
        self.assertEqual(recorte["semana"]["dias"], 5)
        self.assertEqual((recorte["comparacao"]["de"], recorte["comparacao"]["ate"]),
                         ("2026-07-13", "2026-10-02"))
        # 12 semanas de 5 dias, menos o feriado de 07/09.
        self.assertEqual(recorte["comparacao"]["dias"], 59)
        self.assertEqual(recorte["dias_descartados"], ["2026-09-07"])
        self.assertEqual([d["dia"] for d in recorte["dias_incompletos"]], ["2026-10-12"])

    def test_janela_de_3_meses_termina_no_ultimo_dia_completo(self):
        self.assertEqual(self.flow["periodo"]["de"], "2026-08-03")
        self.assertEqual(self.flow["periodo"]["ate"], "2026-10-09")
        self.assertEqual(self.flow["periodo"]["meses_lidos"], ["202608", "202609", "202610"])
        # 44 dias a 10 e a semana a 30; o dia parcial (12/10) fica fora.
        self.assertAlmostEqual(self.fundos[A]["captacao_liquida"], 590 * MI)

    def test_captacao_e_ritmo_da_semana(self):
        self.assertAlmostEqual(self.fundos[A]["captacao_semana"], 150 * MI)
        self.assertEqual(self.fundos[A]["ritmo_semanal"], "acelerando")
        self.assertAlmostEqual(self.fundos[B]["captacao_semana"], 25 * MI)
        self.assertEqual(self.fundos[B]["ritmo_semanal"], "estavel")

    def test_mes_anterior_revalidado_e_o_antes_da_janela_lido_sem_guardar(self):
        chamadas = {c.args[1]: c.kwargs for c in self.principal.call_args_list}
        self.assertEqual(chamadas["202610"], {"force": True, "revalidar": False})
        self.assertEqual(chamadas["202609"], {"force": False, "revalidar": True})
        self.assertEqual(chamadas["202608"], {"force": False, "revalidar": False})
        # A comparacao comeca em 13/07: julho e lido; junho, nao.
        self.assertEqual([c.args[1] for c in self.anterior.call_args_list], ["202607"])

    def test_semana_deduplica_a_retificacao_e_pede_base(self):
        semanas = [Semana(date(2026, 10, 5), (date(2026, 10, 5),))]
        repetido = [
            DailyFlow(cnpj=A, dia=date(2026, 10, 5), captacao=1 * MI),
            DailyFlow(cnpj=A, dia=date(2026, 10, 5), captacao=3 * MI),
        ]
        [semana] = weekly_by_fund(repetido, semanas)[A]
        self.assertAlmostEqual(semana.captacao_liquida, 3 * MI)
        # Menos de 4 semanas com informe na base: sem serie, e nao "acelerando".
        poucas = [SemanaDoFundo(date(2026, 9, 28), 10 * MI, 5), SemanaDoFundo(date(2026, 10, 5), 90 * MI, 5)]
        self.assertEqual(ritmo_da_semana(poucas), "sem_serie")


# --- a oportunidade de grade e o recorte do digest --------------------------------


class TestOportunidadeDeGrade(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from shelf_monitor.report.digest_visao import montar_visao

        cls.visao = montar_visao(_relatorio(), None, "semanal")
        cls.op = cls.visao.oportunidade

    def test_so_gap_real_que_capta_e_acelera_na_semana(self):
        # A e B estao no BTG e fora da XP; B e produto proprio do BTG. O
        # enchimento do BTG e o fundo novo tambem sao gap, mas nao captam. B
        # capta, mas estavel na semana: fica fora da lista (16/10/2026).
        self.assertEqual(self.op["total_reais"], 3 + _ENCHIMENTO)
        self.assertEqual(self.op["total_captando"], 2)
        self.assertTrue(self.op["com_ritmo"])
        self.assertEqual(self.op["total_na_lista"], 1)
        self.assertEqual(self.op["total_produto_do_concorrente"], 0)
        self.assertEqual([f["cnpj"] for f in self.op["linhas"]], [A])
        self.assertEqual(self.op["total_operacionais"], 1)  # a gestora de A esta na XP
        self.assertAlmostEqual(self.op["captacao_semana"], 150 * MI)
        self.assertAlmostEqual(self.op["captacao_3m"], 590 * MI)

    def test_gap_que_a_lista_da_casa_tem_nao_e_oportunidade(self):
        from shelf_monitor.report.digest_visao import montar_visao

        visao = montar_visao(_relatorio(_lista_da_casa(**{A: "faas"})), None, "semanal")
        self.assertNotIn(A, [f["cnpj"] for f in visao.oportunidade["linhas"]])

    def test_recorte_tira_nu_inter_e_o_radar(self):
        recortado = self.visao.report
        self.assertEqual([s.platform for s in recortado.snapshots], ["xp", "btg"])
        # O radar saiu do digest em 17/10/2026: a execucao tem o extrato, o recorte nao.
        self.assertIsNotNone(_relatorio().cvm_snapshot)
        self.assertIsNone(recortado.cvm_snapshot)
        self.assertIsNone(recortado.cvm_diff)
        nomes = [f["nome"] for f in recortado.retail_flow["top_captacao"]]
        self.assertIn("Fundo A FIC FIM", nomes)
        self.assertTrue(all("nu" not in f["vitrines"] for f in recortado.retail_flow["top_captacao"]))

    def test_digest_abre_direto_nas_secoes_com_os_numeros_em_destaque(self):
        """15/10/2026, pedido da area: sem a frase de abertura ("1 fundo capta
        nas vitrines..."), sem a frase do funil ("N gaps nas vitrines...") e sem
        a evolucao semanal da soma. Os numeros em destaque dizem a oportunidade."""
        from shelf_monitor.report.digest import render_html, render_markdown

        report = _relatorio()
        md = render_markdown(report, None, "semanal")
        self.assertIn(
            "execução de 14/10/2026 08:00\n\n## Oportunidade de grade e gaps\n\n"
            "- **1** fundo para a grade: gap acelerando na semana\n"
            "- **+150,0 MM** captação dele na semana\n"
            "- **+590,0 MM** captação dele em 3 meses\n",
            md,
        )
        for fora in ("capta nas vitrines", "gaps nas vitrines", "Evolução semanal",
                     "acima do ritmo das 12"):
            self.assertNotIn(fora, md)
        # A ordem das secoes e o que saiu do semanal.
        ordem = ["## Oportunidade de grade e gaps", "## Ritmo de captação",
                 "## Vitrines: entradas e saídas", "## Fontes"]
        posicoes = [md.index(titulo) for titulo in ordem]
        self.assertEqual(posicoes, sorted(posicoes))
        for fora in ("## Sumário", "Ticket de entrada", "Captação dos fundos de vitrine",
                     "Rentabilidade", "Nu (Nubank)", "Banco Inter", "NU NOVO FIC", "XP NOVO FIC",
                     "Radar CVM", "BTG NOVO FIC", "radar de registros"):
            self.assertNotIn(fora, md)
        html = render_html(report, None, "semanal")
        texto = _texto(html)
        self.assertIn("fundo para a grade: gap acelerando na semana", texto)
        for fora in ("A captação dos gaps", "MM/dia", "gaps nas vitrines", "capta nas vitrines"):
            self.assertNotIn(fora, texto)
        # A previa do e-mail: as casas e a semana, como o cabecalho.
        previa = re.search(r'<div style="display:none;[^"]*">([^<]*)', html).group(1)
        self.assertTrue(
            previa.startswith("BTG Pactual e Itaú · semana de 05/10/2026 a 09/10/2026")
        )
        self.assertIn("Fundo A FIC FIM", texto)
        self.assertNotIn("BTG Proprio B FIC FIM 12.", texto)

    def test_assunto_abre_com_a_grade(self):
        from shelf_monitor.report.mailer import build_subject

        assunto = build_subject(_relatorio())
        self.assertIn("1 fundo para a grade (+150,0 MM na semana)", assunto)
        self.assertIn("1 entrada / 0 saídas", assunto)
        # Sem o radar no digest (17/10/2026), o assunto nao conta registro na CVM.
        self.assertNotIn("CVM", assunto)


class TestOportunidadePorConcorrente(unittest.TestCase):
    """O Itau sempre ao lado do BTG (pedido de 14/10/2026): resumo e lista por
    concorrente, e nao um ranking unico que o BTG ocupa inteiro."""

    @classmethod
    def setUpClass(cls):
        from shelf_monitor.report.digest_visao import montar_visao

        cls.report = _relatorio(com_itau=True)
        cls.op = montar_visao(cls.report, None, "semanal").oportunidade

    def test_resumo_por_concorrente(self):
        por_casa = {c["platform_id"]: c for c in self.op["por_casa"]}
        self.assertEqual(list(por_casa), ["btg", "itau"])
        # A conta nas duas casas; E so no Itau. B e produto proprio do BTG.
        self.assertEqual([f["cnpj"] for f in por_casa["btg"]["linhas"]], [A])
        self.assertEqual({f["cnpj"] for f in por_casa["itau"]["linhas"]}, {A, E})
        self.assertAlmostEqual(por_casa["btg"]["captacao_semana"], 150 * MI)
        self.assertAlmostEqual(por_casa["itau"]["captacao_semana"], 170 * MI)
        self.assertEqual(self.op["total_transferiveis"], 2)

    def test_digest_lista_os_primeiros_de_cada_concorrente(self):
        from shelf_monitor.report.digest import render_html, render_markdown

        md = render_markdown(self.report, None, "semanal")
        # Sem a carteira da CVM a coluna e "Gaps", nao "Gaps reais".
        self.assertIn("| Concorrente | Gaps | Acelerando | Para trazer | Semana | 3 meses |", md)
        # O nome da casa e so "Itaú" (pedido de 15/10/2026).
        self.assertIn("| Itaú |", md)
        self.assertNotIn("íon", md)
        # Cada fundo uma vez: A no grupo do BTG (com as duas casas), E no do Itau.
        self.assertIn("**BTG Pactual** - o primeiro", md)
        self.assertIn("**Itaú** - o primeiro", md)
        self.assertEqual(md.count("Fundo A FIC FIM ("), 1)
        self.assertIn("| BTG, Itaú |", md)
        html = render_html(self.report, None, "semanal")
        texto = _texto(html)
        self.assertIn("Os 10 primeiros de cada concorrente · por captação em 3 meses", texto)
        self.assertIn(">Itaú</td></tr>", html)
        # A evolucao semanal com uma miniatura por casa saiu.
        self.assertNotIn("MM/dia", texto)


class TestListaAcelerando(unittest.TestCase):
    """A lista de gaps do digest (pedido de 16/10/2026): so quem acelera ou virou
    entrada na semana, pela captacao em 3 meses, 10 por concorrente, com o
    produto proprio do concorrente marcado e fora dos numeros de cima."""

    @classmethod
    def setUpClass(cls):
        from shelf_monitor.report.digest_visao import montar_visao

        cls.report = _relatorio(com_itau=True, com_proprio=True)
        cls.visao = montar_visao(cls.report, None, "semanal")
        cls.op = cls.visao.oportunidade

    def test_entra_quem_acelera_ou_virou_entrada(self):
        # B capta estavel na semana e fica fora; E virou entrada e entra.
        self.assertEqual(self.op["total_captando"], 4)  # A, B, E e F
        por_casa = {c["platform_id"]: c for c in self.op["por_casa"]}
        self.assertNotIn(B, [f["cnpj"] for f in por_casa["btg"]["linhas"]])
        self.assertIn(E, [f["cnpj"] for f in por_casa["itau"]["linhas"]])
        ritmos = {f["cnpj"]: f["ritmo_semanal"] for f in self.op["linhas"]}
        self.assertEqual(ritmos, {F: "acelerando", A: "acelerando", E: "virou_entrada"})

    def test_ordem_pela_captacao_em_3_meses_com_o_produto_proprio(self):
        # F (produto proprio, +728 MM) passa na frente de A (+590 MM): o
        # esforco nao ordena mais.
        por_casa = {c["platform_id"]: c for c in self.op["por_casa"]}
        self.assertEqual([f["cnpj"] for f in por_casa["btg"]["linhas"]], [F, A])
        self.assertEqual(por_casa["btg"]["total_na_lista"], 2)
        self.assertEqual(por_casa["btg"]["total_transferiveis"], 1)

    def test_numeros_de_cima_so_com_o_que_da_para_trazer(self):
        self.assertEqual(self.op["total_produto_do_concorrente"], 1)
        self.assertEqual(self.op["total_transferiveis"], 2)  # A e E
        self.assertAlmostEqual(self.op["captacao_semana"], 170 * MI)
        self.assertAlmostEqual(self.op["captacao_3m"], 610 * MI)
        por_casa = {c["platform_id"]: c for c in self.op["por_casa"]}
        self.assertAlmostEqual(por_casa["btg"]["captacao_3m"], 590 * MI)

    def test_digest_marca_o_produto_proprio(self):
        from shelf_monitor.report.digest import render_html, render_markdown
        from shelf_monitor.report.mailer import build_subject

        md = render_markdown(self.report, None, "semanal")
        self.assertIn("- **1** produto próprio do concorrente, marcado na lista\n", md)
        linha_f = next(linha for linha in md.splitlines() if "BTG Proprio F FIC FIM (" in linha)
        self.assertTrue(linha_f.endswith("| produto próprio |"))
        self.assertIn("**BTG Pactual** - os 2 primeiros, por captação em 3 meses", md)
        self.assertIn("estão acelerando na semana - captação por dia mais de 25% acima", md)
        texto = _texto(render_html(self.report, None, "semanal"))
        self.assertIn("produto próprio do concorrente, marcado na lista", texto)
        # O assunto segue contando so o que da para trazer.
        self.assertIn("2 fundos para a grade (+170,0 MM na semana)", build_subject(self.report))

    def test_sem_semana_completa_a_lista_nao_filtra_pelo_ritmo(self):
        from shelf_monitor.report.oportunidade import oportunidade_de_grade

        flow = {**self.visao.report.retail_flow, "_recorte": None}
        op = oportunidade_de_grade(self.visao.gaps, flow)
        self.assertFalse(op["com_ritmo"])
        self.assertIn(B, [f["cnpj"] for f in op["linhas"]])
        self.assertIsNone(op["captacao_semana"])


class TestRentabilidadeDeReferencia(unittest.TestCase):
    """A rentabilidade de cada fundo ao lado dele nos gaps e nas captacoes
    (pedido de 16/10/2026): 12 meses, ou 3 meses para o fundo sem esse
    historico, com a distancia para a mediana da categoria e a janela escrita."""

    @staticmethod
    def _bloco(meses: int, de: str, ate: str, base: str, retornos: dict) -> dict:
        """Um bloco de rentabilidade com a forma do `performance._montar`."""
        from shelf_monitor.report.performance import (
            PerformanceCasa,
            RetornoFundo,
            _onde_esta_a_gente,
            _por_categoria,
        )

        fundos = [
            RetornoFundo(cnpj=cnpj, categoria=CATEGORIA, retorno=retorno, pl=PL,
                         cotistas=1000, dias=200, mediana_categoria=mediana)
            for cnpj, (retorno, mediana) in retornos.items()
        ]
        casa = PerformanceCasa(platform_id="btg", label="BTG Pactual",
                               vitrine=len(fundos), com_serie=len(fundos), fundos=fundos)
        return {
            "available": True,
            "janela": {"rotulo": f"{meses} meses", "meses": meses, "de": de, "ate": ate,
                       "base": base},
            "periodo": {"de": de, "ate": ate, "meses_lidos": [], "meses_ausentes": [],
                        "pregoes": 200},
            "participantes": ["btg"],
            "excluded": [],
            "universo": {"fundos_com_serie": len(fundos), "categorias_com_referencia": 1,
                         "categorias_vistas": 1, "amostra_minima": 20},
            "casas": [casa.to_dict()],
            "por_categoria": _por_categoria([casa], minimo=1),
            "onde_esta_a_gente": _onde_esta_a_gente([casa]),
            "_casas": [casa],
        }

    def _relatorio(self) -> RunReport:
        report = _relatorio(com_itau=True)
        # A tem 12 meses de cota; E so 3 (fundo novo); B e C, nenhum.
        report.performance_12m = self._bloco(
            12, "2025-10-31", "2026-10-14", "cota do fechamento de out/2025",
            {A: (13.45, 13.14), B: (-30.58, 10.31)},
        )
        report.performance = self._bloco(
            3, "2026-08-03", "2026-10-14", "cota do primeiro pregão do período",
            {A: (3.2, 3.1), E: (4.12, 4.02), B: (-1.0, 2.0)},
        )
        return report

    def test_tabela_por_cnpj_com_os_dois_prazos(self):
        from shelf_monitor.report.performance import com_rentabilidade, rentabilidade_por_cnpj

        report = self._relatorio()
        tabela = rentabilidade_por_cnpj(report.performance_12m, report.performance)
        self.assertEqual(tabela[A], {"retorno_12m": 13.45, "mediana_12m": 13.14,
                                     "retorno_3m": 3.2, "mediana_3m": 3.1})
        self.assertNotIn("retorno_12m", tabela[E])
        [linha_c] = com_rentabilidade([{"cnpj": C, "nome": "C"}], tabela)
        self.assertIsNone(linha_c["retorno_12m"])
        self.assertEqual(linha_c["nome"], "C")
        # Sem os objetos (bloco lido de JSON), a tabela sai vazia - e nada quebra.
        self.assertEqual(rentabilidade_por_cnpj({"available": True}, None), {})

    def test_digest_mostra_o_retorno_nos_gaps_e_nas_captacoes(self):
        from shelf_monitor.report.digest import render_html, render_markdown

        report = self._relatorio()
        md = render_markdown(report, None, "semanal")
        self.assertIn("| Cotistas | Rent. 12m | Ritmo na semana | Esforço |", md)
        self.assertIn("| 13,45% (+0,31pp) |", md)
        # Sem 12 meses de cota, 3 meses - com o prazo escrito.
        self.assertIn("| 4,12% em 3 meses (+0,10pp) |", md)
        self.assertRegex(
            md, r"Fundo A FIC FIM - .*, rentabilidade de 13,45% em 12 meses "
            r"\(\+0,31pp vs categoria\) - BTG, Itaú"
        )
        prazo = (
            "Rentabilidade: 12 meses, de 31/10/2025 a 14/10/2026 (cota do fechamento de "
            "out/2025); para o fundo sem esse histórico, 3 meses, de 03/08/2026 a "
            "14/10/2026 (cota do primeiro pregão do período), com o prazo na célula. "
            "Cota patrimonial ex-provento; o pp compara com a mediana da categoria no "
            "universo da CVM."
        )
        # A janela vai nas duas secoes: oportunidade e ritmo.
        self.assertEqual(md.count(prazo), 2)
        texto = _texto(render_html(report, None, "semanal"))
        self.assertIn("Rent. 12m", texto)
        self.assertIn("13,45% +0,31pp vs categoria", texto)
        self.assertIn("4,12% em 3 meses · +0,10pp vs categoria", texto)
        self.assertIn(prazo, texto)

    def test_painel_e_planilha_levam_o_retorno(self):
        from openpyxl import load_workbook

        from shelf_monitor.report.aggregate import build_dashboard_data
        from shelf_monitor.report.spreadsheet import render_spreadsheet

        report = self._relatorio()
        dados = build_dashboard_data(report)
        gap_a = next(f for f in dados["gaps"]["funds"] if f["cnpj"] == A)
        self.assertEqual(gap_a["retorno_12m"], 13.45)
        topo = {f["cnpj"]: f for f in dados["flow"]["top_captacao"]}
        self.assertEqual(topo[A]["mediana_12m"], 13.14)
        self.assertTrue(dados["flow"]["prazo_rentabilidade"].startswith("Rentabilidade: 12 meses"))
        self.assertEqual(dados["gaps"]["prazo_rentabilidade"], dados["flow"]["prazo_rentabilidade"])

        from shelf_monitor.report.retail_flow import build_grade_opportunities

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        grade = build_grade_opportunities(report.snapshots, report.retail_flow["_fluxos"])
        caminho = render_spreadsheet(
            config, report.snapshots, retail_flow=report.retail_flow, grade=grade,
            performance=report.performance, performance_12m=report.performance_12m,
        )
        livro = load_workbook(caminho)

        def coluna(aba, titulo: str) -> dict[str, object]:
            """{nome do fundo: valor da coluna}, pela linha de cabecalho da aba."""
            linhas = list(livro[aba].iter_rows(values_only=True))
            topo = next(i for i, linha in enumerate(linhas) if titulo in linha)
            idx = linhas[topo].index(titulo)
            return {linha[0]: linha[idx] for linha in linhas[topo + 1:] if linha and linha[0]}

        gaps_12m = coluna("Gaps XP", "Rentabilidade 12 meses (%)")
        self.assertEqual(gaps_12m["Fundo A FIC FIM"], 13.45)
        self.assertEqual(coluna("Gaps XP", "Rentabilidade 3 meses (%)")["Fundo E FIC FIM"], 4.12)
        captacao = coluna("Captação por fundo", "Rentabilidade 12 meses (%)")
        self.assertEqual(captacao["BTG Proprio B FIC FIM"], -30.58)
        # Retorno negativo aparece com o sinal - o formato da taxa o trocava por "-".
        aba = livro["Captação por fundo"]
        celula = next(
            c for linha in aba.iter_rows() for c in linha if c.value == -30.58
        )
        self.assertEqual(celula.number_format, '0.00"%";-0.00"%";0.00"%"')

        def notas(aba) -> str:
            return " ".join(str(c) for c in next(livro[aba].iter_cols(max_col=1, values_only=True)) if c)

        # Na planilha cada prazo tem a propria coluna: a nota da a janela de cada
        # uma, e a aba que so tem 12 meses so fala de 12 meses.
        doze = ("Rentabilidade 12 meses: de 31/10/2025 a 14/10/2026 "
                "(cota do fechamento de out/2025).")
        tres = ("Rentabilidade 3 meses: de 03/08/2026 a 14/10/2026 "
                "(cota do primeiro pregão do período).")
        self.assertIn(f"{doze} {tres} Cota patrimonial ex-provento", notas("Gaps XP"))
        for aba in ("Captação por fundo", "Oportunidades de grade"):
            self.assertIn(doze, notas(aba))
            self.assertNotIn("Rentabilidade 3 meses", notas(aba))
            self.assertNotIn("com o prazo na célula", notas(aba))
        self.assertIn("Rentabilidade 12 meses (%)", next(
            linha for linha in livro["Oportunidades de grade"].iter_rows(values_only=True)
            if "Fundo" in linha
        ))


class TestRadarSemCancelados(unittest.TestCase):
    """Na parte da CVM, nenhum fundo cancelado (pedido de 14/10/2026). O digest
    nao traz o radar desde 17/10/2026; o dashboard e a planilha seguem com ele."""

    @staticmethod
    def _fundo(cnpj: str, nome: str, casa: str, situacao: str) -> CvmFund:
        return CvmFund(cnpj=cnpj, name=nome, house=casa, matched_by="name",
                       manager_name="GESTORA DO GRUPO", situation=situacao)

    def _extratos(self) -> tuple[CvmSnapshot, CvmSnapshot]:
        normal, cancelado = "Em Funcionamento Normal", "Cancelado"
        antes = CvmSnapshot(
            collected_at=datetime(2026, 10, 7), status=SnapshotStatus.OK, funds=[
                self._fundo("71000000000101", "BTG ATIVO FIC", "btg", normal),
                self._fundo("71000000000202", "BTG VAI CANCELAR FIC", "btg", normal),
                self._fundo("71000000000303", "BTG JA CANCELADO FIC", "btg", cancelado),
                self._fundo("72000000000101", "ITAU SUMIU FIC", "itau", normal),
            ],
        )
        agora = CvmSnapshot(
            collected_at=datetime(2026, 10, 14), status=SnapshotStatus.OK, funds=[
                self._fundo("71000000000101", "BTG ATIVO FIC", "btg", normal),
                self._fundo("71000000000202", "BTG VAI CANCELAR FIC", "btg", cancelado),
                self._fundo("71000000000303", "BTG JA CANCELADO FIC", "btg", cancelado),
                self._fundo("72000000000202", "ITAU NOVO FIC", "itau", "Fase Pré-Operacional"),
            ],
        )
        return antes, agora

    def test_cancelado_sai_da_contagem_e_dos_movimentos(self):
        from shelf_monitor.diff import diff_cvm, radar_sem_cancelados

        antes, agora = self._extratos()
        ativos, diff = radar_sem_cancelados(agora, diff_cvm(agora, antes))
        self.assertEqual([f.name for f in ativos.funds], ["BTG ATIVO FIC", "ITAU NOVO FIC"])
        # A passagem para "Cancelado" e o fundo que sumiu do extrato nao viram movimento.
        self.assertEqual(diff.situation_changes, [])
        self.assertEqual(diff.cancelled, [])
        self.assertEqual([f.name for f in diff.new_registrations], ["ITAU NOVO FIC"])
        # O extrato da coleta segue inteiro: e ele que a proxima execucao compara.
        self.assertEqual(len(agora.funds), 4)

    def test_digest_nao_traz_o_radar(self):
        from shelf_monitor.diff import diff_cvm, radar_sem_cancelados
        from shelf_monitor.report.digest import render_html, render_markdown

        antes, agora = self._extratos()
        report = _relatorio()
        report.cvm_snapshot, report.cvm_diff = radar_sem_cancelados(agora, diff_cvm(agora, antes))
        md = render_markdown(report, None, "semanal")
        for texto in (md, _texto(render_html(report, None, "semanal"))):
            self.assertNotIn("Radar CVM", texto)
            self.assertNotIn("ITAU NOVO FIC", texto)
            self.assertNotIn("Pré-operacionais", texto)
            self.assertNotIn("Cancelado", texto)
            self.assertNotIn("CANCELAR", texto)
            self.assertNotIn("Saíram do cadastro", texto)


class TestLegendaDoRitmo(unittest.TestCase):
    """A legenda do ritmo no fim da secao de captacao (pedido de 17/10/2026): o
    que cada rotulo quer dizer, com a faixa da regra e as duas janelas."""

    def test_legenda_no_markdown_e_no_html(self):
        from shelf_monitor.report.digest import render_html, render_markdown

        report = _relatorio()
        md = render_markdown(report, None, "semanal")
        self.assertIn("**Como ler o ritmo** - a captação líquida por dia útil de agora "
                      "contra a de antes.", md)
        self.assertIn("| Ritmo | O que quer dizer |", md)
        self.assertIn("| Acelerando | já captava e passou a captar mais de 25% acima do "
                      "ritmo de antes |", md)
        for rotulo in ("Virou entrada", "Estável", "Desacelerando", "Virou saída",
                       "Saída persistente", "Sem série"):
            self.assertIn(f"| {rotulo} |", md)
        # A legenda fecha a secao de captacao, antes das vitrines.
        self.assertLess(md.index("Como ler o ritmo"), md.index("## Vitrines"))
        self.assertGreater(md.index("Como ler o ritmo"), md.index("## Ritmo de captação"))
        texto = _texto(render_html(report, None, "semanal"))
        self.assertIn("Como ler o ritmo · a captação líquida por dia útil", texto)
        self.assertIn("Saída persistente perdia dinheiro antes e continua perdendo", texto)


class TestServicoDaCasa(unittest.TestCase):
    def test_fundo_da_casa_sai_com_o_servico(self):
        from shelf_monitor.report.digest import render_html, render_markdown

        report = _relatorio(_lista_da_casa(**{C: "faas"}))
        md = render_markdown(report, None, "semanal")
        self.assertRegex(md, r"Fundo C FIC FIM - .* - XP · FaaS, BTG")
        self.assertIn("XP · FaaS, BTG", _texto(render_html(report, None, "semanal")))

    def test_lista_le_o_servico_e_avisa_o_desconhecido(self):
        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        pasta = config.distribuivel_dir
        pasta.mkdir(parents=True, exist_ok=True)
        (pasta / "casa.csv").write_text(
            "cnpj;nome;Serviço\n"
            "11.000.000/0001-01;Fundo A;Distribuição\n"
            "12000000000101;Fundo B;FaaS\n"
            "13000000000101;Fundo C;Fund as a Service\n"
            "14000000000101;Fundo D;consultoria\n",
            encoding="utf-8",
        )
        lista = BaseDistribuivel.ler(pasta)
        self.assertEqual(lista.rotulo_do_servico("11000000000101"), "Distribuição")
        self.assertEqual(lista.rotulo_do_servico(B), "FaaS")
        self.assertEqual(lista.rotulo_do_servico(C), "FaaS")
        # Servico desconhecido: o fundo fica na lista, sem servico, com aviso.
        self.assertEqual(lista.situacao_de(D), "distribui")
        self.assertIsNone(lista.rotulo_do_servico(D))
        self.assertTrue(any("serviço desconhecido 'consultoria'" in a for a in lista.avisos))
        self.assertEqual(lista.to_dict()["servicos"], {"distribuicao": 1, "faas": 2})


class TestEdicaoMensal(unittest.TestCase):
    def test_primeiro_digest_do_mes_e_o_mensal(self):
        from shelf_monitor.report.digest_visao import EDICAO_MENSAL, EDICAO_SEMANAL, edicao_do_digest

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        hoje = date(2026, 11, 3)
        self.assertEqual(edicao_do_digest(config, hoje), EDICAO_MENSAL)
        # O de outubro nao conta; o de hoje (rodar de novo) tambem nao.
        (config.reports_dir / "digest-2026-10-27.md").write_text("x", encoding="utf-8")
        (config.reports_dir / "digest-2026-11-03.md").write_text("x", encoding="utf-8")
        self.assertEqual(edicao_do_digest(config, hoje), EDICAO_MENSAL)
        self.assertEqual(edicao_do_digest(config, date(2026, 11, 10)), EDICAO_SEMANAL)
        self.assertEqual(edicao_do_digest(None, hoje), EDICAO_SEMANAL)

    def test_render_digest_decide_a_edicao_uma_vez(self):
        from shelf_monitor.report.digest import render_digest

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        report = _relatorio()
        report.performance = {
            "available": True, "casas": [], "janela": {"rotulo": "3 meses", "de": None, "ate": None},
        }
        md, html = render_digest(config, report)
        # Primeiro do mes, os dois arquivos na mesma edicao.
        self.assertIn("edição mensal", md.read_text(encoding="utf-8"))
        self.assertIn("edição mensal", html.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
