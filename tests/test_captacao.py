"""Visao de captacao: de onde vem o saldo de cada vitrine.

Ate 2026-09-24 o bloco dava um numero por casa e uma coluna de exclusivos, e o
leitor nao tinha como saber de onde vinha a diferenca. Medido naquele dia (07 a
09/2026): a vitrine da XP perdia R$ 1,3 bi somando +R$ 1,5 bi nos fundos que
divide com outra casa e -R$ 2,8 bi nos que so ela exibe; o saldo de +R$ 3,0 bi
do BTG vinha do produto proprio (+R$ 3,3 bi); e a diferenca entre as duas estava
em Renda Fixa e Multimercados. Os testes travam:

- a composicao do saldo (compartilhado x so na casa, proprio x terceiros, % do PL);
- a tabela por familia e a tendencia por casa;
- as frases da leitura, iguais em toda superficie;
- as listas com rotulo de casa, ticket por cotista novo e o que a XP perde;
- na previdencia, listas por dinheiro (o cotista e a seguradora) e sem os cortes
  suprimidos;
- no digest (v2, 14/10/2026), so o ritmo e as maiores captacoes e resgates: a
  captacao por casa ficou no dashboard e na planilha.

O informe diario e substituido por registros sinteticos (`ler_filtrado`), com
numeros redondos calculados a mao.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import re
import unittest
from datetime import date, datetime
from unittest import mock

from shelf_monitor.classify import KEY_CATEGORIA
from shelf_monitor.config import Config
from shelf_monitor.models import (
    BLOCO_CNPJ,
    BLOCO_COTISTAS,
    BLOCO_PRODUTO_PROPRIO,
    CollectorStrategy,
    FundOffer,
    RunReport,
    ShelfSnapshot,
    SnapshotStatus,
)
from shelf_monitor.report.retail_flow import (
    DailyFlow,
    FundFlow,
    build_retail_flow,
    leitura_da_captacao,
)

from fixtures_comuns import temp_config

# Renda fixa que nao e caixa: "Duracao Baixa Grau de Investimento", a categoria
# de antes, virou fundo caixa/simples e sai da analise (16/10/2026).
RF = "Renda Fixa Duração Livre Crédito Livre"
MULTI = "Multimercados Livre"
INDEPENDENTE = "GESTORA INDEPENDENTE LTDA"
MI = 1_000_000.0
PL = 1_000 * MI  # PL de cada fundo, constante
_DIAS = (date(2026, 8, 3), date(2026, 8, 28), date(2026, 9, 2), date(2026, 9, 18))

# (cnpj, nome, casas, categoria, gestor, captacao liquida em R$ mi, cotistas ganhos)
_FUNDOS = [
    ("11000000000101", "Compartilhado RF", ["xp", "btg"], RF, INDEPENDENTE, 100, 1000),
    ("12000000000101", "XP Próprio RF", ["xp"], RF, "XP GESTAO DE RECURSOS LTDA", -300, -500),
    ("13000000000101", "XP Multi", ["xp"], MULTI, INDEPENDENTE, -250, -100),
    ("21000000000101", "BTG Próprio RF", ["btg"], RF,
     "BTG PACTUAL ASSET MANAGEMENT S/A DTVM", 500, 2000),
    ("22000000000101", "BTG Multi", ["btg"], MULTI, INDEPENDENTE, 50, 50),
    # Capta sem cotista novo: fora das listas de fundos, dentro das de previdencia.
    ("23000000000101", "BTG Seguradora", ["btg"], MULTI, INDEPENDENTE, 100, 0),
]
# Completam cada vitrine acima de MIN_VITRINE (30), sem fluxo: XP 35 fundos, BTG 36.
_ENCHIMENTO = 32


def _oferta(platform: str, cnpj: str, nome: str, categoria: str, gestor: str) -> FundOffer:
    fundo = FundOffer(
        platform=platform, name=nome, cnpj=cnpj, manager_name=gestor,
        sources=[CollectorStrategy.PLAYWRIGHT.value],
    )
    fundo.extra[KEY_CATEGORIA] = categoria
    return fundo


def _cenario() -> tuple[list[ShelfSnapshot], list[DailyFlow]]:
    fundos = {"xp": [], "btg": []}
    registros: list[DailyFlow] = []
    for cnpj, nome, casas, categoria, gestor, liquido, cotistas in _FUNDOS:
        for casa in casas:
            fundos[casa].append(_oferta(casa, cnpj, nome, categoria, gestor))
        parcela = liquido * MI / len(_DIAS)
        for i, dia in enumerate(_DIAS):
            registros.append(DailyFlow(
                cnpj=cnpj, dia=dia,
                captacao=max(parcela, 0.0), resgates=max(-parcela, 0.0),
                pl=PL, cota=10.0, cotistas=1000 + round(cotistas * i / (len(_DIAS) - 1)),
            ))
    for casa, base in (("xp", 31), ("btg", 32)):
        for i in range(_ENCHIMENTO):
            cnpj = f"{base}{i:012d}"
            fundos[casa].append(_oferta(casa, cnpj, f"{casa} fundo {i}", RF, INDEPENDENTE))
            registros += [
                DailyFlow(cnpj=cnpj, dia=dia, pl=PL, cota=10.0, cotistas=1000)
                for dia in _DIAS
            ]
    snapshots = [
        ShelfSnapshot(
            platform=casa, collected_at=datetime(2026, 9, 19),
            status=SnapshotStatus.OK, strategy=CollectorStrategy.PLAYWRIGHT, funds=lista,
        )
        for casa, lista in fundos.items()
    ]
    return snapshots, registros


class _ConfigSemCotistas(Config):
    blocos_suprimidos = ((BLOCO_COTISTAS, "o cotista e a seguradora"),)


def _captacao(config: Config | None = None) -> tuple[list[ShelfSnapshot], dict]:
    snapshots, registros = _cenario()

    def ler(config, yyyymm, alvo, **_):
        return [r for r in registros if r.dia.strftime("%Y%m") == yyyymm and r.cnpj in alvo]

    with mock.patch("shelf_monitor.report.retail_flow.ler_filtrado", side_effect=ler), \
            mock.patch("shelf_monitor.report.retail_flow.ler_filtrado_sem_guardar",
                       side_effect=ler), \
            mock.patch("shelf_monitor.report.retail_flow.months_back",
                       return_value=["202608", "202609"]):
        flow = build_retail_flow(config or Config(), snapshots)
    return snapshots, flow


def _relatorio(suprimidos: dict[str, str] | None = None) -> RunReport:
    snapshots, flow = _captacao()
    return RunReport(
        generated_at=datetime(2026, 9, 24), snapshots=snapshots,
        retail_flow=flow, suprimidos=dict(suprimidos or {}),
    )


class TestComposicao(unittest.TestCase):
    def setUp(self):
        _, self.flow = _captacao()
        self.casas = {c["platform_id"]: c for c in self.flow["casas"]}

    def test_saldo_e_percentual_do_pl(self):
        xp, btg = self.casas["xp"], self.casas["btg"]
        self.assertAlmostEqual(xp["captacao_liquida"], -450 * MI)
        self.assertAlmostEqual(btg["captacao_liquida"], 750 * MI)
        self.assertAlmostEqual(xp["pl_inicio"], (3 + _ENCHIMENTO) * PL)
        self.assertAlmostEqual(xp["sobre_pl"], -450 / 35_000 * 100, places=3)

    def test_compartilhado_e_so_na_casa_somam_o_saldo(self):
        for casa in self.casas.values():
            self.assertAlmostEqual(
                casa["captacao_compartilhados"] + casa["captacao_exclusivos"],
                casa["captacao_liquida"],
            )
        self.assertAlmostEqual(self.casas["xp"]["captacao_compartilhados"], 100 * MI)
        self.assertAlmostEqual(self.casas["xp"]["captacao_exclusivos"], -550 * MI)

    def test_produto_proprio_e_relativo_a_casa(self):
        xp, btg = self.casas["xp"], self.casas["btg"]
        self.assertEqual((xp["proprios"], btg["proprios"]), (1, 1))
        self.assertAlmostEqual(xp["captacao_propria"], -300 * MI)
        self.assertAlmostEqual(btg["captacao_propria"], 500 * MI)
        self.assertAlmostEqual(xp["captacao_terceiros"], -150 * MI)
        self.assertAlmostEqual(btg["captacao_terceiros"], 250 * MI)
        dono = {f["nome"]: f["proprio_de"] for f in self.flow["fundos"]}
        self.assertEqual(dono["XP Próprio RF"], ["xp"])
        self.assertEqual(dono["Compartilhado RF"], [])

    def test_familia_conta_cada_fundo_uma_vez_no_universo(self):
        familias = {l["categoria"]: l for l in self.flow["por_categoria"]}
        self.assertEqual(set(familias), {"Renda Fixa", "Multimercados"})
        rf = familias["Renda Fixa"]
        # Compartilhado RF entra nas duas casas, uma vez no universo.
        self.assertAlmostEqual(rf["universo"]["captacao_liquida"], 300 * MI)
        self.assertEqual(rf["universo"]["fundos"], 3 + 2 * _ENCHIMENTO)
        self.assertAlmostEqual(rf["casas"]["xp"]["captacao_liquida"], -200 * MI)
        self.assertAlmostEqual(rf["casas"]["btg"]["captacao_liquida"], 600 * MI)
        # A familia que mais capta vem primeiro.
        self.assertEqual(self.flow["por_categoria"][0]["categoria"], "Renda Fixa")

    def test_tendencia_por_casa_cobre_a_vitrine(self):
        por_casa = self.flow["tendencia_por_casa"]
        # XP: 3 fundos + enchimento; BTG: 4 (o da seguradora a mais) + enchimento.
        self.assertEqual(sum(por_casa["xp"].values()), 3 + _ENCHIMENTO)
        self.assertEqual(sum(por_casa["btg"].values()), 4 + _ENCHIMENTO)

    def test_o_que_a_vitrine_da_xp_perde(self):
        nomes = [f["nome"] for f in self.flow["baseline_perdendo"]]
        self.assertEqual(nomes, ["XP Próprio RF", "XP Multi"])


class TestListasSemCotistas(unittest.TestCase):
    def test_fundo_que_capta_sem_cotista_novo_so_entra_quando_cotista_nao_mede_gente(self):
        _, fundos = _captacao()
        _, previdencia = _captacao(_ConfigSemCotistas())
        nomes = lambda f: [x["nome"] for x in f["top_captacao"]]  # noqa: E731
        self.assertNotIn("BTG Seguradora", nomes(fundos))
        self.assertIn("BTG Seguradora", nomes(previdencia))
        self.assertTrue(fundos["com_cotistas"])
        self.assertFalse(previdencia["com_cotistas"])

    def test_entrou_e_saiu(self):
        fundo = FundFlow(
            cnpj="1", nome="x", categoria="Renda Fixa", gestora=None,
            captacao_bruta=10 * MI, pl_inicio=PL, pl_fim=PL,
            cotistas_inicio=1, cotistas_fim=1,
        )
        self.assertFalse(fundo.entrou())
        self.assertTrue(fundo.entrou(com_cotistas=False))
        self.assertFalse(fundo.saiu(com_cotistas=False))


class TestLeitura(unittest.TestCase):
    def test_as_quatro_frases(self):
        _, flow = _captacao()
        frases = leitura_da_captacao(flow)
        self.assertEqual(frases, [
            "Captação líquida dos fundos da vitrine de 03/08/2026 a 18/09/2026: XP Investimentos "
            "-R$ 450 mi (-1,29% do PL de início), BTG Pactual +R$ 750 mi (+2,08%). "
            "Maior captação entre as casas: BTG Pactual.",
            "Fundos que outra casa também exibe: XP Investimentos +R$ 100 mi, BTG "
            "Pactual +R$ 100 mi. Fundos que só a própria casa exibe: XP Investimentos "
            "-R$ 550 mi, BTG Pactual +R$ 650 mi.",
            "Por família de ativo, a diferença entre XP Investimentos e BTG Pactual "
            "(-R$ 1,2 bi) vem de Renda Fixa (-R$ 800 mi) e Multimercados (-R$ 400 mi).",
            "Fundos da gestora da própria casa: XP Investimentos -R$ 300 mi, BTG "
            "Pactual +R$ 500 mi. De terceiros: XP Investimentos -R$ 150 mi, BTG "
            "Pactual +R$ 250 mi.",
        ])

    def test_supressoes_tiram_as_frases_do_corte(self):
        _, flow = _captacao()
        frases = leitura_da_captacao(flow, {BLOCO_CNPJ: "x", BLOCO_PRODUTO_PROPRIO: "y"})
        self.assertEqual(len(frases), 2)
        self.assertFalse(any("só a própria casa" in f for f in frases))
        self.assertFalse(any("gestora da própria casa" in f for f in frases))

    def test_familia_no_sentido_contrario_e_dita(self):
        _, flow = _captacao()
        # A XP passa a ganhar em Multimercados: a familia vira compensacao.
        for linha in flow["por_categoria"]:
            if linha["categoria"] == "Multimercados":
                linha["casas"]["xp"]["captacao_liquida"] = 450 * MI
        frase = next(f for f in leitura_da_captacao(flow) if f.startswith("Por família"))
        self.assertIn("vem de Renda Fixa (-R$ 800 mi)", frase)
        self.assertIn("Multimercados vai no sentido contrário (+R$ 300 mi)", frase)

    def test_sem_baseline_nao_ha_leitura(self):
        _, flow = _captacao()
        flow["casas"] = [c for c in flow["casas"] if c["platform_id"] != "xp"]
        self.assertEqual(leitura_da_captacao(flow), [])


class TestSuperficies(unittest.TestCase):
    def test_digest_md_e_html_publicam_os_mesmos_numeros(self):
        from shelf_monitor.report.digest import _html_ritmo, _md_ritmo
        from shelf_monitor.report.digest_visao import montar_visao

        visao = montar_visao(_relatorio())
        md = "\n".join(_md_ritmo(visao))
        html = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", _html_ritmo(visao)))
        valores = re.findall(r"[+-][\d.]+,\d MM", md)
        self.assertTrue(valores)
        for valor in valores:
            self.assertIn(valor, html)
        # Casa pelo nome curto, nunca o id; ticket por cotista novo nas listas.
        self.assertIn("XP, BTG", md)
        self.assertNotRegex(md, r" - (xp|btg)(,|$)")
        self.assertIn("por cotista novo", md)
        # A captacao por casa e as listas que sairam do semanal (14/10/2026)
        # continuam no dashboard e na planilha, nao no digest.
        self.assertNotIn("Na vitrine da XP Investimentos, perdendo", md)
        self.assertNotIn("| Renda Fixa |", md)
        self.assertNotIn("Só nesta vitrine", html)

    def test_digest_da_previdencia(self):
        from shelf_monitor.previdencia.config import ConfigPrevidencia
        from shelf_monitor.report.digest import _md_ritmo
        from shelf_monitor.report.digest_visao import montar_visao

        relatorio = _relatorio(dict(ConfigPrevidencia.blocos_suprimidos))
        md = "\n".join(_md_ritmo(montar_visao(relatorio)))
        for proibido in ("Compartilhados", "Só nesta vitrine", "Produto próprio", "cotistas"):
            self.assertNotIn(proibido, md)
        self.assertIn("Maiores captações - fundos ganhando dinheiro", md)

    def test_payload_do_dashboard(self):
        from shelf_monitor.report.aggregate import build_dashboard_data

        report = _relatorio()
        flow = build_dashboard_data(report)["flow"]
        self.assertEqual(flow["leitura"], leitura_da_captacao(report.retail_flow))
        self.assertEqual(flow["rotulos"]["btg"], "BTG Pactual")
        self.assertEqual(len(flow["por_categoria"]), 2)
        self.assertIn("xp", flow["tendencia_por_casa"])
        self.assertEqual(flow["baseline_perdendo"][0]["nome"], "XP Próprio RF")

    def test_planilha(self):
        from openpyxl import load_workbook

        from shelf_monitor.report.spreadsheet import render_spreadsheet

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)

        def abrir(suprimidos, nome):
            report = _relatorio(suprimidos)
            return load_workbook(render_spreadsheet(
                config, report.snapshots, generated_at=report.generated_at,
                path=config.reports_dir / nome, retail_flow=report.retail_flow,
                suprimidos=report.suprimidos,
            ))

        def textos(ws):
            return {str(c.value) for linha in ws.iter_rows() for c in linha if c.value is not None}

        fundos = abrir({}, "fundos.xlsx")
        casa = textos(fundos["Captação por casa"])
        self.assertIn("Compartilhados (R$)", casa)
        self.assertIn("Produto próprio (R$)", casa)
        self.assertIn(leitura_da_captacao(_relatorio().retail_flow)[0], casa)
        self.assertIn("Produto próprio de", [c.value for c in fundos["Captação por fundo"][7]])

        prev = abrir({BLOCO_CNPJ: "x", BLOCO_PRODUTO_PROPRIO: "y"}, "prev.xlsx")
        casa = textos(prev["Captação por casa"])
        self.assertNotIn("Compartilhados (R$)", casa)
        self.assertNotIn("Produto próprio (R$)", casa)
        self.assertNotIn("Produto próprio de", [c.value for c in prev["Captação por fundo"][7]])

    def test_deck_ganha_o_slide_de_captacao(self):
        from pptx import Presentation

        from shelf_monitor.report.deck import render_deck

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        deck = Presentation(render_deck(config, _relatorio()))
        slides = [
            [s.text_frame.text for s in slide.shapes if s.has_text_frame]
            for slide in deck.slides
        ]
        # O primeiro quadro de cada slide e o retangulo de fundo, sem texto.
        captacao = next(t for t in slides if "Captação dos fundos de vitrine" in t)
        self.assertIn(leitura_da_captacao(_relatorio().retail_flow)[0], captacao)


if __name__ == "__main__":
    unittest.main()
