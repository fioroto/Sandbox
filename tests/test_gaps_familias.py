"""Bloco de gaps com o mapa de familias e a lista distribuivel.

Costura 2 das duas combinadas na spec (docs/spec-gaps-xp-familias.md): o bloco
de sempre (`build_gap_block`), agora com os dois insumos novos. Snapshots de
fixture, o mapa montado com a CDA em memoria (tests/carteira_apoio.py) e a
lista distribuivel numa pasta temporaria.

Os fundos da fixture reproduzem os casos medidos em 26/09/2026: o espelho da
linha "Advisory" da XP, o feeder de marca propria do BTG sobre master de outra
gestora (BTGP Thunderbolt -> SPX Raptor), a estrategia que a XP tem so em
outra liquidez (Sparta Max D+61 contra D+121) e o fundo sem carteira publicada
que so o nome aproxima.
"""
from __future__ import annotations

import re
import unittest
from datetime import datetime
from pathlib import Path

from shelf_monitor.distribuivel import DISTRIBUI, PODE_DISTRIBUIR, BaseDistribuivel
from shelf_monitor.familias import (
    COBERTO_ESPELHO,
    GAP_ESTRATEGIA,
    OUTRA_LIQUIDEZ,
    SEM_DADO,
    MapaDeFamilias,
)
from shelf_monitor.models import (
    CollectorStrategy,
    FundOffer,
    RunReport,
    ShelfSnapshot,
    SnapshotStatus,
)
from shelf_monitor.report.aggregate import (
    ESFORCO_ABRIR_GESTORA,
    ESFORCO_ATIVAR,
    ESFORCO_EXPOR,
    ESFORCO_LIQUIDEZ,
    ESFORCO_NAO_TRANSFERIVEL,
    ESFORCO_OPERACIONAL,
    build_dashboard_data,
    build_gap_block,
    cnpjs_cobertos,
    frase_do_funil,
    ranquear_gaps,
)

try:
    from carteira_apoio import Cda, classe, mascara
    from fixtures_comuns import temp_config
except ImportError:  # rodando da raiz, como pacote
    from tests.carteira_apoio import Cda, classe, mascara
    from tests.fixtures_comuns import temp_config

ALFA = ("ALFA GESTORA DE RECURSOS", "11111111000191")
SPX = ("SPX GESTAO DE RECURSOS", "22222222000191")
SPARTA = ("SPARTA ADMINISTRADORA", "33333333000191")
DELTA = ("DELTA CAPITAL", "55555555000191")
EPSILON = ("EPSILON INVESTIMENTOS", "66666666000191")
BTG_ASSET = ("BTG PACTUAL ASSET MANAGEMENT S/A DTVM", "44444444000191")

# CNPJs, por papel no caso.
ALFA_MASTER = "10000000000101"
ALFA_ADVISORY = "10000000000202"   # vitrine XP
ALFA_ACCESS = "10000000000303"     # vitrine BTG - espelho coberto
SPX_MASTER = "20000000000101"
SPX_OUTRO = "20000000000202"       # vitrine XP: a SPX ja e distribuida
THUNDERBOLT = "20000000000303"     # vitrine BTG, gestor BTG, investe no SPX master
BTG_SELIC = "40000000000101"       # vitrine BTG, produto proprio
SPARTA_MASTER = "30000000000101"
SPARTA_XP = "30000000000202"       # vitrine XP, D+121
SPARTA_INTER = "30000000000303"    # vitrine Inter, D+61 - outra liquidez
DELTA_XP = "50000000000101"        # vitrine XP, sem carteira publicada
DELTA_INTER = "50000000000202"     # vitrine Inter, sem carteira publicada
EPSILON_RF = "60000000000101"      # vitrine Inter, na lista distribuivel
EPSILON_CP = "60000000000202"      # vitrine Inter, "pode distribuir"
EPSILON_ENVELOPE = "60000000000900"
EPSILON_CLASSE_A = "60000000000301"
EPSILON_CLASSE_B = "60000000000402"


def _cadastro() -> list[dict]:
    rf = "Renda Fixa Duração Livre Crédito Livre"
    return [
        classe(ALFA_MASTER, "ALFA CREDITO MASTER", *ALFA, anbima=rf, dias=(0, 1)),
        classe(ALFA_ADVISORY, "ALFA CREDITO ADVISORY FIC", *ALFA, anbima=rf),
        classe(ALFA_ACCESS, "BTG ACCESS ALFA CREDITO FIC", *BTG_ASSET, anbima=rf),
        classe(SPX_MASTER, "SPX RAPTOR MASTER", *SPX, dias=(0, 1)),
        classe(SPX_OUTRO, "SPX FALCON ACOES", *SPX, anbima="Ações Livre"),
        classe(THUNDERBOLT, "BTGP THUNDERBOLT FIC", *BTG_ASSET, dias=(360, 1)),
        classe(BTG_SELIC, "BTGP TESOURO SELIC", *BTG_ASSET,
               anbima="Renda Fixa Duração Baixa Soberano", dias=(0, 0)),
        classe(SPARTA_MASTER, "SPARTA MAX MASTER", *SPARTA, anbima=rf, dias=(0, 1)),
        classe(SPARTA_XP, "SPARTA MAX A 60 FIC", *SPARTA, anbima=rf, dias=(120, 1)),
        classe(SPARTA_INTER, "SPARTA MAX FIC", *SPARTA, anbima=rf, dias=(60, 1)),
        classe(DELTA_XP, "DELTA MACRO ADVISORY FIC", *DELTA),
        classe(DELTA_INTER, "DELTA MACRO FIC", *DELTA),
        classe(EPSILON_RF, "EPSILON RF", *EPSILON, codigo="777001"),
        classe(EPSILON_CP, "EPSILON CREDITO", *EPSILON),
        classe(EPSILON_CLASSE_A, "EPSILON MULTI A", *EPSILON, cnpj_fundo=EPSILON_ENVELOPE),
        classe(EPSILON_CLASSE_B, "EPSILON MULTI B", *EPSILON, cnpj_fundo=EPSILON_ENVELOPE),
    ]


def _cda() -> Cda:
    cda = Cda("202608")
    for master in (ALFA_MASTER, SPARTA_MASTER):
        (cda.pl(master, 1000.0).debenture(master, 500.0, "CSAN33")
         .debenture(master, 500.0, "AURE13"))
    for feeder, master in (
        (ALFA_ADVISORY, ALFA_MASTER), (ALFA_ACCESS, ALFA_MASTER),
        (THUNDERBOLT, SPX_MASTER),
        (SPARTA_XP, SPARTA_MASTER), (SPARTA_INTER, SPARTA_MASTER),
    ):
        cda.pl(feeder, 100.0).cota(feeder, master, 100.0)
    cda.pl(SPX_MASTER, 1000.0).titulo(SPX_MASTER, 600.0).acao(SPX_MASTER, 400.0, "PETR4")
    cda.pl(SPX_OUTRO, 100.0).acao(SPX_OUTRO, 100.0, "VALE3")
    cda.pl(BTG_SELIC, 100.0).titulo(BTG_SELIC, 100.0)
    cda.pl(EPSILON_RF, 100.0).titulo(EPSILON_RF, 100.0, titulo="NOTAS DO TESOURO NACIONAL SERIE B")
    cda.pl(EPSILON_CP, 100.0).debenture(EPSILON_CP, 100.0, "LIGHC4")
    return cda


def _mapa() -> MapaDeFamilias:
    return MapaDeFamilias.montar(_cadastro(), [_cda().resumo()])


def _fundo(platform: str, cnpj: str | None, nome: str, gestor: str,
           sources=("playwright",), categoria="Multimercados Livre") -> FundOffer:
    fundo = FundOffer(platform=platform, name=nome, cnpj=cnpj, manager_name=gestor)
    fundo.sources = list(sources)
    fundo.extra["classificacao_anbima"] = categoria
    return fundo


def _snap(platform: str, fundos: list[FundOffer], status=SnapshotStatus.OK) -> ShelfSnapshot:
    return ShelfSnapshot(
        platform=platform, collected_at=datetime(2026, 9, 25),
        strategy=CollectorStrategy.PLAYWRIGHT, status=status, funds=fundos,
    )


def _snapshots(xp_fontes=("playwright",)) -> list[ShelfSnapshot]:
    xp = _snap("xp", [
        _fundo("xp", ALFA_ADVISORY, "ALFA CREDITO ADVISORY FIC", ALFA[0], xp_fontes),
        _fundo("xp", SPX_OUTRO, "SPX FALCON ACOES", SPX[0], xp_fontes),
        _fundo("xp", SPARTA_XP, "SPARTA MAX A 60 FIC", SPARTA[0], xp_fontes),
        _fundo("xp", DELTA_XP, "DELTA MACRO ADVISORY FIC", DELTA[0], xp_fontes),
    ])
    btg = _snap("btg", [
        _fundo("btg", ALFA_ACCESS, "BTG ACCESS ALFA CREDITO FIC", BTG_ASSET[0], ("api-json",)),
        _fundo("btg", THUNDERBOLT, "BTGP THUNDERBOLT FIC", BTG_ASSET[0], ("api-json",)),
        _fundo("btg", BTG_SELIC, "BTGP TESOURO SELIC", BTG_ASSET[0], ("api-json",)),
    ])
    inter = _snap("inter", [
        _fundo("inter", SPARTA_INTER, "SPARTA MAX FIC", SPARTA[0]),
        _fundo("inter", DELTA_INTER, "DELTA MACRO FIC", DELTA[0]),
        _fundo("inter", EPSILON_RF, "EPSILON RF", EPSILON[0]),
        _fundo("inter", EPSILON_CP, "EPSILON CREDITO", EPSILON[0]),
    ])
    return [xp, btg, inter]


def _por_cnpj(linhas: list[dict]) -> dict[str, dict]:
    return {f["cnpj"]: f for f in linhas}


class _ComPasta(unittest.TestCase):
    def setUp(self):
        self.config, self.tmp = temp_config()
        self.addCleanup(self.tmp.cleanup)
        self.pasta = self.config.distribuivel_dir
        self.pasta.mkdir(parents=True, exist_ok=True)
        self.mapa = _mapa()

    def escrever(self, nome: str, texto: str, *, bom: bool = False) -> Path:
        caminho = self.pasta / nome
        caminho.write_bytes(("﻿" if bom else "").encode("utf-8") + texto.encode("utf-8"))
        return caminho

    def lista(self) -> BaseDistribuivel:
        return BaseDistribuivel.ler(self.pasta, self.mapa.resolver)


class TestFunil(unittest.TestCase):
    def setUp(self):
        self.bloco = build_gap_block(_snapshots(), familias=_mapa())

    def test_o_funil_fecha(self):
        funil = self.bloco["funil"]
        self.assertEqual(
            funil["aparentes"],
            funil["cobertos"] + funil["outra_liquidez"] + funil["reais"],
        )
        self.assertEqual(self.bloco["gaps_aparentes"], funil["aparentes"])
        self.assertEqual(self.bloco["total_gaps"], funil["reais"])
        self.assertEqual(
            len(self.bloco["funds"]) + len(self.bloco["cobertos"]), funil["aparentes"]
        )

    def test_espelho_do_master_que_a_xp_exibe_sai_da_conta_com_a_prova(self):
        coberto = _por_cnpj(self.bloco["cobertos"])[ALFA_ACCESS]
        self.assertEqual(coberto["veredito"]["tipo"], COBERTO_ESPELHO)
        self.assertEqual(coberto["veredito"]["cobre"], ALFA_ADVISORY)
        self.assertNotIn(ALFA_ACCESS, _por_cnpj(self.bloco["funds"]))
        self.assertEqual(cnpjs_cobertos(self.bloco), {ALFA_ACCESS})

    def test_estrategia_so_em_outra_liquidez_fica_na_lista_mas_nao_na_conta(self):
        gap = _por_cnpj(self.bloco["funds"])[SPARTA_INTER]
        self.assertEqual(gap["veredito"]["tipo"], OUTRA_LIQUIDEZ)
        self.assertEqual(self.bloco["gaps_liquidez"], 1)
        esforco = _por_cnpj(ranquear_gaps(self.bloco["funds"]))[SPARTA_INTER]["esforco"]
        self.assertEqual(esforco, ESFORCO_LIQUIDEZ)

    def test_feeder_de_marca_propria_sobre_master_de_terceiro_e_da_gestora_do_master(self):
        """BTGP Thunderbolt -> SPX Raptor: a SPX ja e distribuida pela XP."""
        gap = _por_cnpj(ranquear_gaps(self.bloco["funds"]))[THUNDERBOLT]
        self.assertEqual(gap["veredito"]["tipo"], GAP_ESTRATEGIA)
        self.assertEqual(gap["gestora_estrategia"], SPX[0])
        self.assertEqual(gap["esforco"], ESFORCO_OPERACIONAL)

    def test_sem_o_mapa_o_mesmo_feeder_saia_como_produto_proprio(self):
        sem_mapa = build_gap_block(_snapshots())
        gap = _por_cnpj(ranquear_gaps(sem_mapa["funds"]))[THUNDERBOLT]
        self.assertEqual(gap["esforco"], ESFORCO_NAO_TRANSFERIVEL)

    def test_master_da_propria_casa_continua_nao_transferivel(self):
        gap = _por_cnpj(ranquear_gaps(self.bloco["funds"]))[BTG_SELIC]
        self.assertEqual(gap["esforco"], ESFORCO_NAO_TRANSFERIVEL)

    def test_nome_parecido_sem_carteira_informa_e_continua_gap_real(self):
        gap = _por_cnpj(self.bloco["funds"])[DELTA_INTER]
        self.assertEqual(gap["veredito"]["tipo"], SEM_DADO)
        self.assertEqual(gap["veredito"]["possivel_equivalente"], DELTA_XP)
        self.assertTrue(gap["veredito"]["gap_real"])

    def test_a_lista_abre_pelo_gap_real_mais_barato_e_a_liquidez_vem_depois(self):
        """Versao de liquidez custa o mesmo que a inclusao operacional, mas a
        estrategia ja esta na prateleira: vem depois dela, antes de abrir
        gestora nova."""
        ordem = [f["cnpj"] for f in ranquear_gaps(self.bloco["funds"])]
        self.assertLess(ordem.index(THUNDERBOLT), ordem.index(SPARTA_INTER))
        self.assertLess(ordem.index(SPARTA_INTER), ordem.index(EPSILON_RF))
        self.assertEqual(ordem[-1], BTG_SELIC)

    def test_gestora_nova_e_abrir_gestora(self):
        gap = _por_cnpj(ranquear_gaps(self.bloco["funds"]))[EPSILON_RF]
        self.assertEqual(gap["esforco"], ESFORCO_ABRIR_GESTORA)

    def test_categorias_contam_gaps_reais_e_cobertos(self):
        categorias = {c["category"]: c for c in self.bloco["categories"]}
        self.assertEqual(sum(c["gap_count"] for c in categorias.values()), self.bloco["total_gaps"])
        self.assertEqual(sum(c["cobertos"] for c in categorias.values()), 1)

    def test_a_frase_do_funil_diz_os_numeros_do_bloco(self):
        frase = frase_do_funil(self.bloco)
        self.assertIn(f"Dos {self.bloco['gaps_aparentes']} gaps aparentes", frase)
        self.assertIn("1 é espelho", frase)
        self.assertIn(f"sobram {self.bloco['total_gaps']} gaps reais", frase)

    def test_ordem_nao_depende_da_ordem_de_coleta(self):
        invertidos = [
            _snap(s.platform, list(reversed(s.funds))) for s in reversed(_snapshots())
        ]
        outro = build_gap_block(invertidos, familias=_mapa())
        self.assertEqual(
            [f["cnpj"] for f in ranquear_gaps(outro["funds"])],
            [f["cnpj"] for f in ranquear_gaps(self.bloco["funds"])],
        )
        self.assertEqual(
            [f["cnpj"] for f in outro["cobertos"]], [f["cnpj"] for f in self.bloco["cobertos"]]
        )


class TestSemCarteira(unittest.TestCase):
    def test_sem_o_mapa_o_bloco_e_o_de_antes_com_o_aviso(self):
        bloco = build_gap_block(_snapshots())
        self.assertEqual(bloco["total_gaps"], bloco["gaps_aparentes"])
        self.assertEqual(bloco["cobertos"], [])
        self.assertFalse(bloco["carteira"]["available"])
        self.assertIn("espelho e equivalência não foram verificados", bloco["carteira"]["aviso"])
        self.assertIsNone(frase_do_funil(bloco))


class TestListaDistribuivel(_ComPasta):
    def _pasta_completa(self):
        self.escrever(
            "a_equipe_rf.csv",
            "cnpj;nome;situacao;data_referencia\n"
            f"{mascara(EPSILON_RF)};EPSILON RF;distribui;15/09/2026\n"
            f"{mascara(EPSILON_CP)};EPSILON CREDITO;pode_distribuir;15/09/2026\n"
            "12.345;LIXO;distribui;\n",
            bom=True,
        )
        self.escrever(
            "b_private.csv",
            "CNPJ,Nome,Situação,Código CVM\n"
            f"{EPSILON_ENVELOPE},EPSILON MULTI,Pode distribuir,\n"
            ",SO O CODIGO,distribui,777001\n"
            f"{mascara(SPX_OUTRO)},SPX FALCON,talvez,\n"
            "99.999.999/0001-99,NAO EXISTE,distribui,\n",
        )

    def test_os_arquivos_sao_unidos_e_resolvidos_para_classe(self):
        self._pasta_completa()
        lista = self.lista()
        self.assertEqual(lista.situacao_de(EPSILON_RF), DISTRIBUI)
        self.assertEqual(lista.situacao_de(EPSILON_CP), PODE_DISTRIBUIR)
        # CNPJ de fundo com duas classes vale para as duas, com aviso.
        self.assertEqual(lista.situacao_de(EPSILON_CLASSE_A), PODE_DISTRIBUIR)
        self.assertEqual(lista.situacao_de(EPSILON_CLASSE_B), PODE_DISTRIBUIR)
        self.assertTrue(any("2 classes" in a for a in lista.avisos))
        self.assertEqual([a.nome for a in lista.arquivos], ["a_equipe_rf.csv", "b_private.csv"])
        # A data da coluna, quando ha; sem a coluna, a do arquivo.
        self.assertEqual(lista.arquivos[0].data_referencia.isoformat(), "2026-09-15")
        self.assertIsNotNone(lista.arquivos[1].data_referencia)

    def test_linha_invalida_sai_com_arquivo_linha_e_motivo(self):
        self._pasta_completa()
        invalidas = {(i.arquivo, i.linha): i.motivo for i in self.lista().invalidas}
        self.assertIn("CNPJ inválido", invalidas[("a_equipe_rf.csv", 4)])
        self.assertIn("situação desconhecida", invalidas[("b_private.csv", 4)])
        self.assertIn("não encontrado no cadastro", invalidas[("b_private.csv", 5)])
        self.assertEqual(len(invalidas), 3)

    def test_repetido_com_situacoes_diferentes_fica_com_distribui(self):
        self.escrever("1.csv", f"cnpj;situacao\n{EPSILON_CP};pode_distribuir\n")
        self.escrever("2.csv", f"cnpj;situacao\n{EPSILON_CP};distribui\n")
        self.escrever("3.csv", f"cnpj;situacao\n{EPSILON_CP};pode_distribuir\n")
        self.assertEqual(self.lista().situacao_de(EPSILON_CP), DISTRIBUI)

    def test_gap_da_vitrine_coberto_pela_lista_vira_expor_ou_ativar(self):
        self._pasta_completa()
        bloco = build_gap_block(_snapshots(), familias=self.mapa, distribuivel=self.lista())
        esforcos = {f["cnpj"]: f["esforco"] for f in ranquear_gaps(bloco["funds"])}
        self.assertEqual(esforcos[EPSILON_RF], ESFORCO_EXPOR)
        self.assertEqual(esforcos[EPSILON_CP], ESFORCO_ATIVAR)
        # Continua gap da VITRINE: o cliente nao o ve. So o esforco muda.
        self.assertIn(EPSILON_RF, _por_cnpj(bloco["funds"]))
        # E abre a lista: expor e o degrau mais barato.
        self.assertEqual(ranquear_gaps(bloco["funds"])[0]["esforco"], ESFORCO_EXPOR)
        # A outra lente mede a vitrine concorrente contra a lista: o que ja
        # esta nela nao e gap aparente dessa lente (7 fundos, 2 na lista).
        lente = bloco["lentes"]["distribuivel"]
        self.assertTrue(lente["available"])
        self.assertEqual(lente["funil"]["aparentes"], 5)

    def test_pasta_vazia_suprime_so_a_lente_distribuivel(self):
        bloco = build_gap_block(_snapshots(), familias=self.mapa, distribuivel=self.lista())
        self.assertTrue(bloco["available"])
        self.assertEqual(bloco["base"], "vitrine")
        lente = bloco["lentes"]["distribuivel"]
        self.assertFalse(lente["available"])
        self.assertIn("nenhum CSV", lente["reason"])

    def test_sem_vitrine_da_xp_a_lista_vira_a_base(self):
        """Dia de bloqueio de CDN: a XP so veio pelo Open Finance."""
        self._pasta_completa()
        snapshots = _snapshots(xp_fontes=("open-finance",))
        sem_lista = build_gap_block(snapshots, familias=self.mapa)
        self.assertFalse(sem_lista["available"])
        self.assertIn("sem vitrine coletada", sem_lista["reason"])

        com_lista = build_gap_block(snapshots, familias=self.mapa, distribuivel=self.lista())
        self.assertTrue(com_lista["available"])
        self.assertEqual(com_lista["base"], "distribuivel")
        self.assertIn("lista distribuível", com_lista["baseline_warning"])
        # Na lista, nao e gap; fora dela, e - com o veredito contra a lista.
        todos = _por_cnpj(com_lista["funds"] + com_lista["cobertos"])
        self.assertNotIn(EPSILON_RF, todos)
        self.assertIn(ALFA_ACCESS, todos)
        self.assertEqual(com_lista["gaps_aparentes"], 5)

    def test_sem_base_de_mercado_a_lista_entra_sem_validacao_com_aviso(self):
        self.escrever("x.csv", f"cnpj;situacao\n{EPSILON_RF};distribui\n")
        lista = BaseDistribuivel.ler(self.pasta, None)
        self.assertEqual(lista.situacao_de(EPSILON_RF), DISTRIBUI)
        self.assertTrue(any("não foram conferidos" in a for a in lista.avisos))


class TestOutrosBlocosUsamOMesmoVeredito(unittest.TestCase):
    def test_oportunidade_de_grade_nao_lista_o_que_a_xp_ja_cobre(self):
        from shelf_monitor.report.retail_flow import FundFlow, build_grade_opportunities

        def fluxo(cnpj: str, nome: str) -> FundFlow:
            return FundFlow(
                cnpj=cnpj, nome=nome, categoria="Multimercados Livre", gestora=None,
                vitrines=["btg"], captacao_bruta=100_000_000.0, resgates=0.0,
                pl_inicio=500_000_000.0, pl_fim=600_000_000.0,
                cotistas_inicio=1000, cotistas_fim=5000, dias=60,
            )

        # O bloco de grade so cruza vitrines com 30 fundos ou mais: enche as
        # tres com os mesmos fundos, que nao viram gap de ninguem.
        snapshots = _snapshots()
        for snapshot in snapshots:
            snapshot.funds += [
                _fundo(snapshot.platform, f"8800000000{n:02d}01", f"COMUM {n}", ALFA[0],
                       snapshot.funds[0].sources)
                for n in range(30)
            ]
        fluxos = [fluxo(ALFA_ACCESS, "BTG ACCESS ALFA"), fluxo(THUNDERBOLT, "THUNDERBOLT")]
        cobertos = cnpjs_cobertos(build_gap_block(snapshots, familias=_mapa()))
        sem = build_grade_opportunities(snapshots, fluxos)
        com = build_grade_opportunities(snapshots, fluxos, cobertos=cobertos)
        self.assertEqual(sem["total"], 2)
        self.assertEqual(com["total"], 1)
        self.assertEqual(com["total_cobertos"], 1)

    def test_consenso_nao_chama_de_lacuna_o_que_a_xp_ja_cobre(self):
        from shelf_monitor.report.competitive import build_xp_decisions

        snapshots = _snapshots()
        # O espelho coberto na vitrine das duas concorrentes: vira "consenso".
        snapshots[2].funds.append(
            _fundo("inter", ALFA_ACCESS, "BTG ACCESS ALFA CREDITO FIC", BTG_ASSET[0])
        )
        for snapshot in snapshots:
            snapshot.funds += [
                _fundo(snapshot.platform, f"8800000000{n:02d}01", f"COMUM {n}", ALFA[0],
                       snapshot.funds[0].sources)
                for n in range(30)
            ]

        def consenso(decisoes: dict) -> list[dict]:
            return [
                d for d in decisoes.get("decisoes") or []
                if d["titulo"].startswith("Fundos que a concorrência exibe em consenso")
            ]

        cobertos = cnpjs_cobertos(build_gap_block(snapshots, familias=_mapa()))
        self.assertEqual(len(consenso(build_xp_decisions(snapshots))), 1)
        self.assertEqual(consenso(build_xp_decisions(snapshots, cobertos=cobertos)), [])


class TestSuperficies(_ComPasta):
    def _relatorio(self) -> RunReport:
        self.escrever("lista.csv", f"cnpj;situacao\n{EPSILON_RF};distribui\n")
        return RunReport(
            generated_at=datetime(2026, 9, 26, 8, 0), snapshots=_snapshots(),
            familias=self.mapa, distribuivel=self.lista(),
        )

    def test_digest_md_e_html_publicam_o_mesmo_funil_e_os_mesmos_cobertos(self):
        from shelf_monitor.report.digest import render_html, render_markdown

        relatorio = self._relatorio()
        markdown = render_markdown(relatorio, self.config)
        html = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", render_html(relatorio, self.config)))
        import html as html_mod
        html = html_mod.unescape(html)
        frase = frase_do_funil(build_gap_block(
            relatorio.snapshots, familias=self.mapa, distribuivel=relatorio.distribuivel
        ))
        self.assertIn(frase, markdown)
        self.assertIn(frase, html)
        self.assertIn("Já cobertos", markdown)
        cnpjs_md = set(re.findall(r"\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}", markdown))
        cnpjs_html = set(re.findall(r"\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}", html))
        self.assertIn(mascara(ALFA_ACCESS), cnpjs_md)
        self.assertEqual(set(), cnpjs_md - cnpjs_html)
        self.assertIn("carteira mensal da CVM (CDA) de ago/2026", markdown)

    def test_dashboard_leva_o_funil_os_cobertos_e_as_notas(self):
        dados = build_dashboard_data(self._relatorio())["gaps"]
        self.assertEqual(dados["frase_funil"], frase_do_funil(dados))
        self.assertEqual([f["cnpj"] for f in dados["cobertos"]], [ALFA_ACCESS])
        self.assertTrue(any("Lista distribuível" in n for n in dados["notas_cobertura"]))

    def test_planilha_ganha_gaps_cobertos_e_familias(self):
        from openpyxl import load_workbook

        from shelf_monitor.report.spreadsheet import render_spreadsheet

        relatorio = self._relatorio()
        caminho = render_spreadsheet(
            self.config, relatorio.snapshots, generated_at=relatorio.generated_at,
            familias=self.mapa, distribuivel=relatorio.distribuivel,
        )
        livro = load_workbook(caminho)
        self.assertIn("Gaps cobertos", livro.sheetnames)
        self.assertIn("Famílias", livro.sheetnames)
        cobertos = livro["Gaps cobertos"]
        self.assertEqual(cobertos.cell(row=2, column=2).value, mascara(ALFA_ACCESS))
        self.assertEqual(cobertos.cell(row=2, column=8).value, mascara(ALFA_ADVISORY))
        gaps = livro["Gaps XP"]
        cabecalho = [c.value for c in gaps[1]]
        self.assertIn("Veredito", cabecalho)
        self.assertIn("Na lista distribuível", cabecalho)
        familias = [
            r[1] for r in livro["Famílias"].iter_rows(min_row=2, values_only=True) if r[1]
        ]
        self.assertIn("ALFA CREDITO MASTER", familias)


if __name__ == "__main__":
    unittest.main()
