"""A taxa em pontos percentuais e a taxa global nos gaps por casa (20/10/2026).

O `maxAdminFee` do Open Finance vem em fracao na XP, no BTG e no Itau (0,02 para
2%) e em percentual na Nu; e, em parte dos fundos, so com a taxa de administracao -
nas vitrines de 14/10/2026, o Itau Small Cap Valuation tinha 0,0005 no Open
Finance e 2,0% no site do Itau (`shelf_monitor.taxas`). Os testes travam:

- a escala decidida por catalogo, no coletor do Open Finance, sem aviso de coleta;
- a taxa da vitrine vencendo a do catalogo no merge, e a do catalogo ficando
  quando o site nao publica;
- o snapshot da versao 1 corrigido na leitura, sem o diff acusar mudanca de taxa;
- a linha do gap com a menor taxa global entre as casas que o exibem;
- a coluna Taxa global nas tabelas de gaps do digest e na aba Gaps XP.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import unittest

from shelf_monitor.merge import merge_offers
from shelf_monitor.models import KEY_TAXA_VITRINE, CollectorStrategy, FundOffer

try:
    from fixtures_comuns import temp_config
    from test_digest_v2 import A, E, _relatorio, _texto
except ImportError:  # rodando da raiz, como pacote
    from tests.fixtures_comuns import temp_config
    from tests.test_digest_v2 import A, E, _relatorio, _texto

OF = CollectorStrategy.OPEN_FINANCE.value
NAVEGADOR = CollectorStrategy.PLAYWRIGHT.value


def _fundo_v1(cnpj: str, taxa: float | None, sources: list[str], **extra) -> dict:
    """Um fundo como o snapshot da versao 1 o gravava."""
    fundo = FundOffer(platform="btg", name=f"Fundo {cnpj[:2]}", cnpj=cnpj,
                      admin_fee_pct=taxa, sources=sources)
    fundo.extra.update(extra)
    return fundo.to_dict()


def _snapshot_v1(fundos: list[dict]) -> dict:
    return {
        "schema_version": 1, "platform": "btg", "collected_at": "2026-09-28T08:00:00",
        "strategy": OF, "status": "ok", "funds": fundos,
    }


class TestEscalaDoCatalogo(unittest.TestCase):
    def test_catalogo_em_fracao_vira_percentual(self):
        """XP, BTG e Itau publicam o maxAdminFee em fracao; a Nu, em percentual."""
        from shelf_monitor.taxas import catalogo_em_percentual

        fracao = [FundOffer(platform="btg", name=str(t), admin_fee_pct=t)
                  for t in (0.02, 0.0242, 0.005, 0.0, None, 2.0)]
        self.assertEqual(catalogo_em_percentual(fracao), 3)
        # Acima de 0,1 ja e percentual: fracao acima de 10% ao ano nao e taxa.
        self.assertEqual([o.admin_fee_pct for o in fracao], [2.0, 2.42, 0.5, 0.0, None, 2.0])
        percentual = [FundOffer(platform="nu", name=str(t), admin_fee_pct=t)
                      for t in (0.15, 0.5, 1.0, 0.05)]
        self.assertEqual(catalogo_em_percentual(percentual), 0)
        self.assertEqual([o.admin_fee_pct for o in percentual], [0.15, 0.5, 1.0, 0.05])

    def test_coletor_do_open_finance_entrega_em_percentual(self):
        from unittest import mock

        from shelf_monitor.collectors.open_finance import OpenFinanceFunds
        from shelf_monitor.config import PLATFORMS, Config

        cliente = OpenFinanceFunds(Config(), PLATFORMS["btg"], session=None)
        pagina = {"data": [
            {"name": f"Fundo {i}", "cnpjNumber": f"{i:014d}", "fees": {"maxAdminFee": taxa}}
            for i, taxa in enumerate(("0.020000", "0.005000", "0.015000"), start=1)
        ], "meta": {"totalPages": 1, "totalRecords": 3}}
        with mock.patch.object(cliente, "_discover_base_url", return_value="https://x"), \
                mock.patch("shelf_monitor.collectors.open_finance.get_json", return_value=pagina):
            ofertas = cliente.fetch()
        self.assertEqual([o.admin_fee_pct for o in ofertas], [2.0, 0.5, 1.5])
        self.assertEqual(cliente.warnings, [])  # aviso marcaria a coleta como parcial


class TestTaxaDaVitrine(unittest.TestCase):
    def test_a_taxa_da_vitrine_vence_a_do_catalogo(self):
        """O Itau Small Cap Valuation: 0,05% no Open Finance (so a de
        administracao), 2,0% no site do Itau."""
        catalogo = FundOffer(platform="itau", name="Small Cap", cnpj="2" * 14,
                             admin_fee_pct=0.05)
        vitrine = FundOffer(platform="itau", name="Small Cap", cnpj="2" * 14,
                            admin_fee_pct=2.0)
        [unido] = merge_offers({OF: [catalogo], NAVEGADOR: [vitrine]})
        self.assertEqual(unido.admin_fee_pct, 2.0)
        self.assertEqual(unido.taxa_da_vitrine, 2.0)
        self.assertEqual(unido.extra["divergencias"]["admin_fee_pct"], {
            "valor_usado": 2.0, "fonte_usada": NAVEGADOR,
            "valor_alternativo": 0.05, "fonte_alternativa": OF,
        })

    def test_taxa_de_site_acima_do_teto_nao_e_taxa(self):
        """O site do BTG publicava 100,0 no Brasil Plural Debentures (28/09/2026):
        fica a do catalogo."""
        catalogo = FundOffer(platform="btg", name="Plural", cnpj="9" * 14, admin_fee_pct=1.1)
        vitrine = FundOffer(platform="btg", name="Plural", cnpj="9" * 14, admin_fee_pct=100.0)
        [unido] = merge_offers({OF: [catalogo], NAVEGADOR: [vitrine]})
        self.assertEqual(unido.admin_fee_pct, 1.1)
        self.assertIsNone(unido.taxa_da_vitrine)
        so_site = FundOffer(platform="btg", name="Z", cnpj="1" * 14, admin_fee_pct=90.0,
                            sources=[NAVEGADOR])
        self.assertIsNone(so_site.taxa_da_vitrine)

    def test_sem_taxa_na_vitrine_fica_a_do_catalogo(self):
        catalogo = FundOffer(platform="btg", name="Y", cnpj="5" * 14, admin_fee_pct=1.2)
        vitrine = FundOffer(platform="btg", name="Y", cnpj="5" * 14)
        [unido] = merge_offers({OF: [catalogo], NAVEGADOR: [vitrine]})
        self.assertEqual(unido.admin_fee_pct, 1.2)
        self.assertIsNone(unido.taxa_da_vitrine)


class TestSnapshotDaVersao1(unittest.TestCase):
    """A coleta gravada antes de 20/10/2026 e corrigida na leitura."""

    def _lido(self):
        from shelf_monitor.models import ShelfSnapshot

        fundos = [
            # JGP Hedge: 0,02 no catalogo, 2,0 no site, guardado como divergencia.
            _fundo_v1("3" * 14, 0.02, [OF, NAVEGADOR], divergencias={"admin_fee_pct": {
                "valor_usado": 0.02, "fonte_usada": OF,
                "valor_alternativo": 2.0, "fonte_alternativa": NAVEGADOR,
            }}),
            # Itau Small Cap: 0,0005 no catalogo (so a de administracao), 2,0 no site.
            _fundo_v1("8" * 14, 0.0005, [OF, NAVEGADOR], divergencias={"admin_fee_pct": {
                "valor_usado": 0.0005, "fonte_usada": OF,
                "valor_alternativo": 2.0, "fonte_alternativa": NAVEGADOR,
            }}),
            _fundo_v1("4" * 14, 0.015, [OF]),
            _fundo_v1("6" * 14, 0.005, [OF]),
            _fundo_v1("7" * 14, 1.5, [NAVEGADOR]),
        ]
        return ShelfSnapshot.from_dict(_snapshot_v1(fundos))

    def test_taxa_em_percentual_e_a_vitrine_na_frente(self):
        snapshot = self._lido()
        taxas = {f.cnpj: f.admin_fee_pct for f in snapshot.funds}
        self.assertEqual(taxas, {
            "3" * 14: 2.0, "8" * 14: 2.0, "4" * 14: 1.5, "6" * 14: 0.5, "7" * 14: 1.5,
        })
        jgp = next(f for f in snapshot.funds if f.cnpj == "3" * 14)
        self.assertEqual(jgp.taxa_da_vitrine, 2.0)
        # Convertida, a do catalogo bate com a do site: nao diverge mais.
        self.assertNotIn("admin_fee_pct", jgp.extra.get("divergencias") or {})
        small_cap = next(f for f in snapshot.funds if f.cnpj == "8" * 14)
        self.assertEqual(small_cap.extra["divergencias"]["admin_fee_pct"], {
            "valor_usado": 2.0, "fonte_usada": NAVEGADOR,
            "valor_alternativo": 0.05, "fonte_alternativa": OF,
        })
        so_catalogo = next(f for f in snapshot.funds if f.cnpj == "4" * 14)
        self.assertIsNone(so_catalogo.taxa_da_vitrine)
        # Lido e regravado, sai na versao nova - e nao e convertido de novo.
        self.assertEqual(snapshot.to_dict()["schema_version"], 2)

    def test_versao_2_nao_e_convertida(self):
        from shelf_monitor.models import ShelfSnapshot

        dados = self._lido().to_dict()
        relido = ShelfSnapshot.from_dict(dados)
        self.assertEqual({f.cnpj: f.admin_fee_pct for f in relido.funds}, {
            "3" * 14: 2.0, "8" * 14: 2.0, "4" * 14: 1.5, "6" * 14: 0.5, "7" * 14: 1.5,
        })

    def test_diff_contra_a_coleta_antiga_nao_acusa_mudanca_de_taxa(self):
        from shelf_monitor.diff import diff_shelf
        from shelf_monitor.models import ShelfSnapshot

        antigo = self._lido()
        novo = ShelfSnapshot.from_dict(antigo.to_dict())
        self.assertEqual(diff_shelf(novo, antigo).changed, [])


class TestTaxaGlobalNosGaps(unittest.TestCase):
    """A no BTG (2,00%) e no Itau (1,90%); E so no Itau, sem taxa no site."""

    def setUp(self):
        self.report = _relatorio(com_itau=True)
        for snapshot in self.report.snapshots:
            for oferta in snapshot.funds:
                if oferta.cnpj == A:
                    oferta.extra[KEY_TAXA_VITRINE] = {"btg": 2.0, "itau": 1.9}[snapshot.platform]

    def test_a_menor_taxa_global_entre_as_casas(self):
        from shelf_monitor.report.digest_visao import montar_visao

        linhas = {f["cnpj"]: f for f in montar_visao(self.report, None, "semanal")
                  .oportunidade["linhas"]}
        self.assertEqual(linhas[A]["taxa_global"], 1.9)
        self.assertIsNone(linhas[E]["taxa_global"])

    def test_digest_traz_a_coluna_e_a_nota(self):
        from shelf_monitor.report.digest import render_html, render_markdown

        md = render_markdown(self.report, None, "semanal")
        # Rent. 12m e Taxa global lado a lado; sem serie de cota, a rentabilidade e "-".
        self.assertRegex(md, r"\| Fundo A FIC FIM .*\| \+49 \| - \| 1,90% \| acelerando \|")
        self.assertRegex(md, r"\| Fundo E FIC FIM .*\| \+49 \| - \| - \| acelerando \|")
        self.assertIn("_Taxa global: a que o site da casa mostra para o fundo, em % ao ano", md)
        texto = _texto(render_html(self.report, None, "semanal"))
        self.assertIn("Rent. 12m Taxa global Esforço", texto)
        self.assertIn("1,90%", texto)

    def test_planilha_traz_a_taxa_global_nos_gaps(self):
        from openpyxl import load_workbook

        from shelf_monitor.report.spreadsheet import render_spreadsheet

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        caminho = render_spreadsheet(config, self.report.snapshots,
                                     retail_flow=self.report.retail_flow)
        aba = load_workbook(caminho)["Gaps XP"]
        cabecalho = [c.value for c in aba[1]]
        coluna = cabecalho.index("Taxa global na vitrine (%)")
        por_cnpj = {
            linha[cabecalho.index("CNPJ")]: linha[coluna]
            for linha in aba.iter_rows(min_row=2, values_only=True) if linha[0]
        }
        self.assertEqual(por_cnpj["11.000.000/0001-01"], 1.9)


if __name__ == "__main__":
    unittest.main()
