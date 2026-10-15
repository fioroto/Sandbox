"""Blocos suprimidos: a previdencia nao publica cotistas nem cruza CNPJ.

Em previdencia o cotista do fundo (FIE) e a seguradora - 438 dos 439 fundos de
vitrine com informe tinham 1 cotista em 23/09 - e cada seguradora distribui pelo
proprio veiculo, entao XP e BTG dividiam 1 CNPJ. Ate a onda 1 da auditoria as
superficies publicavam "ponderada por cotistas", "fundos de muita gente",
"+0 cotistas", "225 fundos que faltam na XP", "inclusao operacional" e
"217 fundos exclusivos" sobre esses numeros.

Os testes rodam o MESMO relatorio duas vezes - com as supressoes da
previdencia e sem nenhuma, como no monitor de fundos - para travar os dois
lados: some na previdencia, fica em fundos.

Os blocos de captacao e rentabilidade vem de `fixtures/previdencia/
blocos_2026-09-23.json` (gerado por `make_blocos.py` a partir do relatorio
real de 23/09, com as listas aparadas).

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import json
import re
import unittest
from datetime import datetime
from pathlib import Path

from shelf_monitor.diff import diff_shelf
from shelf_monitor.models import (
    BLOCO_CNPJ,
    BLOCO_COTISTAS,
    BLOCO_PRODUTO_PROPRIO,
    CollectorStrategy,
    FundOffer,
    RunReport,
    ShelfSnapshot,
    SnapshotStatus,
)
from shelf_monitor.previdencia.config import ConfigPrevidencia
from shelf_monitor.report.competitive import build_xp_decisions
from shelf_monitor.report.digest import render_html, render_markdown

from fixtures_comuns import temp_config

_BLOCOS = json.loads(
    (Path(__file__).parent / "fixtures" / "previdencia" / "blocos_2026-09-23.json")
    .read_text(encoding="utf-8")
)
AGORA = datetime(2026, 9, 23, 22, 0)
PREVIDENCIA = dict(ConfigPrevidencia.blocos_suprimidos)

# O que nao pode aparecer em superficie nenhuma da previdencia.
_COTISTAS = ("Ponderada por cotistas", "PL menos cotistas", "muita gente", " cotistas -")
_CNPJ = ("Só nesta vitrine", "só ela exibe", "inclusão operacional", "faltam na XP")


def _vitrine(platform: str, cnpjs: list[str]) -> ShelfSnapshot:
    fundos = []
    for i, cnpj in enumerate(cnpjs):
        fundo = FundOffer(platform=platform, name=f"{platform.upper()} Prev {i} FIC FIM", cnpj=cnpj)
        fundo.sources = [CollectorStrategy.API_JSON.value]
        fundos.append(fundo)
    return ShelfSnapshot(
        platform=platform, collected_at=AGORA, strategy=CollectorStrategy.API_JSON,
        status=SnapshotStatus.OK, funds=fundos,
    )


def _relatorio(suprimidos: dict[str, str]) -> RunReport:
    # Um CNPJ em comum, como em 23/09: sem a supressao, o resto vira gap e exclusivo.
    xp = _vitrine("xp", ["11111111000111", "22222222000122", "33333333000133"])
    btg = _vitrine("btg", ["11111111000111", "44444444000144", "55555555000155"])
    report = RunReport(
        generated_at=AGORA, snapshots=[xp, btg],
        shelf_diffs=[diff_shelf(xp, None), diff_shelf(btg, None)],
        suprimidos=dict(suprimidos),
    )
    report.performance = json.loads(json.dumps(_BLOCOS["performance"]))
    report.retail_flow = json.loads(json.dumps(_BLOCOS["retail_flow"]))
    report.grade_opportunities = json.loads(json.dumps(_BLOCOS["grade_opportunities"]))
    return report


def _texto(html: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html))


class TestDigest(unittest.TestCase):
    def test_previdencia_nao_publica_cotistas_nem_cnpj(self):
        # Na edicao mensal, que traz a rentabilidade (e os cotistas, nos fundos).
        report = _relatorio(PREVIDENCIA)
        for superficie in (
            render_markdown(report, None, "mensal"),
            _texto(render_html(report, None, "mensal")),
        ):
            for proibido in _COTISTAS + _CNPJ:
                self.assertNotIn(proibido, superficie)
            self.assertIn("Não se aplica a este monitor", superficie)
            self.assertIn("Sem ponderação por cotistas", superficie)
            # O que continua valendo: a ponderada por PL e o fluxo por fundo.
            self.assertIn("Ponderada por PL", superficie)
            self.assertIn("Ritmo de captação", superficie)

    def test_fundos_segue_publicando_tudo(self):
        markdown = render_markdown(_relatorio({}), None, "mensal")
        self.assertIn("Ponderada por cotistas", markdown)
        self.assertIn("PL menos cotistas", markdown)
        self.assertIn("Oportunidade de grade e gaps", markdown)
        self.assertIn("nas vitrines de BTG Pactual", markdown)
        self.assertNotIn("Não se aplica", markdown)


class TestDashboardEDeck(unittest.TestCase):
    def setUp(self):
        self.config, self._tmp = temp_config()
        self.addCleanup(self._tmp.cleanup)

    def _dashboard(self, suprimidos):
        from shelf_monitor.report.dashboard import render_dashboard

        html = render_dashboard(self.config, _relatorio(suprimidos)).read_text(encoding="utf-8")
        inicio = html.index('id="dashboard-data">') + len('id="dashboard-data">')
        return html, json.loads(html[inicio:html.index("</script>", inicio)])

    def test_dashboard_da_previdencia(self):
        html, data = self._dashboard(PREVIDENCIA)
        self.assertEqual(
            set(data["suprimidos"]), {BLOCO_COTISTAS, BLOCO_CNPJ, BLOCO_PRODUTO_PROPRIO}
        )
        self.assertTrue(data["gaps"]["nao_se_aplica"])
        self.assertNotIn("Fundos que faltam na XP", html)
        self.assertNotIn("Fundos de muita gente", html)
        self.assertNotIn('id="chart-exclusive"', html)
        self.assertNotIn('id="perf-expostos"', html)
        self.assertIn("Não se aplica a este monitor", html)

    def test_dashboard_de_fundos_nao_muda(self):
        html, data = self._dashboard({})
        self.assertEqual(data["suprimidos"], {})
        self.assertIn("Fundos que faltam na XP", html)
        self.assertIn("Fundos de muita gente", html)
        self.assertIn('id="chart-exclusive"', html)

    @staticmethod
    def _executiva(html: str) -> str:
        """O texto da visao executiva - o HTML estatico acima dos paineis."""
        return _texto(html[html.index('id="visao-executiva"'):html.index('class="detalhes-head"')])

    def test_visao_executiva_da_previdencia(self):
        """A camada executiva e montada em Python: a supressao tem de valer nela.

        O JS de cada painel ja trocava o bloco pelo motivo; as manchetes, as
        leituras e o placar nao passam por ele.
        """
        html, _ = self._dashboard(PREVIDENCIA)
        executiva = self._executiva(html)
        for proibido in _COTISTAS + _CNPJ + ("só nela",):
            self.assertNotIn(proibido, executiva)
        # O que continua valendo, com o prazo ao lado do numero.
        self.assertIn("Captação líquida em 3 meses", executiva)
        self.assertIn("de 01/07/2026 a 21/09/2026", executiva)
        self.assertIn("Rentabilidade em 3 meses", executiva)
        self.assertIn("de 01/07/2026 a 22/09/2026", executiva)

    def test_visao_executiva_de_fundos(self):
        html, _ = self._dashboard({})
        executiva = self._executiva(html)
        self.assertIn("Fundos que faltam na XP", executiva)
        self.assertIn("Captação líquida em 3 meses", executiva)
        self.assertIn("de 01/07/2026 a 21/09/2026", executiva)
        # A melhor concorrente ao lado do numero da XP, e o placar com ela.
        self.assertIn("Maior: BTG Pactual", executiva)
        self.assertIn("só nela", executiva)
        self.assertIn('<tr class="base"><th scope="row">XP Investimentos</th>', html)

    def test_deck_da_previdencia(self):
        from pptx import Presentation

        from shelf_monitor.report.deck import render_deck

        deck = Presentation(render_deck(self.config, _relatorio(PREVIDENCIA)))
        texto = " ".join(
            s.text_frame.text for slide in deck.slides for s in slide.shapes if s.has_text_frame
        )
        self.assertNotIn("Gaps de prateleira", texto)
        self.assertIn("Não se aplica a este monitor", texto)


class TestPlanilha(unittest.TestCase):
    def setUp(self):
        self.config, self._tmp = temp_config()
        self.addCleanup(self._tmp.cleanup)

    def _planilha(self, suprimidos, nome):
        from openpyxl import load_workbook

        from shelf_monitor.report.spreadsheet import render_spreadsheet

        report = _relatorio(suprimidos)
        return load_workbook(render_spreadsheet(
            self.config, report.snapshots, generated_at=AGORA,
            path=self.config.reports_dir / nome, diffs=report.shelf_diffs,
            retail_flow=report.retail_flow, grade=report.grade_opportunities,
            performance=report.performance, suprimidos=report.suprimidos,
        ))

    @staticmethod
    def _cabecalho(sheet, linha):
        return [c.value for c in sheet[linha] if c.value is not None]

    def test_planilha_da_previdencia(self):
        wb = self._planilha(PREVIDENCIA, "previdencia.xlsx")
        for aba in ("Gaps XP", "Gaps por categoria", "Onde está a gente"):
            self.assertNotIn(aba, wb.sheetnames)
        self.assertNotIn("Só nesta vitrine", self._cabecalho(wb["Resumo"], 4))
        self.assertNotIn("Só nesta vitrine", self._cabecalho(wb["Vitrine varejo"], 1))
        fluxo = self._cabecalho(wb["Captação por fundo"], 7)
        for coluna in ("Exclusivo", "Cotistas fim", "Variação cotistas"):
            self.assertNotIn(coluna, fluxo)
        # Sem desalinhar: a coluna de captacao liquida ainda leva formato de dinheiro.
        liquida = fluxo.index("Captação líquida (R$)") + 1
        self.assertIn("#,##0", wb["Captação por fundo"].cell(row=8, column=liquida).number_format)
        rent = " ".join(
            str(c.value) for linha in wb["Rentabilidade 3 meses"].iter_rows() for c in linha if c.value
        )
        self.assertNotIn("Ponderada por cotistas", rent)
        self.assertIn("Ponderada por PL", rent)
        metodologia = " ".join(
            str(c.value) for linha in wb["Metodologia"].iter_rows() for c in linha if c.value
        )
        self.assertIn("Não se aplica a este monitor", metodologia)
        self.assertNotIn("Onde está a gente -", metodologia)

    def test_planilha_de_fundos_nao_muda(self):
        wb = self._planilha({}, "fundos.xlsx")
        for aba in ("Gaps XP", "Onde está a gente"):
            self.assertIn(aba, wb.sheetnames)
        self.assertIn("Só nesta vitrine", self._cabecalho(wb["Resumo"], 4))
        self.assertIn("Exclusivo", self._cabecalho(wb["Captação por fundo"], 7))


class TestDecisoes(unittest.TestCase):
    def test_sem_decisao_por_cnpj_na_previdencia(self):
        # Vitrines acima de MIN_VITRINE (30), com 1 CNPJ em comum como em 23/09.
        comum = "11111111000111"
        xp = _vitrine("xp", [comum] + [f"{10_000 + i:014d}" for i in range(39)])
        btg = _vitrine("btg", [comum] + [f"{20_000 + i:014d}" for i in range(39)])
        titulos = lambda d: [x["titulo"] for x in d.get("decisoes") or []]  # noqa: E731
        com = titulos(build_xp_decisions([xp, btg]))
        sem = titulos(build_xp_decisions([xp, btg], suprimidos=PREVIDENCIA))
        self.assertTrue(any("exclusivo" in t.lower() for t in com))
        self.assertFalse(any("exclusivo" in t.lower() for t in sem))


if __name__ == "__main__":
    unittest.main()
