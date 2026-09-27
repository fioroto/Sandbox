"""Captacao de ETF mes a mes: fluxo por cotas, estados do mes e as quatro visoes.

Os testes travam:

- a captacao pela variacao de cotas (valorizacao e distribuicao nao sao fluxo),
  a estreia, o informe que falta, o encerramento so em liquidacao e o
  desdobramento com o fator escolhido pelos pares;
- os estados do mes - acelerando, captando, virou saida, saida persistente,
  sem movimento (com o limiar relativo ao PL) e sem dado -, e o mes antes da
  estreia em branco;
- o bloco: mercado, gestoras, segmentos, oportunidades e defesa, e cada ETF
  com o seu segmento, no mes a mes, no ano e em 12 meses;
- o coletor do fechamento (a versao do documento que da informe utilizavel),
  e as superficies: relatorio md/html, planilha, CSV e orquestracao.

O cenario tem numeros redondos, calculados a mao.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import csv
import unittest
from datetime import date, datetime
from unittest import mock

from shelf_monitor.market import etf_captacao_mensal as mensal
from shelf_monitor.market.etf import EtfMetadata
from shelf_monitor.market.etf_captacao_mensal import (
    ajustar_evento_na_cota,
    build_captacao_mensal,
    estado_do_mes,
    fim_do_mes,
    fluxo_do_mes,
    limiar_do_mes,
    meses_fechados,
)
from shelf_monitor.market.fie_balancete import EtfMonthlyPl
from shelf_monitor.market.fundosnet import InformeDiarioEtf

from fixtures_comuns import temp_config

XP = "XP ALLOCATION ASSET MANAGEMENT LTDA"
BTG = "BTG PACTUAL ASSET MANAGEMENT S/A DTVM"
BLACKROCK = "BLACKROCK BRASIL GESTORA DE INVESTIMENTOS LTDA"
ITAU = "ITAU UNIBANCO ASSET MANAGEMENT LTDA."
MI = 1e6


def _informe(cnpj: str, competencia: str, pl: float, cota: float) -> InformeDiarioEtf:
    return InformeDiarioEtf(cnpj=cnpj, data=fim_do_mes(competencia), pl=pl, cota=cota)


class TestMeses(unittest.TestCase):
    def test_fim_do_mes(self):
        self.assertEqual(fim_do_mes("202602"), date(2026, 2, 28))
        self.assertEqual(fim_do_mes("202512"), date(2025, 12, 31))

    def test_so_entra_mes_com_a_entrega_encerrada(self):
        # Em 05/09 o fechamento de agosto ainda esta sendo entregue.
        self.assertEqual(meses_fechados(date(2026, 9, 5), 2), ["202606", "202607"])
        self.assertEqual(meses_fechados(date(2026, 9, 24), 2), ["202607", "202608"])
        self.assertEqual(len(meses_fechados(date(2026, 9, 24))), 14)


class TestFluxoDoMes(unittest.TestCase):
    def test_captacao_e_variacao_de_cotas_pela_cota_media(self):
        # 100 mi de cotas a R$ 100 -> 95 mi a R$ 110: saiu 5 mi x R$ 105.
        f = fluxo_do_mes("202512", _informe("C", "202511", 10_000 * MI, 100.0),
                         _informe("C", "202512", 10_450 * MI, 110.0), True)
        self.assertEqual(f.origem, "cotas")
        self.assertAlmostEqual(f.fluxo, -525 * MI)

    def test_valorizacao_e_distribuicao_nao_sao_fluxo(self):
        alta = fluxo_do_mes("202512", _informe("C", "202511", 1_000 * MI, 10.0),
                            _informe("C", "202512", 1_100 * MI, 11.0), True)
        self.assertAlmostEqual(alta.fluxo, 0.0)
        # Distribuicao: a cota cai junto com o PL, a quantidade de cotas fica.
        proventos = fluxo_do_mes("202512", _informe("C", "202511", 1_000 * MI, 10.0),
                                 _informe("C", "202512", 990 * MI, 9.9), True)
        self.assertAlmostEqual(proventos.fluxo, 0.0)

    def test_estreia_e_o_pl_do_fechamento(self):
        f = fluxo_do_mes("202601", None, _informe("E", "202601", 600 * MI, 10.0), False)
        self.assertEqual((f.origem, f.fluxo), ("estreia", 600 * MI))

    def test_informe_faltando_fica_sem_dado(self):
        antes = fluxo_do_mes("202601", None, _informe("F", "202601", 130 * MI, 10.0), True)
        depois = fluxo_do_mes("202512", _informe("F", "202511", 110 * MI, 10.0), None, True)
        self.assertEqual((antes.origem, antes.fluxo), ("sem_informe", None))
        self.assertEqual((depois.origem, depois.fluxo), ("sem_informe", None))

    def test_encerramento_so_em_liquidacao(self):
        f = fluxo_do_mes("202512", _informe("F", "202511", 110 * MI, 10.0), None, True, encerrado=True)
        self.assertEqual((f.origem, f.fluxo), ("encerramento", -110 * MI))

    def test_desdobramento_pelo_retorno_dos_pares(self):
        """IT NOW S&P500 TRN, 31/12/2025 -> 30/01/2026: cota 411,29 -> 49,97 e
        cotas 3,42 mi -> 30,56 mi. Com os pares a -2,8% o fator e 8."""
        f = fluxo_do_mes(
            "202601",
            _informe("S", "202512", 1_406_622_383.28, 411.2930945),
            _informe("S", "202601", 1_527_084_098.58, 49.9700294),
            True,
        )
        self.assertEqual((f.origem, f.fluxo), ("evento_na_cota", None))
        ajustado = ajustar_evento_na_cota(f, -0.028)
        self.assertEqual((ajustado.origem, ajustado.fator), ("desdobramento", 8))
        # 30,56 mi - 3,42 mi x 8 = 3,2 mi de cotas novas, a ~R$ 50,7.
        self.assertAlmostEqual(ajustado.fluxo / MI, 162.2, places=0)
        self.assertIs(ajustar_evento_na_cota(f, None), f)
        # Pares a +40%: o melhor fator (12) fica a 5,8 pp e o 11 a 6,4 pp - sem
        # um fator claro, o mes continua sem dado.
        self.assertEqual(ajustar_evento_na_cota(f, 0.40).origem, "evento_na_cota")


class TestEstadoDoMes(unittest.TestCase):
    def test_os_quatro_estados_e_o_neutro(self):
        self.assertEqual(estado_do_mes(200 * MI, 100 * MI), "acelerando")
        self.assertEqual(estado_do_mes(50 * MI, 100 * MI), "captando")
        self.assertEqual(estado_do_mes(50 * MI, -100 * MI), "captando")
        self.assertEqual(estado_do_mes(-50 * MI, 100 * MI), "virou_saida")
        self.assertEqual(estado_do_mes(-50 * MI, -10 * MI), "saida_persistente")
        self.assertEqual(estado_do_mes(50_000.0, 100 * MI), "sem_movimento")
        self.assertEqual(estado_do_mes(None, 100 * MI), "sem_dado")

    def test_primeiro_mes_nao_acelera_nem_persiste(self):
        self.assertEqual(estado_do_mes(50 * MI, None), "captando")
        self.assertEqual(estado_do_mes(-50 * MI, None), "virou_saida")

    def test_limiar_relativo_ao_pl(self):
        """BTG em 08/2026: +R$ 0,6 mi liquidos numa gestora de R$ 21 bi."""
        limiar = limiar_do_mes(21_000 * MI)
        self.assertEqual(limiar, 21 * MI)
        self.assertEqual(estado_do_mes(0.6 * MI, 682 * MI, limiar), "sem_movimento")
        self.assertEqual(limiar_do_mes(5 * MI), 100_000.0)
        # Mes anterior abaixo do proprio limiar nao conta como saida.
        self.assertEqual(estado_do_mes(-50 * MI, -0.05 * MI, 1 * MI, 1 * MI), "virou_saida")


def _etf(
    cnpj: str, nome: str, gestor: str, situacao: str = "Em Funcionamento Normal", ticker: str | None = None,
) -> EtfMetadata:
    return EtfMetadata(cnpj=cnpj, name=nome, kind="etf", manager_name=gestor, situation=situacao, ticker=ticker)


def _cenario() -> tuple[dict, dict, list]:
    catalogo = {
        "A": _etf("A", "TREND ETF LBMA OURO CLASSE DE ÍNDICE", XP, ticker="GOLD11"),
        "B": _etf("B", "BTG PACTUAL FUTURO DE OURO B3 FUNDO DE ÍNDICE", BTG),
        "C": _etf("C", "ISHARES IBOVESPA CLASSE DE ÍNDICE", BLACKROCK, ticker="BOVA11"),
        "D": _etf("D", "IT NOW IBOVESPA FUNDO DE ÍNDICE", ITAU),
        "E": _etf("E", "IT NOW IMA-B5 P2 FUNDO DE INDICE", ITAU),
        "F": _etf("F", "TREND ETF IBOVESPA CLASSE DE ÍNDICE", XP),
    }
    meses = ["202510", "202511", "202512", "202601", "202602"]
    serie = {
        "A": [(1000, 10), (1100, 10), (1300, 10), (1250, 10), (1150, 10)],
        "B": [(500, 10), (500, 10), (520, 10), (600, 10), (900, 10)],
        "C": [(10000, 100), (11000, 110), (10450, 110), (9405, 99), (8910, 99)],
        "D": [(2000, 10), (2100, 10), (2300, 10), (2600, 10), (2800, 10)],
        "E": [None, None, None, (600, 10), (750, 10)],
        "F": [(100, 10), (110, 10), None, (130, 10), (140, 10)],
    }
    fechamentos = {m: {} for m in meses}
    for cnpj, pontos in serie.items():
        for mes, ponto in zip(meses, pontos):
            if ponto:
                fechamentos[mes][cnpj] = _informe(cnpj, mes, ponto[0] * MI, ponto[1])
    # O balancete diz que F existia em 12/2025: o mes sem fechamento e falta de
    # informe, nao estreia.
    balancete = [
        EtfMonthlyPl(cnpj="F", etf_name="", ticker=None, category="", manager_name=None,
                     competencia="202512", dt_comptc=date(2025, 12, 31), pl=120 * MI),
    ]
    return catalogo, fechamentos, balancete


class TestBloco(unittest.TestCase):
    def setUp(self):
        catalogo, fechamentos, balancete = _cenario()
        # Tres meses na visao: o primeiro fluxo (11/2025) vira a base dos estados.
        with mock.patch.object(mensal, "MESES_DA_VISAO", 3):
            self.bloco = build_captacao_mensal(catalogo, fechamentos, balancete)
        self.gestoras = {g["gestora"]: g for g in self.bloco["gestoras"]}
        self.segmentos = {s["segmento"]: s for s in self.bloco["segmentos"]}
        self.etfs = {e["cnpj"]: e for e in self.bloco["etfs"]}

    def test_janela_12m_ano_e_base(self):
        b = self.bloco
        self.assertEqual(b["meses"], ["202512", "202601", "202602"])
        self.assertEqual(b["meses_ano"], ["202601", "202602"])
        self.assertEqual(b["mes_base"], "202511")
        self.assertEqual(b["janela_12m"], {"de": "2025-11-30", "ate": "2026-02-28"})
        self.assertEqual(b["janela_ano"], {"de": "2025-12-31", "ate": "2026-02-28"})

    def test_mercado_mes_a_mes(self):
        mercado = self.bloco["mercado"]
        self.assertEqual(
            {m: round(v / MI) for m, v in mercado["fluxo"].items()},
            {"202512": -130, "202601": 930, "202602": 65},
        )
        self.assertEqual(
            mercado["estado"], {"202512": "virou_saida", "202601": "captando", "202602": "captando"}
        )
        self.assertAlmostEqual(mercado["doze_meses"], 865 * MI)
        self.assertAlmostEqual(mercado["no_ano"], 995 * MI)

    def test_competicao_entre_gestoras(self):
        self.assertEqual(
            [g["gestora"] for g in self.bloco["gestoras"]],
            ["Itaú Asset (It Now)", "BTG Pactual", "XP (Trend)", "BlackRock (iShares)"],
        )
        xp = self.gestoras["XP (Trend)"]
        self.assertTrue(xp["foco"])
        self.assertAlmostEqual(xp["doze_meses"], 60 * MI)
        self.assertAlmostEqual(xp["no_ano"], -140 * MI)
        self.assertEqual(
            xp["estado"], {"202512": "acelerando", "202601": "virou_saida", "202602": "saida_persistente"}
        )
        self.assertAlmostEqual(xp["share_doze_meses"], 60 / 865)
        self.assertAlmostEqual(xp["share_pl"], 1290 / 14650)
        self.assertEqual(self.gestoras["BlackRock (iShares)"]["estado"]["202601"], "sem_movimento")

    def test_cada_etf_com_o_seu_segmento(self):
        e = self.etfs["E"]
        self.assertEqual(e["segmento"], "Inflação (IMA-B e Tesouro IPCA)")
        # Antes da estreia o mes fica em branco, e a estreia conta o PL.
        self.assertEqual(e["estado"], {"202512": "fora", "202601": "captando", "202602": "captando"})
        self.assertEqual(e["origem"], {"202601": "estreia"})
        f = self.etfs["F"]
        self.assertEqual(f["estado"], {"202512": "sem_dado", "202601": "sem_dado", "202602": "captando"})
        self.assertEqual(f["origem"], {"202512": "sem informe", "202601": "sem informe"})
        self.assertEqual(self.bloco["cobertura"]["sem_dado"], 2)
        self.assertEqual(self.etfs["B"]["estado"]["202602"], "acelerando")

    def test_segmentos_oportunidades_e_defesa(self):
        ouro = self.segmentos["Ouro"]
        self.assertEqual(ouro["lider_captacao"], "BTG Pactual")
        self.assertAlmostEqual(ouro["foco_doze_meses"], 50 * MI)
        self.assertTrue(ouro["foco_lidera_pl"])
        self.assertEqual([s["segmento"] for s in self.bloco["oportunidades"]],
                         ["Inflação (IMA-B e Tesouro IPCA)"])
        self.assertFalse(self.bloco["oportunidades"][0]["foco_tem_etf"])
        defesa = self.bloco["defesa"][0]
        self.assertEqual(defesa["segmento"], "Ouro")
        self.assertEqual(defesa["alertas"], [
            "a XP (Trend) está em saída persistente",
            "BTG Pactual captou mais em 12 meses: +R$ 400 mi, contra +R$ 50 mi da XP (Trend)",
        ])

    def test_leitura_e_prazos(self):
        texto = " ".join(self.bloco["leitura"])
        self.assertIn("Em 12 meses (dez/2025 a fev/2026), os ETFs captaram +R$ 865 mi líquidos; "
                      "no ano, +R$ 995 mi. Em fev/2026, +R$ 65 mi: captando.", texto)
        self.assertIn("A XP (Trend) captou +R$ 60 mi em 12 meses - 6,9% da captação dos ETFs, "
                      "contra 8,8% do PL - e é a 3ª em captação", texto)
        self.assertIn("Inflação (IMA-B e Tesouro IPCA), +R$ 750 mi com Itaú Asset à frente", texto)
        self.assertIn("Ouro - a XP (Trend) está em saída persistente", texto)
        prazos = mensal.prazos_captacao(self.bloco)
        self.assertIn("12 meses: de dez/2025 a fev/2026 (fechamentos de 30/11/2025 a 28/02/2026)", prazos)
        self.assertIn("no ano: de jan/2026 a fev/2026 (fechamentos de 31/12/2025 a 28/02/2026)", prazos)
        self.assertIn("o estado de dez/2025 compara com nov/2025", prazos)

    def test_sem_dois_fechamentos_diz_por_que(self):
        bloco = build_captacao_mensal({}, {"202608": {}})
        self.assertFalse(bloco["available"])
        self.assertIn("fechamentos", bloco["reason"])


class _Doc:
    def __init__(self, doc_id, ref, situacao="A", fundo="TREND ETF LBMA OURO CLASSE DE ÍNDICE"):
        self.id = doc_id
        self.data_referencia = ref
        self.situacao = situacao
        self.fundo = fundo
        self.cnpj_fundo = None


class _ClienteFalso:
    """Busca e download do Fundos.NET com respostas fixas."""

    def __init__(self, docs, xmls):
        self.docs, self.xmls = docs, xmls
        self.baixados: list[int] = []
        self.buscas = 0

    def fund_types(self):
        return {"ETF": 3}

    def search(self, *_args, **_kwargs):
        self.buscas += 1
        return iter(self.docs)

    def download(self, doc_id):
        self.baixados.append(doc_id)
        return self.xmls[doc_id], "text/xml"


class TestColetor(unittest.TestCase):
    def setUp(self):
        self.config, self.tmp = temp_config()
        self.catalogo = {"A": _etf("11111111000111", "TREND ETF LBMA OURO CLASSE DE ÍNDICE", XP)}
        self.catalogo = {"11111111000111": self.catalogo["A"]}

    def tearDown(self):
        self.tmp.cleanup()

    def test_fica_com_a_versao_que_da_informe_e_nao_rebaixa_o_que_esta_em_disco(self):
        from shelf_monitor.market.etf_sources import collect_fechamentos_mensais
        from shelf_monitor.market.fundosnet import read_informes_csv

        completo = (
            "<DOC_ARQ><CAB_INFORM><DT_COMPT>31/08/2026</DT_COMPT></CAB_INFORM><LISTA_INFORM>"
            "<INFORM><CNPJ_FDO>11111111000111</CNPJ_FDO><VL_QUOTA>52,65</VL_QUOTA>"
            "<VL_PATRIM_LIQ>1760000000,00</VL_PATRIM_LIQ></INFORM></LISTA_INFORM></DOC_ARQ>"
        ).encode("utf-8")
        cabecalho_so = (
            "<DOC_ARQ><CAB_INFORM><DT_COMPT>31/08/2026</DT_COMPT><CNPJ_FDO>11111111000111"
            "</CNPJ_FDO></CAB_INFORM></DOC_ARQ>"
        ).encode("utf-8")
        cliente = _ClienteFalso(
            [
                _Doc(3, date(2026, 9, 1)),                 # depois do fim do mes: fora
                _Doc(2, date(2026, 8, 31), "A"),           # versao ativa sem PL
                _Doc(1, date(2026, 8, 31), "I"),           # versao anterior, completa
                _Doc(0, date(2026, 8, 28), "A"),
            ],
            {2: cabecalho_so, 1: completo, 0: completo},
        )
        with mock.patch("shelf_monitor.market.etf_sources.fundosnet.FundosNetClient", return_value=cliente):
            resumo = collect_fechamentos_mensais(self.config, self.catalogo, meses=["202608"])
            self.assertEqual(resumo["meses"], {"202608": "1 ETF"})
            self.assertEqual(cliente.baixados, [2, 1])
            gravados = read_informes_csv(self.config.etf_fnet_mensal_dir / "202608.csv")
            self.assertEqual([(i.cnpj, i.data, i.pl) for i in gravados],
                             [("11111111000111", date(2026, 8, 31), 1_760_000_000.0)])
            # Mes em disco nao e baixado de novo.
            collect_fechamentos_mensais(self.config, self.catalogo, meses=["202608"])
            self.assertEqual(cliente.buscas, 1)


class TestSuperficies(unittest.TestCase):
    def setUp(self):
        catalogo, fechamentos, balancete = _cenario()
        self.catalogo = catalogo
        with mock.patch.object(mensal, "MESES_DA_VISAO", 3):
            self.bloco = build_captacao_mensal(catalogo, fechamentos, balancete)
        self.config, self.tmp = temp_config()

    def tearDown(self):
        self.tmp.cleanup()

    def test_relatorio_markdown_e_html(self):
        from shelf_monitor.report.etf_captacao import write_etf_captacao

        md_path, html_path = write_etf_captacao(self.config, self.bloco, datetime(2026, 9, 24, 10, 0))
        self.assertEqual(md_path.name, "etf-captacao-2026-09-24.md")
        md = md_path.read_text(encoding="utf-8")
        self.assertIn("# Captação de ETF mês a mês - XP (Trend)", md)
        self.assertIn("Legenda do mês: ▲▲ acelerando; ▲ captando; – sem movimento; ▼ virou saída; "
                      "▼▼ saída persistente; ? sem dado; em branco: antes da estreia.", md)
        self.assertIn("| dez/2025 | -R$ 130 mi | - | ▼ virou saída |", md)
        self.assertIn("| jan/2026 | +R$ 930 mi | +R$ 600 mi | ▲ captando |", md)
        self.assertIn("| 3 | **XP (Trend)** | +R$ 60 mi | 6,9% | 8,8% | -R$ 140 mi | ▲▲ | ▼ | ▼▼ | "
                      "-R$ 90 mi (saída persistente) |", md)
        self.assertIn("**Oportunidades**", md)
        # O ticker abre a linha de cada ETF; sem ticker na B3, sai um traço.
        self.assertIn("| **BOVA11** | ISHARES IBOVESPA | BlackRock | Ibovespa |", md)
        self.assertIn("| - | BTG PACTUAL FUTURO DE OURO B3 | BTG Pactual | Ouro |", md)
        self.assertIn("**Alerta:** a XP (Trend) está em saída persistente", md)
        self.assertIn("[Fundos.NET", md)
        html = html_path.read_text(encoding="utf-8")
        self.assertIn('title="acelerando"', html)
        self.assertIn("<strong>GOLD11</strong>", html)
        self.assertIn("background:#1E7B34", html)
        self.assertIn("href=", html)

    def test_planilha_abre_pela_captacao_mensal(self):
        from openpyxl import load_workbook

        from shelf_monitor.report.etf_captacao import CORES
        from shelf_monitor.report.etf_report import export_etf_excel

        caminho = export_etf_excel([], [], self.catalogo, self.config.reports_dir / "etf.xlsx",
                                   captacao=self.bloco)
        wb = load_workbook(caminho)
        self.assertEqual(wb.sheetnames[:4], [
            "Captação - Competição", "Captação - Segmentos", "Captação - Oport. e defesa",
            "Captação - por ETF",
        ])
        self.assertEqual(wb.active.title, "Captação - Competição")
        ws = wb["Captação - Competição"]
        linha_xp = next(
            r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "XP (Trend)"
        )
        dez = ws.cell(row=linha_xp, column=9)
        self.assertAlmostEqual(dez.value, 200 * MI)
        self.assertEqual(dez.fill.fgColor.rgb[-6:], CORES["acelerando"][0])
        por_etf = wb["Captação - por ETF"]
        linha_a = next(
            r for r in range(1, por_etf.max_row + 1) if por_etf.cell(row=r, column=4).value == "GOLD11"
        )
        self.assertEqual(por_etf.cell(row=linha_a, column=5).value, "TREND ETF LBMA OURO CLASSE DE ÍNDICE")
        # A grade de meses comeca depois do ticker, e a cor acompanha.
        self.assertEqual(por_etf.cell(row=linha_a, column=10).fill.fgColor.rgb[-6:], CORES["virou_saida"][0])

    def test_csv_longo_sem_os_meses_antes_da_estreia(self):
        from shelf_monitor.report.etf_report import export_captacao_csv

        caminho = export_captacao_csv(self.bloco, self.config.reports_dir / "captacao.csv")
        with caminho.open(encoding="utf-8-sig", newline="") as handle:
            linhas = list(csv.DictReader(handle, delimiter=";"))
        de_e = [(l["competencia"], l["estado"], l["origem"]) for l in linhas if l["cnpj"] == "E"]
        self.assertEqual(de_e, [("202601", "captando", "estreia"), ("202602", "captando", "cotas")])
        self.assertEqual({l["ticker"] for l in linhas if l["cnpj"] == "C"}, {"BOVA11"})

    def test_orquestracao_sem_rede_usa_o_que_esta_em_disco(self):
        from shelf_monitor import main as main_mod
        from shelf_monitor.http import SourceUnavailable

        catalogo, fechamentos, _ = _cenario()
        with mock.patch.object(main_mod, "meses_fechados", return_value=sorted(fechamentos)), \
                mock.patch.object(main_mod, "ler_fechamentos", return_value=fechamentos), \
                mock.patch.object(main_mod, "run_etf_pl_pipeline", return_value=([], None, [])), \
                mock.patch.object(main_mod.etf_sources, "collect_fechamentos_mensais",
                                  side_effect=SourceUnavailable("fnet: timeout")), \
                mock.patch.object(main_mod, "print_captacao_mensal") as imprimir:
            with main_mod.console.capture() as saida:
                bloco = main_mod._captacao_mensal_de_etf(self.config, catalogo)
        self.assertTrue(bloco["available"])
        self.assertIn("usa o que está em disco", saida.get())
        imprimir.assert_called_once_with(bloco)

    def test_tickers_chegam_ao_catalogo_e_aos_fluxos(self):
        from types import SimpleNamespace

        from shelf_monitor import main as main_mod
        from shelf_monitor.http import SourceUnavailable
        from shelf_monitor.market import b3_funds

        catalogo = {
            "11111111000111": _etf("11111111000111", "BTG PACTUAL FUTURO DE OURO B3 FUNDO DE ÍNDICE", BTG),
            "22222222000122": _etf("22222222000122", "TREND ETF LBMA OURO CLASSE DE ÍNDICE", XP, ticker="GOLD11"),
        }
        b3_funds.write_tickers_csv(self.config.etf_tickers_path, [
            {"cnpj": "11111111000111", "ticker": "GLDX11", "fonte": "b3", "atualizado_em": "2026-01-01"},
        ])
        fluxo = SimpleNamespace(cnpj="11111111000111", ticker=None)
        with mock.patch.object(main_mod.etf_sources, "collect_b3",
                               side_effect=SourceUnavailable("b3: timeout")):
            with main_mod.console.capture() as saida:
                main_mod._tickers_de_etf(self.config, catalogo, [fluxo])
        # Arquivo de janeiro: vencido, tenta a B3; sem rede, fica com o arquivo.
        self.assertIn("B3 indisponível para os tickers", saida.get())
        self.assertEqual((catalogo["11111111000111"].ticker, fluxo.ticker), ("GLDX11", "GLDX11"))
        self.assertEqual(catalogo["22222222000122"].ticker, "GOLD11")

    def test_console_nao_quebra(self):
        from shelf_monitor.report.etf_report import print_captacao_mensal

        print_captacao_mensal(self.bloco)
        print_captacao_mensal({"available": False, "reason": "sem fechamentos"})


if __name__ == "__main__":
    unittest.main()
