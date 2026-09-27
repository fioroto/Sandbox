"""Competicao em ETFs: segmento pelo nome, gestoras, oportunidades e lancamentos.

A CVM deixa vazios o benchmark e as classificacoes dos ETFs (252 classes FIIM em
2026-09-24), entao o que cada ETF replica sai do nome. Os testes travam:

- a ordem das regras de segmento, nos casos em que o nome engana (bitcoin da
  Hashdex com "NASDAQ", debenture incentivada com "IPCA", S&P 500 em dolar,
  small cap americana);
- o rotulo da gestora e a casa monitorada a que ela pertence;
- o bloco competitivo: participacao, lideranca, ausencia, variacao dentro do
  segmento, lancamentos (pre-operacional, estreia e estreia depois do
  balancete, com o PL do Fundos.NET), cotistas pelo maximo da
  janela e fluxo sem o dia-base;
- as superficies: leitura, relatorio md/html com prazos e fontes, abas da planilha.

O cenario tem numeros redondos, calculados a mao.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import unittest
from datetime import date, datetime
from unittest import mock

from shelf_monitor.market.etf import EtfMetadata
from shelf_monitor.market.etf_competicao import (
    SEM_SEGMENTO,
    build_etf_competicao,
    casa_da_gestora,
    rotulo_da_gestora,
    segmento_do_etf,
)
from shelf_monitor.market.fie_balancete import EtfMonthlyPl
from shelf_monitor.market.fundosnet import InformeDiarioEtf

from fixtures_comuns import temp_config

XP = "XP ALLOCATION ASSET MANAGEMENT LTDA"
BTG = "BTG PACTUAL ASSET MANAGEMENT S/A DTVM"
BLACKROCK = "BLACKROCK BRASIL GESTORA DE INVESTIMENTOS LTDA"
ITAU = "ITAU UNIBANCO ASSET MANAGEMENT LTDA."
BI = 1e9


class TestSegmento(unittest.TestCase):
    CASOS = [
        ("ISHARES IBOVESPA CLASSE DE ÍNDICE - RESPONSABILIDADE LIMITADA", "Ibovespa"),
        ("TREND ETF IBOVESPA B3 ALAVANCADO 2X CLASSE  DE ÍNDICE", "Ibovespa alavancado"),
        ("NU IBOV SMART DIVIDENDOS CLASSE DE ÍNDICE", "Dividendos"),
        ("NU IBOV SMART LOW VOLATILITY B3 CLASSE DE  ÍNDICE", "Fatores e smart beta"),
        ("IT NOW IBOVESPA B3 BR+ FUNDO DE ÍNDICE", "Ibovespa BR+ e variações"),
        # "NASDAQ" no nome, mas o que replica e cripto.
        ("HASHDEX NASDAQ BITCOIN REFERENCE PRICE FUNDO DE ÍNDICE", "Bitcoin"),
        ("HASHDEX NASDAQ CME CRYPTO INDEX FUNDO DE ÍNDICE", "Cesta e temáticos cripto"),
        ("BUENA VISTA HASHDEX GOLD & BITCOIN FUNDO DE ÍNDICE", "Cesta e temáticos cripto"),
        # Renda com opcoes nao compete com o indice subjacente.
        ("BUENA VISTA NASDAQ-100 NEOS HIGH INCOME ETF INDEX (NDXHIEN)", "Renda com opções (high income)"),
        ("TREND ETF NASDAQ 100 CLASSE DE ÍNDICE INVESTIMENTO NO EXTERIOR", "Nasdaq e tecnologia"),
        # Debenture incentivada atrelada a IPCA e credito, nao inflacao.
        ("IT NOW S&P/B3 DEBÊNTURES INCENTIVADAS IPCA FUNDO DE ÍNDICE", "Crédito privado (debêntures e LF)"),
        # Infraestrutura sozinha e debenture incentivada (a categoria da B3); o
        # hibrido do BTG e 80% acao e compete com os setoriais.
        ("TREND ETF INFRAESTRUTURA ESG B3 CLASSE DE  ÍNDICE", "Crédito privado (debêntures e LF)"),
        ("TREND ETF IDEX INFRA TOP LIQUIDEZ CLASSE DE ÍNDICE", "Crédito privado (debêntures e LF)"),
        ("TREND ETF IDEX CDI CORE ESG CLASSE DE ÍNDICE", "Crédito privado (debêntures e LF)"),
        # A familia IDEX B da XP e NTN-B de um vencimento ("TREND NTNB35" na B3).
        ("TREND ETF IDEX B35 CLASSE DE ÍNDICE - RESPONSABILIDADE LIMITADA", "Inflação (IMA-B e Tesouro IPCA)"),
        ("TREND ETF IDEX B60 CLASSE DE ÍNDICE - RESPONSABILIDADE LIMITADA", "Inflação (IMA-B e Tesouro IPCA)"),
        ("BTG PACTUAL TEVA INFRAESTRUTURA HÍBRIDO FUNDO DE ÍNDICE", "Setoriais e temáticos Brasil"),
        ("IT NOW IMA-B5 P2 FUNDO DE INDICE RESPONSABILIDADE LIMITADA", "Inflação (IMA-B e Tesouro IPCA)"),
        ("ETF BRADESCO IMAB5+ - RESPONSABILIDADE LIMITADA", "Inflação (IMA-B e Tesouro IPCA)"),
        # "TREASURY" de titulo publico brasileiro com duration-alvo.
        ("INVESTO ETF MARKETVECTOR BRAZIL TREASURY 760 DAY TARGET DURATION", "Tesouro com duration-alvo"),
        ("B-INDEX ETF TEVA LFT CURTO PRAZO FUNDO DE ÍNDICE", "Tesouro Selic e CDI"),
        ("TREND ETF PREFIXADO DI 2033 CLASSE DE ÍNDICE", "Prefixado (IRF-M e Tesouro pré)"),
        # S&P 500 cotado em dolar continua S&P 500; o indice de dolar e dolar.
        ("TREND ETF S&P 500 FUTURES USD-BRLCLASSE DE  ÍNDICE", "S&P 500"),
        ("TREND ETF S&P / B3 DOLAR CLASSE DE ÍNDICE", "Dólar"),
        # Small cap americana nao e small cap brasileira nem S&P 500.
        ("INVESTO ETF S&P SMALLCAP 600 VALUE CLASSE DE ÍNDICE", "Temáticos internacionais"),
        ("ISHARES BM&FBOVESPA SMALL CAP CLASSE DE  ÍNDICE", "Small caps"),
        ("TREND ETF LBMA OURO CLASSE DE ÍNDICE INVESTIMENTO NO EXTERIOR", "Ouro"),
        ("TREND ETF IFIX-L  CLASSE DE ÍNDICE", "FII (IFIX) e REITs Brasil"),
    ]

    def test_o_que_cada_etf_replica(self):
        for nome, esperado in self.CASOS:
            with self.subTest(nome=nome):
                self.assertEqual(segmento_do_etf(nome)[1], esperado)

    def test_nome_sem_indice_reconhecivel(self):
        self.assertEqual(segmento_do_etf("BB ETF 0000 FUNDO DE ÍNDICE"), SEM_SEGMENTO)
        self.assertEqual(segmento_do_etf(None), SEM_SEGMENTO)


class TestGestora(unittest.TestCase):
    def test_rotulo_curto_com_a_marca(self):
        self.assertEqual(rotulo_da_gestora(XP), "XP (Trend)")
        self.assertEqual(rotulo_da_gestora(ITAU), "Itaú Asset (It Now)")
        self.assertEqual(rotulo_da_gestora("BANCO BRADESCO S.A."), "Bradesco (B-Index)")

    def test_sem_gestor_no_cadastro_usa_a_marca_do_nome(self):
        self.assertEqual(rotulo_da_gestora(None, "QR BLOOMBERG DEFI FUNDO DE ÍNDICE"), "QR Asset")

    def test_gestora_desconhecida_vira_as_duas_primeiras_palavras(self):
        self.assertEqual(rotulo_da_gestora("ALGUMA GESTORA NOVA LTDA"), "Alguma Gestora")

    def test_casa_monitorada(self):
        self.assertEqual(casa_da_gestora(XP), "xp")
        self.assertEqual(casa_da_gestora(BTG), "btg")
        self.assertIsNone(casa_da_gestora(BLACKROCK))


def _etf(cnpj: str, nome: str, gestor: str, situacao: str = "Em Funcionamento Normal") -> EtfMetadata:
    return EtfMetadata(cnpj=cnpj, name=nome, kind="etf", manager_name=gestor, situation=situacao)


def _pl(cnpj: str, competencia: str, valor: float) -> EtfMonthlyPl:
    dia = date(2026, 6, 30) if competencia == "202606" else date(2026, 8, 31)
    return EtfMonthlyPl(
        cnpj=cnpj, etf_name="", ticker=None, category="", manager_name=None,
        competencia=competencia, dt_comptc=dia, pl=valor,
    )


def _cenario() -> tuple[dict, list, list]:
    catalogo = {
        "A": _etf("A", "TREND ETF LBMA OURO CLASSE DE ÍNDICE", XP),
        "B": _etf("B", "BTG PACTUAL FUTURO DE OURO B3 FUNDO DE ÍNDICE", BTG),
        "C": _etf("C", "ISHARES IBOVESPA CLASSE DE ÍNDICE", BLACKROCK),
        "D": _etf("D", "TREND ETF IBOVESPA CLASSE DE ÍNDICE", XP),
        "E": _etf("E", "IT NOW IMA-B5 P2 FUNDO DE INDICE", ITAU),
        "F": _etf("F", "IT NOW NTN-B 2035 FUNDO DE ÍNDICE", ITAU),
        "G": _etf("G", "IT NOW S&P/B3 DEBÊNTURES INCENTIVADAS IPCA FUNDO DE ÍNDICE", ITAU,
                  "Fase Pré-Operacional"),
        "H": _etf("H", "TREND ETF PREFIXADO DI 2033 CLASSE DE ÍNDICE", XP, "Fase Pré-Operacional"),
        "Z": EtfMetadata(cnpj="Z", name="XP FEEDER ETF FIC FIF", kind="feeder", manager_name=XP),
    }
    serie = [
        _pl("A", "202606", 2.0 * BI), _pl("A", "202608", 2.4 * BI),
        _pl("B", "202606", 0.5 * BI), _pl("B", "202608", 0.6 * BI),
        _pl("C", "202606", 14.0 * BI), _pl("C", "202608", 15.0 * BI),
        _pl("D", "202606", 1.0 * BI), _pl("D", "202608", 1.5 * BI),
        _pl("E", "202606", 4.0 * BI), _pl("E", "202608", 4.2 * BI),
        _pl("F", "202608", 1.0 * BI),  # estreia: sem PL em 06/2026
    ]

    def informe(cnpj, dia, cotistas, captacao=0.0, resgates=0.0):
        return InformeDiarioEtf(
            cnpj=cnpj, data=date(2026, 9, dia), pl=1.0, cota=1.0, cotistas=cotistas,
            captacao=captacao, resgates=resgates,
        )

    informes = [
        informe("A", 14, 1000, captacao=100.0),  # dia-base: ja esta no PL
        informe("A", 15, 10, captacao=50.0, resgates=10.0),  # cotistas com erro de origem
        informe("A", 16, 1010, resgates=20.0),
        informe("C", 14, 5000, captacao=999.0),
        informe("C", 16, 5100, captacao=200.0),
        informe("D", 14, 300),
    ]
    return catalogo, serie, informes


class TestBloco(unittest.TestCase):
    def setUp(self):
        catalogo, serie, informes = _cenario()
        self.bloco = build_etf_competicao(catalogo, serie, informes)
        self.gestoras = {g["gestora"]: g for g in self.bloco["gestoras"]}
        self.segmentos = {s["segmento"]: s for s in self.bloco["segmentos"]}

    def test_ranking_e_participacao_das_gestoras(self):
        self.assertEqual(self.bloco["foco"], "XP (Trend)")
        self.assertEqual([g["gestora"] for g in self.bloco["gestoras"]],
                         ["BlackRock (iShares)", "Itaú Asset (It Now)", "XP (Trend)", "BTG Pactual"])
        xp = self.gestoras["XP (Trend)"]
        self.assertEqual((xp["posicao"], xp["etfs"], xp["pre_operacionais"]), (3, 2, 1))
        self.assertAlmostEqual(xp["share"], 3.9 / 24.7)
        # 3,0 de 21,5 em 06/2026 para 3,9 de 24,7 em 08/2026. O feeder nao entra.
        self.assertAlmostEqual(xp["delta_share_pp"], (3.9 / 24.7 - 3.0 / 21.5) * 100)
        self.assertEqual(self.bloco["pl_total"], 24.7 * BI)

    def test_situacao_da_casa_em_cada_segmento(self):
        self.assertEqual(self.segmentos["Ouro"]["foco_status"], "lidera")
        self.assertAlmostEqual(self.segmentos["Ouro"]["foco_share"], 0.8)
        ibov = self.segmentos["Ibovespa"]
        self.assertEqual((ibov["foco_status"], ibov["foco_posicao"], ibov["lider"]),
                         ("presente", 2, "BlackRock (iShares)"))
        # Dentro do segmento a valorizacao e a mesma: 1/15 -> 1,5/16,5.
        self.assertAlmostEqual(ibov["foco_delta_pp"], (1.5 / 16.5 - 1.0 / 15.0) * 100)
        self.assertEqual(self.segmentos["Inflação (IMA-B e Tesouro IPCA)"]["foco_status"], "ausente")
        self.assertEqual(self.segmentos["Prefixado (IRF-M e Tesouro pré)"]["foco_status"], "entrando")

    def test_oportunidades_e_defesa(self):
        self.assertEqual(
            [s["segmento"] for s in self.bloco["oportunidades"]],
            ["Ibovespa", "Inflação (IMA-B e Tesouro IPCA)"],
        )
        self.assertEqual([s["segmento"] for s in self.bloco["defesa"]], ["Ouro"])
        self.assertEqual(self.bloco["defesa"][0]["segundo"], "BTG Pactual")

    def test_lancamentos_pre_operacionais_e_estreias(self):
        lancamentos = [(l["gestora"], l["tipo"][:15], l["segmento"]) for l in self.bloco["lancamentos"]]
        self.assertEqual(lancamentos, [
            ("XP (Trend)", "pré-operacional", "Prefixado (IRF-M e Tesouro pré)"),
            ("Itaú Asset (It Now)", "pré-operacional", "Crédito privado (debêntures e LF)"),
            ("Itaú Asset (It Now)", "começou a opera", "Inflação (IMA-B e Tesouro IPCA)"),
        ])
        # A janela da estreia sai com as datas-base, nao so com os meses.
        self.assertEqual(
            self.bloco["lancamentos"][-1]["tipo"],
            "começou a operar entre os balancetes de 30/06/2026 e 31/08/2026",
        )

    def test_cotistas_pelo_maximo_e_fluxo_sem_o_dia_base(self):
        etfs = {e["cnpj"]: e for e in self.bloco["etfs"]}
        # O "10" de 15/09 e erro de origem: vale o maximo da janela.
        self.assertEqual(etfs["A"]["cotistas"], 1010)
        # 50 - 10 - 20: a captacao de 14/09 ja esta no PL do dia-base.
        self.assertAlmostEqual(etfs["A"]["fluxo"], 20.0)
        self.assertAlmostEqual(etfs["C"]["fluxo"], 200.0)
        self.assertIsNone(etfs["D"]["fluxo"])  # informe unico: nao ha fluxo depois da base
        xp = self.gestoras["XP (Trend)"]
        self.assertEqual(xp["cotistas"], 1310)
        self.assertAlmostEqual(xp["share_cotistas"], 1310 / 6410)

    def test_leitura_com_prazos(self):
        texto = " ".join(self.bloco["leitura"])
        self.assertIn("No PL de ETFs de ago/2026 (R$ 24,7 bi em 6 ETFs", texto)
        self.assertIn("a XP (Trend) é a 3ª gestora, com R$ 3,9 bi (15,8% do mercado)", texto)
        self.assertIn("Lidera 1 segmento: Ouro (80% do segmento)", texto)
        self.assertIn("Não tem ETF em 1 segmento acima de R$ 500 mi", texto)
        self.assertIn("Ibovespa (R$ 16,5 bi), 9,1% e 2ª de 2 gestoras", texto)
        self.assertIn("De jun/2026 a ago/2026, a participação no PL de ETFs foi de 14,0% para 15,8%", texto)
        self.assertIn("Em fase pré-operacional estão 2 ETFs, 1 da XP (Trend)", texto)
        self.assertIn("1.310 cotistas - 20,4% dos cotistas de ETF", texto)

    def test_participacao_competencia_a_competencia(self):
        mensal = self.bloco["participacao_mensal"]
        self.assertEqual(mensal["competencias"], ["202606", "202608"])
        xp = next(g for g in mensal["gestoras"] if g["foco"])
        self.assertAlmostEqual(xp["share"]["202606"], 3.0 / 21.5)
        self.assertAlmostEqual(xp["share"]["202608"], 3.9 / 24.7)
        # A Itau entra com o PL da estreia (F) so na segunda competencia.
        itau = next(g for g in mensal["gestoras"] if g["gestora"].startswith("Itaú"))
        self.assertEqual(itau["pl"], {"202606": 4.0 * BI, "202608": 5.2 * BI})

    def test_sem_pl_o_bloco_diz_por_que(self):
        catalogo, _, _ = _cenario()
        bloco = build_etf_competicao(catalogo, [])
        self.assertFalse(bloco["available"])
        self.assertIn("balancete", bloco["reason"])


class TestEstreiaDepoisDoBalancete(unittest.TestCase):
    """ETF que comecou a operar depois do ultimo balancete.

    Medido em 2026-09-24: 8 ETFs com PL no Fundos.NET em setembro e nenhum
    balancete (R$ 266 mi, 2 da XP) sumiam do radar de lancamentos.
    """

    def setUp(self):
        catalogo, serie, informes = _cenario()
        catalogo["I"] = _etf("I", "TREND ETF BLOOMBERG AGGREGATE US BONDS CLASSE DE ÍNDICE", XP)
        catalogo["I"].ticker = "AGGX11"
        # Pre-operacional no cadastro, mas o informe ja mostra PL: vale o informe.
        catalogo["J"] = _etf(
            "J", "BTG PACTUAL TEVA TESOURO SELIC FUNDO DE ÍNDICE", BTG, "Fase Pré-Operacional"
        )
        # Informe ANTES da data-base e nenhum balancete: lacuna de dado, nao estreia.
        catalogo["K"] = _etf("K", "ISHARES S&P 500 CLASSE DE ÍNDICE", BLACKROCK)
        informes += [
            InformeDiarioEtf(cnpj="I", data=date(2026, 9, 15), pl=4_000_000.0, cota=100.0),
            InformeDiarioEtf(cnpj="I", data=date(2026, 9, 16), pl=5_000_000.0, cota=100.0),
            InformeDiarioEtf(cnpj="J", data=date(2026, 9, 17), pl=80_000_000.0, cota=100.0),
            InformeDiarioEtf(cnpj="K", data=date(2026, 8, 28), pl=9_000_000.0, cota=100.0),
        ]
        self.bloco = build_etf_competicao(catalogo, serie, informes)
        self.segmentos = {s["segmento"]: s for s in self.bloco["segmentos"]}
        self.etfs = {e["cnpj"]: e for e in self.bloco["etfs"]}

    def test_entram_no_radar_com_o_pl_e_a_data_do_ultimo_informe(self):
        recentes = [
            (l["gestora"], l["segmento"], l["pl"], l["data_pl"], l["tipo"])
            for l in self.bloco["lancamentos"] if l["chave"] == "recente"
        ]
        tipo = "começou a operar depois do balancete de 31/08/2026 (Fundos.NET)"
        self.assertEqual(recentes, [
            ("XP (Trend)", "Renda fixa internacional", 5_000_000.0, "2026-09-16", tipo),
            ("BTG Pactual", "Tesouro Selic e CDI", 80_000_000.0, "2026-09-17", tipo),
        ])

    def test_ficam_fora_das_somas_de_pl(self):
        """O PL delas e de outra data e de outra fonte: somar misturaria as duas."""
        self.assertEqual(self.bloco["pl_total"], 24.7 * BI)
        self.assertEqual(self.bloco["etfs_com_pl"], 6)
        xp = next(g for g in self.bloco["gestoras"] if g["foco"])
        self.assertEqual((xp["etfs"], xp["recentes"]), (2, 1))

    def test_informe_vale_mais_que_a_situacao_do_cadastro(self):
        self.assertEqual(self.bloco["pre_operacionais"], 2)  # G e H; J ja opera
        self.assertFalse(self.etfs["J"]["pre_operacional"])
        self.assertTrue(self.etfs["J"]["recente"])

    def test_informe_anterior_a_data_base_nao_e_estreia(self):
        self.assertFalse(self.etfs["K"]["recente"])
        self.assertFalse(any(l["nome"].startswith("ISHARES S&P") for l in self.bloco["lancamentos"]))

    def test_casa_que_acabou_de_estrear_esta_entrando_no_segmento(self):
        segmento = self.segmentos["Renda fixa internacional"]
        self.assertEqual(segmento["foco_status"], "entrando")
        self.assertEqual(segmento["foco_recentes"], [
            "TREND ETF BLOOMBERG AGGREGATE US BONDS CLASSE DE ÍNDICE"
        ])
        selic = self.segmentos["Tesouro Selic e CDI"]
        self.assertEqual([p["gestora"] for p in selic["lancando"]], ["BTG Pactual"])

    def test_superficies_dizem_a_data_e_a_fonte(self):
        from shelf_monitor.market.etf_competicao import prazos_etf
        from shelf_monitor.report.etf_digest import render_etf_html, render_etf_markdown

        texto = " ".join(self.bloco["leitura"])
        self.assertIn(
            "Depois do balancete de 31/08/2026, 2 ETFs começaram a operar, com R$ 85 mi "
            "no último informe do Fundos.NET; 1 da XP (Trend), em Renda fixa internacional.",
            texto,
        )
        self.assertIn(
            "ETFs que começaram a operar depois do balancete: PL do último informe de cada "
            "um no Fundos.NET, com a data ao lado",
            prazos_etf(self.bloco),
        )
        md = render_etf_markdown(self.bloco, datetime(2026, 9, 24, 10, 0))
        self.assertIn("**Começaram a operar depois do balancete de 31/08/2026: 2 ETFs**", md)
        self.assertIn(
            "- **XP (Trend)** - AGGX11 - TREND ETF BLOOMBERG AGGREGATE US BONDS CLASSE DE ÍNDICE "
            "(Renda fixa internacional): R$ 5 mi em 16/09/2026",
            md,
        )
        html = render_etf_html(self.bloco, datetime(2026, 9, 24, 10, 0))
        self.assertIn("R$ 5 mi em 16/09/2026", html)

    def test_um_so_concorda_no_singular(self):
        catalogo, serie, informes = _cenario()
        catalogo["I"] = _etf("I", "TREND ETF BLOOMBERG AGGREGATE US BONDS CLASSE DE ÍNDICE", XP)
        informes.append(InformeDiarioEtf(cnpj="I", data=date(2026, 9, 16), pl=5e6, cota=100.0))
        bloco = build_etf_competicao(catalogo, serie, informes)
        self.assertIn("1 ETF começou a operar, com R$ 5 mi", " ".join(bloco["leitura"]))


class TestSuperficies(unittest.TestCase):
    def setUp(self):
        catalogo, serie, informes = _cenario()
        self.catalogo, self.serie = catalogo, serie
        self.bloco = build_etf_competicao(catalogo, serie, informes)
        self.config, self.tmp = temp_config()

    def tearDown(self):
        self.tmp.cleanup()

    def test_relatorio_markdown_e_html(self):
        from shelf_monitor.report.etf_digest import write_etf_digest

        md_path, html_path = write_etf_digest(self.config, self.bloco, datetime(2026, 9, 24, 10, 0))
        md = md_path.read_text(encoding="utf-8")
        self.assertEqual(md_path.name, "etf-competicao-2026-09-24.md")
        self.assertIn("Prazos: PL: balancete de FIE de ago/2026 (data-base 31/08/2026)", md)
        self.assertIn("comparação: de jun/2026 (data-base 30/06/2026) a ago/2026", md)
        self.assertIn(
            "**Começaram a operar entre os balancetes de 30/06/2026 e 31/08/2026: 1 ETF**", md
        )
        self.assertIn("| 3 | **XP (Trend)** | 2 |", md)
        self.assertIn("| Gestora | jun/2026 | ago/2026 |", md)
        self.assertIn("| **XP (Trend)** | 14,0% | 15,8% |", md)
        self.assertIn("## Oportunidades: onde a XP (Trend) é pequena ou não tem ETF", md)
        self.assertIn("**Ouro** - R$ 3,0 bi, 80,0% do segmento", md)
        self.assertIn("Prefixado (IRF-M e Tesouro pré): **XP**", md)
        self.assertIn("[Balancete de fundos estruturados da CVM](https://dados.cvm.gov.br/dataset/fie-doc-balancete)", md)
        html = html_path.read_text(encoding="utf-8")
        self.assertIn("<strong>Prazos:</strong>", html)
        self.assertIn('href="https://dados.cvm.gov.br/dataset/fie-doc-balancete"', html)

    def test_planilha_abre_pela_competicao(self):
        from openpyxl import load_workbook

        from shelf_monitor.report.etf_report import export_etf_excel

        caminho = export_etf_excel(
            [], [], self.catalogo, self.config.reports_dir / "etf.xlsx",
            pl_series=self.serie, fonte="fnet", competicao=self.bloco,
        )
        wb = load_workbook(caminho)
        self.assertEqual(
            wb.sheetnames[:5],
            ["Competição", "Segmentos", "Oportunidades e defesa", "Lançamentos", "ETFs por segmento"],
        )
        self.assertEqual(wb.active.title, "Competição")
        textos = [c.value for linha in wb["Competição"].iter_rows(max_row=4) for c in linha if c.value]
        self.assertIn("PL: balancete de FIE de ago/2026 (data-base 31/08/2026)", textos)

    def test_console_e_orquestracao(self):
        from shelf_monitor import main as main_mod

        with mock.patch.object(main_mod.fundosnet, "read_fnet_csvs", return_value=[]):
            with mock.patch.object(main_mod, "print_etf_competicao") as imprimir:
                bloco = main_mod._competicao_de_etf(self.config, self.catalogo, self.serie)
        self.assertTrue(bloco["available"])
        self.assertIsNone(bloco["cotistas"]["total"])  # sem informes, sem cotistas
        imprimir.assert_called_once_with(bloco)

    def test_console_nao_quebra(self):
        from shelf_monitor.report.etf_report import print_etf_competicao

        print_etf_competicao(self.bloco)
        print_etf_competicao({"available": False, "reason": "sem balancete"})


if __name__ == "__main__":
    unittest.main()
