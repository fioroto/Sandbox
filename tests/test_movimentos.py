"""Movimento de prateleira: uma definicao, e so quando ha comparacao.

Trava tres coisas da onda 1 da auditoria (docs/auditoria-fluxos.md):

- G3: coleta de casa que declara vitrine e nao trouxe nenhum fundo de vitrine
  e parcial. Em 2026-09-13 XP, BTG e Itau foram gravados OK so com o catalogo.
- G4: vitrine que aparece ou muda de escala e nova linha de base, nao entrada.
  Em 2026-09-19 o digest publicou "828 fundos entraram".
- G7: assunto, console, deck, dashboard e digest contam o mesmo movimento - o
  da vitrine de varejo. Medido no BTG de 13 para 19/09: 236 no diff inteiro,
  174 na vitrine, e cada superficie mostrava um dos dois.

Rode com: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import re
import unittest
from datetime import datetime, timedelta

from shelf_monitor.config import Config
from shelf_monitor.diff import assess_quality, diff_shelf
from shelf_monitor.models import (
    CollectorStrategy,
    FundOffer,
    RunReport,
    ShelfSnapshot,
    SnapshotStatus,
)
from shelf_monitor.report.aggregate import build_dashboard_data
from shelf_monitor.report.digest import _summary_counts, render_html, render_markdown
from shelf_monitor.report.mailer import build_subject
from shelf_monitor.report.shared import movimentos_de_vitrine

from fixtures_comuns import base_cnpj, temp_config

AGORA = datetime(2026, 9, 23, 10, 0)
ONTEM = AGORA - timedelta(days=1)
VITRINE = CollectorStrategy.API_JSON.value
CATALOGO = CollectorStrategy.OPEN_FINANCE.value


def _fundo(platform: str, nome: str, n: int, fonte: str, **campos) -> FundOffer:
    fundo = FundOffer(
        platform=platform, name=nome, cnpj=f"{base_cnpj(platform) + n:014d}", **campos
    )
    fundo.sources = [fonte]
    return fundo


def _serie(platform: str, n: int, fonte: str, inicio: int = 0) -> list[FundOffer]:
    return [
        _fundo(platform, f"{platform.upper()} Fundo {i:03d} FIC FIM", i, fonte)
        for i in range(inicio, inicio + n)
    ]


def _snap(platform: str, fundos: list[FundOffer], quando: datetime) -> ShelfSnapshot:
    return ShelfSnapshot(
        platform=platform, collected_at=quando,
        strategy=CollectorStrategy.API_JSON, status=SnapshotStatus.OK, funds=fundos,
    )


class TestColetaSemVitrine(unittest.TestCase):
    """G3: a coleta que falhou justamente na vitrine nao e OK."""

    def test_casa_com_vitrine_declarada_e_so_catalogo_vira_parcial(self):
        snap = _snap("xp", _serie("xp", 5, CATALOGO), AGORA)
        assess_quality(snap, None, espera_vitrine=True)
        self.assertEqual(snap.status, SnapshotStatus.PARTIAL)
        self.assertIn("nenhum fundo de vitrine", snap.error)

    def test_casa_sem_vitrine_declarada_segue_ok(self):
        # O Nu so declara Open Finance: catalogo sozinho e a coleta completa.
        snap = _snap("nu", _serie("nu", 5, CATALOGO), AGORA)
        assess_quality(snap, None, espera_vitrine=False)
        self.assertEqual(snap.status, SnapshotStatus.OK)

    def test_com_um_fundo_de_vitrine_a_regra_nao_dispara(self):
        fundos = _serie("xp", 5, CATALOGO) + _serie("xp", 1, VITRINE, inicio=5)
        snap = _snap("xp", fundos, AGORA)
        assess_quality(snap, None, espera_vitrine=True)
        self.assertEqual(snap.status, SnapshotStatus.OK)

    def test_a_declaracao_vem_das_estrategias_da_casa(self):
        config = Config()
        self.assertTrue(config.platform("xp").declara_vitrine)
        self.assertTrue(config.platform("btg").declara_vitrine)
        self.assertFalse(config.platform("nu").declara_vitrine)

    def test_a_coleta_grava_a_casa_sem_vitrine_como_parcial(self):
        from shelf_monitor.main import collect_platforms

        class _SoCatalogo:
            def __init__(self, platform_id):
                self.platform_id = platform_id

            def collect(self):
                return _snap(self.platform_id, _serie(self.platform_id, 5, CATALOGO), AGORA)

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        snapshots = collect_platforms(
            config, ["xp", "nu"], construtor=lambda c, pid, parar: _SoCatalogo(pid)
        )
        status = {s.platform: s.status for s in snapshots}
        self.assertEqual(status, {"xp": SnapshotStatus.PARTIAL, "nu": SnapshotStatus.OK})


class TestNovaLinhaDeBase(unittest.TestCase):
    """G4: a coleta que passou a ver a vitrine nao publica entradas."""

    def test_vitrine_que_aparece_nao_vira_entrada(self):
        catalogo = _serie("btg", 3, CATALOGO)
        anterior = _snap("btg", catalogo, ONTEM)
        atual = _snap("btg", catalogo + _serie("btg", 60, VITRINE, inicio=100), AGORA)

        diff = diff_shelf(atual, anterior)
        self.assertTrue(diff.nova_linha_de_base)
        self.assertTrue(diff.is_first_run)
        self.assertEqual((diff.added, diff.removed, diff.changed), ([], [], []))
        self.assertIn("de 0 para 60", diff.suspect_reason)

        report = RunReport(generated_at=AGORA, snapshots=[atual], shelf_diffs=[diff])
        markdown = render_markdown(report)
        self.assertIn("Nova linha de base", markdown)
        self.assertNotIn("fundos entraram", markdown)
        self.assertNotIn("Primeira coleta", markdown)
        self.assertIn("Nova linha de base", render_html(report))
        self.assertIn("nova linha de base: BTG Pactual", build_subject(report))

    def test_crescimento_real_medido_continua_movimento(self):
        # Inter, 31/08 -> 13/09: 399 -> 466 (+67, +17%) - o maior crescimento
        # real medido entre duas coletas com vitrine.
        anterior = _snap("inter", _serie("inter", 399, VITRINE), ONTEM)
        atual = _snap("inter", _serie("inter", 466, VITRINE), AGORA)
        diff = diff_shelf(atual, anterior)
        self.assertFalse(diff.nova_linha_de_base)
        self.assertEqual(len(diff.added), 67)

    def test_vitrine_pequena_dobrando_e_movimento(self):
        anterior = _snap("btg", _serie("btg", 3, VITRINE), ONTEM)
        atual = _snap("btg", _serie("btg", 6, VITRINE), AGORA)
        diff = diff_shelf(atual, anterior)
        self.assertFalse(diff.nova_linha_de_base)
        self.assertEqual(len(diff.added), 3)

    def test_salto_de_escala_vira_linha_de_base(self):
        anterior = _snap("btg", _serie("btg", 100, VITRINE), ONTEM)
        atual = _snap("btg", _serie("btg", 250, VITRINE), AGORA)
        diff = diff_shelf(atual, anterior)
        self.assertTrue(diff.nova_linha_de_base)
        self.assertEqual(diff.added, [])

    def test_casa_rebaseada_fica_fora_da_contagem_das_outras(self):
        xp_antes = _snap("xp", _serie("xp", 3, CATALOGO), ONTEM)
        xp_agora = _snap("xp", _serie("xp", 3, CATALOGO) + _serie("xp", 80, VITRINE, 100), AGORA)
        btg_antes = _snap("btg", _serie("btg", 10, VITRINE), ONTEM)
        btg_agora = _snap("btg", _serie("btg", 12, VITRINE), AGORA)
        report = RunReport(
            generated_at=AGORA, snapshots=[xp_agora, btg_agora],
            shelf_diffs=[diff_shelf(xp_agora, xp_antes), diff_shelf(btg_agora, btg_antes)],
        )
        markdown = render_markdown(report)
        self.assertIn("2 fundos entraram", markdown)
        self.assertIn("XP Investimentos virou nova linha de base", markdown)
        self.assertIn("fica fora da contagem", markdown)
        subject = build_subject(report)
        self.assertIn("2 entradas", subject)
        self.assertIn("nova linha de base: XP Investimentos", subject)


class TestUmaDefinicaoDeMovimento(unittest.TestCase):
    """G7: as cinco superficies contam o mesmo numero."""

    def setUp(self):
        a_antes = _fundo("btg", "BTG Alfa FIC FIM", 1, VITRINE, min_initial=1_000.0)
        saiu = _fundo("btg", "BTG Beta FIC FIM", 2, VITRINE)
        cat_antes = _fundo("btg", "BTG Catálogo FIC FIM", 3, CATALOGO, min_initial=5_000.0)
        anterior = _snap("btg", [a_antes, saiu, cat_antes], ONTEM)

        a_agora = _fundo("btg", "BTG Alfa FIC FIM", 1, VITRINE, min_initial=500.0)
        entrou = _fundo("btg", "BTG Novo Ações FIC FIA", 4, VITRINE)
        restrito = _fundo("btg", "BTG Exclusivo Família FIM", 5, VITRINE)
        cat_agora = _fundo("btg", "BTG Catálogo FIC FIM", 3, CATALOGO, min_initial=1_000.0)
        so_catalogo = _fundo("btg", "BTG Carteira Fechada FIC FIM", 6, CATALOGO)
        self.atual = _snap("btg", [a_agora, entrou, restrito, cat_agora, so_catalogo], AGORA)

        self.diff = diff_shelf(self.atual, anterior)
        self.report = RunReport(
            generated_at=AGORA, snapshots=[self.atual], shelf_diffs=[self.diff]
        )

    def test_o_diff_inteiro_e_maior_que_a_vitrine(self):
        # Premissa do cenario: sem ela o teste nao distingue as duas contas.
        self.assertEqual(
            (len(self.diff.added), len(self.diff.removed), len(self.diff.changed)),
            (3, 1, 2),
        )
        casa = movimentos_de_vitrine(self.diff, self.atual)
        self.assertEqual([f.name for f in casa.entradas], ["BTG Novo Ações FIC FIA"])
        self.assertEqual([f.name for f in casa.saidas], ["BTG Beta FIC FIM"])
        self.assertEqual([c.name for c in casa.alteracoes], ["BTG Alfa FIC FIM"])
        self.assertEqual(casa.so_catalogo, 3)

    def test_digest_e_assunto(self):
        self.assertEqual(_summary_counts(self.report), (1, 1, 1, 3))
        # Vitrine de varejo 2 (o Exclusivo e restrito), catalogo de varejo 4.
        self.assertIn("| BTG Pactual | 2 | 4 | ok | 1 | 1 | 1 |", render_markdown(self.report))
        subject = build_subject(self.report)
        self.assertIn("1 entrada / 1 saída", subject)
        self.assertIn("1 alteração", subject)

    def test_linha_da_casa_concorda_com_o_numero(self):
        # Em 2026-09-24 o digest publicou "1 entradas", "1 saídas" e "1 alterações".
        from shelf_monitor.report.digest import render_html

        esperado = "1 entrada, 1 saída, 1 alteração na vitrine"
        self.assertIn(esperado, render_markdown(self.report))
        # No e-mail a linha da casa virou tabela (secao 5 da linguagem visual):
        # um movimento por linha, com o fundo, e as contagens que restam em
        # texto continuam concordando com o numero.
        html = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", render_html(self.report)))
        self.assertIn("Entrou BTG Novo Ações FIC FIA", html)
        self.assertIn("Saiu BTG Beta FIC FIM", html)
        self.assertIn("Alterou BTG Alfa FIC FIM", html)
        self.assertIn("1 alteração de condições", html)
        self.assertIn("Só no catálogo Open Finance, fora da vitrine: BTG Pactual 3", html)
        for errado in ("1 entradas", "1 saídas", "1 alterações"):
            self.assertNotIn(errado, html)
        self.assertIn("Mais 3 mudanças só no catálogo", render_markdown(self.report))
        self.assertIn(
            "1 fundo entrou, 1 saiu e 1 teve condições alteradas", render_markdown(self.report)
        )

    def test_payload_do_dashboard(self):
        [resumo] = build_dashboard_data(self.report)["diff_summary"]
        self.assertEqual(
            (resumo["added"], resumo["removed"], resumo["changed"], resumo["so_catalogo"]),
            (1, 1, 1, 3),
        )

    def test_resumo_do_console(self):
        from shelf_monitor import main

        with main.console.capture() as saida:
            main.print_summary(self.report)
        linha = next(l for l in saida.get().splitlines() if "BTG Pactual" in l)
        # Fundos (5 coletados), entradas, saidas, alteracoes - so os numeros da
        # linha, qualquer que seja a moldura que o rich desenhe.
        self.assertEqual([n for n in re.split(r"[^0-9]+", linha) if n], ["5", "1", "1", "1"])

    def test_slide_de_movimentos(self):
        from pptx import Presentation

        from shelf_monitor.report.deck import render_deck

        config, tmp = temp_config()
        self.addCleanup(tmp.cleanup)
        deck = Presentation(render_deck(config, self.report))
        tabelas = [
            shape.table
            for slide in deck.slides
            if any(
                s.has_text_frame and "Movimentos desde" in s.text_frame.text
                for s in slide.shapes
            )
            for shape in slide.shapes
            if shape.has_table
        ]
        [tabela] = tabelas
        linha = [tabela.cell(1, c).text for c in range(4)]
        self.assertEqual(linha[:3], ["BTG Pactual", "1", "1"])
        # O exemplo de entrada e da vitrine, nunca do catalogo nem restrito.
        self.assertIn("BTG Novo Ações FIC FIA", linha[3])
        self.assertNotIn("Carteira Fechada", linha[3])
        self.assertNotIn("Exclusivo", linha[3])


if __name__ == "__main__":
    unittest.main()
